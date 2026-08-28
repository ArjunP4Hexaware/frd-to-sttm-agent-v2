"""Redraw slides 6 and 7 of the ACFC program-overview deck (PPTX).

    ~/.virtualenvs/frdsttm/bin/python context/build/build_overview_arch_slides.py [in.pptx] [out.pptx]

Default input : context/ACFC_AI_in_Engineering_Program_Overview_New Draft.pptx
Default output: context/ACFC_AI_in_Engineering_Program_Overview_New Draft v2.pptx

The input deck is never written to. Both slides keep the deck's own chrome —
the blue header band with the flag wave is the slide BACKGROUND, so clearing
the shapes leaves it in place — and match the deck's type: 20pt white
uppercase title in the band, 15pt blue subtitle, Calibri throughout. Palette
and primitives come from acfc_theme (square corners, flat bands, flag red as
the one attention colour).

WHAT THE SLIDES SHOW (Arjun, 2026-08-28): the FINAL architecture of the agent.
Integrations that are designed but not yet built — the SharePoint sync through
Graph, the Entra ID gate, the audit trail, the Collibra register — are drawn
as part of the system; Arjun says in the demo which of them are still to come.
The core (two documents paired by name, four volumes + one table, the one
model call, assess → answer → render, the review app) is `staging` as of
2026-08-28.

Slide 6 — solution architecture, left → right:
    SharePoint (system of record) → Unity Catalog + the job + the review app
    (Databricks) → the draft workbook → a named reviewer → CodeGen.
Slide 7 — data governance architecture, left → right, the previous one-pager's
    shape: four stations in the order a document travels, the audit trail
    beneath them, the two boundary crossings, Collibra as the register on the
    right.
"""
from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

sys.path.insert(0, str(Path(__file__).resolve().parent))
from acfc_theme import (  # noqa: E402
    BLUE, NAVY, SKY, RED, INK, SLATE, MUTED, BAND, HAIR, HAIR2, WHITE,
    TINT_BLUE, TINT_SKY, TINT_RED, MONO, W_IN,
    rect, line, text as _text, arrow, logo, icon,
)

REPO = Path(__file__).resolve().parent.parent.parent
CONTEXT = REPO / "context"
IN_DEFAULT = CONTEXT / "ACFC_AI_in_Engineering_Program_Overview_New Draft.pptx"
OUT_DEFAULT = CONTEXT / "ACFC_AI_in_Engineering_Program_Overview_New Draft v2.pptx"

FONT = "Calibri"          # the deck's face; acfc_theme's Segoe UI is for the standalone one-pagers
L, R = 0.42, W_IN - 0.42
CW = R - L
STRAP_Y = 6.30            # the closing strap; the deck's footer rule sits at ~7.1


def text(s, x, y, w, h, runs, **kw):
    kw.setdefault("font", FONT)
    return _text(s, x, y, w, h, runs, **kw)


def kicker(s, x, y, w, label, *, color=BLUE, size=8, align=PP_ALIGN.LEFT):
    text(s, x, y, w, 0.2, [(label.upper(), {})], size=size, color=color, bold=True,
         spacing=1.4, align=align)


def chip(s, x, y, label, *, color=BLUE, size=7.5):
    """The site's signature control: square, outlined, unfilled."""
    w = 0.062 * len(label) + 0.3
    rect(s, x, y, w, 0.26, fill=None, line=color, line_w=1.0)
    text(s, x, y + 0.045, w, 0.18, [(label.upper(), {})], size=size, color=color,
         bold=True, spacing=0.8, align=PP_ALIGN.CENTER)
    return x + w + 0.1


def clear(slide):
    """Remove every shape; the background (the ACFC chrome) is not a shape."""
    for sh in list(slide.shapes):
        sh._element.getparent().remove(sh._element)


def head(s, title, sub):
    text(s, 0.36, 0.20, 12.4, 0.55, [(title.upper(), {})], size=20, color=WHITE,
         bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 0.36, 0.98, 12.6, 0.34, sub, size=15, color=BLUE, bold=True)


def strap(s, parts):
    rect(s, L, STRAP_Y, CW, 0.42, fill=NAVY)
    text(s, L, STRAP_Y, CW, 0.42, [("     ·     ".join(p.upper() for p in parts), {})],
         size=8.5, color=WHITE, bold=True, spacing=1.2, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE)


# --------------------------------------------------------------------------- #
# slide 6 — solution architecture
# --------------------------------------------------------------------------- #
def solution(s):
    clear(s)
    head(s, "FRD → STTM Agent  ·  Solution architecture",
         "Two documents in, one model call, one workbook out — a person approves every hand-off")

    AX, AW = L, 2.75                 # system of record
    BX, BW = 3.55, 5.95              # Databricks platform
    CX, CW_ = 9.88, R - 9.88         # output · downstream
    TOP = 1.78

    kicker(s, AX, 1.50, AW, "Source systems")
    kicker(s, BX, 1.50, BW, "Databricks platform — where the agent runs")
    kicker(s, CX, 1.50, CW_, "Output  ·  downstream")

    # ---- A: source systems -------------------------------------------------
    # SharePoint — the system of record for the documents themselves
    rect(s, AX, TOP, AW, 1.86, fill=BAND)
    logo(s, "sharepoint", AX + 0.12, TOP + 0.08, 0.48, 0.48)
    text(s, AX + 0.7, TOP + 0.09, AW - 0.84, 0.24, [("Microsoft SharePoint", {})],
         size=10.5, color=INK, bold=True)
    text(s, AX + 0.7, TOP + 0.31, AW - 0.84, 0.2,
         [("system of record — where documents are approved", {})], size=6.8, color=MUTED)
    docs = [
        ("word", "FRD_<source>.docx", "requirements — written by the BSA"),
        ("excel", "VDD_<source>.xlsx", "the vendor's data dictionary"),
        ("excel", "STTM_<source>.xlsx", "approved workbooks — layout only"),
    ]
    y = TOP + 0.62
    for lg, name, desc in docs:
        logo(s, lg, AX + 0.14, y, 0.34, 0.34)
        text(s, AX + 0.6, y - 0.01, AW - 0.7, 0.2, [(name, {"font": MONO})],
             size=7.8, color=BLUE, bold=True)
        text(s, AX + 0.6, y + 0.16, AW - 0.7, 0.2, desc, size=6.8, color=MUTED)
        y += 0.42

    # Jira and Confluence — the requirement's context, read the same way
    others = [
        ("jira", "Jira board / Azure DevOps", "requirement stories and acceptance criteria behind the FRD"),
        ("confluence", "Confluence / other", "source-system notes, run books — same connector pattern"),
    ]
    y = TOP + 1.96
    for lg, name, desc in others:
        rect(s, AX, y, AW, 0.52, fill=BAND)
        logo(s, lg, AX + 0.12, y + 0.05, 0.42, 0.42)
        text(s, AX + 0.66, y + 0.05, AW - 0.8, 0.2, [(name, {})], size=8.5, color=INK, bold=True)
        text(s, AX + 0.66, y + 0.25, AW - 0.8, 0.28, desc, size=6.8, color=MUTED)
        y += 0.62

    # one connector pattern, one identity, for all three
    CY = TOP + 3.2
    rect(s, AX, CY, AW, 1.12, fill=None, line=BLUE)
    logo(s, "entra", AX + 0.12, CY + 0.1, 0.44, 0.44)
    text(s, AX + 0.66, CY + 0.08, AW - 0.8, 0.2, [("One connector layer", {})],
         size=8.5, color=INK, bold=True)
    text(s, AX + 0.66, CY + 0.28, AW - 0.8, 0.54,
         "Entra ID app · Microsoft Graph + Atlassian REST · scoped read-only · one write "
         "folder for approved workbooks.",
         size=6.8, color=SLATE)

    text(s, AX + 0.14, CY + 0.78, AW - 0.28, 0.32,
         [[("00  sync on start-up", {"bold": True, "color": BLUE}),
           ("  —  new FRD_* / VDD_* / STTM_* files are imported, paired by name — "
            "never by similarity.", {})]],
         size=6.8, color=SLATE)

    arrow(s, AX + AW, TOP + 0.55, BX, TOP + 0.55, width=2)

    # ---- B: Databricks platform band ---------------------------------------
    rect(s, BX, TOP, BW, 4.32, fill=BAND)
    P = 0.12
    IX, IW = BX + P, BW - 2 * P

    # Unity Catalog
    rect(s, IX, TOP + P, IW, 1.0, fill=WHITE, line=HAIR2)
    logo(s, "unitycatalog", IX + 0.1, TOP + P + 0.07, 1.15, 0.38)
    text(s, IX + 1.38, TOP + P + 0.1, IW - 1.5, 0.26,
         [("volumes + Delta tables", {"size": 9, "color": MUTED})],
         color=INK)
    text(s, IX + 0.1, TOP + P + 0.52, IW - 0.2, 0.2,
         [("frds  ·  vdds  ·  reference_sttms  ·  output_sttms", {"font": MONO}),
          ("      +  table  ", {"color": MUTED}), ("frd_pairing", {"font": MONO})],
         size=8, color=BLUE, bold=True)
    text(s, IX + 0.1, TOP + P + 0.74, IW - 0.2, 0.26,
         "Every file the agent reads or writes lives here. The review app reads this — "
         "SharePoint is never on the request path.", size=7.5, color=SLATE)

    # standards strip
    SY = TOP + P + 1.1
    rect(s, IX, SY, IW, 0.4, fill=None, line=BLUE)
    icon(s, "braces", IX + 0.1, SY + 0.08, 0.24)
    text(s, IX + 0.44, SY + 0.04, IW - 0.55, 0.34,
         [("ACFC standards as config", {"bold": True, "color": INK}),
          ("  —  ", {"color": MUTED}),
          ("naming_standards.json · engineering_standards.json", {"font": MONO, "color": BLUE}),
          ("  —  catalog and schema per layer, stage type, promoted type, audit columns", {})],
         size=7.5, color=SLATE, anchor=MSO_ANCHOR.MIDDLE)

    # generation job
    JY = SY + 0.53
    rect(s, IX, JY, IW, 1.62, fill=WHITE, line=HAIR2)
    text(s, IX + 0.1, JY + 0.07, IW - 0.2, 0.24,
         [("Generation job", {"bold": True, "size": 10.5}),
          ("    Databricks Jobs · serverless · ", {"size": 8, "color": MUTED}),
          ("frd_sttm_pipeline", {"size": 8, "color": MUTED, "font": MONO})],
         color=INK)
    stages = [
        ("01", "Extract", "docx → markdown; the one Claude call returns the source-level facts. The VDD is parsed by code.", True),
        ("02", "Assess", "Every identifier grounded verbatim in the FRD; sources paired to dictionaries; targets per layer; blockers.", False),
        ("03", "Answer", "ready · needs input · cannot generate. The reviewer settles each question; answers live on run.json.", False),
        ("04", "Render", "Rows from the VDD, targets from the FRD and standards, rules on the row they name — in an approved layout.", False),
    ]
    gap = 0.13
    sw = (IW - 0.2 - 3 * gap) / 4
    for i, (n, name, desc, model) in enumerate(stages):
        x = IX + 0.1 + i * (sw + gap)
        rect(s, x, JY + 0.4, sw, 1.14, fill=TINT_RED if model else TINT_BLUE,
             line=RED if model else None)
        text(s, x + 0.08, JY + 0.45, sw - 0.16, 0.18, [(n, {})], size=7.5,
             color=RED if model else BLUE, bold=True, spacing=1.2)
        text(s, x + 0.08, JY + 0.6, sw - 0.16, 0.22, [(name, {})], size=10, color=INK, bold=True)
        text(s, x + 0.08, JY + 0.83, sw - 0.16, 0.7, desc, size=6.8, color=SLATE)
        if model:
            text(s, x + 0.08, JY + 1.36, sw - 0.16, 0.16, [("the one model call", {})],
                 size=6.5, color=RED, bold=True)
        if i < 3:
            arrow(s, x + sw + 0.01, JY + 0.97, x + sw + gap - 0.01, JY + 0.97, width=1.25)

    # model access + review app
    MY = JY + 1.75
    half = (IW - 0.12) / 2
    rect(s, IX, MY, half, 0.78, fill=WHITE, line=HAIR2)
    logo(s, "claude", IX + 0.1, MY + 0.08, 1.0, 0.22)
    text(s, IX + 1.2, MY + 0.07, half - 1.3, 0.2,
         [("via Foundation Model APIs", {})], size=8.5, color=INK, bold=True)
    text(s, IX + 0.1, MY + 0.34, half - 0.2, 0.44,
         [[("databricks-claude-opus-5", {"font": MONO, "color": BLUE}),
           ("  ·  through AI Gateway: logging, rate limits, guardrails on every call  ·  "
            "inside the client's workspace — no external key", {})]],
         size=7, color=SLATE)
    rx = IX + half + 0.12
    rect(s, rx, MY, half, 0.78, fill=WHITE, line=HAIR2)
    logo(s, "databricks", rx + 0.1, MY + 0.04, 0.56, 0.3)
    text(s, rx + 0.76, MY + 0.07, half - 0.86, 0.2,
         [("Review app", {}), ("    Databricks Apps", {"bold": False, "color": MUTED})],
         size=8.5, color=INK, bold=True)
    text(s, rx + 0.1, MY + 0.34, half - 0.2, 0.44,
         "Pick a source · run (confirm-gated) · answer what the agent will not guess · "
         "re-render · download. Nothing is invented on the reviewer's behalf.",
         size=7, color=SLATE)

    # ---- C: output · downstream --------------------------------------------
    rect(s, CX, TOP, CW_, 1.2, fill=None, line=BLUE)
    logo(s, "excel", CX + 0.12, TOP + 0.1, 0.46, 0.46)
    text(s, CX + 0.68, TOP + 0.12, CW_ - 0.82, 0.26, [("Draft STTM workbook", {})],
         size=10.5, color=INK, bold=True)
    text(s, CX + 0.68, TOP + 0.38, CW_ - 0.82, 0.2,
         [(".xlsx  ·  ACFC layout  ·  open questions flagged", {})], size=7, color=BLUE, bold=True)
    text(s, CX + 0.14, TOP + 0.62, CW_ - 0.28, 0.55,
         "Every cell comes from the FRD, the VDD or the standards — target column names "
         "carried as-is from the dictionary. What cannot be sourced is a question, not a guess.",
         size=7.2, color=SLATE)
    arrow(s, BX + BW, TOP + 0.55, CX, TOP + 0.55, width=2)

    HY = TOP + 1.38
    rect(s, CX, HY, CW_, 1.5, fill=NAVY)
    icon(s, "user", CX + 0.14, HY + 0.14, 0.3)
    text(s, CX + 0.54, HY + 0.13, CW_ - 0.68, 0.26, [("A named reviewer decides", {})],
         size=10.5, color=WHITE, bold=True)
    text(s, CX + 0.14, HY + 0.46, CW_ - 0.28, 1.0,
         "Resolves each open question, edits in Excel if needed, approves — and returns the "
         "workbook to SharePoint. The next sync pairs it: it becomes a layout other sources "
         "may borrow, never content. Nothing moves downstream unapproved.",
         size=7.5, color=WHITE)
    arrow(s, CX + CW_ / 2, TOP + 1.2, CX + CW_ / 2, HY, width=1.5)

    DY = HY + 1.68
    rect(s, CX, DY, CW_, 1.26, fill=BAND)
    kicker(s, CX + 0.14, DY + 0.1, CW_ - 0.28, "Downstream", color=MUTED, size=7)
    icon(s, "code", CX + 0.14, DY + 0.36, 0.28)
    text(s, CX + 0.52, DY + 0.34, CW_ - 0.66, 0.24, [("CodeGen agent", {}), ("   Agent 2", {"bold": False, "color": MUTED})],
         size=9.5, color=INK, bold=True)
    text(s, CX + 0.14, DY + 0.64, CW_ - 0.28, 0.6,
         "Reads the approved workbook — its only input from this agent — and emits the "
         "pipeline code. Unity Catalog carries the lineage from FRD to table.",
         size=7.2, color=SLATE)
    arrow(s, CX + CW_ / 2, HY + 1.5, CX + CW_ / 2, DY, width=1.5)

    strap(s, ["one model call per run", "everything else is deterministic code",
              "nothing is invented", "a person approves"])


# --------------------------------------------------------------------------- #
# slide 7 — data governance architecture
# --------------------------------------------------------------------------- #
def governance(s):
    clear(s)
    head(s, "FRD → STTM Agent  ·  Data governance architecture",
         "Classified data · named people · traceable outputs · a person decides · a record of everything")

    TOP, SH = 1.76, 2.8
    SW, GAP = 2.2, 0.2
    GX = 10.05                      # the register column
    GW = R - GX

    kicker(s, L, 1.50, 6, "The document's journey  →")
    kicker(s, GX, 1.50, GW, "→  Registered in")

    stations = [
        dict(n="01", verb="Enters", q="What data?", logos=["sharepoint", "unitycatalog"],
             title="Classified at the door",
             body="Two documents per source — the FRD and the vendor's data dictionary — arrive "
                  "from SharePoint read-only and land in Unity Catalog volumes. Every volume and "
                  "table carries an owner, a steward, its sensitivity (the dictionary flags PHI "
                  "columns) and a retention rule, set by the data owner.",
             chips="SharePoint read-only · Unity Catalog tags"),
        dict(n="02", verb="Opened", q="Who?", logos=["entra", "databricks"],
             title="Only named people",
             body="Only the client's BSA group — the people already allowed to read this data — "
                  "can open the review app. Each signs in with the company login; nothing runs, "
                  "is decided or is handed out without that name. May run equals may read.",
             chips="Microsoft Entra ID group · Databricks Apps"),
        dict(n="03", verb="Generated", q="Where from?", logos=["claude"], accent=True,
             title="Every cell has a source",
             body="One Claude call reads the FRD — the only one in the system. Every identifier "
                  "it returns must appear word-for-word in the FRD or becomes a question. Column "
                  "names are carried as-is from the dictionary; the agent never opens the "
                  "source's own approved workbook; a missing dictionary blocks the run.",
             chips="verbatim grounding · three sources or a question"),
        dict(n="04", verb="Decided", q="Who decides?", icons=["check"],
             title="A person decides",
             body="The AI only proposes. Code computes the verdict — ready, needs input, cannot "
                  "generate — and lists what it would not guess. A named reviewer answers each "
                  "question (keep · remove · correct); the answers are recorded and survive a "
                  "re-run of the same document. Approval is the reviewer's upload.",
             chips="run.json · human-in-the-loop"),
    ]
    SLOT = {  # (w, h, y-offset) — the wide lockups need width, not a square
        "sharepoint": (0.48, 0.48, 0.0), "entra": (0.48, 0.48, 0.0),
        "unitycatalog": (1.3, 0.46, 0.01), "databricks": (0.8, 0.46, 0.01),
        "claude": (1.5, 0.34, 0.07),
    }
    for i, st in enumerate(stations):
        x = L + i * (SW + GAP)
        acc = st.get("accent", False)
        rect(s, x, TOP, SW, SH, fill=TINT_RED if acc else BAND, line=RED if acc else None)
        c = RED if acc else BLUE
        kicker(s, x + 0.14, TOP + 0.12, SW - 0.28, f"{st['n']}   {st['verb']}", color=c)
        lx = x + 0.14
        for lg in st.get("logos", []):
            w, h, dy = SLOT.get(lg, (0.48, 0.48, 0.0))
            logo(s, lg, lx, TOP + 0.36 + dy, w, h)
            lx += w + 0.16
        for ic in st.get("icons", []):
            icon(s, ic, lx, TOP + 0.36, 0.48)
            lx += 0.64
        kicker(s, x + 0.14, TOP + 0.88, SW - 0.28, st["q"], color=MUTED, size=7)
        text(s, x + 0.14, TOP + 1.06, SW - 0.28, 0.26, [(st["title"], {})], size=10.5,
             color=RED if acc else INK, bold=True)
        text(s, x + 0.14, TOP + 1.36, SW - 0.28, 1.15, st["body"], size=7, color=SLATE)
        line(s, x + 0.14, TOP + SH - 0.34, x + SW - 0.14, TOP + SH - 0.34, color=HAIR2)
        text(s, x + 0.14, TOP + SH - 0.28, SW - 0.28, 0.22, [(st["chips"], {})], size=6.8,
             color=c, bold=True)
        if i < 3:
            arrow(s, x + SW + 0.01, TOP + 1.2, x + SW + GAP - 0.01, TOP + 1.2, width=1.5)

    JR = L + 4 * SW + 3 * GAP        # right edge of the journey (9.82)

    # ---- audit bar -----------------------------------------------------------
    AY, AH = TOP + SH + 0.14, 0.58
    rect(s, L, AY, JR - L, AH, fill=None, line=BLUE)
    icon(s, "shield", L + 0.14, AY + 0.13, 0.32)
    kicker(s, L + 0.58, AY + 0.08, 2.4, "What happened?", color=MUTED, size=7)
    text(s, L + 0.58, AY + 0.26, 2.6, 0.26, [("Nothing is off the record", {})],
         size=10.5, color=INK, bold=True)
    text(s, L + 3.25, AY + 0.07, JR - L - 3.4, 0.36,
         "Every run, every reviewer answer, every render and download is written as a record "
         "in Unity Catalog the moment it happens. The trail reads like a table and cannot be "
         "edited afterwards.", size=7.2, color=SLATE)
    text(s, L + 3.25, AY + 0.4, 4, 0.18, [("Audit trail  ·  Delta run log  ·  a person's name on every action", {})],
         size=6.8, color=BLUE, bold=True)

    # ---- the two boundary crossings -------------------------------------------
    XY, XH = AY + AH + 0.14, 0.74
    text(s, L, XY + 0.04, 2.95, XH,
         [[("Client data crosses the platform boundary in exactly two places", {"bold": True, "color": INK}),
           (" — both under the stations where they happen, both recorded.", {})]],
         size=7.5, color=SLATE)
    bw = (JR - L - 3.15 - 0.16) / 2
    b1 = L + 3.15
    rect(s, b1, XY, bw, XH, fill=TINT_RED, line=RED)
    text(s, b1 + 0.12, XY + 0.06, bw - 0.24, 0.2, [("1   FRD text  →  Claude", {})], size=7.5, color=RED, bold=True)
    text(s, b1 + 0.12, XY + 0.26, bw - 0.24, 0.46,
         "Through the workspace's Foundation Model APIs and AI Gateway — inside the client's "
         "Databricks tenant, logged per call, no external key. Within the client's "
         "data-handling review.", size=6.8, color=SLATE)
    b2 = b1 + bw + 0.16
    rect(s, b2, XY, bw, XH, fill=TINT_BLUE, line=BLUE)
    text(s, b2 + 0.12, XY + 0.06, bw - 0.24, 0.2, [("2   workbook  →  a named person", {})], size=7.5, color=BLUE, bold=True)
    text(s, b2 + 0.12, XY + 0.26, bw - 0.24, 0.46,
         "The download is recorded with the workbook's fingerprint and the reviewer's name. "
         "The reviewer — not the agent — returns the approved workbook to SharePoint.",
         size=6.8, color=SLATE)

    # ---- Collibra register column ------------------------------------------
    GH = XY + XH - TOP
    rect(s, GX, TOP, GW, GH, fill=None, line=NAVY, line_w=1.25)
    logo(s, "collibra", GX + 0.14, TOP + 0.14, GW - 0.28, 0.42)
    kicker(s, GX + 0.14, TOP + 0.68, GW - 0.28, "Governance register", color=MUTED, size=7)
    text(s, GX + 0.14, TOP + 0.88, GW - 0.28, 0.72,
         "The client's data-governance catalog. The agent is registered here as an AI asset — "
         "one place to answer “what is this agent, and is it under control?”",
         size=7.2, color=SLATE)
    uy = TOP + 1.5
    rect(s, GX + 0.14, uy, GW - 0.28, 1.02, fill=BAND)
    logo(s, "unitycatalog", GX + 0.24, uy + 0.08, 1.1, 0.36)
    arrow(s, GX + 1.45, uy + 0.26, GX + 2.0, uy + 0.26, width=1.25)
    text(s, GX + 0.24, uy + 0.44, GW - 0.48, 0.56,
         [[("Collibra harvests Unity Catalog", {"bold": True, "color": INK}),
           (" — owners, sensitivity and descriptions set in Databricks appear here "
            "automatically. Classified once, visible in both.", {})]],
         size=6.8, color=SLATE)
    ky = uy + 1.18
    kicker(s, GX + 0.14, ky, GW - 0.28, "What is registered", color=MUTED, size=7)
    items = ["What the agent is for, and its risk tier", "Who owns it and who stewards it",
             "The data it reads and writes, and where it flows", "The five controls shown to the left"]
    for j, it in enumerate(items):
        yy = ky + 0.26 + j * 0.3
        rect(s, GX + 0.16, yy + 0.05, 0.09, 0.09, fill=SKY)
        text(s, GX + 0.34, yy, GW - 0.48, 0.28, it, size=7.2, color=SLATE)

    arrow(s, JR + 0.01, TOP + 1.2, GX - 0.01, TOP + 1.2, width=1.5)
    arrow(s, JR + 0.01, AY + AH / 2, GX - 0.01, AY + AH / 2, width=1.5)

    strap(s, ["classified data", "named people", "traceable outputs", "a person decides",
              "a record of everything"])


def main(argv):
    src = Path(argv[1]) if len(argv) > 1 else IN_DEFAULT
    out = Path(argv[2]) if len(argv) > 2 else OUT_DEFAULT
    prs = Presentation(str(src))
    solution(prs.slides[5])
    governance(prs.slides[6])
    prs.save(str(out))
    print(out)


if __name__ == "__main__":
    main(sys.argv)
