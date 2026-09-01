"""Build the FRD->STTM client presentation deck (PPTX).

Regenerates docs/FRD_to_STTM_Agent_Demo.pptx -- the 13-slide plain-language
deck used alongside the live demo (docs/DEMO_RUNBOOK.md). Every figure quoted
on the slides comes from docs/LIVE_E2E_2026-08-07.md; update that document
first, then this script, so the two never drift.

    pip install python-pptx
    python scripts/build_presentation.py docs/FRD_to_STTM_Agent_Demo.pptx

Layout is absolute-positioned (no template master), so text length drives
fit: content must stay above y=6.80in and card bodies are sized for the copy
they currently hold. Shorten copy rather than growing cards when editing.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import sys

OUT = sys.argv[1] if len(sys.argv) > 1 else "FRD_to_STTM_Agent.pptx"

# ---------------------------------------------------------------- palette
NAVY   = RGBColor(0x0E, 0x22, 0x33)
INK    = RGBColor(0x1A, 0x24, 0x33)
BODY   = RGBColor(0x3A, 0x4A, 0x5C)
MUTED  = RGBColor(0x6B, 0x7A, 0x8C)
TEAL   = RGBColor(0x00, 0xA1, 0x9A)
BLUE   = RGBColor(0x2E, 0x6E, 0xD1)
AMBER  = RGBColor(0xD9, 0x8A, 0x1F)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
PAPER  = RGBColor(0xF4, 0xF7, 0xF9)
LINE   = RGBColor(0xD8, 0xE0, 0xE7)
NAVY_L = RGBColor(0x1B, 0x3A, 0x52)
PALE   = RGBColor(0xC3, 0xD2, 0xDE)
PALER  = RGBColor(0x9F, 0xB4, 0xC4)

FONT = "Calibri"
FONT_H = "Calibri Light"

W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.85)
CONTENT_W = W - 2 * MARGIN
BOTTOM = Inches(6.80)          # content must not cross this

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]


# ---------------------------------------------------------------- helpers
def slide():
    return prs.slides.add_slide(BLANK)


def rect(s, x, y, w, h, fill=None, line=None, lw=1.0, shape=MSO_SHAPE.RECTANGLE,
         adj=None):
    sh = s.shapes.add_shape(shape, x, y, w, h)
    if adj is not None:
        try:
            sh.adjustments[0] = adj
        except Exception:
            pass
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    if sh.has_text_frame:
        sh.text_frame.word_wrap = True
    return sh


def tbox(s, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    return tf


def para(tf, text, size=16, color=BODY, bold=False, font=FONT, align=PP_ALIGN.LEFT,
         space_before=0, space_after=6, first=False, line_spacing=1.15,
         italic=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.line_spacing = line_spacing
    r = p.add_run()
    r.text = text
    f = r.font
    f.name, f.size, f.bold, f.italic = font, Pt(size), bold, italic
    f.color.rgb = color
    return p


def rich(tf, chunks, size=16, align=PP_ALIGN.LEFT, space_before=0, space_after=6,
         first=False, line_spacing=1.2):
    """chunks: list of (text, color, bold)."""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.line_spacing = line_spacing
    for text, color, bold in chunks:
        r = p.add_run()
        r.text = text
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
    return p


def header(s, kicker, title, sub=None):
    """Standard content-slide header. Returns y where content may start."""
    rect(s, Inches(0), Inches(0), W, H, fill=WHITE)
    rect(s, Inches(0), Inches(0), Inches(0.16), H, fill=TEAL)
    tf = tbox(s, MARGIN, Inches(0.50), CONTENT_W, Inches(0.26))
    para(tf, kicker, size=11.5, color=TEAL, bold=True, first=True, space_after=0)
    tf2 = tbox(s, MARGIN, Inches(0.83), CONTENT_W, Inches(0.55))
    para(tf2, title, size=32, color=INK, font=FONT_H, first=True, space_after=0)
    if sub:
        tf3 = tbox(s, MARGIN, Inches(1.45), CONTENT_W - Inches(1.2), Inches(0.5))
        para(tf3, sub, size=14, color=MUTED, first=True, space_after=0,
             line_spacing=1.2)
        rect(s, MARGIN, Inches(2.10), Inches(1.1), Pt(2.5), fill=LINE)
        return Inches(2.30)
    rect(s, MARGIN, Inches(1.50), Inches(1.1), Pt(2.5), fill=LINE)
    return Inches(1.74)


def footer(s, n, label="FRD → STTM Agent"):
    tf = tbox(s, MARGIN, H - Inches(0.50), CONTENT_W, Inches(0.24))
    p = tf.paragraphs[0]
    p.space_after = 0
    r = p.add_run()
    r.text = label
    r.font.name, r.font.size = FONT, Pt(9.5)
    r.font.color.rgb = MUTED
    tf2 = tbox(s, W - MARGIN - Inches(1.0), H - Inches(0.50), Inches(1.0), Inches(0.24))
    p2 = tf2.paragraphs[0]
    p2.alignment = PP_ALIGN.RIGHT
    p2.space_after = 0
    r2 = p2.add_run()
    r2.text = str(n)
    r2.font.name, r2.font.size, r2.font.bold = FONT, Pt(9.5), True
    r2.font.color.rgb = MUTED


def card(s, x, y, w, h, title, lines, accent=TEAL, title_size=15, body_size=12.5,
         fill=PAPER, num=None, spacing=8):
    rect(s, x, y, w, h, fill=fill, line=LINE, lw=0.75)
    rect(s, x, y, Pt(3.2), h, fill=accent)
    ty = y + Inches(0.24)
    if num:
        tfn = tbox(s, x + Inches(0.28), ty, w - Inches(0.5), Inches(0.24))
        para(tfn, num, size=11, color=accent, bold=True, first=True, space_after=0)
        ty += Inches(0.26)
    tf = tbox(s, x + Inches(0.28), ty, w - Inches(0.56), Inches(0.32))
    para(tf, title, size=title_size, color=INK, bold=True, first=True, space_after=0)
    tf2 = tbox(s, x + Inches(0.28), ty + Inches(0.36), w - Inches(0.56),
               h - (ty - y) - Inches(0.55))
    for i, ln in enumerate(lines):
        para(tf2, ln, size=body_size, color=BODY, first=(i == 0),
             space_after=spacing, line_spacing=1.16)


def metric(s, x, y, w, h, value, label, note=None, accent=TEAL, vsize=32):
    rect(s, x, y, w, h, fill=WHITE, line=LINE, lw=0.75)
    rect(s, x, y, w, Pt(3.2), fill=accent)
    tf = tbox(s, x + Inches(0.22), y + Inches(0.24), w - Inches(0.4), Inches(0.5))
    para(tf, value, size=vsize, color=INK, bold=True, font=FONT_H, first=True,
         space_after=0)
    tf2 = tbox(s, x + Inches(0.22), y + Inches(0.86), w - Inches(0.4), Inches(0.9))
    para(tf2, label, size=11, color=accent, bold=True, first=True, space_after=3)
    if note:
        para(tf2, note, size=10.5, color=MUTED, line_spacing=1.16)


def arrow(s, x, y, w, h, color=LINE):
    return rect(s, x, y, w, h, fill=color, shape=MSO_SHAPE.RIGHT_ARROW)


def notes(s, text):
    tf = s.notes_slide.notes_text_frame
    tf.text = text
    for p in tf.paragraphs:
        for r in p.runs:
            r.font.size = Pt(12)
            r.font.name = FONT


# ================================================================ 1. TITLE
s = slide()
rect(s, Inches(0), Inches(0), W, H, fill=NAVY)
rect(s, Inches(0), Inches(0), Inches(0.16), H, fill=TEAL)
rect(s, W - Inches(4.6), Inches(0), Inches(4.6), H, fill=NAVY_L)
rect(s, W - Inches(4.6), Inches(0), Pt(2), H, fill=TEAL)

tf = tbox(s, MARGIN, Inches(1.65), Inches(7.6), Inches(0.3))
para(tf, "AI IN ENGINEERING PROGRAM   ·   AGENT 2 OF 5", size=12, color=TEAL,
     bold=True, first=True, space_after=0)

tf = tbox(s, MARGIN, Inches(2.20), Inches(7.7), Inches(1.9))
para(tf, "FRD → STTM Agent", size=50, color=WHITE, font=FONT_H, first=True,
     space_after=6, line_spacing=1.0)
para(tf, "It reads an approved FRD document and builds the STTM mapping\n"
         "spreadsheet for you — in about thirty seconds.",
     size=16.5, color=PALE, space_before=8, line_spacing=1.3)

rect(s, MARGIN, Inches(4.55), Inches(1.3), Pt(2.5), fill=TEAL)

tf = tbox(s, MARGIN, Inches(4.95), Inches(7.6), Inches(0.9))
para(tf, "Prepared for AmeriHealth   ·   Hexaware", size=14.5, color=WHITE,
     bold=True, first=True, space_after=5)
para(tf, "Presented by Soham & Arjun   |   Live demo included", size=13,
     color=PALER)

hy = Inches(2.20)
for val, lbl in [("94.1%", "of spreadsheet cells match the\nhand-made answer sheet"),
                 ("0 / 50", "names checked — every single\none is really in the document"),
                 ("~33 s", "start to finish, one paid AI\ncall (about 15 cents)")]:
    tf = tbox(s, W - Inches(4.10), hy, Inches(3.2), Inches(1.1))
    para(tf, val, size=27, color=TEAL, bold=True, font=FONT_H, first=True,
         space_after=3)
    for ln in lbl.split("\n"):
        para(tf, ln, size=11.5, color=PALE, line_spacing=1.15, space_after=0)
    hy += Inches(1.30)

tf = tbox(s, W - Inches(4.10), Inches(6.35), Inches(3.2), Inches(0.4))
para(tf, "Measured on a scrubbed practice document —\nreal run, 7 August 2026",
     size=9.5, color=RGBColor(0x7E, 0x96, 0xA8), first=True, italic=True,
     line_spacing=1.2, space_after=0)

# ================================================================ 2. AGENDA
s = slide()
y = header(s, "PROJECT AGENDA", "What we'll cover today",
           "About thirty minutes. The live run sits in the middle, and we finish with your questions.")

items = [
    ("01", "Where this agent fits", "It's one of five helpers that pass work to each other, one to the next.", TEAL),
    ("02", "The problem we're solving", "Making mapping spreadsheets by hand is slow, uneven, and hard to check later.", TEAL),
    ("03", "How the agent works", "Four simple steps — and what the AI does versus what plain code does.", BLUE),
    ("04", "Why you can trust it", "It shows its work on every value, and it asks a person when it isn't sure.", BLUE),
    ("05", "Live demo", "A document goes in, a spreadsheet comes out — run right here, with an offline backup.", AMBER),
    ("06", "Proof & value for AmeriHealth", "Real measured results, and what they mean for your teams.", AMBER),
    ("07", "Status, roadmap & Q&A", "What's done, what's next, and your questions.", MUTED),
]

col_w = Inches(5.70)
gap = Inches(0.60)
row_h = Inches(1.10)
for i, (num, title, sub, acc) in enumerate(items):
    cx = MARGIN + (col_w + gap) * (i // 4)
    cy = y + Inches(0.06) + row_h * (i % 4)
    tfn = tbox(s, cx, cy + Inches(0.01), Inches(0.58), Inches(0.36))
    para(tfn, num, size=21, color=acc, bold=True, font=FONT_H, first=True,
         space_after=0)
    tf = tbox(s, cx + Inches(0.60), cy, col_w - Inches(0.70), Inches(0.32))
    para(tf, title, size=16.5, color=INK, bold=True, first=True, space_after=3)
    para(tf, sub, size=12.5, color=BODY, line_spacing=1.16, space_after=0)
    rect(s, cx, cy + row_h - Inches(0.22), col_w - Inches(0.25), Pt(0.9), fill=LINE)

footer(s, 2)

# ================================================================ 3. PROGRAM MAP
s = slide()
y = header(s, "WHERE THIS FITS", "The second of five helpers",
           "Each helper hands its finished work to the next one as a file a computer can read — not as an email attachment.")

chain = [
    ("BRD → FRD", "Turns business wishes\ninto a working spec", False),
    ("FRD → STTM", "Makes the mapping\nspreadsheet + contract", True),
    ("CodeGen", "Writes the pipeline\ncode", False),
    ("Code Review", "Checks the generated\ncode", False),
]
bw, bh = Inches(2.62), Inches(1.52)
bx, by = MARGIN, y + Inches(0.10)
for i, (name, desc, active) in enumerate(chain):
    fill = TEAL if active else PAPER
    txt = WHITE if active else INK
    sub = RGBColor(0xD5, 0xEE, 0xEC) if active else BODY
    rect(s, bx, by, bw, bh, fill=fill, line=None if active else LINE, lw=0.75)
    tf = tbox(s, bx + Inches(0.22), by + Inches(0.24), bw - Inches(0.44), Inches(1.0))
    if active:
        para(tf, "YOU ARE HERE", size=9.5, color=RGBColor(0xC8, 0xEC, 0xE9),
             bold=True, first=True, space_after=3)
        para(tf, name, size=18, color=txt, bold=True, font=FONT_H, space_after=4)
    else:
        para(tf, name, size=18, color=txt, bold=True, font=FONT_H, first=True,
             space_after=4)
    for ln in desc.split("\n"):
        para(tf, ln, size=11.5, color=sub, space_after=0, line_spacing=1.15)
    if i < 3:
        arrow(s, bx + bw + Inches(0.10), by + Inches(0.62), Inches(0.34),
              Inches(0.26))
    bx += bw + Inches(0.54)

tf = tbox(s, MARGIN, by + bh + Inches(0.18), CONTENT_W, Inches(0.26))
para(tf, "SQL Optimization runs as a standalone fifth helper alongside this chain.",
     size=12, color=MUTED, first=True, italic=True, space_after=0)

cy = by + bh + Inches(0.58)
cw = (CONTENT_W - Inches(0.5)) / 2
card(s, MARGIN, cy, cw, Inches(2.28), "Upstream — already solved", [
    "This agent and the BRD→FRD agent read the same shared rule file, kept "
    "word-for-word identical in both code bases.",
    "Headings, ID formats and placeholders all come from that file — never typed "
    "into the code.",
    "If the two copies ever differ, the tests fail the build.",
], accent=TEAL, body_size=12.3)
card(s, MARGIN + cw + Inches(0.5), cy, cw, Inches(2.28),
     "Downstream — one missing piece", [
    "This agent already makes the file CodeGen reads. That part works today.",
    "But CodeGen also needs a second file made from the spreadsheet — and no "
    "tool makes it yet.",
    "Building that tool is the most useful next step. We're saying so up front.",
], accent=AMBER, body_size=12.3)
footer(s, 3)

# ================================================================ 4. PROBLEM
s = slide()
y = header(s, "THE PROBLEM", "Making these spreadsheets by hand is the bottleneck",
           "The document is approved and the targets are known — yet someone still types the whole spreadsheet, one cell at a time.")

left_w = Inches(6.15)
items = [
    ("Slow", "A person reads the whole FRD, cross-checks the data dictionary, and types every row of the spreadsheet by hand. It takes days."),
    ("Uneven", "Two people reading the same document make two different spreadsheets. Names and rules drift between feeds and teams."),
    ("Hard to check later", "Months later, nobody remembers which sentence in the document a spreadsheet cell came from. Finding out again is digging."),
    ("It holds everything up", "Code, tests and reviews all wait for this spreadsheet. The slowest manual step sets the pace for everything after it."),
]
cy = y + Inches(0.10)
for title, sub in items:
    rect(s, MARGIN, cy + Inches(0.09), Inches(0.10), Inches(0.10), fill=AMBER,
         shape=MSO_SHAPE.OVAL)
    tf = tbox(s, MARGIN + Inches(0.30), cy, left_w - Inches(0.35), Inches(0.3))
    para(tf, title, size=16.5, color=INK, bold=True, first=True, space_after=3)
    para(tf, sub, size=13, color=BODY, line_spacing=1.18, space_after=0)
    cy += Inches(1.08)

px = MARGIN + left_w + Inches(0.55)
pw = CONTENT_W - left_w - Inches(0.55)
ph = Inches(4.32)
rect(s, px, y + Inches(0.10), pw, ph, fill=NAVY)
rect(s, px, y + Inches(0.10), pw, Pt(3.2), fill=TEAL)
tf = tbox(s, px + Inches(0.40), y + Inches(0.58), pw - Inches(0.80), Inches(3.4))
para(tf, "WHAT WE'RE AUTOMATING", size=11, color=TEAL, bold=True, first=True,
     space_after=14)
para(tf, "The reading and the\ntyping — not the\nthinking.", size=21,
     color=WHITE, font=FONT_H, line_spacing=1.25, space_after=16)
para(tf, "The agent copies out what the document actually says, checks every "
         "value back against the text, and fills in the spreadsheet.",
     size=12.5, color=PALE, line_spacing=1.3, space_after=12)
para(tf, "When the document isn't clear, it stops and asks a person instead of "
         "guessing.",
     size=12.5, color=WHITE, bold=True, line_spacing=1.3, space_after=0)
footer(s, 4)

# ================================================================ 5. HOW IT WORKS
s = slide()
y = header(s, "HOW IT WORKS", "Four steps, one job each",
           "Runs on Databricks or on a plain laptop — the exact same code either way.")

stages = [
    ("01", "Ingest", "read the document",
     ["Opens the FRD file", "Turns it into plain text", "Saves it in a table"],
     TEAL),
    ("02", "Extract", "the one AI step",
     ["Claude reads the text", "Fills in a strict form", "If it can't, it stops\nand says so"],
     BLUE),
    ("03", "Contract build", "check everything",
     ["Checks the form is\nfilled in right", "Checks each value is\nreally in the document",
      "Sets unclear things aside"], BLUE),
    ("04", "Render + eval", "make the spreadsheet",
     ["Double-checks with the\ndata dictionary", "Builds the STTM\nspreadsheet",
      "Scores it against the\nanswer sheet"], TEAL),
]
bw = Inches(2.86)
bx, by = MARGIN, y + Inches(0.10)
bh = Inches(2.86)
for i, (num, name, tag, lines, acc) in enumerate(stages):
    rect(s, bx, by, bw, bh, fill=WHITE, line=LINE, lw=0.9)
    rect(s, bx, by, bw, Pt(4), fill=acc)
    tf = tbox(s, bx + Inches(0.26), by + Inches(0.28), bw - Inches(0.5), Inches(1.0))
    para(tf, num, size=11, color=acc, bold=True, first=True, space_after=3)
    para(tf, name, size=18.5, color=INK, bold=True, font=FONT_H, space_after=2)
    para(tf, tag, size=11.5, color=MUTED, space_after=0, italic=True)
    tf2 = tbox(s, bx + Inches(0.26), by + Inches(1.40), bw - Inches(0.5), Inches(1.3))
    for j, ln in enumerate(lines):
        sub = ln.split("\n")
        para(tf2, "·  " + sub[0], size=11.5, color=BODY, first=(j == 0),
             space_after=0 if len(sub) > 1 else 9, line_spacing=1.1)
        for extra in sub[1:]:
            para(tf2, "    " + extra, size=11.5, color=BODY, space_after=9,
                 line_spacing=1.1)
    if i < 3:
        arrow(s, bx + bw + Inches(0.045), by + Inches(1.30), Inches(0.25),
              Inches(0.26), color=RGBColor(0xC2, 0xCE, 0xD8))
    bx += bw + Inches(0.34)

dy = by + bh + Inches(0.32)
rect(s, MARGIN, dy, CONTENT_W, Inches(1.02), fill=PAPER, line=LINE, lw=0.75)
rect(s, MARGIN, dy, Pt(3.2), Inches(1.02), fill=INK)
tf = tbox(s, MARGIN + Inches(0.32), dy + Inches(0.18), CONTENT_W - Inches(0.7),
          Inches(0.7))
para(tf, "Who does what", size=13.5, color=INK, bold=True,
     first=True, space_after=5)
rich(tf, [("The AI only reads the document and copies things out. ", BODY, False),
          ("Plain, predictable code does everything else", INK, True),
          (" — the checking, the deciding, and the building of the spreadsheet. "
           "The AI never writes a cell the code can't back up.", BODY, False)],
     size=12.5, line_spacing=1.22, space_after=0)
footer(s, 5)

# ================================================================ 6. TRUST
s = slide()
y = header(s, "WHY TRUST IT", "Why you can trust what it makes",
           "Three simple habits — followed by plain code, on every single document.")

cw = (CONTENT_W - Inches(0.6)) / 3
cards = [
    ("01", "It shows its work", TEAL, [
        "Every important value it writes must appear word-for-word in the source "
        "document. Looser text is checked for close match.",
        "On the live run: all 45 strict checks and all 30 softer checks passed.",
        "This rule is never loosened just to make a run pass.",
    ]),
    ("02", "When unsure, it asks", BLUE, [
        "If the document doesn't say which feed a rule belongs to, the agent "
        "doesn't pick one — it sets the question aside.",
        "A plain code check against the data dictionary then confirms whatever "
        "the data itself can prove.",
        "Whatever it can't prove waits for a person.",
    ]),
    ("03", "A person has the last word", AMBER, [
        "A small review app lists each open question with its options and the "
        "text it came from.",
        "The reviewer picks an answer; that answer is saved and the agent never "
        "overrides it.",
        "Your analysts stay in charge.",
    ]),
]
for i, (num, title, acc, lines) in enumerate(cards):
    cx = MARGIN + (cw + Inches(0.3)) * i
    card(s, cx, y + Inches(0.10), cw, Inches(3.22), title, lines, accent=acc,
         title_size=15.5, body_size=12.3, fill=WHITE, num=num, spacing=9)

by2 = y + Inches(3.42)
rect(s, MARGIN, by2, CONTENT_W, Inches(1.05), fill=NAVY)
tf = tbox(s, MARGIN + Inches(0.38), by2 + Inches(0.20), CONTENT_W - Inches(0.8),
          Inches(0.68))
para(tf, "IN ONE SENTENCE", size=10.5, color=TEAL, bold=True,
     first=True, space_after=6)
para(tf, "Every cell in the spreadsheet can be traced back to the sentence it "
         "came from — and nothing the agent is unsure about goes out without a "
         "person looking first.",
     size=15, color=WHITE, line_spacing=1.2, space_after=0)
footer(s, 6)

# ================================================================ 7. DEMO DIVIDER
s = slide()
rect(s, Inches(0), Inches(0), W, H, fill=NAVY)
rect(s, Inches(0), Inches(0), Inches(0.16), H, fill=AMBER)
rect(s, Inches(0), H - Inches(2.45), W, Inches(2.45), fill=NAVY_L)

tf = tbox(s, MARGIN, Inches(1.70), Inches(9.5), Inches(0.3))
para(tf, "SECTION 05", size=12, color=AMBER, bold=True, first=True, space_after=0)
tf = tbox(s, MARGIN, Inches(2.15), Inches(8.2), Inches(1.7))
para(tf, "The Agent Demo", size=48, color=WHITE, font=FONT_H, first=True,
     space_after=10)
para(tf, "One FRD document goes in. One STTM spreadsheet comes out.\nLive, on "
         "this laptop, in about thirty seconds.",
     size=16.5, color=PALE, line_spacing=1.3, space_after=0)

by = H - Inches(1.95)
labels = [("Ingest", "read the document"), ("Extract", "the one AI step"),
          ("Contract build", "check everything"),
          ("Render + eval", "spreadsheet + score")]
bw = Inches(2.72)
bx = MARGIN
for i, (n, d) in enumerate(labels):
    rect(s, bx, by, Inches(0.11), Inches(0.11), fill=AMBER, shape=MSO_SHAPE.OVAL)
    tf = tbox(s, bx, by + Inches(0.28), bw - Inches(0.35), Inches(0.8))
    para(tf, n, size=15.5, color=WHITE, bold=True, first=True, space_after=3)
    para(tf, d, size=11.5, color=PALER, line_spacing=1.15, space_after=0)
    if i < 3:
        rect(s, bx + Inches(0.22), by + Inches(0.045), bw - Inches(0.44), Pt(1.2),
             fill=RGBColor(0x3C, 0x5A, 0x72))
    bx += bw

tf = tbox(s, W - Inches(4.15), Inches(2.35), Inches(3.30), Inches(1.4))
para(tf, "BACKUP PLAN", size=10.5, color=AMBER, bold=True, first=True, space_after=8)
para(tf, "Every demo run is saved. If the internet acts up, we replay a saved "
         "run — the same screens, no internet needed at all.",
     size=12.5, color=PALE, line_spacing=1.3, space_after=0)

# ================================================================ 8. DEMO WATCH
s = slide()
y = header(s, "THE DEMO", "What to watch for",
           "Five moments on the results screen — each one backs up something we've claimed today.")

left_w = Inches(6.3)
beats = [
    ("What it found", "Three data feeds, their file names, target tables and rules. All of it read out of the document — none of it pre-written."),
    ("The gate strip", "\"N unclear things found, M proven by the data, K waiting for a person.\" This is the ask-a-human part, live on screen."),
    ("The verdict tile", "PASS means nothing needs a second look. A warning stays up only when a person really should look at something."),
    ("The score vs. the answer sheet", "94.1% of cells match a hand-made reference. The rest differ on purpose — we planted those differences to test the scoring."),
    ("The spreadsheet itself", "We download and open the file — the thing a person would otherwise spend days making."),
]
cy = y + Inches(0.06)
for i, (t, sub) in enumerate(beats):
    tfn = tbox(s, MARGIN, cy + Inches(0.01), Inches(0.45), Inches(0.32))
    para(tfn, f"0{i+1}", size=15, color=AMBER, bold=True, font=FONT_H, first=True,
         space_after=0)
    tf = tbox(s, MARGIN + Inches(0.50), cy, left_w - Inches(0.55), Inches(0.32))
    para(tf, t, size=16, color=INK, bold=True, first=True, space_after=3)
    para(tf, sub, size=12.5, color=BODY, line_spacing=1.16, space_after=0)
    cy += Inches(0.88)

px = MARGIN + left_w + Inches(0.45)
pw = CONTENT_W - left_w - Inches(0.45)
card(s, px, y + Inches(0.06), pw, Inches(2.06), "It tells you the cost first", [
    "Before running, it says exactly what it will spend: one paid AI call, about "
    "15 cents, about 35 seconds.",
    "Runs differ in how many questions get auto-answered. The spreadsheet and "
    "its 94.1% score haven't varied.",
], accent=BLUE, title_size=15, body_size=12, fill=WHITE, spacing=9)
card(s, px, y + Inches(2.34), pw, Inches(2.06), "No real documents — ever", [
    "The document on screen is scrubbed of anything real and approved for "
    "showing.",
    "In real use, documents and results stay inside your own Databricks "
    "workspace and never leave it.",
], accent=TEAL, title_size=15, body_size=12, fill=WHITE, spacing=9)
footer(s, 8)

# ================================================================ 9. PROOF
s = slide()
y = header(s, "PROOF", "Real numbers from a real run",
           "One scrubbed practice FRD, through the real pipeline, 7 August 2026. Nothing simulated.")

mw = (CONTENT_W - Inches(0.9)) / 4
mets = [
    ("94.1%", "SPREADSHEET MATCH", "3,094 of 3,288 cells match the hand-made answer sheet", TEAL),
    ("0 / 50", "NOTHING MADE UP", "Every name it wrote really appears in the source document", TEAL),
    ("45 / 45", "SHOW-YOUR-WORK CHECKS", "Plus all 30 softer checks — passed without one miss", BLUE),
    ("~33 s", "START TO FINISH", "One paid AI call, about 15 cents per document", BLUE),
]
for i, (v, l, n, a) in enumerate(mets):
    metric(s, MARGIN + (mw + Inches(0.3)) * i, y + Inches(0.10), mw, Inches(1.86),
           v, l, n, accent=a)

cy = y + Inches(2.10)
cw = (CONTENT_W - Inches(0.5)) / 2
card(s, MARGIN, cy, cw, Inches(2.28), "What the numbers mean", [
    "All 194 non-matching cells are differences we planted on purpose, to prove "
    "the scoring notices them. It caught exactly those — nothing more.",
    "It missed nothing, and it invented nothing.",
    "An expert reviewing the output would keep about 96% of it exactly as-is. "
    "The bar we had to clear was 80%.",
], accent=TEAL, title_size=15, body_size=12.3, fill=WHITE, spacing=8)
card(s, MARGIN + cw + Inches(0.5), cy, cw, Inches(2.28), "Built carefully", [
    "77 automatic tests that run with no internet, no AI calls and no cluster.",
    "It has room for a document with roughly 75–90 feeds; this one used just 5% "
    "of that room.",
    "When something is wrong it stops and says so, loudly. It never quietly "
    "retries its way to a nice-looking answer.",
], accent=BLUE, title_size=15, body_size=12.3, fill=WHITE, spacing=8)
footer(s, 9)

# ================================================================ 10. AMERIHEALTH
s = slide()
y = header(s, "THE VALUE CASE", "Why this matters for AmeriHealth")

cw = (CONTENT_W - Inches(0.6)) / 3
rows = [
    [("Days of work become minutes", TEAL, [
        "The spreadsheet is the slowest step between an approved FRD and working "
        "code. The agent makes it in about thirty seconds.",
        "Your people go from typing to checking — the part that actually needs "
        "their judgement.",
      ]),
     ("Easy to audit", BLUE, [
        "Every value points back to the sentence it came from, and that check "
        "runs on every single document.",
        "When an auditor asks \"why is this mapped that way?\", the answer is "
        "already written down.",
      ]),
     ("Same result every time", AMBER, [
        "Two people make two different spreadsheets. The agent makes the same "
        "one every time, from the same rules.",
        "Names and rules stop drifting between teams — which is what lets the "
        "later steps be automated too.",
      ])],
    [("Your data stays with you", TEAL, [
        "Built for Databricks: documents and results live in your own Unity "
        "Catalog storage, and keys live in a locked secret store.",
        "Keeping patient data safe is a design rule here — this demo only ever "
        "uses scrubbed documents.",
      ]),
     ("People stay in charge", BLUE, [
        "The agent never guesses. Unclear things are set aside, proven ones are "
        "confirmed by the data, and the rest goes to a reviewer.",
        "Adopting it means checking a short list of questions — not trusting a "
        "black box.",
      ]),
     ("It pays off down the line", AMBER, [
        "What this agent produces is exactly what CodeGen consumes, which then "
        "feeds Code Review.",
        "Every document processed here speeds up the whole chain, from "
        "requirement to running code.",
      ])],
]
row_h = Inches(2.42)
for r, row in enumerate(rows):
    for i, (title, acc, lines) in enumerate(row):
        card(s, MARGIN + (cw + Inches(0.3)) * i, y + Inches(0.06) + row_h * r,
             cw, Inches(2.26), title, lines, accent=acc, title_size=14.5,
             body_size=11.8, fill=PAPER, spacing=8)
footer(s, 10, "FRD → STTM Agent   ·   Value for AmeriHealth")

# ================================================================ 11. DATABRICKS
s = slide()
y = header(s, "DEPLOYMENT", "Built for your Databricks workspace",
           "Shown on a laptop today; the workspace version is already written into the code.")

cw = (CONTENT_W - Inches(0.55)) / 2
card(s, MARGIN, y + Inches(0.06), cw, Inches(3.40), "Ready today", [
    "One Databricks job runs all four steps in order, on serverless compute — "
    "no machines for you to manage.",
    "Documents and results live in Unity Catalog storage and never leave your "
    "workspace.",
    "The AI key sits in Databricks' locked secret store; there is no key "
    "anywhere in the code.",
    "The review app is packaged the official Databricks way.",
    "If a step fails, the job stops and says so — it never quietly retries past "
    "a real problem.",
], accent=TEAL, title_size=16, body_size=12.2, fill=WHITE, spacing=9)

card(s, MARGIN + cw + Inches(0.55), y + Inches(0.06), cw, Inches(3.40),
     "Still to do — three plumbing items", [
    "Today the demo runs the steps on the laptop; the workspace version will "
    "press the Databricks job button instead.",
    "The app's AI key must come from Databricks' secret store rather than a "
    "local file.",
    "Live progress updates through Databricks' front door should work, but "
    "haven't been tried against a real workspace yet.",
    "None of this changes how the agent reads, checks or builds — it's plumbing "
    "between the same parts.",
], accent=AMBER, title_size=16, body_size=12.2, fill=WHITE, spacing=9)

by = y + Inches(3.72)
rect(s, MARGIN, by, CONTENT_W, Inches(0.84), fill=NAVY)
tf = tbox(s, MARGIN + Inches(0.38), by + Inches(0.24), CONTENT_W - Inches(0.8),
          Inches(0.45))
rich(tf, [("The framing: ", TEAL, True),
          ("built for the workspace, shown on a laptop. What you watched running "
           "is the exact same code the workspace would run.",
           WHITE, False)], size=13.5, first=True, line_spacing=1.2, space_after=0)
footer(s, 11)

# ================================================================ 12. STATUS
s = slide()
y = header(s, "STATUS & NEXT STEPS", "Where we are, and what's next",
           "The honest picture: the agent and its checks are proven; the wiring around them is the remaining work.")

rows = [
    ("Reading quality", "PROVEN", "Real run: 94.1% match, nothing made up, and about 96% of the output keepable exactly as-is.", TEAL),
    ("Checks & asking", "PROVEN", "All 45 strict + 30 softer checks passed, and the ask-a-person flow works end to end.", TEAL),
    ("Hand-off from upstream", "PROVEN", "Both agents read one shared rule file, kept word-for-word identical by tests.", TEAL),
    ("Databricks wiring", "NEXT", "Press the workspace job button from the app, move the key, and test live progress updates.", AMBER),
    ("The file for CodeGen", "NEXT", "Build the tool that makes the extra file CodeGen needs — the program's biggest missing piece.", AMBER),
    ("Bigger documents", "NEXT", "Try a real, larger FRD to confirm the roughly 75–90 feed headroom holds up.", BLUE),
]
cy = y + Inches(0.04)
rh = Inches(0.64)
for i, (label, status, detail, acc) in enumerate(rows):
    if i % 2 == 0:
        rect(s, MARGIN, cy - Inches(0.04), CONTENT_W, rh - Inches(0.02), fill=PAPER)
    tf = tbox(s, MARGIN + Inches(0.28), cy + Inches(0.10), Inches(3.2), Inches(0.32))
    para(tf, label, size=14.5, color=INK, bold=True, first=True, space_after=0)
    pill = rect(s, MARGIN + Inches(3.62), cy + Inches(0.09), Inches(1.0),
                Inches(0.28), fill=acc, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.5)
    ptf = pill.text_frame
    ptf.margin_left = ptf.margin_right = ptf.margin_top = ptf.margin_bottom = 0
    ptf.vertical_anchor = MSO_ANCHOR.MIDDLE
    pp = ptf.paragraphs[0]
    pp.alignment = PP_ALIGN.CENTER
    pr = pp.add_run()
    pr.text = status
    pr.font.name, pr.font.size, pr.font.bold = FONT, Pt(9.5), True
    pr.font.color.rgb = WHITE
    tf2 = tbox(s, MARGIN + Inches(4.95), cy + Inches(0.11), CONTENT_W - Inches(5.2),
               Inches(0.4))
    para(tf2, detail, size=12.3, color=BODY, first=True, space_after=0,
         line_spacing=1.15)
    cy += rh

by = cy + Inches(0.22)
rect(s, MARGIN, by, CONTENT_W, Inches(0.74), fill=NAVY)
rect(s, MARGIN, by, Pt(3.2), Inches(0.74), fill=TEAL)
tf = tbox(s, MARGIN + Inches(0.38), by + Inches(0.22), CONTENT_W - Inches(0.8),
          Inches(0.34))
rich(tf, [("Our ask: ", TEAL, True),
          ("one real (non-sensitive) FRD to test at full size, and a workspace "
           "to install into.", WHITE, False)],
     size=13.5, first=True, line_spacing=1.15, space_after=0)
footer(s, 12)

# ================================================================ 13. Q&A
s = slide()
rect(s, Inches(0), Inches(0), W, H, fill=NAVY)
rect(s, Inches(0), Inches(0), Inches(0.16), H, fill=TEAL)
rect(s, W - Inches(5.0), Inches(0), Inches(5.0), H, fill=NAVY_L)
rect(s, W - Inches(5.0), Inches(0), Pt(2), H, fill=TEAL)

tf = tbox(s, MARGIN, Inches(2.35), Inches(7.0), Inches(0.3))
para(tf, "SECTION 07", size=12, color=TEAL, bold=True, first=True, space_after=0)
tf = tbox(s, MARGIN, Inches(2.80), Inches(7.0), Inches(1.9))
para(tf, "Questions", size=50, color=WHITE, font=FONT_H, first=True, space_after=12)
para(tf, "Happy to run the demo again, open the spreadsheet, or walk through any "
         "part in as much detail as you like.",
     size=15.5, color=PALE, line_spacing=1.3, space_after=0)
rect(s, MARGIN, Inches(4.90), Inches(1.3), Pt(2.5), fill=TEAL)
tf = tbox(s, MARGIN, Inches(5.25), Inches(7.0), Inches(0.8))
para(tf, "Soham  ·  Arjun   |   Hexaware — AI in Engineering", size=13.5,
     color=WHITE, bold=True, first=True, space_after=4)
para(tf, "FRD → STTM Agent   ·   Agent 2 of 5", size=12, color=PALER, space_after=0)

tf = tbox(s, W - Inches(4.50), Inches(1.95), Inches(3.6), Inches(3.9))
para(tf, "LIKELY QUESTIONS", size=10.5, color=TEAL, bold=True, first=True,
     space_after=14)
for q in ["What about much bigger documents?",
          "What happens when the document is unclear?",
          "How do we know it didn't make something up?",
          "How much checking is left for a person to do?",
          "How does it get into our workspace — and who holds the keys?",
          "How big is the missing CodeGen piece?"]:
    para(tf, "—  " + q, size=12.5, color=PALE, space_after=11, line_spacing=1.2)

# ================================================================ notes
SPEAKER_NOTES = [
    # 1 title
    "Open. Introduce yourself and Arjun and split the room's expectation: slides "
    "first, then a live run of the agent, then questions.\n\n"
    "The three numbers on the right are the whole story in miniature — 94.1% "
    "mapping fidelity, zero ungrounded identifiers, about thirty seconds and "
    "fifteen cents per document. All measured on a real live run on 7 August "
    "2026, not a mock, not a projection.\n\n"
    "Say plainly: the FRD in the demo is anonymized fixture material. No real "
    "client document is used at any point today.",
    # 2 agenda
    "Keep this to about forty-five seconds. Signal that the demo (item 05) is the "
    "centre of gravity and that we will finish with open Q&A.\n\n"
    "Flag the two-hander: agree in advance who takes 03/04 (mechanics) and who "
    "takes 06 (value), so the hand-off is not negotiated on stage.",
    # 3 program map
    "The point of this slide is that the agents are chained by contracts, not by "
    "email attachments.\n\n"
    "Upstream is genuinely solved: our label contract is versioned and committed "
    "byte-identically to both repositories, with a round-trip test that fails the "
    "build on drift.\n\n"
    "Downstream, say the gap out loud — CodeGen needs a workbook-derived mapping "
    "contract and no committed tool in the programme produces one yet. Naming it "
    "before anyone asks buys credibility for everything else on the deck.",
    # 4 problem
    "Anchor in the audience's own experience — ask whether their mapping "
    "workbooks are built by hand today; the answer is almost always yes.\n\n"
    "The right-hand panel is the line to land: we automate the reading and the "
    "transcription, not the judgement. Where the FRD is genuinely ambiguous the "
    "agent stops and asks rather than guessing. That framing pre-empts the "
    "\"can we trust it\" question later.",
    # 5 how it works
    "Ninety seconds, no deeper. Walk the four stages left to right; the audience "
    "will see these same four stages light up in the demo, so this slide is the "
    "map for what they are about to watch.\n\n"
    "The band at the bottom matters more than the boxes: the model reads prose, "
    "and deterministic code owns validation, grounding, attribution and "
    "rendering. The model never writes a cell that code cannot justify.\n\n"
    "If asked which model: Claude, called through the Anthropic SDK, with the "
    "schema and prompt living in our repository — not a black-box endpoint.",
    # 6 trust
    "This is the slide for the sceptical engineer and the compliance-minded "
    "stakeholder at the same time.\n\n"
    "Grounding audit: every strict field must appear verbatim in the source "
    "document. On the live run, 45 of 45 strict checks passed. Stress that we "
    "never relax this to make a run go green — a failing audit is a failing run.\n\n"
    "Gating: the agent records an ambiguity rather than choosing; a deterministic "
    "data-dictionary cross-check auto-confirms only what the data proves.\n\n"
    "If pushed on the difference between strict and softer checks: strict fields "
    "are identifiers and must match verbatim; the softer (advisory) fields are "
    "prose and are checked by token overlap.",
    # 7 demo divider
    "Transition slide. Hand over to whoever is driving the laptop.\n\n"
    "Before starting the run, read the cost-confirmation dialog aloud — one "
    "billed call, about $0.15, about thirty-five seconds. The demo tells you what "
    "it costs before it spends anything.\n\n"
    "Contingency: if the network misbehaves, switch to Replay and pick the "
    "post-fix run. Identical results view, zero API calls, works offline. Do not "
    "debug on stage.",
    # 8 demo watch
    "Leave this up while the run executes, then walk the results page top to "
    "bottom in this order.\n\n"
    "Give the gate strip time even when the counts are zero — zero means nothing "
    "needed a human on this run, which is itself the point. If the run shows a "
    "flag that then gets auto-confirmed, slow down: that is the best possible "
    "outcome, because the flag-then-confirm story renders on screen.\n\n"
    "Runs vary in how much the dictionary can auto-confirm. The mapping quality "
    "and the 94.1% eval have not varied across runs — say so before anyone "
    "wonders about it.\n\n"
    "Finish by downloading and opening the workbook. That artifact is the "
    "deliverable an analyst would otherwise hand-build.",
    # 9 proof
    "Come back to slides with the workbook still fresh in their minds.\n\n"
    "The number to defend is 94.1%. Every one of the 194 non-matching cells is a "
    "datatype difference deliberately built into the test fixture — the eval "
    "catches exactly what it should and nothing else. Nothing was missed and "
    "nothing was invented.\n\n"
    "The hallucination sweep checked all 50 identifiers in the output — feed "
    "names, file patterns, catalogs, schemas, tables, load strategies, "
    "requirement ids — for verbatim presence in the source. Zero ungrounded.\n\n"
    "If asked about the keep rate: a proxy-SME review estimated roughly 96% "
    "first-pass keep, against an acceptance bar of 80%.",
    # 10 amerihealth
    "This is the slide to slow down on — it is the reason we are in the room.\n\n"
    "Lead with the top-left card and the bottom-left card: speed, and the fact "
    "that your data never leaves your tenancy. For a regulated payer those two "
    "usually decide the conversation.\n\n"
    "Auditability is the strongest card technically: the traceability is not a "
    "report we generate afterwards, it is a hard gate that runs on every "
    "document.\n\n"
    "Close on the bottom-right: the value is not one workbook, it is a governed "
    "path from requirement to running pipeline, and every FRD processed here "
    "shortens the downstream chain.",
    # 11 deployment
    "Framing line: built for the workspace, currently demoed locally. Say it "
    "early so nobody thinks the local demo is the product.\n\n"
    "Left column is real and committed today. Right column is the honest list of "
    "port work — three items, all plumbing, none of them touching the pipeline "
    "logic or the quality gates.\n\n"
    "If asked about credentials: in the workspace, the key comes from a "
    "Databricks secret scope, and there is no key anywhere in the repository.",
    # 12 status
    "A one-page honest read. Three things proven, three things next.\n\n"
    "Do not soften the two amber rows. The workspace port and the downstream "
    "mapping contract are real work, and being straight about them is what makes "
    "the three green rows believable.\n\n"
    "End on the ask: a real, non-sensitive FRD to validate at scale, and a "
    "workspace target for the deployment port. Get a name and a date if you can.",
    # 13 Q&A
    "Answers to have ready:\n\n"
    "Larger FRDs — the run used 5% of its output budget, giving roughly 75 to 90 "
    "feeds of headroom in one document. Beyond that, per-feed chunked extraction "
    "is the known path. Validating on a real large FRD is our ask.\n\n"
    "Genuine ambiguity — it is gated, not guessed. The dictionary auto-confirms "
    "what it can prove; everything else goes to a reviewer, whose decision is "
    "stored and never re-decided by the heuristic.\n\n"
    "Invented mappings — the grounding audit and the hallucination sweep. Offer "
    "to show the audit output.\n\n"
    "Review workload — on the demo document, at most a handful of items per run, "
    "often zero.\n\n"
    "Keys and tenancy — Databricks secret scope, Unity Catalog volumes, nothing "
    "leaves the workspace.\n\n"
    "CodeGen gap — no committed extractor today; it is scoped work, and we would "
    "size it properly rather than guess on stage.",
]
for i, txt in enumerate(SPEAKER_NOTES):
    notes(prs.slides[i], txt)

prs.save(OUT)
print("wrote", OUT, "-", len(prs.slides._sldIdLst), "slides,",
      len(SPEAKER_NOTES), "notes")
