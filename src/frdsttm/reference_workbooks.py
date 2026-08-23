"""
frdsttm.reference_workbooks — reference STTM workbook parsing + feed matching.

Factored verbatim out of notebooks/04_sttm_render.py (2026-08-22) so the
similarity/corpus modules and the review app backend can read reference
workbooks with EXACTLY the dictionary the render stage uses — one parser,
no drift. 04_sttm_render consumes this module through the
notebooks/_reference_workbooks.py shim (%run in Databricks, plain import
locally), the same pattern as _models / _live_extraction.

Both client dialects are supported: sheet_per_table (FILE_DETAILS +
MAPPING-<TABLE> sheets) and single_sheet (CAQH-style wide sheet). See
SKILL.md "Stage 4" for the roles a parsed dictionary plays.
"""

from __future__ import annotations

import re
import unicodedata

from openpyxl import load_workbook

# --------------------------------------------------------------------------- #
# normalization
# --------------------------------------------------------------------------- #
_UNI = str.maketrans({"\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"',
                      "\u2010": "-", "\u2011": "-", "\u2013": "-", "\u2014": "-",
                      "\u00a0": " "})


def _n(s) -> str:
    if s is None:
        return ""
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", str(s)).translate(_UNI)).strip()


def _nl(s) -> str:
    return _n(s).lower()


_TRUE = {"yes", "y", "true"}


def _flag(s) -> bool:
    return _nl(s) in _TRUE


def _not_null(s) -> bool:
    return "not null" in _nl(s)


# --------------------------------------------------------------------------- #
# source-dictionary parsing
# --------------------------------------------------------------------------- #
# Source-side cell (normalized) that marks a trailing audit-column row.
# Mirrors CodeGen's `extractor.audit_source_markers: ["na"]`.
AUDIT_SOURCE_MARKER = "na"

_HEADER_ALIASES = {
    "source_column": ["database column name", "database name",
                      "client data table column name", "field name"],
    "description": ["description"],
    "sample": ["sample value", "example values"],
    "datatype": ["datatype", "data type"],
    "nullable_raw": ["null check"],
    "phi_raw": ["phi field", "phi/pii field", "pii"],
    "mandatory_raw": ["mandatory field", "mandatory", "mandatory column"],
    "comment": ["comment", "comments"],
    "length": ["length"],
    "fixed_length": ["field length (fixed width)"],
    "fixed_start": ["start position (fixed width)"],
    "fixed_end": ["end position (fixed width)"],
    "segment": ["segment (ex:header,trailer,detail)", "segment"],
    "business_rule": ["business rule"],
}

_TARGET_ALIASES = {
    "catalog": ["catalog"],
    "schema": ["schema"],
    "table": ["tablename", "table name"],
    "column": ["columnname", "column name"],
    "datatype": ["datatype", "data type"],
}


def _map_headers(headers, aliases, offset=0):
    out = {}
    for i, h in enumerate(headers):
        hn = _nl(h)
        if not hn:
            continue
        for field, names in aliases.items():
            if field not in out and any(hn == a or hn.startswith(a) for a in names):
                out[field] = i + offset
    return out


def _section_starts(label_row):
    """Column indices where 'Source', 'Stage Layer', 'Standard Layer' begin."""
    src = stage = std = None
    for i, v in enumerate(label_row):
        vn = _nl(v)
        if vn.startswith("source") and src is None:
            src = i
        elif "stage layer" in vn and stage is None:
            stage = i
        elif "standard layer" in vn and std is None:
            std = i
    return src, stage, std


def _row_vals(ws, r, width):
    return [c.value for c in ws[r][:width]] if r <= ws.max_row else [None] * width


def parse_reference_workbook(path):
    """Returns {"dialect": ..., "feeds": {table_key: feed_dict}, "meta": {...}}.

    feed_dict: {"sheet", "fields": [source-field dicts],
                "ref_targets": [{stage:{...}, standard:{...}} per field]}  # eval only

    Trailing rows whose source column is the audit marker "NA" are kept as
    fields (positional eval and similarity stay unchanged) but flagged
    `audit: True` (2026-08-22): they are the client's ETL audit columns
    (load timestamps, source file name, ...) — a workbook convention, not an
    FRD fact — and 04 derives their target identity from the template's own
    targets instead of the 1:1 source-name rule. CodeGen's `extract-sttm`
    requires at least one such row per mapping sheet.
    """
    wb = load_workbook(path, read_only=True)
    if any(s.startswith("MAPPING-") for s in wb.sheetnames):
        return _parse_sheet_per_table(wb)
    return _parse_single_sheet(wb)


def _parse_sheet_per_table(wb):
    feeds, meta = {}, {}
    if "FILE_DETAILS" in wb.sheetnames:
        ws = wb["FILE_DETAILS"]
        rows = list(ws.iter_rows(values_only=True))
        if rows:
            hdr = [_nl(v) for v in rows[0]]
            meta["file_details"] = [
                {hdr[i]: _n(v) for i, v in enumerate(r) if i < len(hdr) and _n(v)}
                for r in rows[1:] if any(_n(v) for v in r)
            ]
    for name in wb.sheetnames:
        if not name.startswith("MAPPING-"):
            continue
        ws = wb[name]
        width = ws.max_column
        labels = _row_vals(ws, 1, width)
        headers = _row_vals(ws, 2, width)
        src_i, stage_i, std_i = _section_starts(labels)
        if src_i is None or stage_i is None:
            continue
        smap = _map_headers(headers[src_i:stage_i], _HEADER_ALIASES, src_i)
        gmap = _map_headers(headers[stage_i:(std_i or width)], _TARGET_ALIASES, stage_i)
        dmap = _map_headers(headers[std_i:], _TARGET_ALIASES, std_i) if std_i is not None else {}
        recycle_note = next((_n(h) for h in headers if "recycle" in _nl(h)), None)

        fields, ref_targets, table_key = [], [], None
        for r in ws.iter_rows(min_row=3, values_only=True):
            get = lambda m, k: _n(r[m[k]]) if k in m and m[k] < len(r) else ""
            col = get(smap, "source_column")
            if not col:
                continue
            fields.append({
                "source_column": col,
                "description": get(smap, "description"),
                "sample": get(smap, "sample"),
                "datatype": get(smap, "datatype") or "String",
                "nullable": not _not_null(get(smap, "nullable_raw")),
                "phi": _flag(get(smap, "phi_raw")),
                "mandatory": _flag(get(smap, "mandatory_raw")),
                "comment": get(smap, "comment"),
                "segment": "",
                "business_rule": "",
                "audit": _nl(col) == AUDIT_SOURCE_MARKER,
            })
            ref_targets.append({
                "stage": {k: get(gmap, k) for k in _TARGET_ALIASES},
                "standard": {k: get(dmap, k) for k in _TARGET_ALIASES},
            })
            table_key = table_key or _nl(get(gmap, "table"))
        if table_key:
            feeds[table_key] = {"sheet": name, "fields": fields,
                                "ref_targets": ref_targets,
                                "recycle_note": recycle_note}
    return {"dialect": "sheet_per_table", "feeds": feeds, "meta": meta}


def _parse_single_sheet(wb):
    for name in wb.sheetnames:
        ws = wb[name]
        header_r = label_r = None
        for r in range(1, min(ws.max_row, 30) + 1):
            vals = [_nl(v) for v in _row_vals(ws, r, ws.max_column)]
            if any(v.startswith("field name") for v in vals):
                header_r = r
                label_r = r - 1
                break
        if header_r is None:
            continue
        width = ws.max_column
        meta = {}
        for r in range(1, label_r):
            k, v = _n(ws.cell(r, 1).value), _n(ws.cell(r, 2).value)
            if k:
                meta[k] = v
        labels = _row_vals(ws, label_r, width)
        headers = _row_vals(ws, header_r, width)
        src_i, stage_i, std_i = _section_starts(labels)
        smap = _map_headers(headers[src_i:stage_i], _HEADER_ALIASES, src_i)
        gmap = _map_headers(headers[stage_i:std_i], _TARGET_ALIASES, stage_i)
        dmap = _map_headers(headers[std_i:], _TARGET_ALIASES, std_i)

        fields, ref_targets = [], []
        blanks = 0
        for r in ws.iter_rows(min_row=header_r + 1, values_only=True):
            get = lambda m, k: _n(r[m[k]]) if k in m and m[k] < len(r) else ""
            col = get(smap, "source_column")
            if not col:
                blanks += 1
                if blanks > 5:
                    break
                continue
            blanks = 0
            fields.append({
                "source_column": col,
                "description": get(smap, "description") or get(smap, "comment"),
                "sample": get(smap, "sample"),
                "datatype": get(smap, "datatype") or "String",
                "nullable": not _not_null(get(smap, "nullable_raw")),
                "phi": _flag(get(smap, "phi_raw")),
                "mandatory": _flag(get(smap, "mandatory_raw")),
                "comment": get(smap, "comment"),
                "length": get(smap, "length"),
                "fixed_length": get(smap, "fixed_length"),
                "fixed_start": get(smap, "fixed_start"),
                "fixed_end": get(smap, "fixed_end"),
                "segment": get(smap, "segment"),
                "business_rule": get(smap, "business_rule"),
                "audit": _nl(col) == AUDIT_SOURCE_MARKER,
            })
            ref_targets.append({
                "stage": {k: get(gmap, k) for k in _TARGET_ALIASES},
                "standard": {k: get(dmap, k) for k in _TARGET_ALIASES},
            })
        feed_key = _nl(name)
        return {"dialect": "single_sheet",
                "feeds": {feed_key: {"sheet": name, "fields": fields,
                                     "ref_targets": ref_targets,
                                     "recycle_note": None}},
                "meta": meta}
    return {"dialect": "single_sheet", "feeds": {}, "meta": {}}


# --------------------------------------------------------------------------- #
# layout descriptor — the WRITE-side view of a reference workbook (2026-08-22)
# --------------------------------------------------------------------------- #
# The parser above reads a client workbook by header NAME. `layout_of` turns
# the same observation into something 04 can render INTO: which sheet holds
# which feed, where the band / header / first data rows are, which column
# carries which logical field (via the same alias tables), and which columns
# carry nothing the contract knows — those stay blank on fill and are
# reported. This is what retired the two hard-coded output dialects: the
# chosen template's own layout is the dialect, whatever it is.

_FILE_DETAILS_ALIASES = {
    "vendor": ["vendor", "source system", "source"],
    "file_name": ["filename", "file name", "file", "files"],
    "description": ["file description", "description"],
    "location": ["location", "landing location", "file location", "path"],
    "frequency": ["frequency", "file frequency"],
}
_VERSION_HISTORY_ALIASES = {
    "version": ["version"],
    "date": ["date"],
    "author": ["author"],
    "change": ["change description", "description", "change", "changes"],
}
# single-sheet metadata block keys -> contract fields the renderer can fill
_META_ALIASES = {
    "file_name_patterns": ["file(s)", "file", "files", "file name", "filename"],
    "source_system": ["file generator", "vendor", "source system", "source"],
    "landing_location": ["file location", "location", "landing location"],
    "lobs": ["lob", "lobs", "line of business"],
    "frequency": ["file frequency", "frequency"],
    "domain": ["domain"],
    "sub_domain": ["sub-domain", "sub domain", "subdomain"],
    "file_type": ["file type", "format", "file format"],
}


def _column_roles(headers, src_i, stage_i, std_i, width):
    """Per column: {"index", "header", "role"} where role is
    ("source", field) | ("stage", key) | ("standard", key) | ("recycle", None)
    | ("index", None) | None (unmapped: left blank on fill)."""
    smap = _map_headers(headers[src_i:stage_i], _HEADER_ALIASES, src_i)
    gmap = _map_headers(headers[stage_i:(std_i or width)], _TARGET_ALIASES, stage_i)
    dmap = _map_headers(headers[std_i:], _TARGET_ALIASES, std_i) if std_i is not None else {}
    inv = {}
    for f, i in smap.items():
        inv[i] = ("source", f)
    for k, i in gmap.items():
        inv[i] = ("stage", k)
    for k, i in dmap.items():
        inv[i] = ("standard", k)
    for i, h in enumerate(headers):
        if i in inv:
            continue
        hn = _nl(h)
        if "recycle" in hn:
            inv[i] = ("recycle", None)
        elif hn == "#":
            inv[i] = ("index", None)
    return [{"index": i, "header": _n(h), "role": inv.get(i)} for i, h in enumerate(headers)]


def _header_map(headers, aliases):
    return {k: i for k, i in _map_headers(headers, aliases, 0).items()}


def layout_of(path) -> dict:
    """Layout descriptor of one reference workbook (see module comment).

    sheet_per_table → {"dialect", "path", "sheets": {sheet name: {label_row,
    header_row, first_data_row, width, columns, unmapped}}, "file_details":
    {"headers", "columns": {logical: col idx}} | None, "version_history":
    {...} | None}
    single_sheet → {"dialect", "path", "sheet", "label_row", "header_row",
    "first_data_row", "width", "columns", "unmapped", "meta_rows": [(row,
    key, field-or-None)]}
    """
    wb = load_workbook(path, read_only=True)
    out = {"path": str(path)}
    if any(n.startswith("MAPPING-") for n in wb.sheetnames):
        out["dialect"] = "sheet_per_table"
        sheets = {}
        for name in wb.sheetnames:
            if not name.startswith("MAPPING-"):
                continue
            ws = wb[name]
            width = ws.max_column
            labels = _row_vals(ws, 1, width)
            headers = _row_vals(ws, 2, width)
            src_i, stage_i, std_i = _section_starts(labels)
            if src_i is None or stage_i is None:
                continue
            cols = _column_roles(headers, src_i, stage_i, std_i, width)
            sheets[name] = {
                "label_row": 1, "header_row": 2, "first_data_row": 3, "width": width,
                "columns": cols,
                "unmapped": [c["header"] for c in cols if c["role"] is None and c["header"]],
            }
        out["sheets"] = sheets
        out["file_details"] = None
        if "FILE_DETAILS" in wb.sheetnames:
            ws = wb["FILE_DETAILS"]
            hdr = _row_vals(ws, 1, ws.max_column)
            out["file_details"] = {"headers": [_n(h) for h in hdr],
                                   "columns": _header_map(hdr, _FILE_DETAILS_ALIASES)}
        out["version_history"] = None
        if "VERSION_HISTORY" in wb.sheetnames:
            ws = wb["VERSION_HISTORY"]
            hdr = _row_vals(ws, 1, ws.max_column)
            out["version_history"] = {"headers": [_n(h) for h in hdr],
                                      "columns": _header_map(hdr, _VERSION_HISTORY_ALIASES)}
        return out

    out["dialect"] = "single_sheet"
    for name in wb.sheetnames:
        ws = wb[name]
        header_r = None
        for r in range(1, min(ws.max_row, 30) + 1):
            vals = [_nl(v) for v in _row_vals(ws, r, ws.max_column)]
            if any(v.startswith("field name") for v in vals):
                header_r = r
                break
        if header_r is None:
            continue
        label_r = header_r - 1
        width = ws.max_column
        labels = _row_vals(ws, label_r, width)
        headers = _row_vals(ws, header_r, width)
        src_i, stage_i, std_i = _section_starts(labels)
        if src_i is None or stage_i is None:
            continue
        cols = _column_roles(headers, src_i, stage_i, std_i, width)
        meta_rows = []
        for r in range(1, label_r):
            k = _n(ws.cell(r, 1).value)
            if not k:
                continue
            kn = _nl(k)
            # longest matching alias wins: "file generator" must not fall to
            # the "file" alias of file_name_patterns
            hits = [(len(a), f) for f, names in _META_ALIASES.items() for a in names
                    if kn == a or kn.startswith(a)]
            field = max(hits)[1] if hits else None
            meta_rows.append({"row": r, "key": k, "field": field})
        out.update({"sheet": name, "label_row": label_r, "header_row": header_r,
                    "first_data_row": header_r + 1, "width": width, "columns": cols,
                    "unmapped": [c["header"] for c in cols if c["role"] is None and c["header"]],
                    "meta_rows": meta_rows})
        return out
    out.update({"sheet": None, "columns": [], "unmapped": [], "meta_rows": []})
    return out


# --------------------------------------------------------------------------- #
# match contract feeds <-> dictionary feeds
# --------------------------------------------------------------------------- #
def match_feeds(contract, dictionary):
    """Returns {contract feed index: dict feed_key}. Matches on stage tables
    (sheet-per-table) or falls back to the single dictionary feed."""
    out = {}
    dict_keys = list(dictionary["feeds"])
    for i, feed in enumerate(contract["feeds"]):
        tables = [_nl(t) for t in (feed.get("stage_target") or {}).get("tables", [])]
        hit = next((k for k in dict_keys if k in tables), None)
        if hit is None and len(dict_keys) == 1 and len(contract["feeds"]) == 1:
            hit = dict_keys[0]
        if hit is not None:
            out[i] = hit
    return out



# --------------------------------------------------------------------------- #
# loose token bag — the coarse similarity primitive 04's driver used for
# reference pairing, now shared with frdsttm.similarity
# --------------------------------------------------------------------------- #
def loose_tokens(s):
    return {t for t in re.findall(r"[a-z0-9]{4,}", str(s).lower())}
