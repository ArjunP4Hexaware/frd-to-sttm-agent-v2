"""
frdsttm.render — rows from the VDD + targets from the FRD/standards → .xlsx.

* Every ROW comes from the vendor data dictionary. Source column names are
  carried AS-IS into both target layers (neither client standard states a
  column naming rule; nothing is guessed).
* Stage datatype is the standards' stage default; standard datatype is
  promoted from the VENDOR type per the coding standard.
* The client's audit columns (SRC_FILE_NAME, REC_CREATION_TIME, …) are
  appended per segment from the standards, never from a vendor document.
* FRD rules land on the row whose column they name; a rule naming no column
  goes to a source-level cell. Placement is recorded, never dropped.
* The workbook is written INTO an approved STTM's layout when one is
  available (sheets, band labels, headers, styles — structure only); the two
  built-in dialects are the fallback.
"""

from __future__ import annotations

import re
from copy import copy
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from frdsttm import standards as std
from frdsttm.reference_layout import _nl

GENERATOR = "frd-sttm-agent"
_SEGMENT_SUFFIX = {"header": "_HDR", "detail": "_DTL", "trailer": "_TRL"}
_DATA_BEARING_SEGMENTS = {"", "detail", "dtl", "data", "body", "record"}


# --------------------------------------------------------------------------- #
# rows
# --------------------------------------------------------------------------- #
def _table_for_segment(tables, segment):
    suffix = _SEGMENT_SUFFIX.get(_nl(segment))
    if suffix:
        for t in tables:
            if str(t).upper().endswith(suffix):
                return t
    return tables[0] if tables else ""


def _row(f: dict) -> dict:
    req = f.get("required")
    return {
        "source_column": f.get("name"),
        "datatype": f.get("datatype") or "",
        "length": f.get("length") or "",
        "description": f.get("description") or "",
        "sample": f.get("example") or "",
        "phi": f.get("phi"),
        "mandatory": req,
        "nullable": (not req) if req is not None else None,
        "segment": f.get("segment") or "",
        "comment": f.get("allowed_values") or "",
        "business_rule": f.get("notes") or "",
        "fixed_start": f.get("start_position") or "",
        "fixed_end": f.get("end_position") or "",
        "fixed_length": f.get("length") or "",
        "position": f.get("position"),
        "audit": False,
    }


def _target(layer_spec: dict, table: str, column: str, datatype: str) -> dict:
    return {"catalog": layer_spec.get("catalog") or "", "schema": layer_spec.get("schema") or "",
            "table": table, "column": column, "datatype": datatype}


def build_rows(source: dict, feed: dict, vdd_fields: list[dict]) -> tuple[list[dict], dict]:
    """The rendered rows for one source. Returns (rows, notes)."""
    stage, standard = source["layers"]["stage"], source["layers"]["standard"]
    stage_dt = std.stage_default_type()
    rows, unpromoted, segments = [], set(), []
    for f in vdd_fields:
        seg = f.get("segment") or ""
        if seg not in segments:
            segments.append(seg)
        stage_tbl = _table_for_segment(stage["tables"], seg)
        std_tbl = _table_for_segment(standard["tables"], seg) or stage_tbl
        promoted = std.promote_type(f.get("datatype"))
        if promoted is None and f.get("datatype"):
            unpromoted.add(str(f["datatype"]))
        r = _row(f)
        r["stage"] = _target(stage, stage_tbl, f.get("name") or "", stage_dt)
        r["standard"] = _target(standard, std_tbl, f.get("name") or "", promoted or stage_dt)
        rows.append(r)
    n_audit = 0
    for seg in segments or [""]:
        data_bearing = _nl(seg) in _DATA_BEARING_SEGMENTS
        stage_tbl = _table_for_segment(stage["tables"], seg)
        std_tbl = _table_for_segment(standard["tables"], seg) or stage_tbl
        std_cols = {c["name"]: c for c in std.audit_columns("standard", data_bearing)}
        for col in std.audit_columns("stage", data_bearing):
            sc = std_cols.get(col["name"], col)
            r = {**_row({"name": "NA", "segment": seg}), "source_column": "NA", "audit": True}
            r["stage"] = _target(stage, stage_tbl, col["name"], col.get("datatype", "String"))
            r["standard"] = _target(standard, std_tbl, sc["name"], sc.get("datatype", "String"))
            rows.append(r)
            n_audit += 1
    return rows, {"n_rows": len(rows), "n_audit_rows": n_audit,
                  "unpromoted_types": sorted(unpromoted), "segments": [s for s in segments if s]}


# --------------------------------------------------------------------------- #
# rule placement
# --------------------------------------------------------------------------- #
def _column_mentions(text, columns):
    hits = [c for c in columns
            if c and re.search(rf"(?<![A-Za-z0-9_]){re.escape(c)}(?![A-Za-z0-9_])", text, re.I)]
    return sorted(hits, key=len, reverse=True)


def place_rules(feed: dict, rows: list[dict]) -> dict:
    columns = [r["source_column"] for r in rows if not r.get("audit")]
    by_column: dict = {c: [] for c in columns}
    unattributed = []
    for rule in feed.get("validation_rules") or []:
        cols = _column_mentions(rule, columns)
        if cols:
            for c in cols:
                by_column[c].append(rule)
        else:
            unattributed.append(rule)
    recycle = None
    if feed.get("recycle_rule"):
        cols = _column_mentions(feed["recycle_rule"], columns)
        recycle = {"text": feed["recycle_rule"], "column": cols[0] if cols else None}
    return {"by_column": {c: r for c, r in by_column.items() if r},
            "recycle": recycle, "unattributed": unattributed}


def _source_level_text(placement):
    parts = list(placement["unattributed"])
    if placement["recycle"] and placement["recycle"]["column"] is None:
        parts.append(f"Recycle rule: {placement['recycle']['text']}")
    return "\n".join(parts)


# --------------------------------------------------------------------------- #
# built-in renderers (fallback when no approved workbook supplies a layout)
# --------------------------------------------------------------------------- #
_HDR_FILL = PatternFill("solid", fgColor="D9E1F2")
_SEC_FILL = PatternFill("solid", fgColor="BDD7EE")
_BOLD = Font(bold=True)
_SRC_HEADERS = ["Database column Name", "NULL CHECK", "Description", "Sample Value",
                "DataType", "PHI Field", "Mandatory Field", "Comment"]
_TGT_HEADERS = ["Schema", "TableName", "ColumnName", "DataType"]


def _style_row(ws, r, n_cols, fill):
    for c in range(1, n_cols + 1):
        cell = ws.cell(r, c)
        cell.font, cell.fill = _BOLD, fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def _yn(v, yes="Yes", no="No"):
    return "" if v is None else (yes if v else no)


def _render_sheet_per_table(spec, units, out_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "FILE_DETAILS"
    ws.append(["Vendor", "FileName", "File Description", "Location", "Frequency"])
    _style_row(ws, 1, 5, _HDR_FILL)
    for u in units:
        f = u["feed"]
        ws.append([f.get("source_system") or "", "; ".join(f.get("file_name_patterns") or []),
                   _source_level_text(u["placement"]), f.get("landing_location") or "",
                   f.get("frequency") or ""])
    vh = wb.create_sheet("VERSION_HISTORY")
    vh.append(["Version", "Date", "Author", "Change Description"])
    _style_row(vh, 1, 4, _HDR_FILL)
    vh.append(["0.1", datetime.now(timezone.utc).date().isoformat(), GENERATOR,
               f"Auto-generated from {spec.get('source_file', 'FRD')}"])
    n_src, n_tgt = len(_SRC_HEADERS), len(_TGT_HEADERS)
    for u in units:
        rows, placement = u["rows"], u["placement"]
        if not rows:
            continue
        table = u["source"]["layers"]["stage"]["tables"][:1] or [u["feed"].get("feed_name") or "SOURCE"]
        ws = wb.create_sheet(f"MAPPING-{str(table[0]).upper()}"[:31])
        recycle = placement["recycle"]
        n_cols = n_src + n_tgt + 1 + n_tgt + (1 if recycle else 0)
        ws.cell(1, 1, "Source File Layout")
        ws.cell(1, n_src + 1, "Stage Layer")
        ws.cell(1, n_src + n_tgt + 2, "Standard Layer")
        _style_row(ws, 1, n_cols, _SEC_FILL)
        headers = _SRC_HEADERS + _TGT_HEADERS + [""] + _TGT_HEADERS + (["Recycle Flag"] if recycle else [])
        for c, h in enumerate(headers, 1):
            ws.cell(2, c, h)
        _style_row(ws, 2, n_cols, _HDR_FILL)
        for r in rows:
            comment = "\n".join(placement["by_column"].get(r["source_column"], [])
                                + ([r["comment"]] if r.get("comment") else []))
            line = [r["source_column"], _yn(r["nullable"], "NULL", "Not NULL"), r["description"],
                    r["sample"], r["datatype"] or "String", _yn(r["phi"]), _yn(r["mandatory"]), comment,
                    r["stage"]["schema"], r["stage"]["table"], r["stage"]["column"], r["stage"]["datatype"],
                    "", r["standard"]["schema"], r["standard"]["table"], r["standard"]["column"],
                    r["standard"]["datatype"]]
            if recycle:
                line.append(f"Y ( {recycle['text']} )" if recycle["column"] == r["source_column"] else "")
            ws.append(line)
        ws.freeze_panes = "A3"
        for c in range(1, n_cols + 1):
            ws.column_dimensions[get_column_letter(c)].width = 22
    wb.save(out_path)
    return {"dialect": "sheet_per_table", "layout_from": None}


def _render_single_sheet(spec, units, out_path):
    wb = Workbook()
    u = units[0]
    feed, placement, rows = u["feed"], u["placement"], u["rows"]
    ws = wb.active
    ws.title = (feed.get("feed_name") or "mapping")[:31]
    meta = [
        ("File(s)", "\n".join(feed.get("file_name_patterns") or [])),
        ("File Generator", feed.get("source_system") or ""),
        ("File Location", feed.get("landing_location") or ""),
        ("LOB", ",".join(re.sub(r"^REG#\d+\s+", "", x) for x in feed.get("lobs") or [])),
        ("File frequency", feed.get("frequency") or ""),
        ("Domain", (feed.get("domain") or "").upper()),
        ("Sub-Domain", feed.get("sub_domain") or ""),
        ("File type", f"{feed.get('file_format') or ''}"
                      + (f" ({feed.get('delimiter')} delimited)" if feed.get("delimiter") else "")),
    ]
    if placement["recycle"]:
        meta.append(("Recycle rule", placement["recycle"]["text"]))
    if placement["unattributed"]:
        meta.append(("Validation rules (source-level)", "\n".join(placement["unattributed"])))
    for k, v in meta:
        ws.append([k, v])
        ws.cell(ws.max_row, 1).font = _BOLD
    src_h = ["#", "Field Name", "Data Type", "Length", "Field Length\n(fixed width)",
             "Start position\n(fixed width)", "End Position\n(fixed width)", "Segment", "PII",
             "Comments", "Business Rule"]
    tgt_h = ["Catalog", "Schema", "TableName", "ColumnName", "DataType", "Mandatory\nColumn",
             "Primary Key", "Field Description"]
    label_r = ws.max_row + 1
    ws.cell(label_r, 1, "Source Layout")
    ws.cell(label_r, len(src_h) + 1, "Stage Layer")
    ws.cell(label_r, len(src_h) + len(tgt_h) + 1, "Standard Layer")
    n_cols = len(src_h) + 2 * len(tgt_h)
    _style_row(ws, label_r, n_cols, _SEC_FILL)
    for c, h in enumerate(src_h + tgt_h + tgt_h, 1):
        ws.cell(label_r + 1, c, h)
    _style_row(ws, label_r + 1, n_cols, _HDR_FILL)
    for idx, r in enumerate(rows, 1):
        def tgt(t):
            return [t["catalog"], t["schema"], t["table"], t["column"], t["datatype"] or "String",
                    _yn(r["mandatory"], "Yes", ""), "", r["description"]]
        rule = "\n".join(placement["by_column"].get(r["source_column"], [])
                         + ([r["business_rule"]] if r.get("business_rule") else []))
        ws.append([idx, r["source_column"], r["datatype"], r["length"], r["fixed_length"],
                   r["fixed_start"], r["fixed_end"], r["segment"], _yn(r["phi"], "Yes", ""),
                   r["description"] or r["comment"], rule] + tgt(r["stage"]) + tgt(r["standard"]))
    ws.freeze_panes = ws.cell(label_r + 2, 1).coordinate
    for c in range(1, n_cols + 1):
        ws.column_dimensions[get_column_letter(c)].width = 18
    wb.save(out_path)
    return {"dialect": "single_sheet", "layout_from": None}


# --------------------------------------------------------------------------- #
# render INTO an approved workbook's layout (structure only)
# --------------------------------------------------------------------------- #
def _field_value(r, role, placement, sheet_has_comment):
    kind, key = role
    if kind == "source":
        rules = placement["by_column"].get(r["source_column"], [])
        if key == "source_column":
            return r["source_column"]
        if key == "nullable_raw":
            return _yn(r["nullable"], "NULL", "Not NULL")
        if key == "phi_raw":
            return _yn(r["phi"])
        if key == "mandatory_raw":
            return _yn(r["mandatory"])
        if key == "comment":
            return "\n".join(rules + ([r["comment"]] if r.get("comment") else []))
        if key == "business_rule":
            own = [] if sheet_has_comment else rules
            return "\n".join(own + ([r["business_rule"]] if r.get("business_rule") else []))
        if key == "datatype":
            return r.get("datatype") or "String"
        return r.get(key, "") or ""
    if kind in ("stage", "standard"):
        return (r.get(kind) or {}).get(key) or ""
    if kind == "recycle":
        rec = placement["recycle"]
        return f"Y ( {rec['text']} )" if rec and rec["column"] == r["source_column"] else ""
    return ""


def _fill_sheet(ws, sheet_layout, rows, placement):
    first, width, cols = sheet_layout["first_data_row"], sheet_layout["width"], sheet_layout["columns"]
    styles = None
    if ws.max_row >= first:
        styles = [(copy(ws.cell(first, c).font), copy(ws.cell(first, c).fill),
                   copy(ws.cell(first, c).border), copy(ws.cell(first, c).alignment),
                   ws.cell(first, c).number_format) for c in range(1, width + 1)]
        ws.delete_rows(first, ws.max_row - first + 1)
    has_comment = any(c["role"] == ("source", "comment") for c in cols)
    for n, r in enumerate(rows, 1):
        rr = first + n - 1
        for c in cols:
            if c["role"] is None:
                continue
            val = n if c["role"] == ("index", None) else _field_value(r, c["role"], placement, has_comment)
            ws.cell(rr, c["index"] + 1).value = val if val != "" else None
        if styles:
            for ci, (fo, fi, bo, al, nf) in enumerate(styles, 1):
                cell = ws.cell(rr, ci)
                cell.font, cell.fill, cell.border, cell.alignment, cell.number_format = fo, fi, bo, al, nf


def _render_into_sheet_per_table(spec, units, layout, out_path):
    wb = load_workbook(layout["path"])
    sheets = layout["sheets"]
    objs = {n: wb[n] for n in sheets if n in wb.sheetnames}
    lead = next(iter(sheets))
    info = {"dialect": "sheet_per_table", "layout_from": Path(layout["path"]).name,
            "sheets": {}, "unfilled_columns": {}, "removed_sheets": []}
    fd = layout.get("file_details")
    if fd and "FILE_DETAILS" in wb.sheetnames:
        ws = wb["FILE_DETAILS"]
        if ws.max_row >= 2:
            ws.delete_rows(2, ws.max_row - 1)
        for i, u in enumerate(units, 2):
            f = u["feed"]
            vals = {"vendor": f.get("source_system") or "",
                    "file_name": "; ".join(f.get("file_name_patterns") or []),
                    "description": _source_level_text(u["placement"]),
                    "location": f.get("landing_location") or "", "frequency": f.get("frequency") or ""}
            for k, v in vals.items():
                if k in fd["columns"] and v:
                    ws.cell(i, fd["columns"][k] + 1, v)
    vh = layout.get("version_history")
    if vh and "VERSION_HISTORY" in wb.sheetnames:
        ws = wb["VERSION_HISTORY"]
        if ws.max_row >= 2:
            ws.delete_rows(2, ws.max_row - 1)
        vals = {"version": "0.1", "date": datetime.now(timezone.utc).date().isoformat(),
                "author": GENERATOR, "change": f"Auto-generated from {spec.get('source_file', 'FRD')}"}
        for k, v in vals.items():
            if k in vh["columns"]:
                ws.cell(2, vh["columns"][k] + 1, v)
    used = set()
    for i, u in enumerate(units):
        if not u["rows"]:
            continue
        src = lead if lead not in used else next((n for n in sheets if n not in used), lead)
        table = (u["source"]["layers"]["stage"]["tables"] or [u["feed"].get("feed_name") or "SOURCE"])[0]
        title = f"MAPPING-{str(table).upper()}"[:31]
        if src in used:
            ws = wb.copy_worksheet(objs[src])
        else:
            ws = objs[src]
            used.add(src)
        ws.title = title if title not in wb.sheetnames or wb[title] is ws else f"{title[:28]}-{i}"
        _fill_sheet(ws, sheets[src], u["rows"], u["placement"])
        info["sheets"][ws.title] = len(u["rows"])
        if sheets[src]["unmapped"]:
            info["unfilled_columns"][ws.title] = sheets[src]["unmapped"]
    for name in list(sheets):
        if name not in used and name in objs:
            wb.remove(objs[name])
            info["removed_sheets"].append(name)
    wb.save(out_path)
    return info


def _render_into_single_sheet(spec, units, layout, out_path):
    wb = load_workbook(layout["path"])
    ws = wb[layout["sheet"]]
    u = units[0]
    feed, placement = u["feed"], u["placement"]
    info = {"dialect": "single_sheet", "layout_from": Path(layout["path"]).name,
            "sheets": {layout["sheet"]: len(u["rows"])},
            "unfilled_columns": {layout["sheet"]: layout["unmapped"]} if layout["unmapped"] else {},
            "removed_sheets": []}
    for m in layout["meta_rows"]:
        field, val = m["field"], None
        if field == "file_name_patterns":
            val = "\n".join(feed.get("file_name_patterns") or [])
        elif field == "lobs":
            val = ",".join(re.sub(r"^REG#\d+\s+", "", x) for x in feed.get("lobs") or [])
        elif field == "file_type":
            val = (f"{feed.get('file_format') or ''}"
                   + (f" ({feed.get('delimiter')} delimited)" if feed.get("delimiter") else "")) or None
        elif field == "domain":
            val = (feed.get("domain") or "").upper() or None
        elif field is not None:
            val = feed.get(field) or None
        ws.cell(m["row"], 2).value = val
    leftover = _source_level_text(placement)
    if leftover:
        info["source_level_rules"] = leftover
    _fill_sheet(ws, layout, u["rows"], placement)
    # Every other sheet carries the template feed's own content (a reference
    # mapping, a lookup) — structure is borrowed, content never is.
    for name in wb.sheetnames:
        if name != layout["sheet"]:
            del wb[name]
            info["removed_sheets"].append(name)
    wb.save(out_path)
    return info


# --------------------------------------------------------------------------- #
# entry point
# --------------------------------------------------------------------------- #
def render_workbook(spec: dict, sources: list[dict], vdd: dict, out_path: str | Path,
                    layout: dict | None = None, pairing_override: dict | None = None) -> dict:
    """Build every source's rows and write the workbook. Returns render info
    (dialect, layout used, rows per sheet, unfilled template columns, rule
    placement, unpromoted vendor types)."""
    override = pairing_override or {}
    files_by_pattern = {f["file_name_pattern"]: f for f in vdd.get("files", [])}
    units, notes = [], {"unpromoted_types": set(), "rule_placement": []}
    for s in sources:
        feed = spec["feeds"][s["feed_index"]]
        pattern = override.get(s["feed_index"]) or s.get("file")
        sheet = files_by_pattern.get(pattern, {}).get("field_sheet") if pattern else None
        fields = vdd.get("fields", {}).get(sheet, []) if sheet else []
        rows, rn = build_rows(s, feed, fields) if fields else ([], {"unpromoted_types": []})
        placement = place_rules(feed, rows)
        notes["unpromoted_types"].update(rn.get("unpromoted_types", []))
        notes["rule_placement"].append({"source": s["feed_name"], **placement})
        units.append({"source": s, "feed": feed, "rows": rows, "placement": placement, "file": pattern})
    if not any(u["rows"] for u in units):
        raise ValueError("no rows to render — no source is paired with a dictionary file that names columns")
    out_path = str(out_path)
    if layout and layout.get("dialect") == "sheet_per_table" and layout.get("sheets"):
        info = _render_into_sheet_per_table(spec, units, layout, out_path)
    elif layout and layout.get("dialect") == "single_sheet" and layout.get("sheet") and len(units) == 1:
        info = _render_into_single_sheet(spec, units, layout, out_path)
    elif len(units) == 1:
        info = _render_single_sheet(spec, units, out_path)
    else:
        info = _render_sheet_per_table(spec, units, out_path)
    info["n_rows"] = sum(len(u["rows"]) for u in units)
    info["rows_per_source"] = {u["source"]["feed_name"]: len(u["rows"]) for u in units}
    info["files_per_source"] = {u["source"]["feed_name"]: u["file"] for u in units}
    info["unpromoted_types"] = sorted(notes["unpromoted_types"])
    info["rule_placement"] = notes["rule_placement"]
    return info


def preview_rows(spec: dict, sources: list[dict], vdd: dict, pairing_override: dict | None = None,
                 limit: int = 400) -> list[dict]:
    """The same rows the workbook gets, for the app's table (no file written)."""
    override = pairing_override or {}
    files_by_pattern = {f["file_name_pattern"]: f for f in vdd.get("files", [])}
    out = []
    for s in sources:
        pattern = override.get(s["feed_index"]) or s.get("file")
        sheet = files_by_pattern.get(pattern, {}).get("field_sheet") if pattern else None
        fields = vdd.get("fields", {}).get(sheet, []) if sheet else []
        rows, _ = build_rows(s, spec["feeds"][s["feed_index"]], fields) if fields else ([], {})
        out.append({"source": s["feed_name"], "file": pattern, "n_rows": len(rows),
                    "rows": [{"source_column": r["source_column"], "datatype": r["datatype"],
                              "stage": r["stage"], "standard": r["standard"], "audit": r["audit"]}
                             for r in rows[:limit]]})
    return out
