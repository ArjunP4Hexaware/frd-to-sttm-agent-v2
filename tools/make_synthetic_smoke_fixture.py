"""Generate a fully SYNTHETIC local smoke fixture — no client content.

Since 2026-08-22 the repo carries no FRD/STTM material of any kind (see
CLAUDE.md "Fixtures & data rules"), so a fresh clone cannot exercise the
pipeline offline. This tool replaces that capability without re-tracking
documents: it fabricates, under gitignored `local_dev_fixtures/`,

- two FRDs (plain .txt — 01 parses txt/md natively, no python-docx needed),
- two matching reference STTM workbooks (sheet_per_table dialect),
- two extraction JSONs whose every strict field appears VERBATIM in its
  FRD (so 03's grounding audit passes honestly — the same property the
  hand-authored mock specs had),

and prints the exact commands for the offline smoke:

    .venv/bin/python tools/make_synthetic_smoke_fixture.py
    .venv/bin/python notebooks/01_frd_ingest.py
    .venv/bin/python notebooks/03_contract_build.py     # 02 skipped: extractions pre-written
    .venv/bin/python notebooks/04_sttm_render.py

The two documents are deliberately structural twins with disjoint
vocabulary (member-risk vs claim-intake), which makes them a minimal but
real corpus for the template architecture: build the corpus index and the
render stage will pair each FRD with its own workbook, exclude it from
template candidacy, choose the OTHER workbook (or freeform, depending on
thresholds), and still eval against its own — the full cross-validation
path, offline, zero API calls.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "src"))

from openpyxl import Workbook  # noqa: E402

from frdsttm.models import Feed, FrdIngestionSpec, Project, TableTarget  # noqa: E402

LOCAL_ROOT = REPO / "local_dev_fixtures"

# The source-band headers CodeGen's extractor requires (its header_synonyms),
# as every real client workbook in this dialect carries them.
SRC_HDRS = ["Database Column Name", "Description", "Sample Value", "Datatype", "Null Check",
            "PHI Field", "Mandatory Field", "Comment"]
# Stage/standard bands as the client's sheet-per-table dialect has them (no
# Catalog column there — that belongs to the single-sheet dialect).
TGT_HDRS = ["Schema", "TableName", "ColumnName", "Datatype"]
# (column, datatype) — synthetic names; datatypes are the two CodeGen accepts
# for audit columns ("string" / "timestamp").
AUDIT_ROWS = [("SYN_SRC_FILE", "string"), ("SYN_LOAD_TS", "timestamp")]

DOCS = {
    "synthetic_member_risk": {
        "project_id": "9100001",
        "table": "SYN_MEMBER_RISK",
        "pattern": "SYN_MEMBER_RISK_YYYYMMDD.txt",
        "columns": ["MEMBER_ID", "ZIP_CODE", "RISK_SCORE", "EFFECTIVE_DATE", "LOB_CODE"],
        "rule": "Process shall reject the record when MEMBER_ID is NULL and move it to the reject table.",
    },
    "synthetic_claim_intake": {
        "project_id": "9100002",
        "table": "SYN_CLAIM_INTAKE",
        "pattern": "SYN_CLAIM_INTAKE_YYYYMMDD.txt",
        "columns": ["CLAIM_NUMBER", "PROVIDER_NPI", "PAID_AMOUNT", "SERVICE_CODE", "ADJUDICATION_FLAG"],
        "rule": "Process shall fail the file when CLAIM_NUMBER is NULL or duplicated within the batch.",
    },
}


def frd_text(doc_id: str, d: dict) -> str:
    cols = ", ".join(d["columns"])
    return f"""# Functional Requirements Document — {doc_id} (SYNTHETIC)

Project ID: {d['project_id']}
Project Name: {project_name(doc_id)}

## Business Context

This synthetic document exists only to smoke-test the pipeline offline.
Every fact below is fabricated.

## Data Ingestion Requirements

The Synthetic Vendor delivers the pipe (|) delimited text file {d['pattern']}
weekly to /synthetic/landing/{doc_id}. The feed loads the columns {cols}
into the stage table syn_cat.syn_stg.{d['table']} (Truncate and Load) and is
promoted to syn_cat.syn_std.{d['table']} (Truncate and Load).

**REQ-001** {d['rule']}

## Data Quality

{d['rule']}
"""


def project_name(doc_id: str) -> str:
    return f"Synthetic {doc_id.replace('synthetic_', '').replace('_', ' ').title()} Ingestion"


def spec_for(doc_id: str, d: dict) -> FrdIngestionSpec:
    # project_name (strict-grounded), delimiter and load_strategy are stated
    # in frd_text() above so the contract round-trips through CodeGen's
    # `extract-sttm` unpatched — its FrdContract requires all three
    # (non-null name, a delimiter for 'txt', a load_strategy literal).
    return FrdIngestionSpec(
        project=Project(project_id=d["project_id"],
                        project_name=project_name(doc_id),
                        business_context_summary=None),
        feeds=[Feed(
            feed_name=f"{doc_id} feed",
            source_system="Synthetic Vendor",
            file_name_patterns=[d["pattern"]],
            file_format="txt",
            delimiter="|",
            frequency="weekly",
            landing_location=f"/synthetic/landing/{doc_id}",
            stage_target=TableTarget(catalog="syn_cat", schema="syn_stg",
                                     tables=[d["table"]], load_strategy="Truncate and Load"),
            standard_target=TableTarget(catalog="syn_cat", schema="syn_std",
                                        tables=[d["table"]], load_strategy="Truncate and Load"),
            validation_rules=[d["rule"]],
            requirement_ids=["REQ-001"],
        )],
    )


def workbook_for(path: Path, d: dict) -> None:
    wb = Workbook()
    # The client dialect's bookkeeping sheets (CodeGen's extractor requires
    # both; 04's template fill keeps them from the template).
    fd = wb.active
    fd.title = "FILE_DETAILS"
    fd.append(["Vendor", "FileName", "File Description", "Location", "Frequency"])
    fd.append(["Synthetic Vendor", d["pattern"], "", f"/synthetic/landing/{path.stem.replace('.sttm', '')}", "weekly"])
    vh = wb.create_sheet("VERSION_HISTORY")
    vh.append(["Version", "Date", "Author", "Change Description"])
    vh.append(["1.0", "2026-01-01", "synthetic analyst", "hand-built reference"])
    ws = wb.create_sheet(f"MAPPING-{d['table']}"[:31])
    ws.append(["Source File Layout"] + [""] * (len(SRC_HDRS) - 1)
              + ["Stage Layer"] + [""] * (len(TGT_HDRS) - 1)
              + ["Standard Layer"] + [""] * (len(TGT_HDRS) - 1))
    ws.append(SRC_HDRS + TGT_HDRS + TGT_HDRS)
    for col in d["columns"]:
        ws.append([col, f"synthetic {col.lower().replace('_', ' ')}", "", "String",
                   "Not Null", "No", "Yes", ""]
                  + ["syn_stg", d["table"], col, "String"]
                  + ["syn_std", d["table"], col, "String"])
    # Trailing audit rows, as every client workbook in this dialect carries
    # them: source "NA", the ETL audit column named only on the target side.
    # 04 derives these from the template (not 1:1), and CodeGen's
    # `extract-sttm` requires at least one per mapping sheet.
    for col, dtype in AUDIT_ROWS:
        ws.append(["NA", f"synthetic audit column {col.lower()}", "", dtype, "", "", "", ""]
                  + ["syn_stg", d["table"], col, dtype]
                  + ["syn_std", d["table"], col, dtype])
    wb.save(path)


def main() -> None:
    frd_raw = LOCAL_ROOT / "frd_raw"
    reference = LOCAL_ROOT / "sttm_reference"
    extractions = LOCAL_ROOT / "sttm_out" / "extractions"
    for p in (frd_raw, reference, extractions):
        p.mkdir(parents=True, exist_ok=True)

    for doc_id, d in DOCS.items():
        (frd_raw / f"{doc_id}.txt").write_text(frd_text(doc_id, d), encoding="utf-8")
        workbook_for(reference / f"{doc_id}.sttm.xlsx", d)
        (extractions / f"{doc_id}.json").write_text(
            spec_for(doc_id, d).model_dump_json(by_alias=True, indent=2),
            encoding="utf-8")
        print(f"wrote {doc_id}: frd_raw/*.txt, sttm_reference/*.sttm.xlsx, "
              f"sttm_out/extractions/*.json")

    print(json.dumps({"next": [
        ".venv/bin/python notebooks/01_frd_ingest.py",
        ".venv/bin/python notebooks/03_contract_build.py",
        ".venv/bin/python notebooks/04_sttm_render.py",
    ]}, indent=2))


if __name__ == "__main__":
    main()
