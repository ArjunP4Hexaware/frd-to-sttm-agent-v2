"""Score an agent-rendered STTM against the BSA's APPROVED workbook for the same source.

    ~/.virtualenvs/frdsttm/bin/python tools/eval_against_approved.py \\
        <agent .xlsx> <approved STTM_*.xlsx> [--run <run.json>] [--out eval.xlsx]

Both workbooks go through the same reader (``frdsttm.reference_layout.
parse_reference_workbook``), which resolves either client dialect and every
header variant by NAME, so the eval has no parsing of its own: it diffs two
identical structures.

WHAT IS COMPARED
----------------
Tables are matched by their stage table name (the reader's feed key). Inside
a table, rows are keyed by the normalised SOURCE column — except the client's
audit rows (source column ``NA``: LOB, SRC_FILE_NAME, …), which are keyed by
their STAGE column name. Every row lands in one of: matched · agent_only ·
approved_only.

For a matched row, twelve structured fields each get one outcome:

    match           equal after normalisation
    question        the run asked about this attribute and it is unanswered —
                    by design, not a miss (needs --run)
    agent_blank     the BSA filled it, the agent did not
    approved_blank  the agent filled it, the BSA did not
    disagree        both filled, different

Prose (description, sample) is shown side by side and never scored.

A disagreement is NOT automatically an agent error: the BSA may have renamed
a column by judgement, an ACFC standard may be missing from ``standards/``,
or the BSA may be wrong. The per-table sheets and the distinct-disagreement
roll-up exist so a person can tag each cause. The summary states counts, not
a single blended score.

KNOWN BLIND SPOT: the reader coerces a blank source DataType to "String" (and
the renderer never writes a blank one), so "agent_blank" cannot occur for that
field. The summary says so.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

REPO = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO / "src"))

from frdsttm.reference_layout import _nl, parse_reference_workbook  # noqa: E402

FIELDS = [
    ("source.datatype", lambda r: r["field"]["datatype"]),
    ("source.nullable", lambda r: r["field"]["nullable"]),
    ("source.phi", lambda r: r["field"]["phi"]),
    ("source.mandatory", lambda r: r["field"]["mandatory"]),
    ("stage.schema", lambda r: r["targets"]["stage"]["schema"]),
    ("stage.table", lambda r: r["targets"]["stage"]["table"]),
    ("stage.column", lambda r: r["targets"]["stage"]["column"]),
    ("stage.datatype", lambda r: r["targets"]["stage"]["datatype"]),
    ("standard.schema", lambda r: r["targets"]["standard"]["schema"]),
    ("standard.table", lambda r: r["targets"]["standard"]["table"]),
    ("standard.column", lambda r: r["targets"]["standard"]["column"]),
    ("standard.datatype", lambda r: r["targets"]["standard"]["datatype"]),
]
PROSE = [("description", lambda r: r["field"]["description"]),
         ("sample", lambda r: r["field"]["sample"])]
OUTCOMES = ("match", "question", "agent_blank", "approved_blank", "disagree")


# --------------------------------------------------------------------------- #
# rows and cells
# --------------------------------------------------------------------------- #
def row_key(field: dict, targets: dict) -> str:
    """Source column for dictionary rows; stage column for audit rows (whose
    source column is the literal ``NA`` and would collide)."""
    if field.get("audit"):
        return "audit:" + _nl(targets["stage"]["column"])
    return _nl(field["source_column"])


def index_rows(feed: dict) -> dict[str, dict]:
    out = {}
    for field, targets in zip(feed["fields"], feed["ref_targets"]):
        out.setdefault(row_key(field, targets), {"field": field, "targets": targets})
    return out


def _norm(v):
    if isinstance(v, bool):
        return v
    return _nl(v)


def outcome(agent_value, approved_value, questioned: bool = False) -> str:
    a, b = _norm(agent_value), _norm(approved_value)
    if a == b:
        return "match"
    if questioned and a in ("", None):
        return "question"
    if a in ("", None):
        return "agent_blank"
    if b in ("", None):
        return "approved_blank"
    return "disagree"


# --------------------------------------------------------------------------- #
# the run's open questions → (table key, field) pairs the agent chose not to fill
# --------------------------------------------------------------------------- #
_ATTR = {"schema": "schema", "catalog": "catalog", "tables": "table", "table": "table"}
_FIELD_PATH = re.compile(r"(stage|standard)_target\W+(schema|catalog|tables?)")


def questioned_fields(run: dict | None) -> set[tuple[str, str]]:
    """{(stage table key, 'stage.schema'), …} for every UNANSWERED question
    that concerns a target attribute. Tables are resolved through the
    extraction's own stage_target.tables, so the key matches the reader's."""
    if not run:
        return set()
    feeds = (run.get("extraction") or {}).get("feeds") or []
    keys: dict[int, set[str]] = {}
    for i, f in enumerate(feeds):
        tables = (f.get("stage_target") or {}).get("tables") or []
        # sheet_per_table keys by stage table; single_sheet keys by the sheet
        # name, which the renderer sets to the feed name — accept either
        keys[i] = {_nl(t) for t in tables[:1]} | ({_nl(f["feed_name"])} if f.get("feed_name") else set())
    out = set()
    for q in (run.get("assessment") or {}).get("questions") or []:
        if q.get("answer") is not None:
            continue
        ctx = q.get("context") or {}
        layer, attr = None, None
        if q.get("kind") == "target_gap":
            layer, attr = ctx.get("layer"), _ATTR.get(ctx.get("attribute"))
        elif q.get("kind") in ("unverified", "weak_match"):
            m = _FIELD_PATH.search(str(ctx.get("field", "")))
            if m:
                layer, attr = m.group(1), _ATTR.get(m.group(2))
        idx = ctx.get("feed_index")
        if layer and attr and idx in keys:
            out.update((k, f"{layer}.{attr}") for k in keys[idx])
    return out


# --------------------------------------------------------------------------- #
# comparison
# --------------------------------------------------------------------------- #
def pair_tables(agent: dict, approved: dict) -> tuple[list[tuple[str, str, str]], list[str]]:
    """[(label, agent key, approved key)] plus table-level findings. Keys that
    match by name pair directly; if the two sides do not line up 1:1, the
    remainder pairs by sheet order and says so."""
    a_keys, b_keys = list(agent["feeds"]), list(approved["feeds"])
    pairs = [(k, k, k) for k in a_keys if k in approved["feeds"]]
    rest_a = [k for k in a_keys if k not in approved["feeds"]]
    rest_b = [k for k in b_keys if k not in agent["feeds"]]
    findings = []
    for ka, kb in zip(rest_a, rest_b):
        findings.append(f"table key mismatch — agent {ka!r} vs approved {kb!r}: paired by sheet order")
        pairs.append((f"{ka} ~ {kb}", ka, kb))
    for ka in rest_a[len(rest_b):]:
        findings.append(f"table {ka!r} exists only in the agent's workbook")
    for kb in rest_b[len(rest_a):]:
        findings.append(f"table {kb!r} exists only in the approved workbook")
    return pairs, findings


def compare_table(label: str, a_feed: dict, b_feed: dict, questioned: set[str]) -> dict:
    """questioned: the field names (e.g. 'stage.schema') the run left open for this table."""
    ra, rb = index_rows(a_feed), index_rows(b_feed)
    matched = [k for k in ra if k in rb]
    rows, per_field, distinct = [], {f: Counter() for f, _ in FIELDS}, Counter()
    for k in matched:
        a, b = ra[k], rb[k]
        cells = {}
        for name, get in FIELDS:
            if a["field"].get("audit") and name.startswith("source."):
                continue          # an audit row has no source column; its source band is 'NA' by definition
            va, vb = get(a), get(b)
            o = outcome(va, vb, questioned=name in questioned)
            per_field[name][o] += 1
            cells[name] = (o, va, vb)
            if o == "disagree":
                distinct[(name, str(va), str(vb))] += 1
        prose = {name: (get(a), get(b)) for name, get in PROSE}
        rows.append({"key": k, "audit": bool(a["field"].get("audit")), "cells": cells, "prose": prose})
    return {
        "table": label,
        "n_agent": len(ra), "n_approved": len(rb), "n_matched": len(matched),
        "agent_only": [k for k in ra if k not in rb],
        "approved_only": [k for k in rb if k not in ra],
        "per_field": {f: dict(c) for f, c in per_field.items()},
        "distinct_disagreements": sorted(((f, va, vb, n) for (f, va, vb), n in distinct.items()),
                                         key=lambda t: (-t[3], t[0])),
        "rows": rows,
    }


def evaluate(agent_path: str | Path, approved_path: str | Path, run: dict | None = None) -> dict:
    agent, approved = parse_reference_workbook(str(agent_path)), parse_reference_workbook(str(approved_path))
    pairs, findings = pair_tables(agent, approved)
    qf = questioned_fields(run)
    tables = []
    for label, ka, kb in pairs:
        questioned = {f for (t, f) in qf if t == ka}
        tables.append(compare_table(label, agent["feeds"][ka], approved["feeds"][kb], questioned))
    totals = {f: Counter() for f, _ in FIELDS}
    for t in tables:
        for f, c in t["per_field"].items():
            totals[f].update(c)
    return {
        "agent": str(agent_path), "approved": str(approved_path),
        "agent_dialect": agent["dialect"], "approved_dialect": approved["dialect"],
        "findings": findings, "tables": tables,
        "totals": {f: dict(c) for f, c in totals.items()},
        "n_questioned_fields": len(qf),
        "notes": ["source.datatype: a blank is read as 'String' on both sides, so agent_blank "
                  "cannot occur for this field (reader coercion)."],
    }


# --------------------------------------------------------------------------- #
# report
# --------------------------------------------------------------------------- #
_FILL = {"match": "E2F0D9", "question": "FFF2CC", "agent_blank": "FCE4D6",
         "approved_blank": "DDEBF7", "disagree": "F8CBAD"}


def summary_lines(result: dict) -> list[str]:
    lines = [f"agent:    {result['agent']}  ({result['agent_dialect']})",
             f"approved: {result['approved']}  ({result['approved_dialect']})", ""]
    lines += [f"! {f}" for f in result["findings"]]
    for t in result["tables"]:
        lines.append(f"[{t['table']}]  agent rows {t['n_agent']} · approved rows {t['n_approved']} · "
                     f"matched {t['n_matched']} · agent-only {len(t['agent_only'])} · "
                     f"approved-only {len(t['approved_only'])}")
    lines.append("")
    lines.append(f"{'field':<20}" + "".join(f"{o:>15}" for o in OUTCOMES))
    for f, c in result["totals"].items():
        lines.append(f"{f:<20}" + "".join(f"{c.get(o, 0):>15}" for o in OUTCOMES))
    lines.append("")
    n_dis = sum(len(t["distinct_disagreements"]) for t in result["tables"])
    lines.append(f"distinct disagreements to adjudicate: {n_dis}")
    for t in result["tables"]:
        for f, va, vb, n in t["distinct_disagreements"][:10]:
            lines.append(f"  {t['table']} · {f}: agent {va!r} vs approved {vb!r}  ×{n}")
    lines += [""] + [f"note: {n}" for n in result["notes"]]
    return lines


def write_report(result: dict, out_path: str | Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    for line in summary_lines(result):
        ws.append([line])
    ws.column_dimensions["A"].width = 120

    ws = wb.create_sheet("Disagreements")
    ws.append(["table", "field", "agent value", "approved value", "rows", "cause (fill in)"])
    for t in result["tables"]:
        for f, va, vb, n in t["distinct_disagreements"]:
            ws.append([t["table"], f, va, vb, n, ""])
    for col, w in zip("ABCDEF", (28, 20, 30, 30, 8, 40)):
        ws.column_dimensions[col].width = w

    bold = Font(bold=True)
    for t in result["tables"]:
        ws = wb.create_sheet(t["table"][:31].replace("/", "-"))
        hdr = ["row key", "audit"] + [f for f, _ in FIELDS] + [f"agent {p}" for p, _ in PROSE] + [f"approved {p}" for p, _ in PROSE]
        ws.append(hdr)
        for c in ws[1]:
            c.font = bold
        for r in t["rows"]:
            line = [r["key"], "Y" if r["audit"] else ""]
            for f, _ in FIELDS:
                o, va, vb = r["cells"].get(f, ("n/a", "", ""))
                line.append("n/a" if o == "n/a" else o if o == "match" else f"{o}: {va!r} → {vb!r}")
            line += [r["prose"][p][0] for p, _ in PROSE] + [r["prose"][p][1] for p, _ in PROSE]
            ws.append(line)
            for i, (f, _) in enumerate(FIELDS, start=3):
                o = r["cells"].get(f, ("n/a",))[0]
                if o in _FILL:
                    ws.cell(ws.max_row, i).fill = PatternFill("solid", fgColor=_FILL[o])
        for k in t["agent_only"]:
            ws.append([k, "", "agent_only"])
        for k in t["approved_only"]:
            ws.append([k, "", "approved_only"])
        ws.freeze_panes = "C2"
        for i in range(1, len(hdr) + 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = 22
    out_path = Path(out_path)
    wb.save(str(out_path))
    return out_path


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("agent_xlsx")
    ap.add_argument("approved_xlsx")
    ap.add_argument("--run", help="the run's run.json — turns unanswered questions into the 'question' outcome")
    ap.add_argument("--out", default="eval.xlsx")
    args = ap.parse_args(argv)
    run = json.loads(Path(args.run).read_text()) if args.run else None
    result = evaluate(args.agent_xlsx, args.approved_xlsx, run)
    print("\n".join(summary_lines(result)))
    print(f"\nreport: {write_report(result, args.out)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
