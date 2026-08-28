"""Derive a WORKED VENDOR DATA DICTIONARY from an approved STTM's source band.

    ~/.virtualenvs/frdsttm/bin/python scripts/build_vdd_from_sttm.py \\
        <STTM_*.xlsx> [out.xlsx]

Default output: local_dev_fixtures/vdd_raw/VDD_<stem after STTM_>.xlsx

WHAT THIS IS, AND WHAT IT IS NOT
--------------------------------
A vendor data dictionary is supposed to come FROM THE VENDOR. Nobody has one
yet, so this builds a stand-in by reading the SOURCE band of a mapping the
analysts already produced — the band that is itself a transcription of the
vendor's spec (see the 2026-08-25 decision in CLAUDE.md).

That makes the output a **worked example, not evidence**, and the README sheet
it writes says so in the workbook itself. Two things follow, and both are the
point rather than a limitation:

* It can only carry what the STTM carries. For CAQH the source band has no
  PHI flag and no mandatory flag on ANY of its 115 fields, so the generated
  dictionary has neither, and `frdsttm.dictionary` reports both as gaps. That
  is the honest answer: those facts exist only in the vendor's specification,
  which is exactly the argument slide 3 of the input-requirements deck makes.
* It is CIRCULAR for evaluation. A dictionary derived from an STTM must never
  be used to score a run against that same STTM — the source side would be
  grounded in the answer. Use it to exercise the ingestion path and to show a
  vendor what "filled in" looks like; never as a golden input.

The SD dictionary was hand-built on 2026-08-25 and so could not be
reproduced. This script replaces that: both worked examples now come from a
command, and a corrected STTM regenerates its dictionary.

HOW IT READS AN STTM
--------------------
Only the SOURCE band, bounded on the right by the next band label ("Stage
Layer", "STG - Dest 1", …). That boundary is load-bearing: both real workbooks
have a second `DataType` column in the STAGE band, and reading past the
boundary would silently take the TARGET type as if it were the vendor's.
Headers are matched tolerantly (case, punctuation, trailing spaces) because
the two real workbooks disagree on almost every spelling — even between sheets
of the SAME workbook ("Database column Name" vs "Database Name").
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

REPO = Path(__file__).resolve().parent.parent
OUT_DIR = REPO / "local_dev_fixtures" / "vdd_raw"

FONT = "Arial"
INK, CORE_BG, OPT_BG = "1F2933", "1F3A5F", "5B7C99"
GREY, WARN = "8A94A0", "B23B3B"
_thin = Side(style="thin", color="C6CED6")
BOX = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

# The template's column order (scripts/build_dict_template.py). Kept identical
# so ONE parser reads the template, the hand-built example and this output.
FILES_HEADERS = [
    "File Name Pattern", "File Title", "Format", "Delimiter", "Header Row",
    "Encoding", "Delivery Cadence", "Content Description", "Field Sheet",
    "File Generator", "Text Qualifier", "Line Ending", "Multi-Record-Type",
    "Record Type Field", "Record Type Values", "Expected Field Count",
    "Approx Rows per Delivery", "Trailer / Control Record", "Notes",
]
FIELD_HEADERS = [
    "Position", "Field Name", "Data Type", "Length", "Required (Y/N)",
    "Description", "Allowed Values / Range", "Example Value", "PHI/PII (Y/N)",
    "Segment", "Business Name", "Precision", "Scale", "Format / Pattern",
    "Default Value", "Start Position", "End Position", "Notes",
]

#: Labels that mark the START of a TARGET band. Everything from the first of
#: these rightwards is the client's decision, not the vendor's, and must not
#: leak into the dictionary.
TARGET_BAND = ("stage layer", "std", "stg", "standard layer", "dest", "target")

#: source-band header → our key. First match wins.
FIELD_MAP = {
    # Three dialects, and two of them are in the SAME workbook — SD names the
    # source column differently on every sheet.
    "name": ("database column name", "database name",
             "client data table column name", "client data table column",
             "field name", "source column", "column name", "source field"),
    # NOTE "#" is handled by RAW_HEADERS below, not here: norm("#") is the
    # empty string, so a punctuation-only header can never match this table.
    "position": ("no", "sr no", "position", "seq", "ordinal"),
    "datatype": ("data type", "datatype", "source datatype"),
    "length": ("length", "field length fixed width"),
    "description": ("description", "comments", "field description"),
    "sample": ("sample value", "example values", "example value", "sample"),
    # SD's INDIV_RISK sheet puts the vendor's value domain in "Comment"
    # ("integer, valid range 1 - 5"), which is Allowed Values, not a note.
    "allowed": ("comment", "allowed values range", "allowed values", "valid values"),
    "nullcheck": ("null check", "nullcheck", "null"),
    "mandatory": ("mandatory field", "mandatory column", "mandatory", "required"),
    "phi": ("phi pii field", "phi field", "phi pii", "pii", "phi"),
    "segment": ("segment ex header trailer detail", "segment", "record type"),
    "business_rule": ("business rule",),
    "start": ("start position fixed width", "start position"),
    "end": ("end position fixed width", "end position"),
}


#: Headers that survive no normalisation. Both real STTMs number their source
#: rows with a bare "#", and norm() reduces that to "" — so it is matched on
#: the raw, stripped text instead. Missing this silently fell back to row
#: order, which flattened the per-segment restart (Detail began at 9, not 1).
RAW_HEADERS = {"#": "position", "#.": "position", "no.": "position", "s.no": "position"}


def norm(v) -> str:
    if v is None:
        return ""
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(v).lower()).split())


def text(v) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def yn(v) -> str | None:
    """A source band's null/mandatory wording → Y/N, or None when unstated.

    None is preserved rather than defaulted: "the analyst did not record
    whether this is mandatory" and "it is not mandatory" are different facts,
    and only one of them is safe to hand a vendor as if it were their answer.
    """
    s = (text(v) or "").lower()
    if not s:
        return None
    if s.startswith("not null") or s in ("yes", "y", "true", "mandatory"):
        return "Y"
    if s.startswith("null") or s in ("no", "n", "false", "optional"):
        return "N"
    return None


def find_header_row(rows: list) -> tuple[int, dict, int]:
    """(row index, {col: key}, target_band_start_col) for a mapping sheet.

    Scans for the row that names a source column AND a datatype or a
    description — the two shapes both real dialects share.
    """
    for i, row in enumerate(rows[:14]):
        mapping: dict[int, str] = {}
        for j, cell in enumerate(row):
            raw = str(cell).strip().lower() if cell is not None else ""
            if raw in RAW_HEADERS and RAW_HEADERS[raw] not in mapping.values():
                mapping[j] = RAW_HEADERS[raw]
                continue
            n = norm(cell)
            if not n:
                continue
            for key, spellings in FIELD_MAP.items():
                if n in spellings and key not in mapping.values():
                    mapping[j] = key
                    break
        if "name" in mapping.values() and (
            "datatype" in mapping.values() or "description" in mapping.values()
        ):
            return i, mapping, band_boundary(rows, i, mapping)
    raise SystemExit("could not find a source-band header row in the first 14 rows")


def band_boundary(rows: list, header_row: int, mapping: dict) -> int:
    """First column belonging to a TARGET band, or a large number.

    Both real workbooks put a band label ("Stage Layer") one row ABOVE the
    header row, spanning the target columns. Without this cut the second
    `DataType` column — the STAGE type, a client decision — would be read as
    the vendor's source type.
    """
    for probe in (header_row - 1, header_row - 2):
        if probe < 0:
            continue
        for j, cell in enumerate(rows[probe]):
            n = norm(cell)
            if n and any(n.startswith(t) for t in TARGET_BAND):
                return j
    return 10_000


def read_source_band(ws) -> list[dict]:
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    hdr, mapping, cut = find_header_row(rows)
    mapping = {j: k for j, k in mapping.items() if j < cut}
    out = []
    for row in rows[hdr + 1:]:
        rec = {k: None for k in FIELD_MAP}
        for j, key in mapping.items():
            if j < len(row):
                rec[key] = text(row[j])
        if not rec["name"]:
            continue                        # spacer rows
        if rec["name"].strip().upper() in ("NA", "N/A", "-"):
            # The trailing AUDIT rows (LOB, SRC_FILE_NAME, REC_CREATION_TIME,
            # REC_UPDATED_TIME). They are the CLIENT's columns — derived by
            # the pipeline, with no source field — so they must never appear
            # in a document a VENDOR is asked to fill in. Dropping them is
            # what makes the counts match the vendor-side reality (SD:
            # 86/267/46, not 90/271/50).
            continue
        out.append(rec)
    return out


def read_metadata(ws) -> dict:
    """The leading key/value block some dialects put above the header row."""
    meta = {}
    for row in list(ws.iter_rows(max_row=12, values_only=True)):
        k, v = text(row[0] if row else None), text(row[1] if len(row) > 1 else None)
        if k and v and len(k) < 40:
            meta[norm(k)] = v
    return meta


# --------------------------------------------------------------------------- #
# writing
# --------------------------------------------------------------------------- #
def head(ws, headers, n_core=9):
    for i, name in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i, value=name)
        c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=CORE_BG if i <= n_core else OPT_BG)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "A2"


def body(ws, row, values):
    for i, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=i, value=v)
        c.font = Font(name=FONT, size=10, color=INK)
        c.alignment = Alignment(vertical="top", wrap_text=True)
        c.border = BOX


def assign_files(keys: list[str], file_details: dict) -> tuple[dict, list[str]]:
    """{sheet key: file name pattern}, plus notes. Three passes, in order:

    1. **decisive lexical match** — a sheet's tokens overlap exactly one
       unclaimed file's, and beat every runner-up;
    2. **elimination** — one sheet and one file left, so it is the only
       assignment possible. SD needs this: "SD_COMMUNITY_RISK" ↔
       "analytics_package_YYYY_MM.csv" share not one token, and no amount of
       string cleverness recovers a pairing that is not in the names;
    3. **placeholder + note** — anything still ambiguous gets a synthesised
       name and a line in the README telling whoever sends this to a vendor to
       fix it. A WRONG file pattern would send the landing check hunting for a
       file that does not exist, which is worse than an obvious blank.

    A file is claimed at most once: without that, two sheets happily matched
    the same pattern and one feed silently described another feed's file.
    """
    assigned: dict[str, str] = {}
    notes: list[str] = []
    if not file_details:
        return {k: f"{k}.txt" for k in keys}, notes
    remaining, unclaimed = list(keys), list(file_details)

    # GLOBAL best-first, not greedy in sheet order. Sheet order put
    # "sd_community_risk" ahead of "sd_indiv_risk" and let it claim
    # "sd_ind_risk_data_package…" on the single shared token "risk" — which
    # then pushed the individual-risk sheet onto the wrong file by
    # elimination, describing 267 community columns as an individual-risk
    # feed. Ranking every (sheet, file) pair by score and taking the strongest
    # first gives the sheet that actually shares two tokens the claim.
    pairs = sorted(
        ((_overlap(k, f), k, f) for k in remaining for f in unclaimed),
        key=lambda t: (-t[0], t[1], t[2]),
    )
    for score_, key, fname in pairs:
        if score_ < 1 or key not in remaining or fname not in unclaimed:
            continue
        rivals = [sc for sc, k2, f2 in pairs
                  if f2 == fname and k2 != key and k2 in remaining]
        if rivals and max(rivals) == score_:
            continue                     # a tie for this file: decide nothing
        assigned[key] = fname
        unclaimed.remove(fname)
        remaining.remove(key)

    if len(remaining) == 1 and len(unclaimed) == 1:
        assigned[remaining[0]] = unclaimed[0]
        notes.append(f"sheet {remaining[0]!r} was matched to {unclaimed[0]!r} BY ELIMINATION "
                     f"— their names share nothing. Confirm it before sending this to a vendor.")
        remaining, unclaimed = [], []

    for key in remaining:
        assigned[key] = f"{key}.txt"
        notes.append(f"sheet {key!r} could not be matched to a row on FILE_DETAILS; its file "
                     f"name pattern is a PLACEHOLDER — replace it with the real one before "
                     f"sending this to a vendor")
    return assigned, notes


def _overlap(key: str, fname: str) -> int:
    """Tokens the two names share, allowing a truncated one to match.

    Excel cuts sheet names at 31 characters, so the real pairing is
    "sd_community_DEMOGRAPHI" ↔ "DEMOGRAPHIcs_package_YYYY_MM.csv" — a token
    PREFIX, not an equality, and not a leading-character match either (the two
    names start differently).
    """
    # Tokens of 3+ characters. FOUR was too strict and cost a real pairing:
    # it dropped "ind" from "sd_IND_risk_data_package", which is the only
    # token distinguishing the individual-risk file from the community-risk
    # sheet — both then scored 1 on "risk" alone, tied, and fell through to
    # placeholders. Three keeps the discriminator; the tie guard above still
    # refuses to decide when two sheets really are equally close.
    ft = [t for t in re.split(r"[^a-z0-9]+", sheet_key(fname)) if len(t) > 2]
    kt = [t for t in re.split(r"[^a-z0-9]+", key) if len(t) > 2]
    return sum(1 for a in kt for b in ft if a.startswith(b) or b.startswith(a))


def _unused_match_file(key: str, file_details: dict) -> tuple[str, str | None]:
    """(file name pattern, note). Decisive longest-shared-prefix match, else
    a synthesised name and a note saying so."""
    def score(fname: str) -> int:
        """Tokens the two names share, allowing a truncated one to match.

        Sheet names are cut to 31 characters by Excel, so the real pairing is
        "sd_community_DEMOGRAPHI" ↔ "DEMOGRAPHIcs_package_YYYY_MM.csv" — a
        token PREFIX, not an equality. Leading-character comparison misses it
        entirely because the two names start differently.
        """
        ft = [t for t in re.split(r"[^a-z0-9]+", sheet_key(fname)) if len(t) > 3]
        kt = [t for t in re.split(r"[^a-z0-9]+", key) if len(t) > 3]
        return sum(1 for a_ in kt for b_ in ft if a_.startswith(b_) or b_.startswith(a_))
    if not file_details:
        return f"{key}.txt", None
    ranked = sorted(file_details, key=score, reverse=True)
    best = score(ranked[0])
    runner = score(ranked[1]) if len(ranked) > 1 else -1
    if best >= 1 and best > runner:
        return ranked[0], None
    return (f"{key}.txt",
            f"sheet {key!r} could not be matched to a row on FILE_DETAILS "
            f"(best token overlap {best}); its file name pattern is a placeholder — "
            f"replace it with the real one before sending to a vendor")


def sheet_key(name: str) -> str:
    """A field-sheet name Excel accepts and FILES can point at (<=31 chars)."""
    s = re.sub(r"[^A-Za-z0-9_]+", "_", name).strip("_").lower()
    return (s[:31] or "fields")


def build(sttm: Path, out: Path) -> Path:
    wb = load_workbook(str(sttm), read_only=True, data_only=True)
    mapping_sheets = [n for n in wb.sheetnames
                      if norm(n) not in ("file details", "version history", "forreference",
                                         "revision history", "cover")]
    files, fields, notes = [], {}, []

    file_details = {}
    for n in wb.sheetnames:
        if norm(n) == "file details":
            rows = [list(r) for r in wb[n].iter_rows(values_only=True)]
            hdr = [norm(c) for c in rows[0]]
            for r in rows[1:]:
                rec = dict(zip(hdr, [text(c) for c in r]))
                if rec.get("filename"):
                    file_details[rec["filename"]] = rec

    collected = []          # (key, sheet name, band, metadata)
    for name in mapping_sheets:
        ws = wb[name]
        try:
            band = read_source_band(ws)
        except SystemExit:
            notes.append(f"sheet {name!r} has no readable source band and was skipped")
            continue
        if not band:
            continue
        key = sheet_key(name.replace("MAPPING-", ""))
        collected.append((key, name, band, read_metadata(ws)))
        fields[key] = band

    # File names are assigned ONCE, across all sheets, so a file can be claimed
    # only once and an unresolvable pairing is named rather than guessed.
    inline = {k: (m.get("file s") or m.get("files") or "").split("\n")[0].strip()
              for k, _n, _b, m in collected}
    needs_lookup = [k for k, v in inline.items() if not v]
    assigned, assign_notes = assign_files(needs_lookup, file_details)
    notes.extend(assign_notes)

    for key, name, band, meta in collected:

        # -- one FILES row per mapping sheet. A multi-record file (Header /
        # Detail / Trailer in one stream) is still ONE file, so the segments
        # go on the field rows and the record-type facts go here.
        segs = [s for s in dict.fromkeys(f["segment"] for f in band if f["segment"])]
        pattern = inline[key] or assigned.get(key, f"{key}.txt")
        fd = file_details.get(pattern, {})
        file_type = meta.get("file type") or ""
        delim = None
        m = re.search(r"\(([^)]*)\)", file_type)
        if m and "delim" in file_type.lower():
            delim = m.group(1).replace("pipe delimited", "").strip() or None
        counts = ", ".join(f"{s}={sum(1 for f in band if f['segment'] == s)}" for s in segs)

        files.append([
            pattern,
            fd.get("filedescription") or fd.get("filename") or name.replace("MAPPING-", ""),
            "Delimited text" if ("delim" in file_type.lower() or not file_type) else file_type,
            delim,
            "N" if segs else None,   # a multi-record stream has no column-name row
            None,                    # encoding: not in either STTM
            fd.get("frequency") or meta.get("file frequency"),
            fd.get("filedescription"),
            key,
            meta.get("file generator") or fd.get("vendor"),
            None, None,
            "Y" if len(segs) > 1 else "N",
            "Position 1" if len(segs) > 1 else None,
            " / ".join(segs) if len(segs) > 1 else None,
            counts or len(band),
            None,
            "Y" if any(norm(s) == "trailer" for s in segs) else None,
            "DERIVED FROM THE APPROVED STTM — confirm every cell with the vendor.",
        ])

    # ------------------------------------------------------------------ write
    out_wb = Workbook()
    out_wb.remove(out_wb.active)

    missing = sorted({
        label for key, label in (("phi", "PHI/PII"), ("mandatory", "Required"),
                                 ("sample", "Example Value"))
        if all(f[key] is None for rows in fields.values() for f in rows)
    })

    ws = out_wb.create_sheet("README")
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 108
    rows = [
        ("Vendor Data Dictionary", f"DERIVED WORKED EXAMPLE — {out.stem}"),
        ("PROVENANCE — READ THIS",
         f"Generated by scripts/build_vdd_from_sttm.py from the SOURCE band of "
         f"{sttm.name}. It is NOT the vendor's specification: it is a transcription of "
         f"what analysts had already written down. Treat it as a worked example of a "
         f"filled-in dictionary, never as evidence of what the vendor actually sends."),
        ("Do not use it to score a run",
         "It is circular: a dictionary derived from an STTM must never be used to "
         "evaluate a run against that same STTM, because the source side would be "
         "grounded in the answer."),
        ("What it contains",
         "FILES: one row per file the vendor delivers. Then one sheet per file, one row "
         "per source column: position, name, data type, length, description, segment."),
        ("What it deliberately does NOT contain",
         "Anything the CLIENT decides — target catalog, schema, table, column names, "
         "datatypes, load strategy. Those come from the FRD and the standards contracts."),
    ]
    if missing:
        rows.append((
            "NOT RECOVERABLE from this STTM",
            f"{', '.join(missing)} — blank on every field, because the source band does "
            f"not record them. This is the gap, not a defect of the generator: those "
            f"facts exist only in the vendor's spec. The agent will report them as "
            f"missing and gate the affected cells rather than guess."))
    for note in notes:
        rows.append(("Note", note))
    for i, (k, v) in enumerate(rows, start=1):
        c = ws.cell(row=i, column=1, value=k)
        c.font = Font(name=FONT, size=10, bold=True,
                      color=WARN if "NOT" in k or "Do not" in k else INK)
        c.alignment = Alignment(vertical="top")
        c2 = ws.cell(row=i, column=2, value=v)
        c2.font = Font(name=FONT, size=10, color=INK)
        c2.alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[i].height = 46

    ws = out_wb.create_sheet("FILES")
    head(ws, FILES_HEADERS)
    for i, row in enumerate(files, start=2):
        body(ws, i, row)
    for col, w in zip("ABCDEFGHIJKLMNOPQRS",
                      [40, 26, 16, 11, 11, 11, 22, 40, 22, 18, 13, 12, 17, 16, 26, 24, 20, 20, 46]):
        ws.column_dimensions[col].width = w

    for key, band in fields.items():
        s = out_wb.create_sheet(key)
        head(s, FIELD_HEADERS)
        seen: dict[str | None, int] = {}
        for i, f in enumerate(band, start=2):
            seg = f["segment"]
            seen[seg] = seen.get(seg, 0) + 1
            # The template's rule: Position is the ordinal WITHIN ITS RECORD
            # TYPE and restarts at 1 for each. Prefer the source's own number;
            # fall back to a PER-SEGMENT counter, never to the sheet row.
            body(s, i, [
                int(f["position"]) if (f["position"] or "").isdigit() else seen[seg],
                f["name"], f["datatype"], f["length"],
                yn(f["mandatory"]) or yn(f["nullcheck"]),
                f["description"], f["allowed"], f["sample"], yn(f["phi"]),
                f["segment"], None, None, None, None, None,
                f["start"], f["end"],
                f["business_rule"],
            ])
        for col, w in zip("ABCDEFGHIJKLMNOPQR",
                          [9, 34, 14, 14, 13, 60, 26, 16, 12, 12, 20, 10, 8, 16, 14, 13, 12, 30]):
            s.column_dimensions[col].width = w
        s.freeze_panes = "C2"

    out.parent.mkdir(parents=True, exist_ok=True)
    out_wb.save(str(out))
    wb.close()
    total = sum(len(v) for v in fields.values())
    print(f"{out}\n  {len(files)} file(s), {total} column(s) "
          f"across {len(fields)} field sheet(s)")
    if missing:
        print(f"  NOT recoverable from this STTM: {', '.join(missing)}")
    return out


if __name__ == "__main__":
    if len(sys.argv) < 2:
        raise SystemExit(__doc__.strip().splitlines()[2].strip())
    src = Path(sys.argv[1])
    stem = re.sub(r"^STTM[_ -]*", "", src.stem)
    dest = Path(sys.argv[2]) if len(sys.argv) > 2 else OUT_DIR / f"VDD_{stem}.xlsx"
    build(src, dest)
