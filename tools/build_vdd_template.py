"""Build the unified Vendor Data Dictionary template.

The template is NOT versioned (Arjun, 2026-08-27): one file name, one
current shape, git history is the record of what changed. A version in the
file name means every bump changes the path the FRD's `Source Data
Dictionary` row and the decks point at.

Backward-compatible with the existing DICT_*.xlsx workbooks: the first 9
columns of FILES and of every field sheet are identical in name and order, so
one parser reads the worked example and anything returned on this template.
Verified against a headerless, pipe-delimited, multi-record-type source
(Header/Detail/Trailer record types, per-segment position restart,
SQL-flavoured type vocabulary, non-numeric lengths).
"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

FONT = "Arial"
INK, CORE_BG, OPT_BG = "1F2933", "1F3A5F", "5B7C99"
FILL_ME, EX_GREY, RULE_BG = "FFF2A8", "8A94A0", "EEF2F6"
thin = Side(style="thin", color="C6CED6")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def h(ws, row, headers, n_core, comments=None):
    for i, name in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=name)
        c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=CORE_BG if i <= n_core else OPT_BG)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX
        if comments and name in comments:
            c.comment = Comment(comments[name], "Template")
    ws.row_dimensions[row].height = 36


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def row_from(headers, cells):
    """Turn a header-keyed dict into a list in header order.

    Example rows are written BY POSITION (see ``body``), and the parser
    recognises them only by the marker in the LAST column, Notes. A header
    inserted without its value used to shift every later cell one column
    left, un-marking the row — the blank template then parsed as a real
    file (found 2026-08-27). Keying by header makes that shift impossible: a missing
    key is a blank, and a key that is not a header is an error here, not a
    mystery downstream.
    """
    unknown = set(cells) - set(headers)
    if unknown:
        raise KeyError(f"not a column of this sheet: {sorted(unknown)}")
    return [cells.get(col) for col in headers]


def body(ws, row, values, italic=False, color=INK):
    for i, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=i, value=v)
        c.font = Font(name=FONT, size=10, italic=italic, color=color)
        c.alignment = Alignment(vertical="top", wrap_text=True)
        c.border = BOX


def blanks(ws, first, last, ncol):
    for row in range(first, last + 1):
        for i in range(1, ncol + 1):
            c = ws.cell(row=row, column=i)
            c.border = BOX
            c.font = Font(name=FONT, size=10)
            c.alignment = Alignment(vertical="top", wrap_text=True)


def dv_list(ws, ref, formula, strict=False):
    """strict=False -> the list is a SUGGESTION; any other value is still accepted."""
    dv = DataValidation(type="list", formula1=formula, allow_blank=True,
                        showErrorMessage=strict)
    ws.add_data_validation(dv)
    dv.add(ref)


# Suggestions, not a gate: real vendor specs say varchar / datetime2 / Alpha Numeric.
DATATYPES = '"String,Int,BigInt,Decimal,Date,Timestamp,Boolean,varchar,char,numeric,datetime"'
YN, YNU = '"Y,N"', '"Y,N,Unknown"'
FORMATS = '"Delimited text,Fixed width,Excel,JSON,XML,Parquet"'
LAST = 500

wb = Workbook()

# ================================================================== README
ws = wb.active
ws.title = "README"
widths(ws, {"A": 30, "B": 96, "C": 44})
ws.sheet_view.showGridLines = False

ws["A1"] = "Vendor Data Dictionary — unified template"
ws["A1"].font = Font(name=FONT, size=15, bold=True, color=CORE_BG)
ws.merge_cells("A1:B1")
ws.row_dimensions[1].height = 26

r = 3
ws.cell(row=r, column=1, value="ABOUT THIS FILE").font = Font(name=FONT, size=11, bold=True, color=CORE_BG)
r += 1
for k, v in [
    ("Purpose",
     "Describes the CONTENTS of the data files you send us: every file, and every field in every "
     "file. It is the source-side input to our automated Source-to-Target Mapping. Without it we "
     "cannot map your data without guessing, and we will not guess."),
    ("Who fills it in",
     "The data provider (the party that produces the files). You know what your own fields mean; "
     "we do not."),
    ("What it is NOT",
     "It is not a mapping document and it says nothing about our warehouse. No target table names, "
     "no target column names, no load rules — those are ours to decide, from this."),
    ("How to return it",
     "One workbook per delivery. Name the file VDD_<delivery name>.xlsx (DICT_ is also accepted, for anyone already issued the earlier template). Return the .xlsx "
     "itself, "
     "not a PDF or a screenshot."),
]:
    ws.cell(row=r, column=1, value=k).font = Font(name=FONT, size=10, bold=True)
    c = ws.cell(row=r, column=2, value=v)
    c.font = Font(name=FONT, size=10)
    c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.row_dimensions[r].height = 30
    r += 1

r += 1
ws.cell(row=r, column=1, value="DELIVERY IDENTIFICATION").font = Font(name=FONT, size=11, bold=True, color=CORE_BG)
ws.cell(row=r, column=2, value="Yellow cells: please complete.").font = Font(name=FONT, size=9, italic=True, color=EX_GREY)
r += 1
for label, hint in [
    ("Vendor / data provider", "the organisation that produces the files"),
    ("Delivery name",          "e.g. Coordination of Benefits — weekly"),
    ("Dictionary version",     "e.g. 1.0"),
    ("Effective date",         "first delivery this version describes (YYYY-MM-DD)"),
    ("Prepared by",            "name and role"),
    ("Contact for questions",  "email"),
]:
    ws.cell(row=r, column=1, value=label).font = Font(name=FONT, size=10, bold=True)
    c = ws.cell(row=r, column=2)
    c.fill = PatternFill("solid", fgColor=FILL_ME)
    c.border = BOX
    c.font = Font(name=FONT, size=10)
    ws.cell(row=r, column=3, value=hint).font = Font(name=FONT, size=9, italic=True, color=EX_GREY)
    r += 1

r += 1
ws.cell(row=r, column=1, value="DIVISION OF LABOUR").font = Font(name=FONT, size=11, bold=True, color=CORE_BG)
ws.cell(row=r, column=2, value="Two documents describe this data load. Please do not restate the other one.").font = Font(name=FONT, size=9, italic=True, color=EX_GREY)
r += 1
ws.cell(row=r, column=1, value="Our FRD supplies").font = Font(name=FONT, size=10, bold=True)
c = ws.cell(row=r, column=2, value=(
    "Object/data format · Target schema · Target table name · Domain and subdomain · "
    "Load strategy (Stage / Standard / Consumption) · Archive schedule · ADLS location · "
    "Inbound file folder path. All of this is ours to decide — none of it belongs in this "
    "workbook, and we will not ask you for it."))
c.font = Font(name=FONT, size=10); c.alignment = Alignment(vertical="top", wrap_text=True)
c.fill = PatternFill("solid", fgColor=RULE_BG); c.border = BOX
ws.row_dimensions[r].height = 46
r += 1
ws.cell(row=r, column=1, value="This document supplies").font = Font(name=FONT, size=10, bold=True)
c = ws.cell(row=r, column=2, value=(
    "Everything about the FILE ITSELF and the FIELDS INSIDE IT: names, positions, types, "
    "lengths, null rules, allowed values, descriptions, PHI/PII. Only you can supply these — "
    "they are not in our FRD and they cannot be read out of the data."))
c.font = Font(name=FONT, size=10); c.alignment = Alignment(vertical="top", wrap_text=True)
c.fill = PatternFill("solid", fgColor=RULE_BG); c.border = BOX
ws.row_dimensions[r].height = 40
r += 1
ws.cell(row=r, column=1, value="Stated in both").font = Font(name=FONT, size=10, bold=True)
c = ws.cell(row=r, column=2, value=(
    "File name pattern, format/delimiter and delivery cadence appear in both documents on "
    "purpose. We reconcile them and query any disagreement — that check is the point, so "
    "please fill them in even though we hold a copy."))
c.font = Font(name=FONT, size=10); c.alignment = Alignment(vertical="top", wrap_text=True)
c.fill = PatternFill("solid", fgColor=RULE_BG); c.border = BOX
ws.row_dimensions[r].height = 40
r += 1
ws.cell(row=r, column=1, value="Where this is filed").font = Font(name=FONT, size=10, bold=True)
c = ws.cell(row=r, column=2, value=(
    "The completed workbook is referenced from the FRD's Structural Metadata row "
    "'Source Data Dictionary'. That row is the link between the two documents."))
c.font = Font(name=FONT, size=10); c.alignment = Alignment(vertical="top", wrap_text=True)
ws.row_dimensions[r].height = 30
r += 2

ws.cell(row=r, column=1, value="HOW TO FILL IT IN").font = Font(name=FONT, size=11, bold=True, color=CORE_BG)
r += 1
for s in [
    "1. Complete the yellow cells above.",
    "2. FILES sheet — one row per physical file you deliver. Put a short key in 'Field Sheet' "
    "(e.g. cob_weekly).",
    "3. For each row on FILES, copy FIELDS_TEMPLATE, rename the copy to exactly that 'Field Sheet' "
    "key, and list every field in that file.",
    "4. If a file carries SEVERAL RECORD TYPES (a header record, detail records, a trailer), list "
    "them all on the one sheet: fill in 'Segment', and restart 'Position' at 1 for each record "
    "type. See FIELDS_EXAMPLE_MULTI.",
    "5. Delete the grey example rows once you have replaced them.",
    "6. CODE_SETS — define a reused code list once and reference it as @<Code Set ID> in "
    "'Allowed Values / Range'.",
    "7. CHANGE_LOG — record what changed whenever you re-issue this file.",
]:
    c = ws.cell(row=r, column=2, value=s)
    c.font = Font(name=FONT, size=10)
    c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.row_dimensions[r].height = 28
    r += 1

r += 1
ws.cell(row=r, column=1, value="THE RULES").font = Font(name=FONT, size=11, bold=True, color="9B2C2C")
r += 1
for rule in [
    "LEAVE IT BLANK RATHER THAN GUESS. A blank cell becomes a question we ask you. A wrong cell "
    "becomes a wrong column in production that nobody catches.",
    "DESCRIPTIONS ARE THE POINT. We can infer a type from your file. We can never infer what a "
    "field MEANS. If you fill in one optional thing, make it Description.",
    "FLAG PHI/PII GENEROUSLY. If you are unsure whether a field is identifying, mark it Y. "
    "Over-flagging costs us a masking rule; under-flagging is a breach.",
    "DATA TYPE IS LOAD-BEARING. Our engineering standard builds the warehouse column from "
    "the SOURCE column's data type — not from a sample of the values. A wrong type here "
    "becomes a wrong column in our warehouse.",
]:
    c = ws.cell(row=r, column=2, value=rule)
    c.font = Font(name=FONT, size=10)
    c.fill = PatternFill("solid", fgColor=RULE_BG)
    c.alignment = Alignment(vertical="top", wrap_text=True)
    c.border = BOX
    ws.row_dimensions[r].height = 32
    r += 1

r += 1
ws.cell(row=r, column=1, value="COLUMN LEGEND").font = Font(name=FONT, size=11, bold=True, color=CORE_BG)
r += 1
ws.cell(row=r, column=1, value="Dark blue header").font = Font(name=FONT, size=10, bold=True, color=CORE_BG)
ws.cell(row=r, column=2, value="REQUIRED — we need this for every field.").font = Font(name=FONT, size=10)
r += 1
ws.cell(row=r, column=1, value="Light blue header").font = Font(name=FONT, size=10, bold=True, color=OPT_BG)
ws.cell(row=r, column=2, value="OPTIONAL — but 'Segment' becomes REQUIRED for a multi-record-type file, "
        "and Start/End Position become REQUIRED for a fixed-width file.").font = Font(name=FONT, size=10)
r += 2

ws.cell(row=r, column=1, value="FIELD DEFINITIONS").font = Font(name=FONT, size=11, bold=True, color=CORE_BG)
r += 1
for k, v in [
    ("Position",
     "Ordinal of the field WITHIN ITS RECORD TYPE, in delivery order. For a file with one record "
     "type this is just its position in the row. For a multi-record-type file, restart at 1 for "
     "each record type. Required even when the file has a column-header row — it is how we detect "
     "a layout change."),
    ("Field Name",
     "The name as it appears in the column-header row. For a file WITH NO HEADER ROW, the name "
     "from your published file spec. Never leave this blank — we will not invent one."),
    ("Data Type",
     "Your own type vocabulary is fine (varchar, datetime2, Alpha Numeric, int). The dropdown is a "
     "suggestion, not a restriction — type over it."),
    ("Length",
     "Maximum character length. Text is acceptable where that is the truth ('no max length'). For "
     "Decimal, also give Precision and Scale."),
    ("Required (Y/N)",
     "Y = always populated. N = may be empty. This drives our reject rules, so answer for the "
     "CONTRACT, not for whatever last week's file happened to contain."),
    ("Description",
     "What the field means, in a sentence. Spell out abbreviations. Say the unit."),
    ("Allowed Values / Range",
     "Full code list ('M=Male, F=Female, U=Unknown'), a numeric range ('1-5, 1=lowest'), or "
     "@<Code Set ID>."),
    ("Example Value",
     "One realistic value. Use a SYNTHETIC value for anything PHI/PII — never real member data."),
    ("PHI/PII (Y/N)",
     "Y if the field identifies a person or is health information about one."),
    ("Segment",
     "REQUIRED for multi-record-type files: which record this field belongs to. Use Header / "
     "Detail / Trailer, or your own labels — just be consistent. Leave blank for a single-record "
     "file."),
    ("Business Name",
     "A human-readable name where the technical one is cryptic."),
    ("Precision / Scale",
     "Decimal only: total digits, and digits after the point. Decimal(10,2) is Precision 10, "
     "Scale 2."),
    ("Format / Pattern",
     "Required for dates and shaped codes: CCYYMMDD, MM/DD/YYYY, HH24:MI:SS, or a regex."),
    ("Default Value",
     "What appears when the field is not populated: a literal, or 'empty string', or 'null'."),
    ("Start / End Position",
     "REQUIRED for fixed-width files: 1-based character offsets, inclusive. Leave blank for "
     "delimited files."),
    ("Notes",
     "Anything that does not fit above — a caveat, a known data-quality issue, a deprecation."),
]:
    ws.cell(row=r, column=1, value=k).font = Font(name=FONT, size=10, bold=True)
    c = ws.cell(row=r, column=2, value=v)
    c.font = Font(name=FONT, size=10)
    c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.row_dimensions[r].height = 30
    r += 1


# =================================================================== FILES
ws = wb.create_sheet("FILES")
FILES_HEADERS = [
    # --- core 9: identical to the existing DICT_ workbooks ---
    "File Name Pattern", "File Title", "Format", "Delimiter", "Header Row",
    "Encoding", "Delivery Cadence", "Content Description", "Field Sheet",
    # --- optional ---
    "Null Representation",
    "File Generator", "Text Qualifier", "Line Ending", "Multi-Record-Type",
    "Record Type Field", "Record Type Values", "Expected Field Count",
    "Approx Rows per Delivery", "Trailer / Control Record", "Notes",
]
h(ws, 1, FILES_HEADERS, n_core=9, comments={
    "Header Row":
        "Does row 1 of the file contain COLUMN NAMES?\n\n"
        "This is NOT about a header RECORD. A file can have a 'Header' record type "
        "(a control record) and still have no column-name row -- in that case answer N "
        "here and Y under 'Multi-Record-Type'.",
    "Null Representation":
        "How a MISSING value appears in the file: an empty field, the literal text NULL, "
        "\\N, spaces, 0000-00-00 ...\n\n"
        "Needed because the FRD's data-quality rules reject on NULL ('reject the record "
        "when MEMBER_ID is NULL'). In a delimited file 'null' is a convention, not a fact, "
        "and guessing it wrong either rejects good records or loads bad ones.",
    "Record Type Field":
        "Multi-record-type files: which field distinguishes the record types "
        "(e.g. 'position 1').",
    "Record Type Values":
        "The literal values that mark each record type, e.g. HDR=..., DTL=..., TRL='TRAILR'.",
    "Expected Field Count":
        "Per record type where the file has several, e.g. Header=8, Detail=101, Trailer=6.",
})
widths(ws, {"A": 34, "B": 24, "C": 16, "D": 11, "E": 11, "F": 11, "G": 20,
            "H": 38, "I": 20, "J": 22, "K": 16, "L": 13, "M": 12, "N": 17,
            "O": 18, "P": 30, "Q": 26, "R": 20, "S": 20, "T": 30})
FILES_EXAMPLES = [
    {
        "File Name Pattern": "demographics_package_CCYY_MM.csv",
        "File Title": "Community Demographics",
        "Format": "Delimited text", "Delimiter": ",", "Header Row": "Y",
        "Encoding": "UTF-8", "Delivery Cadence": "Twice yearly",
        "Content Description": "One row per ZIP code: population and community characteristics.",
        "Field Sheet": "demographics_package",
        "Null Representation": "empty field",
        "File Generator": "Example Vendor Inc.", "Text Qualifier": '"', "Line Ending": "LF",
        "Multi-Record-Type": "N",
        "Expected Field Count": 86, "Approx Rows per Delivery": "41000",
        "Trailer / Control Record": "N",
        "Notes": "Example: single record type, has column headers — delete once replaced.",
    },
    {
        "File Name Pattern": "CCYYMMDD_<payer>_COBReport.txt",
        "File Title": "Coordination of Benefits",
        "Format": "Delimited text", "Delimiter": "|", "Header Row": "N",
        "Encoding": "UTF-8", "Delivery Cadence": "Weekly",
        "Content Description": "Header record, one detail record per COB result, trailer with counts.",
        "Field Sheet": "cob_weekly",
        "Null Representation": "spaces",
        "File Generator": "Example Vendor Inc.", "Text Qualifier": "none", "Line Ending": "CRLF",
        "Multi-Record-Type": "Y", "Record Type Field": "Position 1",
        "Record Type Values": "Header = 'FILEHD', Detail = 'COBDTL', Trailer = 'TRAILR'",
        "Expected Field Count": "Header=8, Detail=101, Trailer=6",
        "Approx Rows per Delivery": "120000",
        "Trailer / Control Record": "Y",
        "Notes": "Example: NO column-header row; field names come from this dictionary only "
                 "— delete once replaced.",
    },
]
for i, ex in enumerate(FILES_EXAMPLES):
    body(ws, 2 + i, row_from(FILES_HEADERS, ex), italic=True, color=EX_GREY)
blanks(ws, 4, 40, len(FILES_HEADERS))
dv_list(ws, "C2:C40", FORMATS)
dv_list(ws, "E2:E40", YN)
dv_list(ws, "N2:N40", YN)   # Multi-Record-Type (shifted by Null Representation)
dv_list(ws, "S2:S40", YN)   # Trailer / Control Record
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(FILES_HEADERS))}40"

# ============================================================ field sheets
FIELD_HEADERS = [
    # --- core 9: identical to the existing DICT_ workbooks ---
    "Position", "Field Name", "Data Type", "Length", "Required (Y/N)",
    "Description", "Allowed Values / Range", "Example Value", "PHI/PII (Y/N)",
    # --- optional ---
    "Key / Uniqueness",
    "Segment", "Business Name", "Precision", "Scale", "Format / Pattern",
    "Default Value", "Start Position", "End Position", "Notes",
]
FIELD_COMMENTS = {
    "Position": "Ordinal WITHIN ITS RECORD TYPE. Restart at 1 for each record type "
                "on a multi-record-type file.",
    "Key / Uniqueness":
        "Does this field identify a record? PK for the primary key, PK2/PK3 for the "
        "second and third parts of a composite key, U for merely unique, blank otherwise."
        "\n\nNeeded whenever the feed loads with Update Else Insert / Upsert: the match "
        "key decides which existing row a delivery updates. The FRD's own Business Key / "
        "Primary Key rows read 'NA' on the real documents, so the vendor is the only one "
        "who knows.",
    "Segment":  "REQUIRED for multi-record-type files (Header / Detail / Trailer). "
                "Leave blank for a single-record file.",
    "Data Type": "Your own vocabulary is fine — varchar, datetime2, Alpha Numeric, int. "
                 "The dropdown is a suggestion; type over it.",
    "Length": "Text is acceptable where that is the truth, e.g. 'no max length'.",
    "Start Position": "Fixed-width files only: 1-based, inclusive.",
    "End Position": "Fixed-width files only: 1-based, inclusive.",
}
FIELD_WIDTHS = {"A": 9, "B": 30, "C": 13, "D": 13, "E": 13, "F": 50, "G": 30,
                "H": 16, "I": 12, "J": 15, "K": 11, "L": 22, "M": 10, "N": 8,
                "O": 16, "P": 14, "Q": 13, "R": 12, "S": 28}


def field_sheet(title, rows):
    s = wb.create_sheet(title)
    h(s, 1, FIELD_HEADERS, n_core=9, comments=FIELD_COMMENTS)
    widths(s, FIELD_WIDTHS)
    for i, ex in enumerate(rows):
        body(s, 2 + i, row_from(FIELD_HEADERS, ex), italic=True, color=EX_GREY)
    blanks(s, 2 + len(rows), LAST, len(FIELD_HEADERS))
    dv_list(s, f"C2:C{LAST}", DATATYPES)
    dv_list(s, f"E2:E{LAST}", YN)
    dv_list(s, f"I2:I{LAST}", YNU)
    s.freeze_panes = "C2"
    s.auto_filter.ref = f"A1:{get_column_letter(len(FIELD_HEADERS))}{LAST}"
    return s


# ------ FIELDS_TEMPLATE: single record type, the common case
field_sheet("FIELDS_TEMPLATE", [
    {"Position": 1, "Field Name": "member_id", "Data Type": "String", "Length": 20,
     "Required (Y/N)": "Y",
     "Description": "Unique identifier for the member, as supplied by the health plan.",
     "Allowed Values / Range": "Free text", "Example Value": "M000123456",
     "PHI/PII (Y/N)": "Y", "Key / Uniqueness": "PK", "Business Name": "Member ID",
     "Notes": "Example row — delete once replaced."},
    {"Position": 2, "Field Name": "svc_from_dt", "Data Type": "Date",
     "Required (Y/N)": "Y",
     "Description": "First date of service on the claim line.",
     "Allowed Values / Range": "1900-01-01 to current date", "Example Value": "20240115",
     "PHI/PII (Y/N)": "Y", "Key / Uniqueness": "PK2", "Business Name": "Service From Date",
     "Format / Pattern": "CCYYMMDD"},
    {"Position": 3, "Field Name": "financial_strain_score", "Data Type": "Int",
     "Required (Y/N)": "N",
     "Description": "Individual financial strain risk score. Higher means greater strain.",
     "Allowed Values / Range": "1-5 (1 = little or none, 5 = severe)", "Example Value": "3",
     "PHI/PII (Y/N)": "N", "Business Name": "Financial Strain Score",
     "Default Value": "empty string"},
    {"Position": 4, "Field Name": "paid_amt", "Data Type": "Decimal",
     "Required (Y/N)": "N",
     "Description": "Amount paid by the plan for this claim line, in US dollars.",
     "Allowed Values / Range": "0.00 to 9999999.99", "Example Value": "142.50",
     "PHI/PII (Y/N)": "N", "Business Name": "Paid Amount",
     "Precision": 10, "Scale": 2, "Default Value": "0.00"},
    {"Position": 5, "Field Name": "gender_cd", "Data Type": "String", "Length": 1,
     "Required (Y/N)": "N",
     "Description": "Member gender as reported at enrolment.",
     "Allowed Values / Range": "@GENDER", "Example Value": "F",
     "PHI/PII (Y/N)": "Y", "Business Name": "Gender Code", "Default Value": "U"},
])

# ------ FIELDS_EXAMPLE_MULTI: headerless, pipe-delimited, Header/Detail/Trailer
field_sheet("FIELDS_EXAMPLE_MULTI", [
    {"Position": 1, "Field Name": "File Format Version", "Data Type": "varchar", "Length": 4,
     "Required (Y/N)": "Y",
     "Description": "Version of the file specification this delivery conforms to.",
     "Allowed Values / Range": "Free text", "Example Value": "0210",
     "PHI/PII (Y/N)": "N", "Segment": "Header",
     "Notes": "Whole sheet is an example — delete once replaced."},
    {"Position": 2, "Field Name": "Payer ID", "Data Type": "varchar", "Length": 4,
     "Required (Y/N)": "Y",
     "Description": "Identifier assigned to the payer by the reporting entity.",
     "Allowed Values / Range": "Free text", "Example Value": "1234",
     "PHI/PII (Y/N)": "N", "Segment": "Header"},
    {"Position": 3, "Field Name": "Date of Extract", "Data Type": "datetime2",
     "Required (Y/N)": "Y",
     "Description": "Date on which the file was produced.",
     "Allowed Values / Range": "Valid calendar date", "Example Value": "20240115",
     "PHI/PII (Y/N)": "N", "Segment": "Header", "Format / Pattern": "CCYYMMDD",
     "Notes": "Position restarts at 1 below, because Detail is a different record type."},
    {"Position": 1, "Field Name": "Member ID", "Data Type": "varchar", "Length": 80,
     "Required (Y/N)": "Y",
     "Description": "Plan-assigned member identifier.",
     "Allowed Values / Range": "Free text", "Example Value": "M000123456",
     "PHI/PII (Y/N)": "Y", "Key / Uniqueness": "PK", "Segment": "Detail"},
    {"Position": 2, "Field Name": "Relationship", "Data Type": "varchar", "Length": 2,
     "Required (Y/N)": "N",
     "Description": "Relationship of the member to the subscriber.",
     "Allowed Values / Range": "@RELATIONSHIP", "Example Value": "01",
     "PHI/PII (Y/N)": "N", "Segment": "Detail"},
    {"Position": 3, "Field Name": "Termination Date", "Data Type": "datetime2",
     "Required (Y/N)": "N",
     "Description": "Date the other coverage terminated. Empty where coverage is active.",
     "Allowed Values / Range": "Valid calendar date", "Example Value": "20241231",
     "PHI/PII (Y/N)": "N", "Segment": "Detail", "Format / Pattern": "CCYYMMDD",
     "Default Value": "empty string"},
    {"Position": 1, "Field Name": "Record Type", "Data Type": "Alpha Numeric", "Length": 6,
     "Required (Y/N)": "Y",
     "Description": "Static text identifying the trailer record.",
     "Allowed Values / Range": "Always 'TRAILR'", "Example Value": "TRAILR",
     "PHI/PII (Y/N)": "N", "Segment": "Trailer"},
    {"Position": 2, "Field Name": "Record Count", "Data Type": "numeric",
     "Length": "no max length", "Required (Y/N)": "Y",
     "Description": "Number of Detail records in this file. Used to validate completeness.",
     "Allowed Values / Range": "0 or greater", "Example Value": "118432",
     "PHI/PII (Y/N)": "N", "Segment": "Trailer",
     "Notes": "Control total — we reconcile against this on load."},
])

# =============================================================== CODE_SETS
ws = wb.create_sheet("CODE_SETS")
h(ws, 1, ["Code Set ID", "Value", "Meaning", "Notes"], n_core=3)
widths(ws, {"A": 20, "B": 16, "C": 46, "D": 40, "F": 60})
for i, (cs, v, m, n) in enumerate([
    ("GENDER", "F", "Female", "Example rows — delete once replaced."),
    ("GENDER", "M", "Male", None),
    ("GENDER", "U", "Unknown / not reported", None),
    ("RELATIONSHIP", "01", "Spouse", None),
    ("RELATIONSHIP", "19", "Child", None),
]):
    body(ws, 2 + i, [cs, v, m, n], italic=True, color=EX_GREY)
blanks(ws, 7, 200, 4)
ws.freeze_panes = "A2"
ws.auto_filter.ref = "A1:D200"
ws["F2"] = "Reference a code set from a field's 'Allowed Values / Range' cell as @GENDER."
ws["F2"].font = Font(name=FONT, size=10, italic=True, color=EX_GREY)
ws["F2"].alignment = Alignment(wrap_text=True, vertical="top")

# =============================================================== CHANGE_LOG
ws = wb.create_sheet("CHANGE_LOG")
h(ws, 1, ["Version", "Date", "Changed By", "File / Field", "Change", "Reason"], n_core=5)
widths(ws, {"A": 11, "B": 13, "C": 20, "D": 30, "E": 52, "F": 40})
body(ws, 2, ["1.0", "2026-01-15", "A. Vendor", "(all)", "Initial issue.",
             "Example row — delete once replaced."], italic=True, color=EX_GREY)
blanks(ws, 3, 120, 6)
ws.freeze_panes = "A2"

OUT = sys.argv[1] if len(sys.argv) > 1 else "templates/VDD_TEMPLATE.xlsx"
wb.save(OUT)

# --------------------------------------------------------------- self-check
# The blank template must parse as NO files: every FILES example row carries
# the marker in Notes, so the parser drops and counts them all. If a column
# is ever added without threading it through row_from (or the marker moves
# out of Notes), the examples survive as real files and this refuses to
# ship the workbook. Cheaper than finding out from a vendor.
sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))
from frdsttm.dictionary import TEMPLATE_EXAMPLE_MARKER, parse_dictionary_workbook  # noqa: E402

_expected = sum(TEMPLATE_EXAMPLE_MARKER in (ex.get("Notes") or "").lower()
                for ex in FILES_EXAMPLES)
assert _expected == len(FILES_EXAMPLES), "every FILES example row must carry the marker"
_parsed = parse_dictionary_workbook(OUT)
_dropped = [p for p in _parsed["problems"]
            if p["kind"] == "template_example_rows" and p["file"] is None]
assert _parsed["n_files"] == 0, (
    f"{_parsed['n_files']} example row(s) parsed as real files — a column shifted?")
assert _dropped and _dropped[0]["detail"].startswith(f"{_expected} FILES row"), (
    f"expected {_expected} dropped FILES example rows, got: {_dropped}")
print("wrote", OUT, f"(self-check: {_expected} example rows dropped, 0 files)")
