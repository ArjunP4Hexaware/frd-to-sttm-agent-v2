# Databricks notebook source
# MAGIC %md
# MAGIC # 03 — STTM Render: dictionary cross-check → derive → render → golden-pair eval
# MAGIC
# MAGIC Phase 5 of the FRD→STTM pipeline. Consumes the Phase 4 contracts and the
# MAGIC reference STTM workbooks, which play **two deliberately separated roles**:
# MAGIC
# MAGIC 1. Their **source-layout sections** stand in for the vendor source
# MAGIC    dictionary (the production input that will arrive as its own artifact).
# MAGIC 2. Their **stage/standard sections are ground truth for the eval only** —
# MAGIC    never an input to generation. Target mappings are DERIVED from
# MAGIC    contract + dictionary + explicit rules, then diffed against the
# MAGIC    reference to show exactly which decisions the pipeline can and cannot
# MAGIC    make on its own.
# MAGIC
# MAGIC Steps: **(a)** parse source dictionaries (both workbook dialects,
# MAGIC header-name-driven — the sheets disagree on column names/order);
# MAGIC **(b)** apply any saved review-app resolutions
# MAGIC (`_provenance.human_resolutions`) first — a human resolution is
# MAGIC authoritative and is never re-decided below — then resolve whatever's
# MAGIC left of the Phase 4 attribution ambiguities by dictionary column
# MAGIC membership (zip_code/member_id) and dictionary recycle markers;
# MAGIC **(c)** derive stage/standard mappings (1:1 names, String at stage,
# MAGIC segment-suffix table routing HDR/DTL/TRL); **(d)** render the STTM
# MAGIC workbook INTO the chosen template's own layout — its sheets, band
# MAGIC labels, headers, widths and styles, whatever dialect it is
# MAGIC (`layout_of` + `render_into_template`, 2026-08-22; a template column
# MAGIC the contract cannot fill stays blank and is reported in
# MAGIC `_provenance.template_fill`; the two built-in dialects remain only as
# MAGIC the freeform fallback when no template matched),
# MAGIC placing every FRD-stated validation/recycle rule on the row whose
# MAGIC column it names (Comment / Recycle Flag / Business Rule) or, when it
# MAGIC names none, in a feed-level cell — recorded in
# MAGIC `_provenance.rule_placement`, never dropped (2026-08-22);
# MAGIC **(e)** cell-level eval against the reference workbook.

# COMMAND ----------

# MAGIC %pip install openpyxl
# MAGIC dbutils.library.restartPython()

# COMMAND ----------

# Local-mode support: see notebooks/01_frd_ingest.py's parameter cell for the
# full explanation of IS_DATABRICKS / _param(). Databricks execution below is
# unchanged from before this cell existed.
import os
from pathlib import Path

IS_DATABRICKS = "dbutils" in globals()
LOCAL_ROOT = Path(__file__).resolve().parent.parent / "local_dev_fixtures" if not IS_DATABRICKS else None


def _param(name: str, default: str) -> str:
    if IS_DATABRICKS:
        dbutils.widgets.text(name, default)
        return dbutils.widgets.get(name)
    return os.environ.get(name.upper(), default)


CATALOG = _param("catalog", "arjun_workspace")
SCHEMA = _param("schema", "sttm_agent")
CONTRACTS_TABLE_NAME = _param("contracts_table", "frd_contracts")
RUNS_TABLE_NAME = _param("runs_table", "frd_sttm_runs")
REFERENCE_VOLUME = _param("reference_volume", "sttm_reference")
OUT_VOLUME = _param("out_volume", "sttm_out")
# Provenance (docs/AI_GOVERNANCE.md): who asked for this render, the app's
# run label, and the Databricks job run id — job parameters in the two job
# ymls, env vars TRIGGERED_BY / RUN_LABEL / JOB_RUN_ID from the app's
# subprocess env. "manual" / "" on a hand run: a recorded fact, not a gap.
TRIGGERED_BY = _param("triggered_by", "manual")
RUN_LABEL = _param("run_label", "")
JOB_RUN_ID = _param("job_run_id", "")

if IS_DATABRICKS:
    CONTRACTS_TABLE = f"{CATALOG}.{SCHEMA}.{CONTRACTS_TABLE_NAME}"
    RUNS_TABLE = f"{CATALOG}.{SCHEMA}.{RUNS_TABLE_NAME}"
    REFERENCE_DIR = f"/Volumes/{CATALOG}/{SCHEMA}/{REFERENCE_VOLUME}"
    OUT_ROOT = f"/Volumes/{CATALOG}/{SCHEMA}/{OUT_VOLUME}"
else:
    CONTRACTS_TABLE = CONTRACTS_TABLE_NAME
    RUNS_TABLE = RUNS_TABLE_NAME
    REFERENCE_DIR = str(LOCAL_ROOT / REFERENCE_VOLUME)
    OUT_ROOT = str(LOCAL_ROOT / OUT_VOLUME)
RENDERED_DIR = f"{OUT_ROOT}/rendered"
CONTRACTS_DIR = f"{OUT_ROOT}/contracts"
REPORTS_DIR = f"{OUT_ROOT}/reports"
EXTRACTIONS_DIR = f"{OUT_ROOT}/extractions"   # 02's extraction_meta sidecars (read-only here)
print(f"contracts: {CONTRACTS_TABLE}\nreference: {REFERENCE_DIR}\nrendered:  {RENDERED_DIR}")

# COMMAND ----------

# MAGIC %md
# MAGIC ## Core (pure Python — unit-testable as-is)

# COMMAND ----------

import hashlib
import json
import re
import shutil
import tempfile
import unicodedata
from copy import copy
from datetime import datetime, timezone

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

GENERATOR = "frd-sttm-agent phase5 v0.1"


# COMMAND ----------

# MAGIC %run ./_reference_workbooks

# COMMAND ----------

# Reference-workbook parsing (both dialects), feed matching and the shared
# normalization helpers were factored VERBATIM into
# `src/frdsttm/reference_workbooks.py` (2026-08-22) so the similarity/corpus
# modules and the review app read dictionaries with exactly this stage's
# parser -- one parser, no drift. `%run` above is a Databricks-only magic --
# inert when this file executes as a plain script, so the names would never
# get defined locally without this explicit import.
if not IS_DATABRICKS:
    from _reference_workbooks import (  # noqa: F401
        _n,
        _nl,
        layout_of,
        loose_tokens,
        match_feeds,
        parse_reference_workbook,
    )

# Importable in both modes: the %run shim above (Databricks) and the local
# import (script mode) both put src/ on sys.path first.
from frdsttm import standards as _std  # noqa: E402
from frdsttm.corpus import (  # noqa: E402
    frd_features_from_index,
    load_corpus_index,
    own_reference_for,
    reference_features_from_index,
)
from frdsttm.similarity import (  # noqa: E402
    contract_features,
    decide_templates,
    frd_features,
    merge_dictionaries,
    merge_features,
    thresholds_from,
    workbook_features,
)

# --------------------------------------------------------------------------- #
# human resolution wiring — apply saved review-app resolutions
# (`_provenance.human_resolutions`, written by review_app_react/) ahead of
# the automatic paths below. Matching is by each record's `ambiguity_id`
# against the `id` on a GatedAmbiguity object stored directly in
# `_provenance.ambiguities` / `_provenance.grounding.advisory_flagged` (see
# notebooks/_models.py) -- 03_contract_build.py emits
# `id`/`kind`/`has_candidates`/`candidates`/`context` on the ambiguity itself
# now, so this file reads those fields directly instead of reverse-parsing
# display text (the regex-parsing this replaced was deleted from the review
# app's backend along with this change). Precedence: a human resolution,
# once structurally
# applicable, is authoritative and is never re-decided by the automatic
# dictionary cross-check below; when nothing has been reviewed for a given
# ambiguity, behavior is byte-for-byte the pre-existing automatic path (see
# resolve_attribution()'s `already_settled` param, empty by default).
#
# Branching is has_candidates-driven, uniform across all three kinds (not
# kind-name special-cased): a candidate-having ambiguity (attribution,
# disagreement) requires resolution_type == "candidate_pick" with a
# chosen_candidate still present in the ambiguity's current candidates to
# apply structurally; "none_of_these" is a deliberate, valid terminal state
# that leaves it gated for whatever automatic fallback exists (attribution
# only -- disagreement has none); "free_text" on a candidate-having
# ambiguity is malformed relative to this policy and is never applied,
# exactly like today's attribution-only refusal, now generalized to name
# whichever kind triggered it. A candidate-free ambiguity (advisory_
# grounding) has no candidate_pick to make, so "free_text" with a `rationale`
# IS the resolution, applied exactly as today's override_value was.
# --------------------------------------------------------------------------- #

# Verbatim port of 03_contract_build.py's norm() -- not imported, since that
# script runs pipeline driver code (table reads/writes) at module import
# time (see review_app_react/backend/ambiguity_parsing.py's docstring for
# the same constraint on the review app's own FastAPI backend, which can't
# import 03_contract_build.py as a library either). Regrouping ambiguous
# rules here must reproduce
# byte-identical `id`s to what 03_contract_build.py stored in
# `_provenance.ambiguities`, since that id is the human_resolutions lookup
# key -- if 03's norm() or _ambiguity_id() changes, this must change with it.
_ATTR_UNICODE_MAP = str.maketrans({
    "‘": "'", "’": "'", "“": '"', "”": '"',
    "‐": "-", "‑": "-", "‒": "-", "–": "-",
    "—": "-", "―": "-", " ": " ",
})


def _attr_norm(text: str) -> str:
    t = unicodedata.normalize("NFKC", text).translate(_ATTR_UNICODE_MAP)
    t = re.sub(r"[*_#|`]", " ", t)
    t = re.sub(r"\s+", " ", t)
    return t.strip().lower()


def _ambiguity_id(kind: str, text: str, context: dict) -> str:
    """Verbatim port of 03_contract_build.py's _ambiguity_id()."""
    key = json.dumps({"kind": kind, "text": text, "context": context}, sort_keys=True)
    return f"{kind}-{hashlib.sha1(key.encode('utf-8')).hexdigest()[:12]}"


def _attribution_groups(contract):
    """Reproduces attribution_check()'s grouping (03_contract_build.py) from
    the contract's CURRENT feed state, keyed by the same `id` 03 would
    produce for an unchanged contract. Must run before resolve_attribution()
    mutates validation_rules/recycle_rule, since it needs each feed's
    original (still-duplicated) rule set. If a human_resolutions record's
    ambiguity_id isn't a key in the returned dict, the underlying rule set
    has changed since this ambiguity was gated -- the same staleness
    condition the pre-id-based version of this function detected via text
    matching. Returns {id: {rule, feed_indices, feed_names}}."""
    seen = {}
    for i, feed in enumerate(contract["feeds"]):
        rules = list(feed.get("validation_rules") or [])
        if feed.get("recycle_rule"):
            rules = rules + [feed["recycle_rule"]]
        for rule in rules:
            entry = seen.setdefault(_attr_norm(rule), {"rule": rule, "feed_indices": [], "feed_names": []})
            entry["feed_indices"].append(i)
            entry["feed_names"].append(feed["feed_name"])
    groups = {}
    for entry in seen.values():
        if len(entry["feed_indices"]) > 1:
            text = (f"rule applied to {len(entry['feed_names'])} feeds ({', '.join(entry['feed_names'])}) — "
                    f"attribution unconfirmed pending source dictionary: {entry['rule'][:120]!r}")
            context = {"feed_names": entry["feed_names"], "feed_indices": entry["feed_indices"], "rule": entry["rule"]}
            aid = _ambiguity_id("attribution", text, context)
            groups[aid] = {"rule": entry["rule"], "feed_indices": entry["feed_indices"],
                           "feed_names": entry["feed_names"]}
    return groups


# Feed-scoped advisory fields we're confident map 1:1 onto a free-text
# override: plain Optional[str] fields (direct overwrite) vs List[str]
# fields (replace the one flagged element, found by exact value match --
# never the whole list, which would silently drop every other entry).
_ADVISORY_SCALAR_FIELDS = {"recycle_rule", "history_backfill", "archive_retention",
                          "phi_pii_notes", "landing_location", "sttm_reference"}
_ADVISORY_LIST_FIELDS = {"validation_rules", "load_windows_sla"}


def _apply_advisory_override(contract, context, rationale):
    """Returns (applied, target, reason_not_applied). `context` is the
    GatedAmbiguity.context 03_contract_build.py stored on this
    advisory_grounding ambiguity -- feed_index/field are only present when
    the flagged path is feed-scoped (see grounding_audit()'s advisory()
    closure); top-level paths (in_scope/out_of_scope/system_interfaces/
    open_items, acd(name), project.*) carry `path` only and are
    intentionally out of scope here, unchanged from today -- no real
    fixture exercises them and their shapes vary too much to generalize
    with confidence."""
    idx, field = context.get("feed_index"), context.get("field")
    if idx is None or field is None:
        return False, None, f"advisory path shape not recognized for structural write-back: {context.get('path')!r}"
    if idx >= len(contract["feeds"]):
        return False, None, f"feed index {idx} out of range for this contract"
    feed = contract["feeds"][idx]
    target = f"feeds[{idx}].{field}"
    if field in _ADVISORY_SCALAR_FIELDS:
        feed[field] = rationale
        return True, target, None
    if field in _ADVISORY_LIST_FIELDS:
        original_value = context.get("original_value")
        values = list(feed.get(field) or [])
        for j, v in enumerate(values):
            if v == original_value:
                values[j] = rationale
                feed[field] = values
                return True, f"{target}[{j}]", None
        return False, None, f"original flagged value no longer present in {target}"
    return False, None, f"advisory field {field!r} not in the known scalar/list write-back set"


def _apply_disagreement_override(contract, field, value):
    """The only disagreement source today (enrich() in 03_contract_build.py)
    is project.project_id. Only called for a structural candidate_pick now
    (see apply_human_resolutions() below) -- free text is no longer applied
    for this kind, since it has real candidates (agent value vs. content
    value) and this policy requires a pick from candidate-having ambiguities."""
    if field == "project.project_id":
        contract.setdefault("project", {})["project_id"] = value
        return True, "project.project_id", None
    return False, None, f"disagreement field {field!r} has no known target mapping"


def apply_human_resolutions(contract):
    """Applies `_provenance.human_resolutions` before the automatic paths
    below ever run. Branches on `ambiguity.has_candidates`, not on kind name
    (see module-level comment above for the full policy):

    - has_candidates + resolution_type == "candidate_pick" (chosen_candidate
      still among the ambiguity's current candidates): apply structurally
      per kind -- attribution keeps the rule only on the chosen feed and
      strips it from the others in its cross-feed group; disagreement writes
      the chosen value into the one scalar field named in context.field.
    - has_candidates + resolution_type == "none_of_these": NOT applied, by
      design -- a deliberate rejection of every listed candidate, left gated
      for whatever automatic fallback exists (attribution's dictionary
      cross-check; disagreement has none).
    - has_candidates + anything else (resolution_type == "free_text", or a
      chosen_candidate that no longer matches): NOT applied -- a structural
      pick was required and not provided. This generalizes what was
      previously an attribution-only refusal to every candidate-having kind
      (disagreement included, a real behavior change -- see summary).
    - not has_candidates (advisory_grounding today): resolution_type ==
      "free_text" with a `rationale` IS the resolution -- there's nothing
      else it could be. Applied via the kind's known write-back path.

    Staleness: if a resolution's `candidates_snapshot` no longer matches the
    ambiguity's current `candidates`, the entry is flagged `stale` in the
    audit *before* the branches above run -- a stale candidate_pick is still
    attempted if chosen_candidate remains valid, but the staleness is always
    visible in resolution_audit either way.

    Returns (settled_rules, audit) where settled_rules is a
    {(feed_index, normalized_rule_text)} set resolve_attribution() must treat
    as already-decided (never remove, regardless of dictionary evidence), and
    audit is the list of resolution_audit entries this call produced. Mutates
    `contract` and `_provenance.ambiguities` / `_provenance.grounding.
    advisory_flagged` / `_provenance.resolution_audit` in place.
    """
    prov = contract.setdefault("_provenance", {})
    audit = prov.setdefault("resolution_audit", [])
    human_resolutions = prov.get("human_resolutions") or []
    if not human_resolutions:
        return set(), audit

    ambiguities_by_id = {a["id"]: a for a in prov.get("ambiguities", [])}
    advisory_by_id = {a["id"]: a for a in prov.get("grounding", {}).get("advisory_flagged", [])}
    attribution_groups_by_id = _attribution_groups(contract)
    settled_rules = set()

    for hr in human_resolutions:
        aid = hr["ambiguity_id"]
        kind = hr.get("kind")
        resolution_type = hr.get("resolution_type")
        chosen = hr.get("chosen_candidate")
        rationale = hr.get("rationale")
        snapshot = hr.get("candidates_snapshot") or []

        ambiguity = ambiguities_by_id.get(aid) or advisory_by_id.get(aid)
        entry = {
            "ambiguity_id": aid, "kind": kind,
            "ambiguity_text": ambiguity["text"] if ambiguity else None,
            "resolution_source": "human", "resolution_type": resolution_type,
            "chosen_candidate": chosen, "rationale": rationale,
            "applied": False, "target": None, "reason_not_applied": None,
            "stale": False, "stale_detail": None,
        }

        if ambiguity is None:
            entry["reason_not_applied"] = (
                "no gated ambiguity with this id found in the current contract "
                "(may have already been cleared by a prior resolution, or the "
                "contract was regenerated since this resolution was recorded)")
            audit.append(entry)
            continue

        current_candidates = ambiguity["candidates"]
        if snapshot and set(snapshot) != set(current_candidates):
            entry["stale"] = True
            entry["stale_detail"] = (
                f"candidates changed since this resolution was recorded: "
                f"was {snapshot}, now {current_candidates}")

        if ambiguity["has_candidates"]:
            if resolution_type == "candidate_pick" and chosen in current_candidates:
                if kind == "attribution":
                    group = attribution_groups_by_id.get(aid)
                    if group is None:
                        entry["reason_not_applied"] = (
                            "no matching cross-feed rule group found in the current contract "
                            "(validation_rules/recycle_rule may have changed since this "
                            "ambiguity was gated)")
                    else:
                        rule_norm = _attr_norm(group["rule"])
                        for i in group["feed_indices"]:
                            feed = contract["feeds"][i]
                            if feed["feed_name"] == chosen:
                                settled_rules.add((i, rule_norm))
                                continue
                            feed["validation_rules"] = [
                                r for r in feed.get("validation_rules", []) if _attr_norm(r) != rule_norm]
                            if feed.get("recycle_rule") and _attr_norm(feed["recycle_rule"]) == rule_norm:
                                feed["recycle_rule"] = None
                        entry["applied"] = True
                        entry["target"] = f"feeds[*].validation_rules/recycle_rule (kept only on {chosen!r})"
                elif kind == "disagreement":
                    field = ambiguity["context"].get("field")
                    entry["applied"], entry["target"], entry["reason_not_applied"] = \
                        _apply_disagreement_override(contract, field, chosen)
                else:
                    entry["reason_not_applied"] = f"ambiguity kind {kind!r} has candidates but no known structural write-back path"
            elif resolution_type == "none_of_these":
                entry["reason_not_applied"] = (
                    "reviewer explicitly rejected every listed candidate -- left gated for "
                    + ("the automatic dictionary cross-check to decide feed membership"
                       if kind == "attribution" else "manual resolution; no automatic fallback for this kind"))
            else:
                bad_pick = f"chosen_candidate {chosen!r} not among {current_candidates}" if chosen is not None else "free text instead of a pick"
                entry["reason_not_applied"] = (
                    f"structural pick required and not provided ({bad_pick}) -- this ambiguity "
                    f"has known candidates, so a free-text rationale alone has no confident "
                    f"structural mapping; rationale recorded, "
                    + ("automatic dictionary cross-check still decides feed membership"
                       if kind == "attribution" else "still gated"))
        else:
            if resolution_type == "free_text" and rationale:
                if kind == "advisory_grounding":
                    entry["applied"], entry["target"], entry["reason_not_applied"] = \
                        _apply_advisory_override(contract, ambiguity["context"], rationale)
                else:
                    entry["reason_not_applied"] = f"ambiguity kind {kind!r} has no candidates but no known free-text write-back path"
            else:
                entry["reason_not_applied"] = (
                    "no rationale present (unexpected -- this ambiguity has no candidates, "
                    "free text is the only possible resolution)")

        audit.append(entry)

    # A human explicitly reviewed every ambiguity/advisory flag with a saved
    # resolution, whether or not it was structurally applicable -- clear it
    # from the "needs review" lists the same way resolve_attribution() below
    # clears automatically-resolved ones. Which value ended up in the
    # contract (human vs. automatic-fallback) is exactly what
    # `resolution_audit` records.
    resolved_ids = {hr["ambiguity_id"] for hr in human_resolutions}
    prov["ambiguities"] = [a for a in prov.get("ambiguities", []) if a["id"] not in resolved_ids]
    if prov.get("grounding", {}).get("advisory_flagged"):
        prov["grounding"]["advisory_flagged"] = [
            a for a in prov["grounding"]["advisory_flagged"] if a["id"] not in resolved_ids]
    if (contract.get("status") == "PASS_WITH_FLAGS" and not prov["ambiguities"]
            and not prov.get("grounding", {}).get("advisory_flagged")):
        contract["status"] = "PASS"

    return settled_rules, audit


# --------------------------------------------------------------------------- #
# attribution resolution
# --------------------------------------------------------------------------- #
_COL_TOKEN = re.compile(r"\b([A-Za-z][A-Za-z0-9_]{2,})\s+column\b", re.I)


def _quote_rule(rule: str, limit: int = 200) -> str:
    """Rule text for a resolution message: full up to `limit` chars, else cut
    at a word boundary with an ellipsis — never mid-word (the old hard
    `rule[:90]` slice produced quotes like \"...moving it to the reje\")."""
    rule = rule.strip()
    if len(rule) <= limit:
        return rule
    cut = rule[:limit].rsplit(" ", 1)[0].rstrip()
    return f"{cut}…"


def resolve_attribution(contract, dictionary, feed_match, already_settled=frozenset()):
    """Clears cross-feed rule ambiguities using dictionary column membership.
    Mutates the contract; returns list of resolution strings.

    `already_settled` is a {(feed_index, normalized_rule_text)} set of rules
    apply_human_resolutions() has already decided by human candidate-pick --
    those are kept unconditionally, never re-evaluated against the
    dictionary, so a human decision can never be silently overridden by the
    automatic cross-check. Empty by default, which reproduces the exact
    pre-existing behavior when nothing has been human-resolved."""
    cols_by_feed = {i: {_nl(f["source_column"]) for f in dictionary["feeds"][k]["fields"]}
                    for i, k in feed_match.items()}
    recycle_feeds = {i for i, k in feed_match.items()
                     if dictionary["feeds"][k]["recycle_note"]}
    all_cols = set().union(*cols_by_feed.values()) if cols_by_feed else set()
    resolutions = []

    for i, feed in enumerate(contract["feeds"]):
        if i not in cols_by_feed:
            continue
        kept = []
        for rule in feed.get("validation_rules", []):
            if (i, _attr_norm(rule)) in already_settled:
                kept.append(rule)
                continue
            cands = {_nl(c) for c in _COL_TOKEN.findall(rule)} & all_cols
            if cands and not (cands & cols_by_feed[i]):
                resolutions.append(
                    f"removed rule from {feed['feed_name']!r} — column(s) "
                    f"{sorted(cands)} not in its source dictionary: {_quote_rule(rule)!r}")
                continue
            if "recycle" in _nl(rule) and i not in recycle_feeds and recycle_feeds:
                resolutions.append(
                    f"removed recycle rule from {feed['feed_name']!r} — its dictionary "
                    f"carries no recycle marker: {_quote_rule(rule)!r}")
                continue
            kept.append(rule)
        feed["validation_rules"] = kept
        if feed.get("recycle_rule") and recycle_feeds and i not in recycle_feeds \
                and (i, _attr_norm(feed["recycle_rule"])) not in already_settled:
            resolutions.append(
                f"cleared recycle_rule on {feed['feed_name']!r} — dictionary recycle "
                f"marker present only on {[contract['feeds'][j]['feed_name'] for j in sorted(recycle_feeds)]}")
            feed["recycle_rule"] = None

    prov = contract.setdefault("_provenance", {})

    # D2 fix (docs/LIVE_E2E_2026-08-07.md): an attribution ambiguity whose
    # candidate feeds already EQUAL the dictionary-confirmed set is settled
    # by the same evidence a removal uses -- the extraction agrees with the
    # dictionary, there is just nothing left to remove. Before this pass, a
    # dictionary-correct extraction stayed gated (zero removals -> the
    # clearing below never ran) while a spread-across-all-feeds extraction
    # cleared to PASS. Confirm-and-clear those individually; ambiguities the
    # dictionary cannot adjudicate (no known column token, or a partial
    # overlap) keep their human gate.
    n_removals = len(resolutions)
    confirmed_ids = set()
    for amb in prov.get("ambiguities", []):
        if amb.get("kind") != "attribution":
            continue
        ctx = amb.get("context") or {}
        rule = ctx.get("rule") or amb.get("text", "")
        cand = set(ctx.get("feed_indices") or [])
        if not cand:
            continue
        cols = {_nl(c) for c in _COL_TOKEN.findall(rule)} & all_cols
        if cols:
            dict_set = {i for i, cs in cols_by_feed.items() if cols & cs}
        elif "recycle" in _nl(rule):
            dict_set = set(recycle_feeds)
        else:
            continue
        if dict_set and cand == dict_set:
            confirmed_ids.add(amb["id"])
            resolutions.append(
                f"confirmed rule attribution on "
                f"{[contract['feeds'][j]['feed_name'] for j in sorted(cand)]} "
                f"— candidates match dictionary column membership exactly: "
                f"{_quote_rule(rule)!r}")

    prov["attribution_resolutions"] = resolutions
    if n_removals:
        # Removals re-attribute every rule deterministically -- the
        # pre-existing behavior clears all attribution ambiguities.
        prov["ambiguities"] = [a for a in prov.get("ambiguities", [])
                               if a["kind"] != "attribution"]
    elif confirmed_ids:
        prov["ambiguities"] = [a for a in prov.get("ambiguities", [])
                               if a["id"] not in confirmed_ids]
    if resolutions:
        if contract.get("status") == "PASS_WITH_FLAGS" and not prov["ambiguities"] \
                and not prov.get("grounding", {}).get("advisory_flagged"):
            contract["status"] = "PASS"
    return resolutions


# --------------------------------------------------------------------------- #
# derive target mappings (generation — reference targets are never read here)
# --------------------------------------------------------------------------- #
_SEGMENT_SUFFIX = {"header": "_HDR", "detail": "_DTL", "trailer": "_TRL"}


def _table_for_segment(tables, segment):
    seg = _nl(segment)
    suffix = _SEGMENT_SUFFIX.get(seg)
    if suffix:
        for t in tables:
            if t.upper().endswith(suffix):
                return t
    return tables[0] if tables else ""


# --------------------------------------------------------------------------- #
# standards fill — the client's target-side rules, where the FRD is silent
# --------------------------------------------------------------------------- #
# Until 2026-08-26 a catalog/schema the FRD did not state came out null, and the
# renderer then leaned on whichever reference workbook matched to make the
# output look complete. That is the implicit borrow docs/THREE_INPUT_ARCHITECTURE
# .md §5 exists to remove: it only works when the FRD being rendered is already
# mapped, which makes the accuracy figure self-referential.
#
# Now the gap is filled from contracts/naming_standards.json instead -- the
# client's own vocabulary, versioned and hashed into this run's provenance.
# Three rules hold, and they are the whole point:
#   1. the FRD ALWAYS wins. A stated catalog/schema is never overwritten.
#   2. a fill is recorded in `_provenance.standards_fill` with the contract's
#      own confidence for that derivation (STATED / OBSERVED /
#      OBSERVED_SINGLE_PAIR), so a reviewer can tell a client rule from an
#      inference this repo made.
#   3. an underivable value stays NULL and raises a gated ambiguity naming the
#      feed and what was missing. `abbreviate("domains", "sdoh")` returning None
#      is a real answer -- the SD FRD's own domain is absent from the client's
#      table -- and the correct response is to ask, never to invent.


def _standards_fill_target(feed, layer, target, prov):
    """Fill catalog/schema for one layer from the standards contract.

    Mutates `target` only where the FRD stated nothing; appends one record per
    filled or unfillable cell to `prov`. Returns the list of attribute names
    that could not be resolved, for the caller to gate on.
    """
    unresolved = []
    domain = feed.get("domain")
    for attr, derive, status_key in (
        ("catalog", lambda: _std.catalog_for(layer), "catalog"),
        ("schema", lambda: _std.schema_for(layer, domain), "schema"),
    ):
        if target.get(attr):
            continue                      # rule 1: the FRD always wins
        value = derive()
        if value:
            target[attr] = value
            prov.append({
                "feed": feed.get("feed_name"), "layer": layer, "attribute": attr,
                "value": value, "source": "standards",
                "confidence": _std.derivation_status(status_key),
                "standards_version": _std.NAMING_VERSION,
            })
        else:
            unresolved.append(attr)
            prov.append({
                "feed": feed.get("feed_name"), "layer": layer, "attribute": attr,
                "value": None, "source": "unresolved",
                "confidence": _std.derivation_status(status_key),
                "standards_version": _std.NAMING_VERSION,
                "reason": (
                    f"domain {domain!r} is not in the client's naming standards"
                    if attr == "schema" and _std.abbreviate("domains", domain) is None
                    else f"the naming standards declare no {attr} for layer {layer!r}"
                ),
            })
    return unresolved


def apply_standards_targets(contract, feed_match):
    """Fill every matched feed's unstated catalog/schema from the standards.

    Runs BEFORE `derive_field_mappings`, so the per-field rows it writes carry
    the filled values rather than nulls. Gated ambiguities are appended to
    `_provenance.ambiguities` in the same shape the rest of the stage uses.
    """
    prov = []
    gated = []
    for i in feed_match:
        feed = contract["feeds"][i]
        for layer, key in (("stage", "stage_target"), ("standard", "standard_target")):
            target = feed.get(key)
            if target is None:
                continue
            unresolved = _standards_fill_target(feed, layer, target, prov)
            for attr in unresolved:
                text = (
                    f"{layer} {attr} for feed {feed.get('feed_name') or i!r} is stated "
                    f"neither in the FRD nor derivable from the client's naming standards "
                    f"(domain: {feed.get('domain')!r})"
                )
                gated.append({
                    "id": _ambiguity_id("standards_gap", text,
                                        {"feed": feed.get("feed_name"), "layer": layer,
                                         "attribute": attr}),
                    "kind": "standards_gap",
                    "text": text,
                    "context": {"feed": feed.get("feed_name"), "layer": layer,
                                "attribute": attr, "domain": feed.get("domain")},
                    "options": sorted(_std.known_terms("domains")) if attr == "schema" else [],
                })
    p = contract.setdefault("_provenance", {})
    p["standards_fill"] = prov
    p["standards"] = {
        "naming_version": _std.NAMING_VERSION,
        "engineering_version": _std.ENGINEERING_VERSION,
        "standards_sha256": _std.standards_sha256(),
        "column_rules_sourced": _std.column_rules_are_sourced(),
    }
    if gated:
        p.setdefault("ambiguities", []).extend(gated)
    return contract


def derive_field_mappings(contract, dictionary, feed_match):
    """Attach derived stage/standard mappings to each matched feed's fields.
    Rules (the pipeline's explicit, contestable defaults):
      - stage: 1:1 column name, datatype from the template's stage target for
        that row (else String), schema/catalog from contract -- and, since
        2026-08-26, from the client's naming standards where the FRD stated
        none (`apply_standards_targets`, which must run first) -- table chosen
        by record segment suffix (HDR/DTL/TRL) when present.
      - standard: same 1:1 name; schema/catalog from contract when stated,
        else the standards fill, else null and gated (the FRD may genuinely
        not state it, and the client's vocabulary may not cover its domain);
        datatype from the template's standard target (else the stage
        datatype).
      - AUDIT rows (source "NA", `audit: True` from the dictionary parser,
        2026-08-22): the template's ETL audit columns — a workbook
        convention, not an FRD fact, so there is no 1:1 source name to
        derive from. Their column name and datatype are taken from the
        template's own stage/standard targets (the one deliberate use of a
        reference's target side as input: with exclude-own on, that template
        is a DIFFERENT pair, so this copies the client convention, never the
        ground truth). Without this the rendered workbook carried audit rows
        whose stage ColumnName was "NA" — unrecoverable downstream
        (CodeGen's `extract-sttm` rejects them). Tables still come from the
        contract's own targets.
    """
    for i, key in feed_match.items():
        feed = contract["feeds"][i]
        stg = feed.get("stage_target") or {}
        std = feed.get("standard_target") or {}
        ref_feed = dictionary["feeds"][key]
        derived = []
        for f, rt in zip(ref_feed["fields"], ref_feed.get("ref_targets") or
                         [{"stage": {}, "standard": {}}] * len(ref_feed["fields"])):
            stage_tbl = _table_for_segment(stg.get("tables", []), f.get("segment"))
            std_tbl = (_table_for_segment(std.get("tables", []), f.get("segment"))
                       or stage_tbl)
            if f.get("audit"):
                stage_col = rt["stage"].get("column") or rt["standard"].get("column") or f["source_column"]
                std_col = rt["standard"].get("column") or stage_col
                stage_dt = rt["stage"].get("datatype") or rt["standard"].get("datatype") or "String"
                std_dt = rt["standard"].get("datatype") or stage_dt
            else:
                stage_col = std_col = f["source_column"]
                # Datatypes come from the template's own targets for this row
                # (2026-08-25). Until then every derived row said "String",
                # which put "String" where the analyst had Decimal(10,2)/Int
                # on 158 standard-layer cells of one real workbook. The row
                # set itself is already the template's, so its per-row
                # datatypes are the same kind of input as its column list --
                # a client convention, not the FRD's ground truth.
                stage_dt = rt["stage"].get("datatype") or "String"
                std_dt = rt["standard"].get("datatype") or stage_dt
            derived.append({
                **f,
                "stage": {"catalog": stg.get("catalog"), "schema": stg.get("schema"),
                          "table": stage_tbl, "column": stage_col, "datatype": stage_dt},
                "standard": {"catalog": std.get("catalog"), "schema": std.get("schema"),
                             "table": std_tbl, "column": std_col, "datatype": std_dt},
            })
        feed["fields"] = derived
    return contract


# --------------------------------------------------------------------------- #
# rule placement — where the FRD's stated rules land in the workbook
# --------------------------------------------------------------------------- #
# Until 2026-08-22 neither renderer wrote `validation_rules` / `recycle_rule`
# anywhere: the rules survived in the contract JSON but the .xlsx a reviewer
# approves (and the workbook CodeGen's `extract-sttm` reads — its Layer-2
# input is the verbatim rule text) carried none of them. Silent loss.
#
# Placement is deterministic and explainable, never a guess:
#   - a rule is attributed to every rendered field whose source column name
#     appears as a whole word in the rule text  → that row's Comment cell
#     (per-table dialect: CodeGen maps the `Comment` header to `value_spec`)
#     or Business Rule cell (single-sheet dialect);
#   - the recycle rule is attributed to the longest such column name
#     → per-table dialect: that row's `Recycle Flag` cell as `Y ( <verbatim> )`,
#     the shape CodeGen's parse_recycle_text digests (applies_to = that row);
#   - anything that names no rendered column is NOT pinned to an arbitrary
#     row (a wrong applies_to generates wrong code downstream) — it goes to a
#     feed-level, human-visible cell: FILE_DETAILS › File Description
#     (per-table) or the metadata block (single-sheet).
# Every placement is recorded in `_provenance.rule_placement` and the
# phase5 report, so an unattributed rule is visible, not silent.


def _column_mentions(text, columns):
    """Rendered source columns (original spelling) mentioned as whole words
    in `text`, longest first. Case-insensitive; `_` is a word character so
    MEMBER_ID does not match inside MEMBER_ID_2."""
    hits = []
    for col in columns:
        if col and re.search(rf"(?<![A-Za-z0-9_]){re.escape(col)}(?![A-Za-z0-9_])", text, re.I):
            hits.append(col)
    return sorted(hits, key=len, reverse=True)


def place_feed_rules(feed):
    """Decide where one feed's FRD-stated rules land. Pure: reads
    `feed["fields"]` (derived), `validation_rules`, `recycle_rule`.
    Returns {"by_column": {source_column: [rule, ...]},   # rendered order
             "recycle": {"text", "column" | None} | None,
             "unattributed": [rule, ...]}"""
    # audit rows (source "NA") are never rule targets
    columns = [f["source_column"] for f in (feed.get("fields") or []) if not f.get("audit")]
    by_column = {c: [] for c in columns}
    unattributed = []
    for rule in feed.get("validation_rules") or []:
        cols = _column_mentions(rule, columns)
        if cols:
            for c in cols:
                by_column[c].append(rule)
        else:
            unattributed.append(rule)
    recycle = None
    if feed.get("recycle_rule"):
        cols = _column_mentions(feed["recycle_rule"], columns)
        recycle = {"text": feed["recycle_rule"], "column": cols[0] if cols else None}
    return {"by_column": {c: r for c, r in by_column.items() if r},
            "recycle": recycle, "unattributed": unattributed}


def _feed_level_rule_text(placement):
    """The human-visible feed-level cell: unattributed rules, plus the
    recycle rule when it could not be pinned to a row."""
    parts = list(placement["unattributed"])
    rec = placement["recycle"]
    if rec and rec["column"] is None:
        parts.append(f"Recycle rule: {rec['text']}")
    return "\n".join(parts)


# --------------------------------------------------------------------------- #
# renderers
# --------------------------------------------------------------------------- #
_HDR_FILL = PatternFill("solid", fgColor="D9E1F2")
_SEC_FILL = PatternFill("solid", fgColor="BDD7EE")
_BOLD = Font(bold=True)


def _style_row(ws, r, n_cols, fill):
    for c in range(1, n_cols + 1):
        cell = ws.cell(r, c)
        cell.font = _BOLD
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")


# Row-2 source-band headers. Every one of these MUST be a header CodeGen's
# `extract-sttm` knows (its config `header_synonyms`, normalized): that
# extractor raises on an unknown source-band header, and `Comment` is the one
# it maps to `value_spec` (the per-field Layer-2 input). Pinned by
# tests/test_render_rules.py. The Recycle Flag column sits BEYOND the
# standard band (as the golden workbook parks it) so it is outside every band.
_PER_TABLE_SRC_HEADERS = ["Database column Name", "NULL CHECK", "Description", "Sample Value",
                          "DataType", "PHI Field", "Mandatory Field", "Comment"]
_PER_TABLE_TGT_HEADERS = ["Schema", "TableName", "ColumnName", "DataType"]
_RECYCLE_HEADER = "Recycle Flag"


def render_sheet_per_table(contract, out_path, placements=None):
    """Demo (CV) dialect: FILE_DETAILS + VERSION_HISTORY + one MAPPING-* sheet per feed.
    `placements` — one `place_feed_rules()` result per feed (computed when None)."""
    if placements is None:
        placements = [place_feed_rules(feed) for feed in contract["feeds"]]
    wb = Workbook()
    ws = wb.active
    ws.title = "FILE_DETAILS"
    ws.append(["Vendor", "FileName", "File Description", "Location", "Frequency"])
    _style_row(ws, 1, 5, _HDR_FILL)
    for feed, placement in zip(contract["feeds"], placements):
        ws.append([feed.get("source_system") or "",
                   "; ".join(feed.get("file_name_patterns", [])),
                   _feed_level_rule_text(placement),
                   feed.get("landing_location") or "", feed.get("frequency") or ""])

    vh = wb.create_sheet("VERSION_HISTORY")
    vh.append(["Version", "Date", "Author", "Change Description"])
    _style_row(vh, 1, 4, _HDR_FILL)
    vh.append(["0.1", datetime.now(timezone.utc).date().isoformat(), GENERATOR,
               f"Auto-generated from {contract.get('generated_from_frd', 'FRD')}"])

    src_headers, tgt_headers = _PER_TABLE_SRC_HEADERS, _PER_TABLE_TGT_HEADERS
    n_src, n_tgt = len(src_headers), len(tgt_headers)
    for feed, placement in zip(contract["feeds"], placements):
        fields = feed.get("fields") or []
        if not fields:
            continue
        table = (feed.get("stage_target") or {}).get("tables", [feed["feed_name"]])[0]
        ws = wb.create_sheet(f"MAPPING-{table.upper()}"[:31])
        recycle = placement["recycle"]
        recycle_col = n_src + n_tgt + 1 + n_tgt + 1 if recycle else None
        n_cols = recycle_col or (n_src + n_tgt + 1 + n_tgt)
        ws.cell(1, 1, "Source File Layout")
        ws.cell(1, n_src + 1, "Stage Layer")
        ws.cell(1, n_src + n_tgt + 2, "Standard Layer")
        _style_row(ws, 1, n_cols, _SEC_FILL)
        headers = src_headers + tgt_headers + [""] + tgt_headers
        if recycle:
            headers = headers + [_RECYCLE_HEADER]
        for c, h in enumerate(headers, 1):
            ws.cell(2, c, h)
        _style_row(ws, 2, n_cols, _HDR_FILL)
        for f in fields:
            comment = "\n".join(placement["by_column"].get(f["source_column"], [])
                                + ([f["comment"]] if f.get("comment") else []))
            row = [
                f["source_column"],
                "Not NULL" if not f["nullable"] else "NULL",
                f.get("description", ""), f.get("sample", ""), f.get("datatype", "String"),
                "Yes" if f["phi"] else "No", "Yes" if f["mandatory"] else "No",
                comment,
                f["stage"]["schema"], f["stage"]["table"], f["stage"]["column"], f["stage"]["datatype"],
                "",
                f["standard"]["schema"], f["standard"]["table"], f["standard"]["column"], f["standard"]["datatype"],
            ]
            if recycle:
                row.append(f"Y ( {recycle['text']} )"
                           if recycle["column"] == f["source_column"] else "")
            ws.append(row)
        ws.freeze_panes = "A3"
        for c in range(1, n_cols + 1):
            ws.column_dimensions[get_column_letter(c)].width = 22
    wb.save(out_path)
    return out_path


def render_single_sheet(contract, out_path, sheet_name="mapping", placements=None):
    """CAQH dialect: metadata block + one wide sheet (source | stage | standard).
    `placements` — one `place_feed_rules()` result per feed (computed when None);
    only the first feed is rendered in this dialect."""
    if placements is None:
        placements = [place_feed_rules(feed) for feed in contract["feeds"]]
    wb = Workbook()
    feed = contract["feeds"][0]
    placement = placements[0]
    ws = wb.active
    ws.title = sheet_name[:31]
    meta_rows = [
        ("File(s)", "\n".join(feed.get("file_name_patterns", []))),
        ("File Generator", feed.get("source_system") or ""),
        ("File Location", feed.get("landing_location") or ""),
        ("LOB", ",".join(re.sub(r"^REG#\d+\s+", "", l) for l in feed.get("lobs", []))),
        ("File frequency", feed.get("frequency") or ""),
        ("Domain", (feed.get("domain") or "").upper()),
        ("Sub-Domain", (feed.get("sub_domain") or "")),
        ("File type", f"{feed.get('file_format') or ''}"
                      + (f" ({feed.get('delimiter')} delimited)" if feed.get("delimiter") else "")),
    ]
    # Feed-level rules this dialect has no column for: the recycle rule
    # (always — there is no Recycle Flag column here) and every rule that
    # names no rendered column. Human-visible in the metadata block.
    if placement["recycle"]:
        meta_rows.append(("Recycle rule", placement["recycle"]["text"]))
    if placement["unattributed"]:
        meta_rows.append(("Validation rules (feed-level)", "\n".join(placement["unattributed"])))
    for k, v in meta_rows:
        ws.append([k, v])
        ws.cell(ws.max_row, 1).font = _BOLD
    src_headers = ["#", "Field Name", "Data Type", "Length",
                   "Field Length\n(fixed width)", "Start position\n(fixed width)",
                   "End Position\n(fixed width)", "Segment", "PII", "Comments", "Business Rule"]
    tgt_headers = ["Catalog", "Schema", "TableName", "ColumnName", "DataType",
                   "Mandatory\nColumn", "Primary Key", "Field Description"]
    label_r = ws.max_row + 1
    ws.cell(label_r, 1, "Source Layout")
    ws.cell(label_r, len(src_headers) + 1, "Stage Layer")
    ws.cell(label_r, len(src_headers) + len(tgt_headers) + 1, "Standard Layer")
    n_cols = len(src_headers) + 2 * len(tgt_headers)
    _style_row(ws, label_r, n_cols, _SEC_FILL)
    for c, h in enumerate(src_headers + tgt_headers + tgt_headers, 1):
        ws.cell(label_r + 1, c, h)
    _style_row(ws, label_r + 1, n_cols, _HDR_FILL)

    for idx, f in enumerate(contract["feeds"][0].get("fields") or [], 1):
        def tgt(t):
            return [t.get("catalog") or "", t.get("schema") or "", t.get("table") or "",
                    t.get("column") or "", t.get("datatype") or "String",
                    "Yes" if f["mandatory"] else "", "", f.get("description", "")]
        business_rule = "\n".join(placement["by_column"].get(f["source_column"], [])
                                  + ([f["business_rule"]] if f.get("business_rule") else []))
        ws.append([idx, f["source_column"], f.get("datatype", ""), f.get("length", ""),
                   f.get("fixed_length", ""), f.get("fixed_start", ""), f.get("fixed_end", ""),
                   f.get("segment", ""), "Yes" if f["phi"] else "", f.get("comment", ""),
                   business_rule]
                  + tgt(f["stage"]) + tgt(f["standard"]))
    ws.freeze_panes = ws.cell(label_r + 2, 1).coordinate
    for c in range(1, n_cols + 1):
        ws.column_dimensions[get_column_letter(c)].width = 18
    wb.save(out_path)
    return out_path


def render_contract(contract, dictionary, out_path, layout=None, feed_match=None):
    """Render the workbook and record, per feed, where every FRD-stated rule
    landed (`_provenance.rule_placement`) — so an unattributed rule is
    visible in the contract and report, never silent.

    With a `layout` (from `layout_of(<template workbook>)`, 2026-08-22) the
    workbook is rendered INTO that template — its sheets, bands, headers,
    widths and styles, whatever dialect it is — and `_provenance.template_fill`
    records what the template had that the contract could not fill. Without
    one (freeform: no template matched) the two built-in dialects remain the
    last-resort fallback."""
    placements = [place_feed_rules(feed) for feed in contract["feeds"]]
    contract.setdefault("_provenance", {})["rule_placement"] = [
        {"feed": feed.get("feed_name"),
         "dialect": dictionary["dialect"],
         "by_column": p["by_column"],
         "recycle": p["recycle"],
         "unattributed": p["unattributed"]}
        for feed, p in zip(contract["feeds"], placements)]
    if layout is not None and layout.get("dialect") == "sheet_per_table" and layout.get("sheets"):
        return render_into_template(contract, dictionary, layout, out_path, placements,
                                    feed_match=feed_match or {})
    if layout is not None and layout.get("dialect") == "single_sheet" and layout.get("sheet"):
        return render_into_single_sheet_template(contract, layout, out_path, placements)
    if dictionary["dialect"] == "sheet_per_table":
        return render_sheet_per_table(contract, out_path, placements=placements)
    sheet = next(iter(dictionary["feeds"].values()))["sheet"]
    return render_single_sheet(contract, out_path, sheet_name=sheet, placements=placements)


# --------------------------------------------------------------------------- #
# template fill — render INTO the chosen reference workbook (2026-08-22)
# --------------------------------------------------------------------------- #
# The template's layout is the dialect. Its sheets, band labels, headers,
# column widths and cell styles are kept; its data rows (another FRD's) are
# removed and ours written under the SAME headers via the logical roles the
# parser recovered (`layout_of`). A template column the contract knows
# nothing about (e.g. "Owner", "Start position") stays BLANK and is listed in
# `_provenance.template_fill.unfilled_columns` — the ad-lib is the shape,
# never a cell value. Sheets the template has for feeds we do not render are
# removed; feeds the template has no sheet for get a copy of the lead sheet.

def _copy_row_style(ws, src_row, dst_row, width):
    for c in range(1, width + 1):
        a, b = ws.cell(src_row, c), ws.cell(dst_row, c)
        b.font, b.fill, b.border = copy(a.font), copy(a.fill), copy(a.border)
        b.alignment, b.number_format = copy(a.alignment), a.number_format


def _field_value(f, role, placement, sheet_has_comment):
    """The cell value for one derived field under one logical column role."""
    kind, key = role
    if kind == "source":
        rules = placement["by_column"].get(f["source_column"], [])
        if key == "source_column":
            return f["source_column"]
        # Template-sourced fields carry the analyst's verbatim cell text
        # (`*_raw`, frdsttm.reference_workbooks, 2026-08-25); write it back
        # as-is so the draft keeps the workbook's own spelling ("NO", "NOT
        # NULL", "NA", double spaces). Fields built without a template
        # (freeform, synthetic) have no raw text and fall back to the
        # derived spelling.
        if key == "description":
            return f["description_raw"] if f.get("description_raw") else f.get("description", "")
        if key == "sample":
            return f.get("sample", "")
        if key == "datatype":
            return f.get("datatype", "String")
        if key == "nullable_raw":
            if "nullable_raw" in f:
                return f["nullable_raw"]
            return "NULL" if f.get("nullable") else "Not NULL"
        if key == "phi_raw":
            if "phi_raw" in f:
                return f["phi_raw"]
            return "Yes" if f.get("phi") else "No"
        if key == "mandatory_raw":
            if "mandatory_raw" in f:
                return f["mandatory_raw"]
            return "Yes" if f.get("mandatory") else "No"
        if key == "comment":
            return "\n".join(rules + ([f["comment"]] if f.get("comment") else []))
        if key == "business_rule":
            # rules go to Comment when the sheet has one; else here
            own = [] if sheet_has_comment else rules
            return "\n".join(own + ([f["business_rule"]] if f.get("business_rule") else []))
        return f.get(key, "") or ""
    if kind in ("stage", "standard"):
        return (f.get(kind) or {}).get(key) or ""
    if kind == "recycle":
        rec = placement["recycle"]
        return f"Y ( {rec['text']} )" if rec and rec["column"] == f["source_column"] else ""
    return ""


def _fill_sheet(ws, sheet_layout, fields, placement, audit_style_from=None):
    """Replace a template sheet's data rows with `fields`, under its headers."""
    first = sheet_layout["first_data_row"]
    width = sheet_layout["width"]
    cols = sheet_layout["columns"]
    # remember the first data row's styles before clearing
    style_row = first if ws.max_row >= first else None
    styles = None
    if style_row:
        styles = [(copy(ws.cell(style_row, c).font), copy(ws.cell(style_row, c).fill),
                   copy(ws.cell(style_row, c).border), copy(ws.cell(style_row, c).alignment),
                   ws.cell(style_row, c).number_format) for c in range(1, width + 1)]
    if ws.max_row >= first:
        ws.delete_rows(first, ws.max_row - first + 1)
    has_comment = any(c["role"] == ("source", "comment") for c in cols)
    r = first
    for n, f in enumerate(fields, 1):
        for c in cols:
            role = c["role"]
            if role is None:
                continue
            val = n if role == ("index", None) else _field_value(f, role, placement, has_comment)
            ws.cell(r, c["index"] + 1).value = val if val != "" else None
        if styles:
            for ci, (fo, fi, bo, al, nf) in enumerate(styles, 1):
                cell = ws.cell(r, ci)
                cell.font, cell.fill, cell.border, cell.alignment, cell.number_format = fo, fi, bo, al, nf
        r += 1


def render_into_template(contract, dictionary, layout, out_path, placements, feed_match):
    from openpyxl import load_workbook as _load
    wb = _load(layout["path"])
    sheets = layout["sheets"]
    sheet_objs = {name: wb[name] for name in sheets if name in wb.sheetnames}
    lead_sheet_name = next(iter(sheets))
    fill = {"template": Path(layout["path"]).name, "dialect": "sheet_per_table",
            "sheets": {}, "unfilled_columns": {}, "created_sheets": [], "removed_sheets": [],
            "file_details": None, "version_history": None}

    # FILE_DETAILS: our feeds under the template's own header row
    fd = layout.get("file_details")
    if fd and "FILE_DETAILS" in wb.sheetnames:
        ws = wb["FILE_DETAILS"]
        if ws.max_row >= 2:
            ws.delete_rows(2, ws.max_row - 1)
        cols = fd["columns"]
        for i, (feed, placement) in enumerate(zip(contract["feeds"], placements), 2):
            vals = {"vendor": feed.get("source_system") or "",
                    "file_name": "; ".join(feed.get("file_name_patterns", [])),
                    "description": _feed_level_rule_text(placement),
                    "location": feed.get("landing_location") or "",
                    "frequency": feed.get("frequency") or ""}
            for k, v in vals.items():
                if k in cols and v:
                    ws.cell(i, cols[k] + 1, v)
        fill["file_details"] = {"filled": sorted(cols), "unfilled": [h for h in fd["headers"]
                                                                    if h and _nl(h) not in
                                                                    {_nl(fd["headers"][i]) for i in cols.values()}]}
    else:
        # The template has no FILE_DETAILS sheet. It carries contract FACTS
        # (vendor, file pattern, location, frequency) plus the feed-level
        # rules' only home, and CodeGen's extractor requires it — so create
        # a minimal one, flagged as created rather than copied.
        ws = wb.create_sheet("FILE_DETAILS", 0)
        ws.append(["Vendor", "FileName", "File Description", "Location", "Frequency"])
        for feed, placement in zip(contract["feeds"], placements):
            ws.append([feed.get("source_system") or "", "; ".join(feed.get("file_name_patterns", [])),
                       _feed_level_rule_text(placement), feed.get("landing_location") or "",
                       feed.get("frequency") or ""])
        fill["created_sheets"].append("FILE_DETAILS (template had none; minimal sheet — CodeGen requires it)")

    vh = layout.get("version_history")
    if vh and "VERSION_HISTORY" in wb.sheetnames:
        ws = wb["VERSION_HISTORY"]
        if ws.max_row >= 2:
            ws.delete_rows(2, ws.max_row - 1)
        cols = vh["columns"]
        vals = {"version": "0.1", "date": datetime.now(timezone.utc).date().isoformat(),
                "author": GENERATOR,
                "change": f"Auto-generated from {contract.get('generated_from_frd', 'FRD')} "
                          f"into the layout of {Path(layout['path']).name}"}
        for k, v in vals.items():
            if k in cols:
                ws.cell(2, cols[k] + 1, v)
        fill["version_history"] = {"filled": sorted(cols)}
    else:
        ws = wb.create_sheet("VERSION_HISTORY", 1 if "FILE_DETAILS" in wb.sheetnames else 0)
        ws.append(["Version", "Date", "Author", "Change Description"])
        ws.append(["0.1", datetime.now(timezone.utc).date().isoformat(), GENERATOR,
                   f"Auto-generated from {contract.get('generated_from_frd', 'FRD')} "
                   f"into the layout of {Path(layout['path']).name}"])
        fill["created_sheets"].append("VERSION_HISTORY (template had none; minimal sheet — CodeGen requires it)")

    # which template sheet each rendered feed uses
    used = set()
    for i, feed in enumerate(contract["feeds"]):
        fields = feed.get("fields") or []
        if not fields:
            continue
        key = feed_match.get(i)
        src_sheet = (dictionary["feeds"].get(key) or {}).get("sheet") if key is not None else None
        if src_sheet not in sheets or src_sheet in used:
            src_sheet = lead_sheet_name if lead_sheet_name not in used else next(
                (n for n in sheets if n not in used), lead_sheet_name)
        table = (feed.get("stage_target") or {}).get("tables", [feed["feed_name"]])[0]
        title = f"MAPPING-{table.upper()}"[:31]
        if src_sheet in used:
            ws = wb.copy_worksheet(sheet_objs[src_sheet])
            fill["created_sheets"].append(f"{title} (copy of {src_sheet})")
        else:
            ws = sheet_objs[src_sheet]
            used.add(src_sheet)
        ws.title = title if title not in wb.sheetnames or wb[title] is ws else f"{title[:28]}-{i}"
        sl = sheets[src_sheet]
        _fill_sheet(ws, sl, fields, placements[i])
        fill["sheets"][ws.title] = {"from": src_sheet, "rows": len(fields)}
        if sl["unmapped"]:
            fill["unfilled_columns"][ws.title] = sl["unmapped"]
    # template sheets for feeds we did not render carry another FRD's rows: remove
    for name in list(sheets):
        if name not in used and name in sheet_objs:
            wb.remove(sheet_objs[name])
            fill["removed_sheets"].append(name)
    contract.setdefault("_provenance", {})["template_fill"] = fill
    wb.save(out_path)
    return out_path


def render_into_single_sheet_template(contract, layout, out_path, placements):
    from openpyxl import load_workbook as _load
    wb = _load(layout["path"])
    ws = wb[layout["sheet"]]
    feed, placement = contract["feeds"][0], placements[0]
    fill = {"template": Path(layout["path"]).name, "dialect": "single_sheet",
            "sheets": {layout["sheet"]: {"rows": len(feed.get("fields") or [])}},
            "unfilled_columns": {layout["sheet"]: layout["unmapped"]} if layout["unmapped"] else {},
            "unfilled_meta": [], "created_sheets": [], "removed_sheets": []}
    # metadata block: fill the keys we know, clear the rest (they are another feed's)
    for m in layout["meta_rows"]:
        field = m["field"]
        val = None
        if field == "file_name_patterns":
            val = "\n".join(feed.get("file_name_patterns", []))
        elif field == "lobs":
            val = ",".join(re.sub(r"^REG#\d+\s+", "", l) for l in feed.get("lobs", []))
        elif field == "file_type":
            val = (f"{feed.get('file_format') or ''}"
                   + (f" ({feed.get('delimiter')} delimited)" if feed.get("delimiter") else "")) or None
        elif field == "domain":
            val = (feed.get("domain") or "").upper() or None
        elif field is not None:
            val = feed.get(field) or None
        else:
            fill["unfilled_meta"].append(m["key"])
        ws.cell(m["row"], 2).value = val   # .value=: None must CLEAR the template's old value
    # feed-level rules: the recycle rule and unattributed rules get rows of
    # their own in the metadata block only if the template has such keys;
    # otherwise they go to the report (recorded), never invented into a row
    leftover = _feed_level_rule_text(placement)
    if leftover:
        fill["feed_level_rules_unplaced"] = leftover
    _fill_sheet(ws, layout, feed.get("fields") or [], placement)
    for name in wb.sheetnames:
        if name != layout["sheet"] and name.startswith("MAPPING-"):
            del wb[name]
            fill["removed_sheets"].append(name)
    contract.setdefault("_provenance", {})["template_fill"] = fill
    wb.save(out_path)
    return out_path


# --------------------------------------------------------------------------- #
# golden-pair eval: derived mappings vs reference targets
# --------------------------------------------------------------------------- #
def evaluate_against_reference(contract, dictionary, feed_match):
    report = {"feeds": [], "totals": {"cells": 0, "match": 0}}
    for i, key in feed_match.items():
        feed = contract["feeds"][i]
        ref = dictionary["feeds"][key]["ref_targets"]
        diffs, cells, match = [], 0, 0
        for f, rt in zip(feed.get("fields") or [], ref):
            for layer in ("stage", "standard"):
                for attr in ("schema", "table", "column", "datatype"):
                    ref_v = _nl(rt[layer].get(attr))
                    if not ref_v:
                        continue
                    cells += 1
                    der_v = _nl(f[layer].get(attr))
                    if der_v == ref_v:
                        match += 1
                    elif len(diffs) < 8:
                        diffs.append(f"{f['source_column']}.{layer}.{attr}: "
                                     f"derived={der_v!r} ref={ref_v!r}")
        pct = round(100 * match / cells, 1) if cells else 100.0
        report["feeds"].append({"feed": feed["feed_name"], "cells": cells,
                                "match": match, "pct": pct, "sample_diffs": diffs})
        report["totals"]["cells"] += cells
        report["totals"]["match"] += match
    t = report["totals"]
    t["pct"] = round(100 * t["match"] / t["cells"], 1) if t["cells"] else 100.0
    return report


def evaluate_cross_reference(contract, own_dictionary, feed_match_own):
    """Eval against the document's OWN reference workbook when the render was
    driven by a DIFFERENT template (exclude-own-reference cross-validation,
    2026-08-22). `evaluate_against_reference` zips rows positionally, which
    is only valid when the dictionary that derived the fields IS the eval
    reference; here rows are aligned by source column name instead, and the
    reference defines the denominator — a reference column the render never
    produced counts as unmatched cells, not as ignored.
    """
    report = {"feeds": [], "totals": {"cells": 0, "match": 0}}

    def _key(source_column, audit, stage_column):
        # Audit rows all share source "NA" — align them by target column
        # instead, or every audit row would collide on one key.
        base = _nl(source_column)
        return f"{base}:{_nl(stage_column)}" if audit else base

    for i, key in feed_match_own.items():
        feed = contract["feeds"][i]
        derived_by_col = {_key(f.get("source_column"), f.get("audit"),
                               (f.get("stage") or {}).get("column")): f
                          for f in (feed.get("fields") or [])}
        ref_feed = own_dictionary["feeds"][key]
        diffs, cells, match = [], 0, 0
        for src_f, rt in zip(ref_feed["fields"], ref_feed["ref_targets"]):
            col = _key(src_f.get("source_column"), src_f.get("audit"),
                       rt["stage"].get("column"))
            derived = derived_by_col.get(col)
            for layer in ("stage", "standard"):
                for attr in ("schema", "table", "column", "datatype"):
                    ref_v = _nl(rt[layer].get(attr))
                    if not ref_v:
                        continue
                    cells += 1
                    der_v = _nl((derived or {}).get(layer, {}).get(attr)) if derived else ""
                    if der_v == ref_v:
                        match += 1
                    elif len(diffs) < 8:
                        diffs.append(f"{col}.{layer}.{attr}: "
                                     f"derived={der_v!r} ref={ref_v!r}"
                                     + ("" if derived else " (column not rendered)"))
        pct = round(100 * match / cells, 1) if cells else 100.0
        report["feeds"].append({"feed": feed["feed_name"], "cells": cells,
                                "match": match, "pct": pct, "sample_diffs": diffs})
        report["totals"]["cells"] += cells
        report["totals"]["match"] += match
    t = report["totals"]
    t["pct"] = round(100 * t["match"] / t["cells"], 1) if t["cells"] else 100.0
    return report

# COMMAND ----------

# MAGIC %md
# MAGIC ## Driver: contract × reference workbook → resolve → derive → render → eval

# COMMAND ----------

from pathlib import Path


_loose_tokens = loose_tokens  # factored into frdsttm.reference_workbooks


refs = sorted(Path(REFERENCE_DIR).glob("*.xlsx"))
if not refs:
    # Pre-2026-08-22 this was a hard assert. Freeform mode makes an empty
    # reference library a WARNED state, not a fatal one: every document
    # renders best-effort from its contract alone, flagged, with no eval.
    print(f"WARNING: no reference workbooks in {REFERENCE_DIR} — every "
          f"document renders freeform (no dictionary, no eval). Import "
          f"reference STTMs via the review app's corpus bootstrap.")

# --------------------------------------------------------------------------- #
# Template architecture (2026-08-22, docs/TEMPLATE_ARCHITECTURE.md).
# The reference pick is a computed decision with exactly three modes:
#   single   — one workbook is the template (dialect + dictionary + eval ref)
#   amalgam  — top-k workbooks merge, first-wins per table key
#   freeform — nothing matched; best-effort render from the contract, FLAGGED
# Scores come from frdsttm.similarity over the corpus index the app's
# bootstrap wrote next to the workbooks; without an index, features are
# computed from the contract JSON (weaker, still deterministic). A document's
# OWN paired workbook is excluded from candidacy (cross-validation, decided
# 2026-08-22) and used for eval only — exclude_own_reference=0 restores the
# old self-referential behavior for debugging, never for demos.
# --------------------------------------------------------------------------- #
EXCLUDE_OWN = _param("exclude_own_reference", "1").strip().lower() in ("1", "true", "yes")
TEMPLATE_THRESHOLDS = thresholds_from(_param)

corpus_index = load_corpus_index(REFERENCE_DIR)  # None when absent; raises when corrupt
dicts_by_name = {p.name: parse_reference_workbook(str(p)) for p in refs}
wb_feats = reference_features_from_index(corpus_index) if corpus_index else {}
wb_feats = {n: f for n, f in wb_feats.items() if n in dicts_by_name}
for _name, _d in dicts_by_name.items():
    if _name not in wb_feats:
        wb_feats[_name] = workbook_features(_name, _d)

_FREEFORM_DICT = {"dialect": "sheet_per_table", "feeds": {}, "meta": {}, "feed_sources": {}}


def _own_reference(doc_id):
    """The doc's own paired workbook: corpus pairing first, exact-stem naming
    convention (<doc_id>.sttm.xlsx / <doc_id>.xlsx) as the index-less guard."""
    own = own_reference_for(corpus_index, doc_id) if corpus_index else None
    if own is None:
        lowered = {n: Path(n).stem.lower() for n in dicts_by_name}
        for n, stem in lowered.items():
            if stem in (f"{doc_id.lower()}.sttm", doc_id.lower()):
                own = n
                break
    return own if own in dicts_by_name else None


def _own_pair(doc_id):
    """The corpus pairing record for this doc (reference, score, matched_by),
    or None — the index-less fallback in _own_reference has no record."""
    return (corpus_index.get("pairs", {}) or {}).get(doc_id) if corpus_index else None


def _target_features(doc_id, contract):
    """What the template scorer sees for this document: the FRD-markdown
    features the index holds (identifiers + prose) MERGED with the
    extraction's own table vocabulary (frdsttm.similarity.contract_features).
    Calibrated 2026-08-25 on the two real pairs: real FRDs do not enumerate
    columns, so markdown-only features scored 0.11 against the document's
    own STTM; the feed keys and stage/standard table names the extraction
    carries are what the workbook actually shares."""
    feat = frd_features_from_index(corpus_index, doc_id) if corpus_index else None
    if feat is None:
        feat = frd_features(doc_id, json.dumps(contract, ensure_ascii=False))
    return merge_features(feat, contract_features(doc_id, contract))

# Read contracts from the CONTRACTS_DIR JSON files, not the frd_contracts
# Delta table. The table is a snapshot from 03_contract_build.py's run, at
# which point `_provenance.human_resolutions` doesn't exist yet -- the
# review app (review_app_react/) writes reviewer resolutions into the
# same-named JSON file under CONTRACTS_DIR, in both local mode
# (plain filesystem) and Databricks mode (UC volume, which is just as
# directly Path-readable -- see OUT_ROOT/CONTRACTS_DIR above). Reading the
# table here would silently make every human resolution invisible to this
# notebook; this was confirmed empirically while building the human-
# resolution wiring below (docs/STANDUP_NOTES.md has the before/after).
_contract_paths = sorted(Path(CONTRACTS_DIR).glob("*.contract.json"))
assert _contract_paths, f"No contract JSONs found in {CONTRACTS_DIR}"

contracts = {}
for _p in _contract_paths:
    doc_id = _p.name[: -len(".contract.json")]
    contracts[doc_id] = json.loads(_p.read_text(encoding="utf-8"))
print(f"{len(contracts)} contract(s), {len(refs)} reference workbook(s)")

Path(RENDERED_DIR).mkdir(parents=True, exist_ok=True)
Path(CONTRACTS_DIR).mkdir(parents=True, exist_ok=True)
Path(REPORTS_DIR).mkdir(parents=True, exist_ok=True)

runs = []
for doc_id, contract in contracts.items():
    own_ref = _own_reference(doc_id)
    exclude = {own_ref} if (EXCLUDE_OWN and own_ref) else set()
    # The corpus's exact-name pairing (FRD_<name> <-> STTM_<name>) is
    # definitive in pair_corpus, so it is definitive here too (2026-08-25):
    # a name-paired own STTM is the template by identity, not by score --
    # unless exclude-own is on, in which case cross-validation wins.
    pair = _own_pair(doc_id)
    pinned = own_ref if (own_ref and not exclude and pair
                         and pair.get("matched_by") == "name") else None
    if wb_feats:
        decision = decide_templates(_target_features(doc_id, contract), wb_feats,
                                    TEMPLATE_THRESHOLDS, exclude=exclude, pinned=pinned)
    else:
        decision = {"mode": "freeform", "selections": [], "ranked": [],
                    "pinned": None, "thresholds": {}}
    order = [sel["reference"] for sel in decision["selections"]]
    if decision["mode"] == "freeform":
        dictionary = dict(_FREEFORM_DICT)
    elif len(order) == 1:
        dictionary = dicts_by_name[order[0]]
    else:
        dictionary = merge_dictionaries(dicts_by_name, order)
    fm = match_feeds(contract, dictionary)
    if not fm and decision["mode"] != "freeform":
        # The scored template did not structurally match any feed. Demote to
        # freeform and KEEP rendering — the pre-2026-08-22 behavior (silent
        # SKIP) left the document with no workbook at all.
        decision = {**decision, "mode": "freeform", "demoted_from": order}
        order = []
        dictionary = dict(_FREEFORM_DICT)
        fm = {}

    human_settled, human_audit = apply_human_resolutions(contract)
    human_audit = list(human_audit)  # snapshot -- resolution_audit gets automatic entries appended next
    resolutions = resolve_attribution(contract, dictionary, fm, already_settled=human_settled)
    contract["_provenance"]["resolution_audit"].extend(
        {"ambiguity_id": None, "kind": None, "ambiguity_text": None, "resolution_source": "automatic",
         "resolution_type": None, "applied": True, "target": None, "detail": r} for r in resolutions)
    # Standards BEFORE derivation: the per-field rows must carry the filled
    # catalog/schema, not the nulls the FRD left. Anything the client's
    # vocabulary cannot resolve stays null and is gated, never invented.
    apply_standards_targets(contract, fm)
    derive_field_mappings(contract, dictionary, fm)
    out_xlsx = str(Path(RENDERED_DIR) / f"{doc_id}.sttm.xlsx")
    # Render INTO the lead template's own layout (2026-08-22): its sheets,
    # headers and styles are the dialect. Freeform keeps the built-in fallback.
    template_layout = layout_of(str(Path(REFERENCE_DIR) / order[0])) if order else None
    # openpyxl's save() writes in a way (seek + partial rewrites) that the UC
    # volumes FUSE mount rejects with OSError [Errno 5], observed 2026-08-24
    # on the first live App run. Write to local scratch first, then copy to
    # the volume so the transfer is a plain sequential write. Local mode
    # (RENDERED_DIR not under /Volumes/) keeps the direct write.
    if out_xlsx.startswith("/Volumes/"):
        Path(RENDERED_DIR).mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as _tmp:
            _tmp_path = _tmp.name
        try:
            render_contract(contract, dictionary, _tmp_path, layout=template_layout, feed_match=fm)
            shutil.copyfile(_tmp_path, out_xlsx)
        finally:
            try:
                os.unlink(_tmp_path)
            except OSError:
                pass
    else:
        render_contract(contract, dictionary, out_xlsx, layout=template_layout, feed_match=fm)

    # Eval precedence: the doc's OWN workbook is the ground truth whenever it
    # exists (name-aligned cross eval — valid even though the dictionary that
    # drove the render was a different template). Otherwise the positional
    # eval against the chosen template, matching pre-2026-08-22 semantics.
    # Freeform with no own reference has nothing honest to eval against.
    ev = None
    ev_reference = None
    if own_ref is not None:
        fm_own = match_feeds(contract, dicts_by_name[own_ref])
        if fm_own:
            ev = evaluate_cross_reference(contract, dicts_by_name[own_ref], fm_own)
            ev_reference = own_ref
    if ev is None and fm:
        ev = evaluate_against_reference(contract, dictionary, fm)
        ev_reference = " + ".join(order)

    contract["_provenance"]["template_decision"] = {
        "mode": decision["mode"],
        "selections": decision["selections"],
        "ranked": decision["ranked"][:5],
        "own_reference": own_ref,
        "own_excluded": bool(exclude),
        "pinned": decision.get("pinned"),
        "eval_reference": ev_reference,
        "feed_sources": dictionary.get("feed_sources", {}),
        "thresholds": decision.get("thresholds", {}),
        **({"demoted_from": decision["demoted_from"]}
           if "demoted_from" in decision else {}),
    }

    (Path(CONTRACTS_DIR) / f"{doc_id}.contract.v2.json").write_text(
        json.dumps(contract, indent=2, ensure_ascii=False), encoding="utf-8")

    n_human_applied = sum(1 for a in human_audit if a["applied"])
    template_desc = {
        "single": f"template: {order[0]}" if order else "template: (none)",
        "amalgam": "amalgam of: " + ", ".join(order),
        "freeform": "FREEFORM — no template matched"
                    + (f" (demoted from {', '.join(decision['demoted_from'])})"
                       if "demoted_from" in decision else ""),
    }[decision["mode"]]
    eval_line = (
        f"**Golden-pair eval vs {ev_reference}: {ev['totals']['pct']}%** "
        f"({ev['totals']['match']}/{ev['totals']['cells']} target cells match the reference)"
        if ev else
        "**No eval** — no paired reference workbook for this document"
    )
    lines = [f"# STTM render — {doc_id}", "",
             f"**Status: {contract['status']}** | dialect: {dictionary['dialect']} | "
             f"{template_desc}",
             eval_line, ""]
    lines += ["## Template decision",
              f"- mode: **{decision['mode']}**"
              + (f" | own reference {own_ref} excluded from candidacy" if exclude else "")]
    for r in decision["ranked"][:5]:
        marker = " (excluded — own reference)" if r["excluded"] else ""
        chosen = " ← chosen" if r["reference"] in order else ""
        lines.append(f"- {r['reference']}: score {r['score']:.3f} "
                     f"(columns {r['components']['columns']:.2f}, tables "
                     f"{r['components']['tables']:.2f}, prose "
                     f"{r['components']['tokens']:.2f}){marker}{chosen}")
    if dictionary.get("feed_sources") and decision["mode"] == "amalgam":
        lines += [f"- sheet {k!r} from {v}" for k, v in
                  sorted(dictionary["feed_sources"].items())]
    lines += [""]
    if human_audit:
        lines += ["## Resolution audit (human vs. automatic)"]
        for a in human_audit:
            status = "applied" if a["applied"] else f"NOT applied — {a['reason_not_applied']}"
            stale_note = " [STALE: candidates changed since this resolution was recorded]" if a.get("stale") else ""
            text_preview = (a.get("ambiguity_text") or "")[:100]
            lines.append(f"- [human] {a['kind']} ({a.get('resolution_type')}): {text_preview!r} -> {status}{stale_note}"
                         + (f" (target: {a['target']})" if a["target"] else ""))
        lines += [""]
    if resolutions:
        lines += ["## Attribution resolutions (dictionary cross-check, automatic)"] + \
                 [f"- {r}" for r in resolutions] + [""]
    tf = contract["_provenance"].get("template_fill")
    if tf:
        lines += [f"## Template fill — rendered into the layout of {tf['template']} ({tf['dialect']})"]
        for sheet, info in tf.get("sheets", {}).items():
            lines.append(f"- sheet {sheet}: {info.get('rows', 0)} rows"
                         + (f" (layout from {info['from']})" if info.get("from") else ""))
        for sheet, cols in tf.get("unfilled_columns", {}).items():
            lines.append(f"- {sheet}: template columns left EMPTY — the FRD does not state them: "
                         + ", ".join(cols))
        if tf.get("unfilled_meta"):
            lines.append("- metadata keys left empty: " + ", ".join(tf["unfilled_meta"]))
        if tf.get("feed_level_rules_unplaced"):
            lines.append("- feed-level rules with no home in this layout (see Rule placement): "
                         + _quote_rule(tf["feed_level_rules_unplaced"], 120))
        for c in tf.get("created_sheets", []):
            lines.append(f"- created sheet: {c}")
        for r in tf.get("removed_sheets", []):
            lines.append(f"- removed template sheet (another FRD's rows): {r}")
        lines += [""]
    placements = contract["_provenance"].get("rule_placement") or []
    if any(p["by_column"] or p["recycle"] or p["unattributed"] for p in placements):
        lines += ["## Rule placement (where the FRD's rules landed in the workbook)"]
        for p in placements:
            for col, rules in p["by_column"].items():
                for r in rules:
                    lines.append(f"- **{p['feed']}** › row {col}: {_quote_rule(r, 120)!r}")
            rec = p["recycle"]
            if rec:
                lines.append(f"- **{p['feed']}** › recycle rule "
                             + (f"→ row {rec['column']} (Recycle Flag cell)"
                                if rec["column"] and p["dialect"] == "sheet_per_table" else
                                f"→ metadata block" if p["dialect"] != "sheet_per_table" else
                                "→ NOT attributed to a column — FILE_DETAILS › File Description only; "
                                "CodeGen will see no recycle spec for this feed")
                             + f": {_quote_rule(rec['text'], 120)!r}")
            for r in p["unattributed"]:
                lines.append(f"- **{p['feed']}** › NOT attributed (names no rendered column) → "
                             f"feed-level cell only: {_quote_rule(r, 120)!r}")
        lines += [""]
    if ev:
        lines += ["## Per-feed eval"]
        for f in ev["feeds"]:
            lines.append(f"- **{f['feed']}**: {f['pct']}% ({f['match']}/{f['cells']})")
            for d in f["sample_diffs"]:
                lines.append(f"    - {d}")
    (Path(REPORTS_DIR) / f"{doc_id}.phase5.md").write_text("\n".join(lines), encoding="utf-8")

    print(f"{contract['status']:6s} {doc_id}: [{decision['mode']}] "
          f"eval {str(ev['totals']['pct']) + '%' if ev else 'n/a'}, "
          f"{n_human_applied}/{len(human_audit)} human resolution(s) applied, "
          f"{len(resolutions)} attribution resolution(s) -> {Path(out_xlsx).name}")
    # Provenance columns: 02's sidecar says which model/prompt/input bytes
    # produced the extraction this render descends from; the render's own
    # sha256 lets a workbook a reviewer uploads to SharePoint be matched
    # back to exactly this row. An older artifact set without the sidecar
    # yields NULLs, never a failure.
    _meta_path = Path(EXTRACTIONS_DIR) / f"{doc_id}.extraction_meta.json"
    try:
        _meta = json.loads(_meta_path.read_text(encoding="utf-8")) if _meta_path.is_file() else {}
    except (OSError, json.JSONDecodeError):
        _meta = {}
    _usage = _meta.get("usage") or {}
    runs.append({"doc_id": doc_id, "status": contract["status"],
                 "dialect": dictionary["dialect"],
                 "run_label": RUN_LABEL,
                 "triggered_by": TRIGGERED_BY,
                 "job_run_id": JOB_RUN_ID,
                 "provider": _meta.get("provider"),
                 "model": _meta.get("model"),
                 "system_prompt_sha256": _meta.get("system_prompt_sha256"),
                 "input_tokens": int(_usage.get("input_tokens") or 0),
                 "output_tokens": int(_usage.get("output_tokens") or 0),
                 "frd_sha256": _meta.get("content_sha256"),
                 "rendered_sha256": hashlib.sha256(Path(out_xlsx).read_bytes()).hexdigest(),
                 # Which revision of the client's naming/engineering standards
                 # decided this render's target side (docs/AI_GOVERNANCE.md).
                 "standards_sha256": _std.standards_sha256(),
                 # comma-joined template list, or the mode marker when none
                 "reference": ", ".join(order) if order else "(freeform)",
                 "template_mode": decision["mode"],
                 "eval_reference": ev_reference or "",
                 "n_resolutions": len(resolutions),
                 "n_human_resolutions": len(human_audit),
                 "n_human_resolutions_applied": n_human_applied,
                 # -1.0 / 0 are the explicit "no eval" sentinels: the columns
                 # stay non-nullable in both writers, and a dashboard can
                 # filter eval_pct >= 0 without NULL semantics.
                 "eval_pct": float(ev["totals"]["pct"]) if ev else -1.0,
                 "eval_cells": ev["totals"]["cells"] if ev else 0,
                 "rendered_path": out_xlsx,
                 "run_at": datetime.now(timezone.utc)})

if IS_DATABRICKS:
    from pyspark.sql import types as T

    rschema = T.StructType([
        T.StructField("doc_id", T.StringType(), False),
        T.StructField("status", T.StringType(), False),
        T.StructField("dialect", T.StringType(), False),
        T.StructField("run_label", T.StringType(), False),
        T.StructField("triggered_by", T.StringType(), False),
        T.StructField("job_run_id", T.StringType(), False),
        T.StructField("provider", T.StringType(), True),
        T.StructField("model", T.StringType(), True),
        T.StructField("system_prompt_sha256", T.StringType(), True),
        T.StructField("input_tokens", T.IntegerType(), False),
        T.StructField("output_tokens", T.IntegerType(), False),
        T.StructField("frd_sha256", T.StringType(), True),
        T.StructField("rendered_sha256", T.StringType(), False),
        T.StructField("standards_sha256", T.StringType(), True),
        T.StructField("reference", T.StringType(), False),
        T.StructField("template_mode", T.StringType(), False),
        T.StructField("eval_reference", T.StringType(), False),
        T.StructField("n_resolutions", T.IntegerType(), False),
        T.StructField("n_human_resolutions", T.IntegerType(), False),
        T.StructField("n_human_resolutions_applied", T.IntegerType(), False),
        T.StructField("eval_pct", T.DoubleType(), False),
        T.StructField("eval_cells", T.IntegerType(), False),
        T.StructField("rendered_path", T.StringType(), False),
        T.StructField("run_at", T.TimestampType(), False),
    ])
    # APPEND, never overwrite (governance pass): the runs table is the audit
    # trail of every render — a re-render after human resolutions is a new
    # row (run_label ends in ":rerender"), not a rewrite of the last one.
    # mergeSchema lets a table written before the provenance columns existed
    # keep its history and gain the columns.
    spark.createDataFrame(runs, rschema).write.mode("append").option(
        "mergeSchema", "true").saveAsTable(RUNS_TABLE)
    display(spark.table(RUNS_TABLE).orderBy("run_at", ascending=False))
else:
    from _local_tables import append_table, read_table

    append_table(LOCAL_ROOT / "warehouse", CATALOG, SCHEMA, RUNS_TABLE_NAME, runs)
    for _r in read_table(LOCAL_ROOT / "warehouse", CATALOG, SCHEMA, RUNS_TABLE_NAME):
        print(_r)

# COMMAND ----------

# MAGIC %md
# MAGIC ## Reading the eval
# MAGIC The %% is not a score to maximize blindly — the diff list is the point.
# MAGIC Expected gaps with current derivation rules: **standard-layer type casts**
# MAGIC (String vs Decimal/int — an analyst decision the FRD doesn't state) and,
# MAGIC for the single-sheet dialect, **target column renaming conventions**
# MAGIC (e.g. TPL_* codes) plus the standard schema the FRD omits. Those gaps are
# MAGIC the honest backlog: naming-convention config per client, and a typed
# MAGIC gold-layer dictionary input.


# COMMAND ----------

if not IS_DATABRICKS:
    # Local-mode exit guard (2026-08-22). On the py3.14 venv the interpreter can
    # deadlock at SHUTDOWN in C finalizers (deltalake/pyarrow) after every line
    # above has run and every artifact is written and closed — observed on 03
    # as a >13-minute hang at 0% CPU, while the same file exits instantly under
    # runpy. The review app's local-mode runner waits on process exit, so a
    # hang here is a failed demo run. Exit explicitly: nothing in these stages
    # relies on atexit handlers. Never reached in Databricks (no process to
    # exit — the notebook task returns normally).
    import os as _os
    import sys as _sys
    _sys.stdout.flush()
    _sys.stderr.flush()
    _os._exit(0)
