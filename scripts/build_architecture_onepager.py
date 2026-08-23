"""Build the FRD-to-STTM one-slide system-architecture deck (PPTX).

    .venv/bin/pip install python-pptx pillow     # deck-building only, not runtime deps
    .venv/bin/python scripts/build_architecture_onepager.py [output.pptx]

Default output: context/FRD_to_STTM_Agent_System_Architecture.pptx

ONE slide: how the agent actually works, end to end — system of record →
scheduled read-only sync → the four-stage generation pipeline (the single
model call sits in stage 02) → human-in-the-loop review → outputs and the
downstream agents — over the Unity Catalog working store, with the design
doctrine and the color legend along the bottom.

Visual system (2026-08-22): a "Palantir-like" dark, data-dense look. The
palette is Palantir's own open-source Blueprint design system
(github.com/palantir/blueprint, packages/colors/_colors.scss) — near-black
canvas `111418`, dark-gray cards `1c2127`/`252a31`, hairlines `383e47`,
gray text ramp `5f6b7c`→`c5cbd3`, and its accent families used as a
SEMANTIC code, not decoration:
    cerulean `3fa6da`  = the LLM (exactly one place)
    turquoise `13c9ba` = deterministic code (sync, parse, audit, gate, render)
    gold `f0b726`      = a human (review, resolution, upload)
    indigo `9881f3`    = external systems (SharePoint, downstream agents)
    blue `4c90f0`      = storage (Unity Catalog)
Typography: Calibri (body) + Courier New (mono ids / tracked labels) —
both ship with Office and render true-to-width in LibreOffice QA.
Background is a generated PNG (dot grid + soft vignette) written next to the
other deck assets; connectors are thin hairlines with small triangle heads.
Everything is absolute-positioned on a 13.333" × 7.5" canvas; text length
drives fit — shorten copy rather than growing boxes.

Every statement on the slide is true of `staging` as of 2026-08-22 (see
frd-to-sttm-master-context-document.md §11 for what is proven offline vs
unproven live); the slide claims no live deployment.
"""
from __future__ import annotations

import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFilter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

REPO = Path(__file__).resolve().parent.parent
ASSETS = REPO / "scripts" / "deck_assets"
ICONS = ASSETS / "icons_dark"
OUT_DEFAULT = REPO / "context" / "FRD_to_STTM_Agent_System_Architecture.pptx"

# --- Blueprint palette (hex without '#') -----------------------------------
BLACK = "111418"
DG1, DG2, DG3, DG4, DG5 = "1c2127", "252a31", "2f343c", "383e47", "404854"
G1, G2, G3, G4, G5 = "5f6b7c", "738091", "8f99a8", "abb3bf", "c5cbd3"
LG5 = "f6f7f9"
WHITE = "ffffff"
CERULEAN, CERULEAN_DIM = "3fa6da", "0f6894"
TURQ, TURQ_DIM = "13c9ba", "007067"
GOLD, GOLD_DIM = "f0b726", "866103"
INDIGO, INDIGO_DIM = "9881f3", "634dbf"
BLUE, BLUE_DIM = "4c90f0", "215db0"
RED = "e76a6e"
GREEN = "32a467"

SANS = "Calibri"
MONO = "Courier New"

W_IN, H_IN = 13.333, 7.5


# --------------------------------------------------------------------------- #
# helpers
# --------------------------------------------------------------------------- #
def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h.upper())


def make_background(path: Path) -> Path:
    """Near-black canvas, fine dot grid, soft vignette — generated, not shipped."""
    w, h = 2560, 1440
    img = Image.new("RGB", (w, h), "#" + BLACK)
    glow = Image.new("RGB", (w, h), "#" + BLACK)
    gd = ImageDraw.Draw(glow)
    gd.ellipse([-400, -500, 1500, 900], fill="#161c24")
    gd.ellipse([1500, 700, 3200, 1900], fill="#141a21")
    glow = glow.filter(ImageFilter.GaussianBlur(260))
    img = Image.blend(img, glow, 0.9)
    d = ImageDraw.Draw(img)
    step = 40
    for y in range(step, h, step):
        for x in range(step, w, step):
            d.point((x, y), fill="#262c35")
    path.parent.mkdir(parents=True, exist_ok=True)
    img.save(path, optimize=True)
    return path


def add_text(slide, x, y, w, h, runs, *, size=10, color=G5, font=SANS, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None, line_spacing=None,
             margin=0.0):
    """runs: str | list of (text, {overrides}) | list of paragraphs (list of runs).
    A paragraph list is detected when the first element is itself a list."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = anchor
    paragraphs = runs
    if isinstance(runs, str):
        paragraphs = [[(runs, {})]]
    elif runs and not isinstance(runs[0], list):
        paragraphs = [runs]
    for pi, para in enumerate(paragraphs):
        p = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing:
            p.line_spacing = line_spacing
        for text, ov in para:
            r = p.add_run()
            r.text = text
            f = r.font
            f.name = ov.get("font", font)
            f.size = Pt(ov.get("size", size))
            f.bold = ov.get("bold", bold)
            f.italic = ov.get("italic", False)
            f.color.rgb = rgb(ov.get("color", color))
            sp = ov.get("spacing", spacing)
            if sp is not None:
                rPr = r._r.get_or_add_rPr()
                rPr.set("spc", str(int(sp * 100)))
    return tb


def add_rect(slide, x, y, w, h, *, fill=DG1, line=DG4, line_w=0.75, radius=None, shadow=False):
    shp_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius is not None else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shp_type, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius is not None:
        s.adjustments[0] = radius
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(line_w)
    if not shadow:
        # kill the theme's default outer shadow
        spPr = s._element.spPr
        spPr.append(spPr.makeelement(qn("a:effectLst"), {}))
    s.text_frame.text = ""
    return s


def add_line(slide, x1, y1, x2, y2, *, color=G1, width=0.75, head=False, tail=False, dash=None):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if dash:
        ln.append(ln.makeelement(qn("a:prstDash"), {"val": dash}))
    if tail:
        ln.append(ln.makeelement(qn("a:headEnd"), {"type": "triangle", "w": "sm", "len": "sm"}))
    if head:
        ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "sm", "len": "sm"}))
    return c


def add_dot(slide, cx, cy, d, color):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2), Inches(cy - d / 2), Inches(d), Inches(d))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(color)
    s.line.fill.background()
    spPr = s._element.spPr
    spPr.append(spPr.makeelement(qn("a:effectLst"), {}))
    return s


def add_icon(slide, name, x, y, size):
    p = ICONS / f"{name}.png"
    if p.exists():
        slide.shapes.add_picture(str(p), Inches(x), Inches(y), Inches(size), Inches(size))
        return True
    return False


def crosshair(slide, cx, cy, arm=0.07, color=DG5):
    add_line(slide, cx - arm, cy, cx + arm, cy, color=color, width=0.5)
    add_line(slide, cx, cy - arm, cx, cy + arm, color=color, width=0.5)


def label(slide, x, y, w, text, *, color=G3, size=7.5, align=PP_ALIGN.LEFT):
    """Tracked uppercase mono label — the deck's section-header idiom."""
    return add_text(slide, x, y, w, 0.2, [(text.upper(), {})], size=size, color=color,
                    font=MONO, spacing=1.6, align=align)


def card(slide, x, y, w, h, *, accent, title, idx=None, icon=None, fill=DG1):
    """Dark card with a thin hairline border, an accent dot + mono index on
    the top-left, and the title. Returns the y where body copy may start."""
    add_rect(slide, x, y, w, h, fill=fill, line=DG4, radius=0.035)
    add_dot(slide, x + 0.2, y + 0.22, 0.09, accent)
    tx = x + 0.33
    if idx is not None:
        add_text(slide, tx, y + 0.13, 0.5, 0.2, [(idx, {})], size=8, color=accent, font=MONO, spacing=1.2)
        tx += 0.36
    add_text(slide, tx, y + 0.11, w - (tx - x) - 0.15, 0.25, [(title, {})], size=10.5, color=LG5, bold=True)
    if icon:
        add_icon(slide, icon, x + w - 0.42, y + 0.12, 0.26)
    return y + 0.42


def body(slide, x, y, w, h, lines, *, size=8.5, color=G4, line_spacing=1.08):
    """Bulleted body copy: each item is str or (str, accent_color) — colored
    lead phrase before ' — ' or ': ' when given."""
    paras = []
    for item in lines:
        if isinstance(item, tuple):
            text, acc = item
            for sep in (" — ", ": "):
                if sep in text:
                    head_, rest = text.split(sep, 1)
                    paras.append([("▸ ", {"color": DG5, "size": size - 1}),
                                  (head_, {"color": acc, "bold": True}),
                                  (sep + rest, {})])
                    break
            else:
                paras.append([("▸ ", {"color": DG5, "size": size - 1}), (text, {"color": acc})])
        else:
            paras.append([("▸ ", {"color": DG5, "size": size - 1}), (item, {})])
    return add_text(slide, x, y, w, h, paras, size=size, color=color, line_spacing=line_spacing)


def chip(slide, x, y, text, *, color=G4, border=DG5, w=None, size=7, fill=None):
    w = w or (0.085 * len(text) + 0.24)
    add_rect(slide, x, y, w, 0.24, fill=fill, line=border, line_w=0.6, radius=0.5)
    add_text(slide, x, y + 0.045, w, 0.18, [(text.upper(), {})], size=size, color=color, font=MONO,
             spacing=1.0, align=PP_ALIGN.CENTER)
    return x + w + 0.1


# --------------------------------------------------------------------------- #
# the slide
# --------------------------------------------------------------------------- #
def build(out: Path) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(W_IN)
    prs.slide_height = Inches(H_IN)
    s = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    bg = make_background(ASSETS / "onepager_bg.png")
    s.shapes.add_picture(str(bg), 0, 0, prs.slide_width, prs.slide_height)

    # corner registration marks
    for cx, cy in ((0.3, 0.3), (W_IN - 0.3, 0.3), (0.3, H_IN - 0.3), (W_IN - 0.3, H_IN - 0.3)):
        crosshair(s, cx, cy)

    # ----- header -----------------------------------------------------------
    label(s, 0.55, 0.38, 8, "Hexaware  ·  AmeriHealth Caritas  ·  AI-in-Engineering program  ·  agent 01 / 03", color=G2)
    add_text(s, 0.55, 0.58, 7.5, 0.5, [("FRD → STTM Agent", {"color": WHITE, "bold": True, "size": 24}),
                                       ("   system architecture", {"color": G3, "size": 14})], size=24)
    add_text(s, 0.55, 1.02, 8.3, 0.25,
             [("An approved Functional Requirements Document becomes an audited Source-to-Target Mapping workbook "
               "and a machine-readable feed contract — every fact checked against the source text, every ambiguity "
               "routed to a person.", {})], size=9.5, color=G4)
    # status chips (right, two rows, right-aligned to the margin)
    row1 = [("read-only SharePoint", INDIGO, INDIGO_DIM), ("1 model call / document", CERULEAN, CERULEAN_DIM)]
    row2 = [("verdicts computed in code", TURQ, TURQ_DIM), ("human-in-the-loop", GOLD, GOLD_DIM)]
    for yy, row in ((0.6, row1), (0.92, row2)):
        widths = [0.072 * len(t) + 0.26 for t, _, _ in row]
        cx = W_IN - 0.55 - sum(widths) - 0.1 * (len(row) - 1)
        for (t, col, bd), w in zip(row, widths):
            cx = chip(s, cx, yy, t, color=col, border=bd, w=w)

    # ----- main band geometry ----------------------------------------------
    TOP, H = 1.62, 3.55
    # columns: A source of record · B sync · C pipeline · D human · E downstream
    A = (0.55, 1.9)
    B = (2.72, 1.9)
    C = (4.89, 4.55)
    D = (9.71, 1.62)
    E = (11.58, 1.2)

    # column headers (tracked labels above cards)
    label(s, A[0], TOP - 0.27, A[1], "system of record", color=INDIGO)
    label(s, B[0], TOP - 0.27, B[1], "00 · continuous sync", color=TURQ)
    label(s, C[0], TOP - 0.27, C[1], "generation pipeline · one Databricks job, four stages", color=TURQ)
    label(s, D[0], TOP - 0.27, D[1], "human in the loop", color=GOLD)
    label(s, E[0], TOP - 0.27, E[1], "downstream", color=INDIGO)

    # --- A: SharePoint -------------------------------------------------------
    y = card(s, A[0], TOP, A[1], H, accent=INDIGO, title="SharePoint library", icon="folder")
    body(s, A[0] + 0.2, y, A[1] - 0.35, 2.0, [
        ("FRD folder — approved requirements documents (.docx / .pdf / .md)", INDIGO),
        ("STTM folder — approved mapping workbooks (.xlsx), uploaded by reviewers", INDIGO),
        "The only place anything is authored or approved. The agent never writes here "
        "(Entra ID grant: Sites.Selected, read).",
    ], size=8.2)
    add_rect(s, A[0] + 0.2, TOP + H - 0.92, A[1] - 0.4, 0.74, fill=DG2, line=DG4, radius=0.06)
    add_text(s, A[0] + 0.3, TOP + H - 0.86, A[1] - 0.6, 0.65,
             [[("every approved pair is a template", {"color": LG5, "bold": True})],
              [("exemplar for extraction · layout + dictionary for render · ground truth for eval", {})]],
             size=7.3, color=G4, line_spacing=1.05)

    # --- B: sync -------------------------------------------------------------
    y = card(s, B[0], TOP, B[1], H, accent=TURQ, title="Sync job", idx="00", icon="refresh")
    body(s, B[0] + 0.2, y, B[1] - 0.35, 2.2, [
        ("scheduled — every 15 min, or “Sync now” from the app", TURQ),
        ("incremental — Graph id · eTag · modified · size; only new or changed files move", TURQ),
        ("pairs — FRD ↔ STTM by exact name, then deterministic similarity; verdict + score per pair", TURQ),
        ("indexes — corpus_index.json, content hash per document", TURQ),
        "Zero model calls. Empty library warns; a Graph refusal fails loudly.",
    ], size=7.9)
    add_rect(s, B[0] + 0.2, TOP + H - 0.7, B[1] - 0.4, 0.52, fill=DG2, line=DG4, radius=0.06)
    add_text(s, B[0] + 0.3, TOP + H - 0.65, B[1] - 0.6, 0.45,
             [[("lands in Unity Catalog", {"color": BLUE, "bold": True})],
              [("frd_raw · sttm_reference", {"font": MONO, "size": 6.8}), ("  + index, manifest", {"size": 6.8})]],
             size=7.5, color=G4, line_spacing=1.05)

    # --- C: pipeline container -----------------------------------------------
    add_rect(s, C[0], TOP, C[1], H, fill=DG1, line=DG4, radius=0.02)
    add_text(s, C[0] + 0.2, TOP + 0.1, C[1] - 0.4, 0.22,
             [("frd_sttm_pipeline", {"font": MONO, "color": TURQ, "size": 8.5}),
              ("   Jobs API (the app) or bundle CLI · serverless · insulated runs", {"color": G3, "size": 7.3})], size=8)
    # four stage nodes
    sg = 0.12
    sw = (C[1] - 0.36 - 3 * sg) / 4
    sx0 = C[0] + 0.18
    sy = TOP + 0.42
    sh = 2.16
    stages = [
        ("01", "Ingest", TURQ, [
            "docx / pdf / md / txt → markdown",
            "labels from the versioned FRD label contract",
            "req-ids BR · REQ · FR · SRQ · NFR",
            "→ frd_documents",
        ]),
        ("02", "Extract", CERULEAN, [
            "THE one Claude call per document",
            "schema in prompt · streamed · client-side validation",
            "few-shot exemplars retrieved from the corpus",
            "→ extraction JSON",
        ]),
        ("03", "Contract", TURQ, [
            "validate (extra = forbid) → enrich",
            "grounding audit: strict fields verbatim; prose ≥ 0.75 overlap",
            "gate: attribution · disagreement · advisory",
            "→ contract JSON",
        ]),
        ("04", "Render", TURQ, [
            "template: single · amalgam · freeform (flagged)",
            "derive mappings; FRD rules on the rows they name",
            "eval vs the doc’s OWN approved STTM",
            "→ .sttm.xlsx · contract.v2 · report",
        ]),
    ]
    for i, (idx, title, acc, lines) in enumerate(stages):
        x = sx0 + i * (sw + sg)
        add_rect(s, x, sy, sw, sh, fill=DG2, line=DG4, radius=0.05)
        add_dot(s, x + 0.17, sy + 0.19, 0.08, acc)
        add_text(s, x + 0.28, sy + 0.1, 0.5, 0.2, [(idx, {})], size=8, color=acc, font=MONO, spacing=1.2)
        add_text(s, x + 0.12, sy + 0.3, sw - 0.2, 0.25, [(title, {})], size=10.5, color=LG5, bold=True)
        paras = [[(("▸ " if j else "■ "), {"color": (acc if j == 0 else DG5), "size": 6.5}),
                  (t, {"color": (LG5 if j == 0 else G4), "bold": j == 0})]
                 for j, t in enumerate(lines)]
        add_text(s, x + 0.12, sy + 0.6, sw - 0.2, sh - 0.66, paras, size=7.1, color=G4, line_spacing=1.05)
        if i < 3:
            add_line(s, x + sw + 0.01, sy + 0.36, x + sw + sg - 0.01, sy + 0.36, color=G2, width=0.9, head=True)
    # gate strip under stages
    gy = sy + sh + 0.12
    add_text(s, C[0] + 0.2, gy + 0.02, 1.7, 0.2, [("gate · computed in code", {})], size=7, color=G3, font=MONO, spacing=1.0)
    gx = C[0] + 1.9
    for txt, col in (("PASS", GREEN), ("PASS_WITH_FLAGS", GOLD), ("FAIL", RED)):
        gx = chip(s, gx, gy - 0.02, txt, color=col, border=DG5, size=6.5, w=0.068 * len(txt) + 0.22)
    add_text(s, C[0] + 0.2, gy + 0.32, C[1] - 0.4, 0.5,
             [[("grounding is the quality gate", {"color": LG5, "bold": True}),
               (" — a strict field the FRD does not contain verbatim fails the run; prose below the overlap threshold "
                "is flagged for review; nothing is relaxed to make a run pass. In 04, human resolutions apply first, "
                "the automatic cross-check second; every non-application is recorded with a reason.", {})]],
             size=7.3, color=G4, line_spacing=1.05)

    # --- D: human in the loop --------------------------------------------------
    y = card(s, D[0], TOP, D[1], H, accent=GOLD, title="Review App", icon="user")
    body(s, D[0] + 0.18, y, D[1] - 0.3, 2.5, [
        ("pick — unmapped FRDs → Generate (billed, confirm-gated); mapped FRDs → view approved STTM, or regenerate-anyway with auto golden-pair eval", GOLD),
        ("resolve — every gated item as a card: candidate pick · none of these · free text", GOLD),
        ("re-render — stage 04 only; nothing re-extracted, nothing billed", GOLD),
        ("approve — download, last edits, upload to the STTM folder; next sync pairs it", GOLD),
    ], size=7.8)
    add_rect(s, D[0] + 0.18, TOP + H - 0.62, D[1] - 0.36, 0.42, fill=DG2, line=GOLD_DIM, radius=0.06)
    add_text(s, D[0] + 0.26, TOP + H - 0.57, D[1] - 0.5, 0.35,
             [[("no publish button exists", {"color": GOLD, "bold": True})],
              [("the hand-off is a person, by design", {})]], size=7.2, color=G4, line_spacing=1.0)

    # --- E: downstream ---------------------------------------------------------
    add_rect(s, E[0], TOP, E[1], H, fill=DG1, line=DG4, radius=0.035)
    add_dot(s, E[0] + 0.18, TOP + 0.22, 0.09, INDIGO)
    add_text(s, E[0] + 0.3, TOP + 0.11, E[1] - 0.35, 0.25, [("Outputs", {})], size=10.5, color=LG5, bold=True)
    items = [
        ("STTM workbook", ".xlsx · client dialect", TURQ, "sheet"),
        ("Feed contract", "contract.json", TURQ, "braces"),
        ("Mapping contract", "via extract-sttm", INDIGO, "braces"),
        ("CodeGen", "agent 02 / 03", INDIGO, "code"),
        ("Code Review", "agent 03 / 03", INDIGO, "check"),
    ]
    iy = TOP + 0.5
    for j, (t, sub, col, icon) in enumerate(items):
        add_rect(s, E[0] + 0.1, iy, E[1] - 0.2, 0.5, fill=DG2, line=(INDIGO_DIM if col == INDIGO else DG4), radius=0.08)
        add_icon(s, icon, E[0] + 0.16, iy + 0.14, 0.2)
        add_text(s, E[0] + 0.4, iy + 0.06, E[1] - 0.5, 0.42,
                 [[(t, {"color": LG5, "bold": True, "size": 7.4})], [(sub, {"size": 6.5, "color": G3})]],
                 size=7.3, line_spacing=1.0)
        if j < len(items) - 1:
            add_line(s, E[0] + E[1] / 2, iy + 0.5, E[0] + E[1] / 2, iy + 0.6, color=G1, width=0.75, head=(j >= 1))
        iy += 0.6

    # ----- connectors between columns ------------------------------------------
    mid = TOP + 1.35
    add_line(s, A[0] + A[1], mid, B[0], mid, color=G2, width=1.0, head=True)            # SharePoint → sync
    add_line(s, B[0] + B[1], mid, C[0], mid, color=G2, width=1.0, head=True)            # sync → pipeline
    add_line(s, C[0] + C[1], mid, D[0], mid, color=G2, width=1.0, head=True)            # pipeline → review
    add_line(s, D[0] + D[1], mid, E[0], mid, color=G2, width=1.0, head=True)            # review → outputs
    # labels on connectors
    add_text(s, A[0] + A[1] + 0.02, mid - 0.22, B[0] - A[0] - A[1] - 0.04, 0.2, [("read", {})], size=6.5, color=G2, font=MONO, align=PP_ALIGN.CENTER)
    add_text(s, B[0] + B[1] + 0.02, mid - 0.22, C[0] - B[0] - B[1] - 0.04, 0.2, [("frd", {})], size=6.5, color=G2, font=MONO, align=PP_ALIGN.CENTER)
    add_text(s, C[0] + C[1] + 0.02, mid - 0.22, D[0] - C[0] - C[1] - 0.04, 0.2, [("gated", {})], size=6.5, color=G2, font=MONO, align=PP_ALIGN.CENTER)
    add_text(s, D[0] + D[1] + 0.02, mid - 0.22, E[0] - D[0] - D[1] - 0.04, 0.2, [("ok", {})], size=6.5, color=G2, font=MONO, align=PP_ALIGN.CENTER)
    # the human return loop: review → (reviewer uploads) → SharePoint — dashed gold, routed under the band
    ly = TOP + H + 0.13
    add_line(s, D[0] + 0.5, TOP + H, D[0] + 0.5, ly, color=GOLD_DIM, width=1.0, dash="dash")
    add_line(s, D[0] + 0.5, ly, A[0] + 0.5, ly, color=GOLD_DIM, width=1.0, dash="dash")
    add_line(s, A[0] + 0.5, ly, A[0] + 0.5, TOP + H, color=GOLD_DIM, width=1.0, dash="dash", head=True)
    add_text(s, 3.3, ly + 0.03, 7.6, 0.2,
             [("reviewer uploads the approved STTM  →  the next sync pulls it back and pairs it  →  it becomes a template", {})],
             size=6.8, color=GOLD, font=MONO, align=PP_ALIGN.CENTER)
    # app ↔ job (trigger) small annotation
    add_line(s, D[0] + 0.9, TOP, D[0] + 0.9, TOP - 0.06, color=DG5, width=0.5)

    # ----- data layer -------------------------------------------------------------
    DY, DH = 5.58, 0.74
    add_rect(s, B[0], DY, (D[0] + D[1]) - B[0], DH, fill=DG1, line=BLUE_DIM, radius=0.02)
    add_dot(s, B[0] + 0.2, DY + 0.2, 0.09, BLUE)
    add_text(s, B[0] + 0.33, DY + 0.1, 5.5, 0.22, [("Unity Catalog", {"color": LG5, "bold": True, "size": 10}),
                                                   ("  working store · the app mirrors it, never owns it", {"color": G3, "size": 7.5})], size=10)
    add_icon(s, "database", D[0] + D[1] - 0.45, DY + 0.12, 0.24)
    cols = [
        ("volume  frd_raw", "source FRDs — the sync lands them here", BLUE),
        ("volume  sttm_reference", "approved STTMs · corpus_index.json · sync_manifest.json", BLUE),
        ("volume  sttm_out_app/<run>", "extractions · contracts · rendered · reports", BLUE),
        ("delta  frd_documents · frd_contracts · frd_sttm_runs", "stage summaries per run", BLUE),
    ]
    cw = ((D[0] + D[1]) - B[0] - 0.4) / 4
    for j, (t, sub, col) in enumerate(cols):
        x = B[0] + 0.2 + j * cw
        add_text(s, x, DY + 0.36, cw - 0.12, 0.4,
                 [[(t, {"font": MONO, "color": BLUE, "size": 7})], [(sub, {"size": 6.8, "color": G4})]],
                 size=7, line_spacing=1.0)

    # ----- footer: doctrine + legend -------------------------------------------------
    FY = 6.5
    add_line(s, 0.55, FY, W_IN - 0.55, FY, color=DG4, width=0.5)
    add_text(s, 0.55, FY + 0.1, 6.6, 0.5,
             [[("the model proposes", {"color": CERULEAN, "bold": True}), ("   ·   ", {"color": DG5}),
               ("deterministic code audits and decides", {"color": TURQ, "bold": True}), ("   ·   ", {"color": DG5}),
               ("a person resolves", {"color": GOLD, "bold": True})],
              [("Grounding checks, gate status, pairing verdicts, template mode and eval percentages are all computed in "
                "code; the model makes exactly one call per document. Fail loud in both directions — no silent mock, no "
                "silent live, no fallback to a stale copy.", {"size": 7.5, "color": G4})]],
             size=9, line_spacing=1.1)
    # legend
    lx = 7.55
    add_text(s, lx, FY + 0.1, 0.7, 0.2, [("legend", {})], size=7, color=G2, font=MONO, spacing=1.4)
    for k, (name, col) in enumerate((("LLM", CERULEAN), ("deterministic code", TURQ), ("human", GOLD),
                                     ("external system", INDIGO), ("storage", BLUE))):
        row, colm = divmod(k, 3)
        x = lx + 0.7 + colm * 1.55
        yy = FY + 0.12 + row * 0.22
        add_dot(s, x + 0.05, yy + 0.08, 0.08, col)
        add_text(s, x + 0.16, yy, 1.4, 0.2, [(name, {})], size=7.5, color=G4)
    # version tag
    add_text(s, W_IN - 5.0, FY + 0.62, 4.45, 0.2,
             [("staging · 2026-08-22 · verified offline · live deploy pending", {})], size=6.5, color=G1, font=MONO,
             align=PP_ALIGN.RIGHT)

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else OUT_DEFAULT
    print(build(target))
