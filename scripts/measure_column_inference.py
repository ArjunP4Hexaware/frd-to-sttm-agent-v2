"""Measure whether a Claude call can infer TARGET COLUMN NAMES for a real feed.

    ~/.virtualenvs/frdsttm/bin/python scripts/measure_column_inference.py [--dry-run]
    ~/.virtualenvs/frdsttm/bin/python scripts/measure_column_inference.py --provider databricks

`--provider databricks` needs NO Anthropic key: it goes through the workspace's
own Foundation Model APIs on the workspace credential, the same seam
`frdsttm.live_extraction.build_live_client` gives stage 02. The model is the
same Claude; what changes is the front door, the credential and WHO IS BILLED
(the Databricks workspace, not an Anthropic account).

THE QUESTION (Arjun, 2026-08-27). `contracts/naming_standards.json` records that
"TPL_" + UPPER_SNAKE(source) reproduces 22 of CAQH's 115 target column names —
a STRING RULE fails. That measurement says nothing about whether a MODEL can do
it, because the names are not a transformation of the source: MEME, SBSB and GRP
are Facets column names, and the CAQH FRD names Facets as the incumbent. A
frontier model plausibly knows that vocabulary. This script replaces the
assertion with a number.

WHAT IS MEASURED, AND WHY IT IS NOT JUST ACCURACY
-------------------------------------------------
The agent's goal is NOT a perfect STTM. It is to carry the FRD and the VDD as
far as they go, IMPLEMENT what it can settle, SURFACE what it cannot, and leave
the rest to the reviewer after download. Under that goal a 60% hit rate whose
misses are flagged beats a 75% hit rate that is confident about everything —
because the first can be shipped with 40% gated and the second poisons the
workbook silently.

So every condition reports two numbers:

  accuracy     exact matches / rows attempted
  CALIBRATION  accuracy WITHIN each confidence band

Calibration is the one that decides whether this is usable. A band that is
~100% correct is a band the agent can fill in; anything below is a band it must
gate. A model that cannot separate the two is not usable at any accuracy.

THE FOUR CONDITIONS
-------------------
  1 blind      source name/type/length/description/segment only
  2 facets     + told the incumbent warehouse is Facets
  3 fewshot    + the 8 HEADER rows with their real answers, then predict the
               101 DETAIL rows. This is the one to watch: it is not invention,
               it is generalising a vocabulary from examples, and it maps to a
               real case — a NEW feed in an ALREADY-MAPPED family.
  4 catalog    + the list of real target names to choose from. This simulates
               the term catalog (information_schema over PR_STD/PR_DLK). It is
               EASIER THAN REALITY — the real catalog has thousands of columns,
               not 115 — so read it as an upper bound on matching, not a
               forecast.

Ground truth is the CAQH STTM's own stage band. Nothing is written back and no
contract is modified: this script only prints.
"""

from __future__ import annotations

import json
import os
import re
import sys
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
STTM = REPO / "sample_documents" / "STTM_STG_STD_PaymentIntegrity_TPL_CAQH_To_DL_Ingestion_1005034.xlsx"

MODEL = "claude-opus-5"
PRICE_IN, PRICE_OUT = 5.00, 25.00        # $ per 1M tokens, Opus 5 (first-party
#                                          list price; the Databricks route is
#                                          billed by the workspace, so the
#                                          dollar figures printed under
#                                          --provider databricks are an
#                                          ANTHROPIC-EQUIVALENT estimate, not
#                                          the invoice.

SCHEMA = {
    "type": "object",
    "properties": {
        "predictions": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "source": {"type": "string"},
                    "target": {"type": "string"},
                    "confidence": {"type": "string", "enum": ["high", "medium", "low"]},
                },
                "required": ["source", "target", "confidence"],
                "additionalProperties": False,
            },
        }
    },
    "required": ["predictions"],
    "additionalProperties": False,
}

JSON_INSTRUCTION = """

Reply with JSON ONLY — no prose, no code fence — in exactly this shape:
{"predictions": [{"source": "<the source field name, copied verbatim>",
                  "target": "<the target column name>",
                  "confidence": "high" | "medium" | "low"}, ...]}
One entry per source field, in the order given."""

BASE = """You are naming columns in a healthcare payer data warehouse.

For each SOURCE field below, give the TARGET column name it would be loaded
into in the stage layer.

Return one prediction per source field, with a confidence:
  high   - you are confident this is the exact name the warehouse uses
  medium - plausible, but you would not bet the load on it
  low    - a guess

Be honest with confidence. A wrong name that is marked high is worse than a
right name marked low: downstream, a high-confidence name is written into the
workbook without review, and a name that does not match an existing column
breaks the load."""

FACETS = """
The warehouse's existing columns follow the naming conventions of Facets
(the TriZetto/Cognizant payer administration system), which is the incumbent
system for this data."""


def read_ground_truth() -> list[dict]:
    from openpyxl import load_workbook
    wb = load_workbook(str(STTM), read_only=True, data_only=True)
    rows = [list(r) for r in wb["caqh"].iter_rows(values_only=True)]
    hdr = next(i for i, r in enumerate(rows)
               if any(str(c).strip() == "Field Name" for c in r if c))
    cols = {str(c).strip().split("\n")[0]: j for j, c in enumerate(rows[hdr]) if c}
    out = []
    for r in rows[hdr + 1:]:
        name = r[cols["Field Name"]] if cols.get("Field Name") is not None else None
        target = r[cols["ColumnName"]] if cols.get("ColumnName") is not None else None
        if not name or not target:
            continue
        out.append({
            "source": str(name).strip(),
            "target": str(target).strip(),
            "datatype": str(r[cols["Data Type"]] or "").strip(),
            "length": str(r[cols["Length"]] or "").strip(),
            "segment": str(r[cols["Segment"]] or "").strip(),
            "description": " ".join(str(r[cols["Comments"]] or "").split())[:180],
        })
    wb.close()
    return out


def field_lines(rows: list[dict]) -> str:
    return "\n".join(
        f"{i + 1}. {r['source']} | {r['datatype']} | len {r['length'] or '-'} "
        f"| {r['segment']} | {r['description']}"
        for i, r in enumerate(rows)
    )


def build(condition: str, rows: list[dict]) -> tuple[str, list[dict]]:
    """(prompt, rows the model must answer)."""
    header = [r for r in rows if r["segment"].lower() == "header"]
    detail = [r for r in rows if r["segment"].lower() == "detail"]

    if condition == "blind":
        return BASE + "\n\nSOURCE FIELDS:\n" + field_lines(rows), rows
    if condition == "facets":
        return BASE + FACETS + "\n\nSOURCE FIELDS:\n" + field_lines(rows), rows
    if condition == "fewshot":
        examples = "\n".join(f"  {r['source']}  ->  {r['target']}" for r in header)
        return (BASE + FACETS +
                "\n\nHere are the names this same feed already uses for its HEADER "
                "record. They are the ground truth for that segment:\n" + examples +
                "\n\nNow name the DETAIL record's fields, following the same "
                "conventions.\n\nSOURCE FIELDS:\n" + field_lines(detail)), detail
    if condition == "catalog":
        catalog = sorted({r["target"] for r in rows})
        examples = "\n".join(f"  {r['source']}  ->  {r['target']}" for r in header)
        return (BASE + FACETS +
                "\n\nHEADER record names already in use (ground truth):\n" + examples +
                "\n\nThe warehouse contains exactly these column names. Choose the "
                "one each source field maps to — do not invent a name that is not "
                "on this list:\n" + "\n".join(f"  {c}" for c in catalog) +
                "\n\nSOURCE FIELDS:\n" + field_lines(detail)), detail
    raise SystemExit(f"unknown condition {condition!r}")


def parse_predictions(text: str) -> list[dict]:
    """Pull the predictions array out of a reply, fence or no fence.

    Defensive because the Databricks route has no schema enforcement: a run
    that costs money must not be thrown away over a stray ```json wrapper or a
    sentence before the object.
    """
    t = text.strip()
    if t.startswith("```"):
        t = re.sub(r"^```[a-zA-Z]*\s*", "", t)
        t = re.sub(r"```\s*$", "", t).strip()
    try:
        return json.loads(t)["predictions"]
    except Exception:  # noqa: BLE001 — fall through to a brace scan
        pass
    start = t.find("{")
    while start != -1:
        depth, i = 0, start
        while i < len(t):
            depth += t[i] == "{"
            depth -= t[i] == "}"
            if depth == 0:
                try:
                    return json.loads(t[start:i + 1])["predictions"]
                except Exception:  # noqa: BLE001
                    break
            i += 1
        start = t.find("{", start + 1)
    return []


def score(preds: list[dict], truth: list[dict]) -> dict:
    by_source = {r["source"].strip().lower(): r["target"] for r in truth}
    bands: dict[str, list[bool]] = {"high": [], "medium": [], "low": []}
    hits = 0
    unmatched = 0
    for p in preds:
        want = by_source.get(str(p.get("source", "")).strip().lower())
        if want is None:
            unmatched += 1
            continue
        ok = str(p.get("target", "")).strip().upper() == want.strip().upper()
        hits += ok
        bands.setdefault(p.get("confidence", "low"), []).append(ok)
    return {
        "n": len(truth), "answered": len(preds) - unmatched, "hits": hits,
        "accuracy": hits / len(truth) if truth else 0.0,
        "bands": {k: (sum(v), len(v)) for k, v in bands.items() if v},
        "unmatched": unmatched,
    }


def main() -> int:
    dry = "--dry-run" in sys.argv
    rows = read_ground_truth()
    print(f"ground truth: {len(rows)} CAQH source fields "
          f"({sum(1 for r in rows if r['segment'].lower()=='header')} header / "
          f"{sum(1 for r in rows if r['segment'].lower()=='detail')} detail / "
          f"{sum(1 for r in rows if r['segment'].lower()=='trailer')} trailer)\n")

    conditions = ["blind", "facets", "fewshot", "catalog"]
    if dry:
        for c in conditions:
            prompt, target = build(c, rows)
            print(f"  {c:9s} prompt {len(prompt):6d} chars (~{len(prompt)//4:5d} tok), "
                  f"{len(target):3d} rows to answer")
        print("\n--dry-run: nothing sent, nothing billed.")
        return 0

    provider = "databricks" if "--provider" in sys.argv and \
        sys.argv[sys.argv.index("--provider") + 1] == "databricks" else "anthropic"
    sys.path.insert(0, str(REPO / "src"))
    from frdsttm.live_extraction import build_live_client, databricks_model_name
    client = build_live_client(provider, max_retries=2)
    model = databricks_model_name(MODEL) if provider == "databricks" else MODEL
    print(f"provider: {provider}   model: {model}\n")
    spend = 0.0
    results = {}

    for c in conditions:
        prompt, target = build(c, rows)
        kwargs = dict(model=model, max_tokens=16000,
                      messages=[{"role": "user", "content": prompt + JSON_INSTRUCTION}])
        if provider != "databricks":
            # The Databricks Foundation Model APIs proxy accepts the Messages
            # API shape but rejects output_config ("Extra inputs are not
            # permitted", 2026-08-27), so structured outputs are first-party
            # only. On that route the JSON contract lives in the prompt and
            # `parse_predictions` does the defensive parsing.
            kwargs["thinking"] = {"type": "adaptive"}
            kwargs["output_config"] = {"format": {"type": "json_schema", "schema": SCHEMA}}
        resp = client.messages.create(**kwargs)
        text = "".join(b.text for b in resp.content if b.type == "text")
        preds = parse_predictions(text)
        if not preds:
            print(f"  {c}: no predictions parsed; raw head: {text[:220]!r}")
        s = score(preds, target)
        cost = (resp.usage.input_tokens * PRICE_IN
                + resp.usage.output_tokens * PRICE_OUT) / 1_000_000
        spend += cost
        results[c] = (s, cost, resp.usage)
        print(f"  {c:9s} {s['hits']:3d}/{s['n']:3d} = {s['accuracy']*100:5.1f}%   "
              f"bands " + " ".join(f"{k}:{h}/{n}" for k, (h, n) in s["bands"].items())
              + f"   ${cost:.3f}")

    print(f"\n{'='*72}\nCALIBRATION — the number that decides usability\n{'='*72}")
    for c, (s, _cost, _u) in results.items():
        line = []
        for band in ("high", "medium", "low"):
            if band in s["bands"]:
                h, n = s["bands"][band]
                line.append(f"{band} {h}/{n} = {h/n*100:.0f}%")
        print(f"  {c:9s} " + "   ".join(line))
    print(f"\ntotal spend: ${spend:.2f}")
    print("\nRead it this way: a confidence band that is ~100% correct is a band the\n"
          "agent may FILL IN. Anything below that is a band it must GATE — which is\n"
          "still a win, because a named gap is what the reviewer finishes after\n"
          "download. What would make this unusable is a high band that is wrong.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
