"""Build the FRD->STTM client presentation deck (PPTX).

Regenerates docs/FRD_to_STTM_Agent_Demo.pptx -- the 13-slide deck used
alongside the live demo (docs/DEMO_RUNBOOK.md). Every figure quoted on the
slides comes from docs/LIVE_E2E_2026-08-07.md; update that document first,
then this script, so the two never drift.

    pip install python-pptx
    python scripts/build_presentation.py docs/FRD_to_STTM_Agent_Demo.pptx

Layout is absolute-positioned (no template master), so text length drives
fit: content must stay above y=6.80in and card bodies are sized for the copy
they currently hold. Shorten copy rather than growing cards when editing.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import sys

OUT = sys.argv[1] if len(sys.argv) > 1 else "FRD_to_STTM_Agent.pptx"

# ---------------------------------------------------------------- palette
NAVY   = RGBColor(0x0E, 0x22, 0x33)
INK    = RGBColor(0x1A, 0x24, 0x33)
BODY   = RGBColor(0x3A, 0x4A, 0x5C)
MUTED  = RGBColor(0x6B, 0x7A, 0x8C)
TEAL   = RGBColor(0x00, 0xA1, 0x9A)
BLUE   = RGBColor(0x2E, 0x6E, 0xD1)
AMBER  = RGBColor(0xD9, 0x8A, 0x1F)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
PAPER  = RGBColor(0xF4, 0xF7, 0xF9)
LINE   = RGBColor(0xD8, 0xE0, 0xE7)
NAVY_L = RGBColor(0x1B, 0x3A, 0x52)
PALE   = RGBColor(0xC3, 0xD2, 0xDE)
PALER  = RGBColor(0x9F, 0xB4, 0xC4)

FONT = "Calibri"
FONT_H = "Calibri Light"

W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.85)
CONTENT_W = W - 2 * MARGIN
BOTTOM = Inches(6.80)          # content must not cross this

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]


# ---------------------------------------------------------------- helpers
def slide():
    return prs.slides.add_slide(BLANK)


def rect(s, x, y, w, h, fill=None, line=None, lw=1.0, shape=MSO_SHAPE.RECTANGLE,
         adj=None):
    sh = s.shapes.add_shape(shape, x, y, w, h)
    if adj is not None:
        try:
            sh.adjustments[0] = adj
        except Exception:
            pass
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    if sh.has_text_frame:
        sh.text_frame.word_wrap = True
    return sh


def tbox(s, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    return tf


def para(tf, text, size=16, color=BODY, bold=False, font=FONT, align=PP_ALIGN.LEFT,
         space_before=0, space_after=6, first=False, line_spacing=1.15,
         italic=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.line_spacing = line_spacing
    r = p.add_run()
    r.text = text
    f = r.font
    f.name, f.size, f.bold, f.italic = font, Pt(size), bold, italic
    f.color.rgb = color
    return p


def rich(tf, chunks, size=16, align=PP_ALIGN.LEFT, space_before=0, space_after=6,
         first=False, line_spacing=1.2):
    """chunks: list of (text, color, bold)."""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.line_spacing = line_spacing
    for text, color, bold in chunks:
        r = p.add_run()
        r.text = text
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
    return p


def header(s, kicker, title, sub=None):
    """Standard content-slide header. Returns y where content may start."""
    rect(s, Inches(0), Inches(0), W, H, fill=WHITE)
    rect(s, Inches(0), Inches(0), Inches(0.16), H, fill=TEAL)
    tf = tbox(s, MARGIN, Inches(0.50), CONTENT_W, Inches(0.26))
    para(tf, kicker, size=11.5, color=TEAL, bold=True, first=True, space_after=0)
    tf2 = tbox(s, MARGIN, Inches(0.83), CONTENT_W, Inches(0.55))
    para(tf2, title, size=32, color=INK, font=FONT_H, first=True, space_after=0)
    if sub:
        tf3 = tbox(s, MARGIN, Inches(1.45), CONTENT_W - Inches(1.2), Inches(0.5))
        para(tf3, sub, size=14, color=MUTED, first=True, space_after=0,
             line_spacing=1.2)
        rect(s, MARGIN, Inches(2.10), Inches(1.1), Pt(2.5), fill=LINE)
        return Inches(2.30)
    rect(s, MARGIN, Inches(1.50), Inches(1.1), Pt(2.5), fill=LINE)
    return Inches(1.74)


def footer(s, n, label="FRD → STTM Agent"):
    tf = tbox(s, MARGIN, H - Inches(0.50), CONTENT_W, Inches(0.24))
    p = tf.paragraphs[0]
    p.space_after = 0
    r = p.add_run()
    r.text = label
    r.font.name, r.font.size = FONT, Pt(9.5)
    r.font.color.rgb = MUTED
    tf2 = tbox(s, W - MARGIN - Inches(1.0), H - Inches(0.50), Inches(1.0), Inches(0.24))
    p2 = tf2.paragraphs[0]
    p2.alignment = PP_ALIGN.RIGHT
    p2.space_after = 0
    r2 = p2.add_run()
    r2.text = str(n)
    r2.font.name, r2.font.size, r2.font.bold = FONT, Pt(9.5), True
    r2.font.color.rgb = MUTED


def card(s, x, y, w, h, title, lines, accent=TEAL, title_size=15, body_size=12.5,
         fill=PAPER, num=None, spacing=8):
    rect(s, x, y, w, h, fill=fill, line=LINE, lw=0.75)
    rect(s, x, y, Pt(3.2), h, fill=accent)
    ty = y + Inches(0.24)
    if num:
        tfn = tbox(s, x + Inches(0.28), ty, w - Inches(0.5), Inches(0.24))
        para(tfn, num, size=11, color=accent, bold=True, first=True, space_after=0)
        ty += Inches(0.26)
    tf = tbox(s, x + Inches(0.28), ty, w - Inches(0.56), Inches(0.32))
    para(tf, title, size=title_size, color=INK, bold=True, first=True, space_after=0)
    tf2 = tbox(s, x + Inches(0.28), ty + Inches(0.36), w - Inches(0.56),
               h - (ty - y) - Inches(0.55))
    for i, ln in enumerate(lines):
        para(tf2, ln, size=body_size, color=BODY, first=(i == 0),
             space_after=spacing, line_spacing=1.16)


def metric(s, x, y, w, h, value, label, note=None, accent=TEAL, vsize=32):
    rect(s, x, y, w, h, fill=WHITE, line=LINE, lw=0.75)
    rect(s, x, y, w, Pt(3.2), fill=accent)
    tf = tbox(s, x + Inches(0.22), y + Inches(0.24), w - Inches(0.4), Inches(0.5))
    para(tf, value, size=vsize, color=INK, bold=True, font=FONT_H, first=True,
         space_after=0)
    tf2 = tbox(s, x + Inches(0.22), y + Inches(0.86), w - Inches(0.4), Inches(0.9))
    para(tf2, label, size=11, color=accent, bold=True, first=True, space_after=3)
    if note:
        para(tf2, note, size=10.5, color=MUTED, line_spacing=1.16)


def arrow(s, x, y, w, h, color=LINE):
    return rect(s, x, y, w, h, fill=color, shape=MSO_SHAPE.RIGHT_ARROW)


def notes(s, text):
    tf = s.notes_slide.notes_text_frame
    tf.text = text
    for p in tf.paragraphs:
        for r in p.runs:
            r.font.size = Pt(12)
            r.font.name = FONT


# ================================================================ 1. TITLE
s = slide()
rect(s, Inches(0), Inches(0), W, H, fill=NAVY)
rect(s, Inches(0), Inches(0), Inches(0.16), H, fill=TEAL)
rect(s, W - Inches(4.6), Inches(0), Inches(4.6), H, fill=NAVY_L)
rect(s, W - Inches(4.6), Inches(0), Pt(2), H, fill=TEAL)

tf = tbox(s, MARGIN, Inches(1.65), Inches(7.6), Inches(0.3))
para(tf, "AI IN ENGINEERING PROGRAM   ·   AGENT 2 OF 5", size=12, color=TEAL,
     bold=True, first=True, space_after=0)

tf = tbox(s, MARGIN, Inches(2.20), Inches(7.7), Inches(1.9))
para(tf, "FRD → STTM Agent", size=50, color=WHITE, font=FONT_H, first=True,
     space_after=6, line_spacing=1.0)
para(tf, "From an approved Functional Requirements Document to a governed\n"
         "Source-to-Target Mapping workbook — in about thirty seconds.",
     size=16.5, color=PALE, space_before=8, line_spacing=1.3)

rect(s, MARGIN, Inches(4.55), Inches(1.3), Pt(2.5), fill=TEAL)

tf = tbox(s, MARGIN, Inches(4.95), Inches(7.6), Inches(0.9))
para(tf, "Prepared for AmeriHealth   ·   Hexaware", size=14.5, color=WHITE,
     bold=True, first=True, space_after=5)
para(tf, "Presented by Soham & Arjun   |   Live demo included", size=13,
     color=PALER)

hy = Inches(2.20)
for val, lbl in [("94.1%", "cell-level mapping fidelity\nvs. the golden reference"),
                 ("0 / 50", "ungrounded identifiers\nin the hallucination sweep"),
                 ("~33 s", "full pipeline, one billed\nmodel call (about $0.15)")]:
    tf = tbox(s, W - Inches(4.10), hy, Inches(3.2), Inches(1.1))
    para(tf, val, size=27, color=TEAL, bold=True, font=FONT_H, first=True,
         space_after=3)
    for ln in lbl.split("\n"):
        para(tf, ln, size=11.5, color=PALE, line_spacing=1.15, space_after=0)
    hy += Inches(1.30)

tf = tbox(s, W - Inches(4.10), Inches(6.35), Inches(3.2), Inches(0.4))
para(tf, "Measured on the anonymized demo pair — live\nend-to-end run, 7 August 2026",
     size=9.5, color=RGBColor(0x7E, 0x96, 0xA8), first=True, italic=True,
     line_spacing=1.2, space_after=0)

# ================================================================ 2. AGENDA
s = slide()
y = header(s, "PROJECT AGENDA", "What we'll cover today",
           "About thirty minutes, with the live agent run in the middle and open Q&A at the end.")

items = [
    ("01", "Where this agent fits", "The five-agent programme, and the hand-offs on either side of us.", TEAL),
    ("02", "The problem we're solving", "Mapping work today: slow, inconsistent, and hard to audit after the fact.", TEAL),
    ("03", "How the agent works", "Four stages, and the split between what the model does and what code does.", BLUE),
    ("04", "Trust architecture", "Verbatim grounding audit, ambiguity gating, human-in-the-loop review.", BLUE),
    ("05", "Live demo", "An FRD in, an STTM workbook out — run on the spot, with an offline backup.", AMBER),
    ("06", "Proof & value for AmeriHealth", "Measured results from the live run, and what they mean for delivery.", AMBER),
    ("07", "Status, roadmap & Q&A", "What is proven, what is still port work, and your questions.", MUTED),
]

col_w = Inches(5.70)
gap = Inches(0.60)
row_h = Inches(1.10)
for i, (num, title, sub, acc) in enumerate(items):
    cx = MARGIN + (col_w + gap) * (i // 4)
    cy = y + Inches(0.06) + row_h * (i % 4)
    tfn = tbox(s, cx, cy + Inches(0.01), Inches(0.58), Inches(0.36))
    para(tfn, num, size=21, color=acc, bold=True, font=FONT_H, first=True,
         space_after=0)
    tf = tbox(s, cx + Inches(0.60), cy, col_w - Inches(0.70), Inches(0.32))
    para(tf, title, size=16.5, color=INK, bold=True, first=True, space_after=3)
    para(tf, sub, size=12.5, color=BODY, line_spacing=1.16, space_after=0)
    rect(s, cx, cy + row_h - Inches(0.22), col_w - Inches(0.25), Pt(0.9), fill=LINE)

footer(s, 2)

# ================================================================ 3. PROGRAM MAP
s = slide()
y = header(s, "WHERE THIS FITS", "Second agent in a five-agent programme",
           "Each agent consumes the previous one's approved output through a versioned, machine-readable contract.")

chain = [
    ("BRD → FRD", "Business requirements\nto functional spec", False),
    ("FRD → STTM", "Mapping workbook\n+ feed contract", True),
    ("CodeGen", "Pipeline code from\nthe mapping contract", False),
    ("Code Review", "Automated review of\nthe generated code", False),
]
bw, bh = Inches(2.62), Inches(1.52)
bx, by = MARGIN, y + Inches(0.10)
for i, (name, desc, active) in enumerate(chain):
    fill = TEAL if active else PAPER
    txt = WHITE if active else INK
    sub = RGBColor(0xD5, 0xEE, 0xEC) if active else BODY
    rect(s, bx, by, bw, bh, fill=fill, line=None if active else LINE, lw=0.75)
    tf = tbox(s, bx + Inches(0.22), by + Inches(0.24), bw - Inches(0.44), Inches(1.0))
    if active:
        para(tf, "YOU ARE HERE", size=9.5, color=RGBColor(0xC8, 0xEC, 0xE9),
             bold=True, first=True, space_after=3)
        para(tf, name, size=18, color=txt, bold=True, font=FONT_H, space_after=4)
    else:
        para(tf, name, size=18, color=txt, bold=True, font=FONT_H, first=True,
             space_after=4)
    for ln in desc.split("\n"):
        para(tf, ln, size=11.5, color=sub, space_after=0, line_spacing=1.15)
    if i < 3:
        arrow(s, bx + bw + Inches(0.10), by + Inches(0.62), Inches(0.34),
              Inches(0.26))
    bx += bw + Inches(0.54)

tf = tbox(s, MARGIN, by + bh + Inches(0.18), CONTENT_W, Inches(0.26))
para(tf, "SQL Optimization runs as a standalone fifth agent alongside this chain.",
     size=12, color=MUTED, first=True, italic=True, space_after=0)

cy = by + bh + Inches(0.58)
cw = (CONTENT_W - Inches(0.5)) / 2
card(s, MARGIN, cy, cw, Inches(2.28), "Upstream — solved, and versioned", [
    "This agent and the BRD→FRD agent share one versioned label contract, "
    "committed byte-identically to both repos.",
    "Section headings, requirement-ID families and placeholders are read from "
    "that file, never hardcoded.",
    "A round-trip test suite fails the build if the two copies drift apart.",
], accent=TEAL, body_size=12.3)
card(s, MARGIN + cw + Inches(0.5), cy, cw, Inches(2.28),
     "Downstream — the gap we're flagging", [
    "This agent emits the FRD feed contract that CodeGen consumes. That half "
    "works today.",
    "CodeGen also needs a workbook-derived mapping contract, and no committed "
    "tool produces it yet.",
    "Closing this is the highest-value next step for the programme — and we are "
    "not hiding it.",
], accent=AMBER, body_size=12.3)
footer(s, 3)

# ================================================================ 4. PROBLEM
s = slide()
y = header(s, "THE PROBLEM", "Source-to-target mapping is the bottleneck",
           "The FRD is approved and the target model is known — yet the mapping workbook is still assembled by hand, one cell at a time.")

left_w = Inches(6.15)
items = [
    ("Slow", "An analyst reads the whole FRD, cross-references the data dictionary, and hand-builds every row of the workbook. Days of effort per document."),
    ("Inconsistent", "Two analysts reading the same FRD produce two different workbooks. Naming, load strategies and rule attribution drift between feeds and teams."),
    ("Hard to audit", "Six months later, nobody can say which sentence in the FRD justified a given mapping cell. Reconstructing that is manual archaeology."),
    ("A blocker for everything downstream", "Code generation, testing and review all wait on the workbook. The slowest manual step sets the pace for the whole delivery chain."),
]
cy = y + Inches(0.10)
for title, sub in items:
    rect(s, MARGIN, cy + Inches(0.09), Inches(0.10), Inches(0.10), fill=AMBER,
         shape=MSO_SHAPE.OVAL)
    tf = tbox(s, MARGIN + Inches(0.30), cy, left_w - Inches(0.35), Inches(0.3))
    para(tf, title, size=16.5, color=INK, bold=True, first=True, space_after=3)
    para(tf, sub, size=13, color=BODY, line_spacing=1.18, space_after=0)
    cy += Inches(1.08)

px = MARGIN + left_w + Inches(0.55)
pw = CONTENT_W - left_w - Inches(0.55)
ph = Inches(4.32)
rect(s, px, y + Inches(0.10), pw, ph, fill=NAVY)
rect(s, px, y + Inches(0.10), pw, Pt(3.2), fill=TEAL)
tf = tbox(s, px + Inches(0.40), y + Inches(0.58), pw - Inches(0.80), Inches(3.4))
para(tf, "WHAT WE ARE AUTOMATING", size=11, color=TEAL, bold=True, first=True,
     space_after=14)
para(tf, "The reading and the\ntranscription — not\nthe judgement.", size=21,
     color=WHITE, font=FONT_H, line_spacing=1.25, space_after=16)
para(tf, "The agent extracts what the document actually says, checks every value "
         "back against the source text, and builds the workbook.",
     size=12.5, color=PALE, line_spacing=1.3, space_after=12)
para(tf, "Where the document is genuinely ambiguous, it stops and asks a human "
         "rather than guessing.",
     size=12.5, color=WHITE, bold=True, line_spacing=1.3, space_after=0)
footer(s, 4)

# ================================================================ 5. HOW IT WORKS
s = slide()
y = header(s, "HOW IT WORKS", "Four stages, one job each",
           "Runs as a Databricks job on serverless compute, or as four plain scripts — the same committed code either way.")

stages = [
    ("01", "Ingest", "FRD .docx → markdown",
     ["Parses the approved FRD", "Section labels from the\nshared contract",
      "Lands in a Delta table"], TEAL),
    ("02", "Extract", "the one model call",
     ["Claude reads the FRD\nagainst a strict schema", "Streaming, validated\nclient-side",
      "Fails loudly, never\nsilently retries"], BLUE),
    ("03", "Contract build", "validate · audit · gate",
     ["Schema validation and\nregex enrichment", "Verbatim grounding audit",
      "Ambiguities gated\nfor review"], BLUE),
    ("04", "Render + eval", "the deliverable",
     ["Data-dictionary\ncross-check", "Renders the client\nSTTM workbook",
      "Cell-level eval vs.\nreference"], TEAL),
]
bw = Inches(2.86)
bx, by = MARGIN, y + Inches(0.10)
bh = Inches(2.86)
for i, (num, name, tag, lines, acc) in enumerate(stages):
    rect(s, bx, by, bw, bh, fill=WHITE, line=LINE, lw=0.9)
    rect(s, bx, by, bw, Pt(4), fill=acc)
    tf = tbox(s, bx + Inches(0.26), by + Inches(0.28), bw - Inches(0.5), Inches(1.0))
    para(tf, num, size=11, color=acc, bold=True, first=True, space_after=3)
    para(tf, name, size=18.5, color=INK, bold=True, font=FONT_H, space_after=2)
    para(tf, tag, size=11.5, color=MUTED, space_after=0, italic=True)
    tf2 = tbox(s, bx + Inches(0.26), by + Inches(1.40), bw - Inches(0.5), Inches(1.3))
    for j, ln in enumerate(lines):
        sub = ln.split("\n")
        para(tf2, "·  " + sub[0], size=11.5, color=BODY, first=(j == 0),
             space_after=0 if len(sub) > 1 else 9, line_spacing=1.1)
        for extra in sub[1:]:
            para(tf2, "    " + extra, size=11.5, color=BODY, space_after=9,
                 line_spacing=1.1)
    if i < 3:
        arrow(s, bx + bw + Inches(0.045), by + Inches(1.30), Inches(0.25),
              Inches(0.26), color=RGBColor(0xC2, 0xCE, 0xD8))
    bx += bw + Inches(0.34)

dy = by + bh + Inches(0.32)
rect(s, MARGIN, dy, CONTENT_W, Inches(1.02), fill=PAPER, line=LINE, lw=0.75)
rect(s, MARGIN, dy, Pt(3.2), Inches(1.02), fill=INK)
tf = tbox(s, MARGIN + Inches(0.32), dy + Inches(0.18), CONTENT_W - Inches(0.7),
          Inches(0.7))
para(tf, "Deliberate division of labour", size=13.5, color=INK, bold=True,
     first=True, space_after=5)
rich(tf, [("The model reads prose and scattered requirement tables. ", BODY, False),
          ("Deterministic code owns everything else", INK, True),
          (" — validation, regex-able facts, grounding checks, rule attribution "
           "and the rendering of the workbook itself.", BODY, False)],
     size=12.5, line_spacing=1.22, space_after=0)
footer(s, 5)

# ================================================================ 6. TRUST
s = slide()
y = header(s, "TRUST ARCHITECTURE", "Why you can put this in front of an auditor",
           "Three mechanisms — all deterministic, all running on every single document.")

cw = (CONTENT_W - Inches(0.6)) / 3
cards = [
    ("01", "Verbatim grounding audit", TEAL, [
        "Every strict field the model produces must appear verbatim in the source "
        "document; prose fields are checked by token overlap.",
        "On the live run: 45 of 45 strict and 30 of 30 advisory checks passed.",
        "This is the quality gate. It is never relaxed to make a run pass.",
    ]),
    ("02", "Ambiguity gating", BLUE, [
        "Where the FRD does not say which feed a rule belongs to, the agent "
        "records a gated ambiguity instead of choosing.",
        "A deterministic data-dictionary cross-check then auto-confirms whatever "
        "the data itself proves.",
        "Anything it cannot prove waits for a person.",
    ]),
    ("03", "Human-in-the-loop review", AMBER, [
        "A review app shows each gated item with its candidates and source "
        "context.",
        "The reviewer picks or overrides, and that decision is written into the "
        "contract — never re-decided by the heuristic.",
        "Your analysts stay the decision-makers.",
    ]),
]
for i, (num, title, acc, lines) in enumerate(cards):
    cx = MARGIN + (cw + Inches(0.3)) * i
    card(s, cx, y + Inches(0.10), cw, Inches(3.22), title, lines, accent=acc,
         title_size=15.5, body_size=12.3, fill=WHITE, num=num, spacing=9)

by2 = y + Inches(3.42)
rect(s, MARGIN, by2, CONTENT_W, Inches(1.05), fill=NAVY)
tf = tbox(s, MARGIN + Inches(0.38), by2 + Inches(0.20), CONTENT_W - Inches(0.8),
          Inches(0.68))
para(tf, "THE RESULT, IN ONE SENTENCE", size=10.5, color=TEAL, bold=True,
     first=True, space_after=6)
para(tf, "Every value in the workbook is traced back to the sentence that "
         "justified it, and nothing the agent is unsure of reaches the workbook "
         "unreviewed.",
     size=15, color=WHITE, line_spacing=1.2, space_after=0)
footer(s, 6)

# ================================================================ 7. DEMO DIVIDER
s = slide()
rect(s, Inches(0), Inches(0), W, H, fill=NAVY)
rect(s, Inches(0), Inches(0), Inches(0.16), H, fill=AMBER)
rect(s, Inches(0), H - Inches(2.45), W, Inches(2.45), fill=NAVY_L)

tf = tbox(s, MARGIN, Inches(1.70), Inches(9.5), Inches(0.3))
para(tf, "SECTION 05", size=12, color=AMBER, bold=True, first=True, space_after=0)
tf = tbox(s, MARGIN, Inches(2.15), Inches(8.2), Inches(1.7))
para(tf, "The Agent Demo", size=48, color=WHITE, font=FONT_H, first=True,
     space_after=10)
para(tf, "One approved FRD in. One STTM workbook out.\nRun live, on this laptop, "
         "in about thirty seconds.",
     size=16.5, color=PALE, line_spacing=1.3, space_after=0)

by = H - Inches(1.95)
labels = [("Ingest", "docx → markdown"), ("Extract", "one live model call"),
          ("Contract build", "validate · audit · gate"),
          ("Render + eval", "workbook + score")]
bw = Inches(2.72)
bx = MARGIN
for i, (n, d) in enumerate(labels):
    rect(s, bx, by, Inches(0.11), Inches(0.11), fill=AMBER, shape=MSO_SHAPE.OVAL)
    tf = tbox(s, bx, by + Inches(0.28), bw - Inches(0.35), Inches(0.8))
    para(tf, n, size=15.5, color=WHITE, bold=True, first=True, space_after=3)
    para(tf, d, size=11.5, color=PALER, line_spacing=1.15, space_after=0)
    if i < 3:
        rect(s, bx + Inches(0.22), by + Inches(0.045), bw - Inches(0.44), Pt(1.2),
             fill=RGBColor(0x3C, 0x5A, 0x72))
    bx += bw

tf = tbox(s, W - Inches(4.15), Inches(2.35), Inches(3.30), Inches(1.4))
para(tf, "BACKUP PLAN", size=10.5, color=AMBER, bold=True, first=True, space_after=8)
para(tf, "Every demo run is saved. If the network misbehaves we replay a recorded "
         "run — the same screens, zero API calls, fully offline.",
     size=12.5, color=PALE, line_spacing=1.3, space_after=0)

# ================================================================ 8. DEMO WATCH
s = slide()
y = header(s, "THE DEMO", "What to watch for",
           "Five moments on the results screen — each one is a claim we can defend.")

left_w = Inches(6.3)
beats = [
    ("Extraction summary", "Three feeds, their file patterns, target tables and rules. All of it came out of the document — nothing is templated."),
    ("The gate strip", "\"N detected, M auto-confirmed against the data dictionary, K awaiting human review.\" The human-in-the-loop centrepiece."),
    ("The verdict tile", "A dictionary-correct extraction earns a PASS. A flag survives only when a human genuinely needs to look at it."),
    ("Eval vs. the golden workbook", "94.1% of cells match a hand-built reference. Every remaining cell is a difference deliberately built into the fixture."),
    ("The workbook itself", "We download and open the .xlsx — the artifact a mapping analyst would otherwise hand-build."),
]
cy = y + Inches(0.06)
for i, (t, sub) in enumerate(beats):
    tfn = tbox(s, MARGIN, cy + Inches(0.01), Inches(0.45), Inches(0.32))
    para(tfn, f"0{i+1}", size=15, color=AMBER, bold=True, font=FONT_H, first=True,
         space_after=0)
    tf = tbox(s, MARGIN + Inches(0.50), cy, left_w - Inches(0.55), Inches(0.32))
    para(tf, t, size=16, color=INK, bold=True, first=True, space_after=3)
    para(tf, sub, size=12.5, color=BODY, line_spacing=1.16, space_after=0)
    cy += Inches(0.88)

px = MARGIN + left_w + Inches(0.45)
pw = CONTENT_W - left_w - Inches(0.45)
card(s, px, y + Inches(0.06), pw, Inches(2.06), "Honesty about the run", [
    "The demo states its cost before it spends anything: one billed call, roughly "
    "$0.15, about thirty-five seconds.",
    "Runs vary in how much the dictionary can auto-confirm. The mapping quality "
    "and the 94.1% eval have not varied.",
], accent=BLUE, title_size=15, body_size=12, fill=WHITE, spacing=9)
card(s, px, y + Inches(2.34), pw, Inches(2.06), "Data handling — non-negotiable", [
    "No real client document is used in this demo. The FRD on screen is "
    "anonymized and cleared for viewing.",
    "In the deployed shape, documents and outputs stay inside your own Databricks "
    "workspace and Unity Catalog volumes.",
], accent=TEAL, title_size=15, body_size=12, fill=WHITE, spacing=9)
footer(s, 8)

# ================================================================ 9. PROOF
s = slide()
y = header(s, "PROOF", "Measured on a live end-to-end run",
           "Anonymized demo FRD → rendered workbook, on the committed pipeline, 7 August 2026. Not a mock run.")

mw = (CONTENT_W - Inches(0.9)) / 4
mets = [
    ("94.1%", "MAPPING FIDELITY", "3,094 of 3,288 cells match the golden reference workbook", TEAL),
    ("0 / 50", "HALLUCINATION SWEEP", "Every identifier traced verbatim to the FRD or the workbook", TEAL),
    ("45 / 45", "STRICT GROUNDING", "Plus 30 of 30 advisory checks — the quality gate passed clean", BLUE),
    ("~33 s", "END TO END", "One billed model call, about $0.15 per document", BLUE),
]
for i, (v, l, n, a) in enumerate(mets):
    metric(s, MARGIN + (mw + Inches(0.3)) * i, y + Inches(0.10), mw, Inches(1.86),
           v, l, n, accent=a)

cy = y + Inches(2.10)
cw = (CONTENT_W - Inches(0.5)) / 2
card(s, MARGIN, cy, cw, Inches(2.28), "What the numbers mean", [
    "Every one of the 194 non-matching cells is a datatype difference deliberately "
    "built into the test fixture. The eval catches exactly what it should.",
    "No mapping was missed, and none was invented.",
    "A proxy-SME review estimated a first-pass keep rate of roughly 96% — against "
    "an acceptance bar of 80%.",
], accent=TEAL, title_size=15, body_size=12.3, fill=WHITE, spacing=8)
card(s, MARGIN + cw + Inches(0.5), cy, cw, Inches(2.28), "Engineering posture", [
    "77 offline unit tests — no network, no model calls and no cluster needed to "
    "run them.",
    "Headroom of roughly 75–90 feeds in one document; this run used 5% of its "
    "output budget.",
    "The pipeline fails loudly on schema violations, refusals or truncation — it "
    "never retries into a plausible answer.",
], accent=BLUE, title_size=15, body_size=12.3, fill=WHITE, spacing=8)
footer(s, 9)

# ================================================================ 10. AMERIHEALTH
s = slide()
y = header(s, "THE VALUE CASE", "Why this matters for AmeriHealth")

cw = (CONTENT_W - Inches(0.6)) / 3
rows = [
    [("Days of work become minutes", TEAL, [
        "The hand-built workbook is the slowest step between an approved FRD and "
        "working pipeline code. The agent produces it in about thirty seconds.",
        "Analysts move from transcription to review — the part that needs their "
        "expertise.",
      ]),
     ("Auditable by construction", BLUE, [
        "Every value is traced verbatim to the sentence that justified it, and "
        "that check is a hard gate on every run.",
        "When an auditor asks why a field maps the way it does, the answer is "
        "recorded, not reconstructed.",
      ]),
     ("Consistent across teams", AMBER, [
        "Two analysts reading one FRD produce two workbooks. The agent produces "
        "the same structure every time.",
        "Naming, load strategies and rule attribution stop drifting — which is "
        "what makes downstream automation possible.",
      ])],
    [("Your data stays with you", TEAL, [
        "Built for Databricks: documents and outputs live in your Unity Catalog "
        "volumes, and credentials come from a secret scope.",
        "The PHI posture is a design constraint — this demo runs only on "
        "anonymized material.",
      ]),
     ("Humans stay in control", BLUE, [
        "The agent never guesses. It flags an ambiguity, auto-confirms only what "
        "the data dictionary proves, and routes the rest to a reviewer.",
        "Adoption means reviewing a short list of open questions, not trusting a "
        "black box.",
      ]),
     ("It compounds downstream", AMBER, [
        "The contract this agent emits is what CodeGen consumes, which in turn "
        "feeds automated code review.",
        "Every FRD processed here shortens the whole chain: a governed path from "
        "requirement to running pipeline.",
      ])],
]
row_h = Inches(2.42)
for r, row in enumerate(rows):
    for i, (title, acc, lines) in enumerate(row):
        card(s, MARGIN + (cw + Inches(0.3)) * i, y + Inches(0.06) + row_h * r,
             cw, Inches(2.26), title, lines, accent=acc, title_size=14.5,
             body_size=11.8, fill=PAPER, spacing=8)
footer(s, 10, "FRD → STTM Agent   ·   Value for AmeriHealth")

# ================================================================ 11. DATABRICKS
s = slide()
y = header(s, "DEPLOYMENT", "Built for your Databricks workspace",
           "Demoed locally today; the workspace shape is already defined in the repository.")

cw = (CONTENT_W - Inches(0.55)) / 2
card(s, MARGIN, y + Inches(0.06), cw, Inches(3.40), "In place today", [
    "An asset bundle defines a single job chaining all four stages with explicit "
    "dependencies, on serverless compute.",
    "Unity Catalog volumes hold raw FRDs, reference workbooks and outputs — client "
    "documents never leave the workspace.",
    "Anthropic credentials are read from a Databricks secret scope; no key exists "
    "in the repository.",
    "The review app is packaged as a Databricks App on the officially supported "
    "stack.",
    "Orchestration retries are deliberately disabled — deterministic errors must "
    "surface, not be retried away.",
], accent=TEAL, title_size=16, body_size=12.2, fill=WHITE, spacing=9)

card(s, MARGIN + cw + Inches(0.55), y + Inches(0.06), cw, Inches(3.40),
     "The remaining port work — named, not hidden", [
    "The demo runs pipeline stages as local processes; in the workspace it "
    "triggers the bundle job and reads outputs from Unity Catalog.",
    "The review app's API key must arrive through a Databricks Apps secret "
    "resource rather than a local environment file.",
    "Live progress streaming through the Apps proxy is expected to work but has "
    "not been exercised against a real workspace.",
    "None of this touches the pipeline logic or the quality gates — it is plumbing "
    "between the same components.",
], accent=AMBER, title_size=16, body_size=12.2, fill=WHITE, spacing=9)

by = y + Inches(3.72)
rect(s, MARGIN, by, CONTENT_W, Inches(0.84), fill=NAVY)
tf = tbox(s, MARGIN + Inches(0.38), by + Inches(0.24), CONTENT_W - Inches(0.8),
          Inches(0.45))
rich(tf, [("The framing: ", TEAL, True),
          ("built for the workspace, currently demoed locally. The four stages you "
           "watch running are the same committed code a workspace job would run.",
           WHITE, False)], size=13.5, first=True, line_spacing=1.2, space_after=0)
footer(s, 11)

# ================================================================ 12. STATUS
s = slide()
y = header(s, "STATUS & NEXT STEPS", "Where we are, and what we would do next",
           "An honest read: the pipeline and its quality gates are proven; the surrounding integration is the work ahead.")

rows = [
    ("Extraction quality", "PROVEN", "Live end-to-end run: 94.1% fidelity, zero hallucinations, ~96% estimated first-pass keep rate.", TEAL),
    ("Grounding & gating", "PROVEN", "45/45 strict and 30/30 advisory checks, with human review wired end-to-end and unit-tested.", TEAL),
    ("Upstream hand-off", "PROVEN", "Versioned label contract shared byte-identically with the BRD→FRD agent, guarded by tests.", TEAL),
    ("Workspace deployment", "NEXT", "Trigger the bundle job from the app, source secrets via an Apps resource, exercise streaming.", AMBER),
    ("Downstream contract", "NEXT", "Emit the workbook-derived mapping contract CodeGen needs — the programme's biggest open gap.", AMBER),
    ("Scale validation", "NEXT", "Run against a real, larger FRD to confirm the 75–90 feed headroom and tune chunking if needed.", BLUE),
]
cy = y + Inches(0.04)
rh = Inches(0.64)
for i, (label, status, detail, acc) in enumerate(rows):
    if i % 2 == 0:
        rect(s, MARGIN, cy - Inches(0.04), CONTENT_W, rh - Inches(0.02), fill=PAPER)
    tf = tbox(s, MARGIN + Inches(0.28), cy + Inches(0.10), Inches(3.2), Inches(0.32))
    para(tf, label, size=14.5, color=INK, bold=True, first=True, space_after=0)
    pill = rect(s, MARGIN + Inches(3.62), cy + Inches(0.09), Inches(1.0),
                Inches(0.28), fill=acc, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.5)
    ptf = pill.text_frame
    ptf.margin_left = ptf.margin_right = ptf.margin_top = ptf.margin_bottom = 0
    ptf.vertical_anchor = MSO_ANCHOR.MIDDLE
    pp = ptf.paragraphs[0]
    pp.alignment = PP_ALIGN.CENTER
    pr = pp.add_run()
    pr.text = status
    pr.font.name, pr.font.size, pr.font.bold = FONT, Pt(9.5), True
    pr.font.color.rgb = WHITE
    tf2 = tbox(s, MARGIN + Inches(4.95), cy + Inches(0.11), CONTENT_W - Inches(5.2),
               Inches(0.4))
    para(tf2, detail, size=12.3, color=BODY, first=True, space_after=0,
         line_spacing=1.15)
    cy += rh

by = cy + Inches(0.22)
rect(s, MARGIN, by, CONTENT_W, Inches(0.74), fill=NAVY)
rect(s, MARGIN, by, Pt(3.2), Inches(0.74), fill=TEAL)
tf = tbox(s, MARGIN + Inches(0.38), by + Inches(0.22), CONTENT_W - Inches(0.8),
          Inches(0.34))
rich(tf, [("Our ask: ", TEAL, True),
          ("a real, non-sensitive FRD to validate at scale, and a workspace target "
           "for the deployment port.", WHITE, False)],
     size=13.5, first=True, line_spacing=1.15, space_after=0)
footer(s, 12)

# ================================================================ 13. Q&A
s = slide()
rect(s, Inches(0), Inches(0), W, H, fill=NAVY)
rect(s, Inches(0), Inches(0), Inches(0.16), H, fill=TEAL)
rect(s, W - Inches(5.0), Inches(0), Inches(5.0), H, fill=NAVY_L)
rect(s, W - Inches(5.0), Inches(0), Pt(2), H, fill=TEAL)

tf = tbox(s, MARGIN, Inches(2.35), Inches(7.0), Inches(0.3))
para(tf, "SECTION 07", size=12, color=TEAL, bold=True, first=True, space_after=0)
tf = tbox(s, MARGIN, Inches(2.80), Inches(7.0), Inches(1.9))
para(tf, "Questions", size=50, color=WHITE, font=FONT_H, first=True, space_after=12)
para(tf, "Happy to go back into the demo, open the workbook, or walk through any "
         "part of the pipeline in detail.",
     size=15.5, color=PALE, line_spacing=1.3, space_after=0)
rect(s, MARGIN, Inches(4.90), Inches(1.3), Pt(2.5), fill=TEAL)
tf = tbox(s, MARGIN, Inches(5.25), Inches(7.0), Inches(0.8))
para(tf, "Soham  ·  Arjun   |   Hexaware — AI in Engineering", size=13.5,
     color=WHITE, bold=True, first=True, space_after=4)
para(tf, "FRD → STTM Agent   ·   Agent 2 of 5", size=12, color=PALER, space_after=0)

tf = tbox(s, W - Inches(4.50), Inches(1.95), Inches(3.6), Inches(3.9))
para(tf, "LIKELY QUESTIONS", size=10.5, color=TEAL, bold=True, first=True,
     space_after=14)
for q in ["How does it behave on an FRD far larger than the demo?",
          "What happens when the document is genuinely ambiguous?",
          "How do we know it hasn't invented a mapping?",
          "What does the review workload look like per document?",
          "How does this reach our workspace, and who holds the keys?",
          "What is the effort to close the CodeGen hand-off gap?"]:
    para(tf, "—  " + q, size=12.5, color=PALE, space_after=11, line_spacing=1.2)

# ================================================================ notes
SPEAKER_NOTES = [
    # 1 title
    "Open. Introduce yourself and Arjun and split the room's expectation: slides "
    "first, then a live run of the agent, then questions.\n\n"
    "The three numbers on the right are the whole story in miniature — 94.1% "
    "mapping fidelity, zero ungrounded identifiers, about thirty seconds and "
    "fifteen cents per document. All measured on a real live run on 7 August "
    "2026, not a mock, not a projection.\n\n"
    "Say plainly: the FRD in the demo is anonymized fixture material. No real "
    "client document is used at any point today.",
    # 2 agenda
    "Keep this to about forty-five seconds. Signal that the demo (item 05) is the "
    "centre of gravity and that we will finish with open Q&A.\n\n"
    "Flag the two-hander: agree in advance who takes 03/04 (mechanics) and who "
    "takes 06 (value), so the hand-off is not negotiated on stage.",
    # 3 program map
    "The point of this slide is that the agents are chained by contracts, not by "
    "email attachments.\n\n"
    "Upstream is genuinely solved: our label contract is versioned and committed "
    "byte-identically to both repositories, with a round-trip test that fails the "
    "build on drift.\n\n"
    "Downstream, say the gap out loud — CodeGen needs a workbook-derived mapping "
    "contract and no committed tool in the programme produces one yet. Naming it "
    "before anyone asks buys credibility for everything else on the deck.",
    # 4 problem
    "Anchor in the audience's own experience — ask whether their mapping "
    "workbooks are built by hand today; the answer is almost always yes.\n\n"
    "The right-hand panel is the line to land: we automate the reading and the "
    "transcription, not the judgement. Where the FRD is genuinely ambiguous the "
    "agent stops and asks rather than guessing. That framing pre-empts the "
    "\"can we trust it\" question later.",
    # 5 how it works
    "Ninety seconds, no deeper. Walk the four stages left to right; the audience "
    "will see these same four stages light up in the demo, so this slide is the "
    "map for what they are about to watch.\n\n"
    "The band at the bottom matters more than the boxes: the model reads prose, "
    "and deterministic code owns validation, grounding, attribution and "
    "rendering. The model never writes a cell that code cannot justify.\n\n"
    "If asked which model: Claude, called through the Anthropic SDK, with the "
    "schema and prompt living in our repository — not a black-box endpoint.",
    # 6 trust
    "This is the slide for the sceptical engineer and the compliance-minded "
    "stakeholder at the same time.\n\n"
    "Grounding audit: every strict field must appear verbatim in the source "
    "document. On the live run, 45 of 45 strict checks passed. Stress that we "
    "never relax this to make a run go green — a failing audit is a failing run.\n\n"
    "Gating: the agent records an ambiguity rather than choosing; a deterministic "
    "data-dictionary cross-check auto-confirms only what the data proves.\n\n"
    "If pushed on the difference between strict and advisory: strict fields are "
    "identifiers and must match verbatim; advisory fields are prose and are "
    "checked by token overlap.",
    # 7 demo divider
    "Transition slide. Hand over to whoever is driving the laptop.\n\n"
    "Before starting the run, read the cost-confirmation dialog aloud — one "
    "billed call, about $0.15, about thirty-five seconds. The demo tells you what "
    "it costs before it spends anything.\n\n"
    "Contingency: if the network misbehaves, switch to Replay and pick the "
    "post-fix run. Identical results view, zero API calls, works offline. Do not "
    "debug on stage.",
    # 8 demo watch
    "Leave this up while the run executes, then walk the results page top to "
    "bottom in this order.\n\n"
    "Give the gate strip time even when the counts are zero — zero means nothing "
    "needed a human on this run, which is itself the point. If the run shows a "
    "flag that then gets auto-confirmed, slow down: that is the best possible "
    "outcome, because the flag-then-confirm story renders on screen.\n\n"
    "Runs vary in how much the dictionary can auto-confirm. The mapping quality "
    "and the 94.1% eval have not varied across runs — say so before anyone "
    "wonders about it.\n\n"
    "Finish by downloading and opening the workbook. That artifact is the "
    "deliverable an analyst would otherwise hand-build.",
    # 9 proof
    "Come back to slides with the workbook still fresh in their minds.\n\n"
    "The number to defend is 94.1%. Every one of the 194 non-matching cells is a "
    "datatype difference deliberately built into the test fixture — the eval "
    "catches exactly what it should and nothing else. Nothing was missed and "
    "nothing was invented.\n\n"
    "The hallucination sweep checked all 50 identifiers in the output — feed "
    "names, file patterns, catalogs, schemas, tables, load strategies, "
    "requirement ids — for verbatim presence in the source. Zero ungrounded.\n\n"
    "If asked about the keep rate: a proxy-SME review estimated roughly 96% "
    "first-pass keep, against an acceptance bar of 80%.",
    # 10 amerihealth
    "This is the slide to slow down on — it is the reason we are in the room.\n\n"
    "Lead with the top-left card and the bottom-left card: speed, and the fact "
    "that your data never leaves your tenancy. For a regulated payer those two "
    "usually decide the conversation.\n\n"
    "Auditability is the strongest card technically: the traceability is not a "
    "report we generate afterwards, it is a hard gate that runs on every "
    "document.\n\n"
    "Close on the bottom-right: the value is not one workbook, it is a governed "
    "path from requirement to running pipeline, and every FRD processed here "
    "shortens the downstream chain.",
    # 11 deployment
    "Framing line: built for the workspace, currently demoed locally. Say it "
    "early so nobody thinks the local demo is the product.\n\n"
    "Left column is real and committed today. Right column is the honest list of "
    "port work — three items, all plumbing, none of them touching the pipeline "
    "logic or the quality gates.\n\n"
    "If asked about credentials: in the workspace, the key comes from a "
    "Databricks secret scope, and there is no key anywhere in the repository.",
    # 12 status
    "A one-page honest read. Three things proven, three things next.\n\n"
    "Do not soften the two amber rows. The workspace port and the downstream "
    "mapping contract are real work, and being straight about them is what makes "
    "the three green rows believable.\n\n"
    "End on the ask: a real, non-sensitive FRD to validate at scale, and a "
    "workspace target for the deployment port. Get a name and a date if you can.",
    # 13 Q&A
    "Answers to have ready:\n\n"
    "Larger FRDs — the run used 5% of its output budget, giving roughly 75 to 90 "
    "feeds of headroom in one document. Beyond that, per-feed chunked extraction "
    "is the known path. Validating on a real large FRD is our ask.\n\n"
    "Genuine ambiguity — it is gated, not guessed. The dictionary auto-confirms "
    "what it can prove; everything else goes to a reviewer, whose decision is "
    "stored and never re-decided by the heuristic.\n\n"
    "Invented mappings — the grounding audit and the hallucination sweep. Offer "
    "to show the audit output.\n\n"
    "Review workload — on the demo document, at most a handful of items per run, "
    "often zero.\n\n"
    "Keys and tenancy — Databricks secret scope, Unity Catalog volumes, nothing "
    "leaves the workspace.\n\n"
    "CodeGen gap — no committed extractor today; it is scoped work, and we would "
    "size it properly rather than guess on stage.",
]
for i, txt in enumerate(SPEAKER_NOTES):
    notes(prs.slides[i], txt)

prs.save(OUT)
print("wrote", OUT, "-", len(prs.slides._sldIdLst), "slides,",
      len(SPEAKER_NOTES), "notes")
