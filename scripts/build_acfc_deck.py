"""Build the ACFC overview deck for this agent (PPTX).

    pip install python-pptx
    python scripts/build_acfc_deck.py [output.pptx]

Self-contained: shared slide helpers (deck_lib) + this agent's content in one file so the
deck can be regenerated without cross-repo dependencies. Layout is absolute-positioned;
text length drives fit — shorten copy rather than growing boxes.
"""
# ---------------------------------------------------------------------------
# deck_lib — shared slide-building helpers (python-pptx, absolute layout, 16:9)
# Every deck in the AI-in-Engineering program uses this same visual system so
# the five agent decks read as one family. Copy is short by design: layout is
# absolute-positioned, so text length drives fit. Shorten copy, don't grow boxes.
# ---------------------------------------------------------------------------
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

# ------------------------------------------------------------------ palette
NAVY = RGBColor(0x0E, 0x22, 0x33)
NAVY_L = RGBColor(0x1B, 0x3A, 0x52)
INK = RGBColor(0x1A, 0x24, 0x33)
BODY = RGBColor(0x3A, 0x4A, 0x5C)
MUTED = RGBColor(0x6B, 0x7A, 0x8C)
TEAL = RGBColor(0x00, 0xA1, 0x9A)
TEAL_D = RGBColor(0x00, 0x7C, 0x76)
BLUE = RGBColor(0x2E, 0x6E, 0xD1)
AMBER = RGBColor(0xD9, 0x8A, 0x1F)
RED = RGBColor(0xC2, 0x3B, 0x3B)
GREEN = RGBColor(0x1E, 0x7A, 0x46)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PAPER = RGBColor(0xF4, 0xF7, 0xF9)
LINE = RGBColor(0xD8, 0xE0, 0xE7)
PALE = RGBColor(0xC3, 0xD2, 0xDE)
TEAL_BG = RGBColor(0xE3, 0xF5, 0xF4)
BLUE_BG = RGBColor(0xE6, 0xEF, 0xFB)
AMBER_BG = RGBColor(0xFB, 0xF1, 0xE1)
GREY_BG = RGBColor(0xEE, 0xF1, 0xF4)

KIND_COLOR = {  # pipeline step kinds
    "llm": (BLUE, BLUE_BG, "LLM step (Claude)"),
    "code": (TEAL_D, TEAL_BG, "Deterministic code"),
    "human": (AMBER, AMBER_BG, "Human decision"),
    "data": (MUTED, GREY_BG, "Artifact / data"),
}

FONT = "Calibri"
FONT_H = "Calibri Light"

W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.7)
CONTENT_W = W - 2 * MARGIN
TOP = Inches(1.45)  # content starts below title band
BOTTOM = Inches(6.75)  # content must stay above footer band


class Deck:
    def __init__(self, agent_label, program_label="Hexaware × ACFC · AI-in-Engineering"):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self.blank = self.prs.slide_layouts[6]
        self.agent_label = agent_label
        self.program_label = program_label
        self.n = 0

    # ------------------------------------------------------------ primitives
    def slide(self, dark=False):
        s = self.prs.slides.add_slide(self.blank)
        self.n += 1
        if dark:
            rect(s, 0, 0, W, H, fill=NAVY)
        return s

    def save(self, path):
        self.prs.save(path)

    def footer(self, s, dark=False):
        col = PALE if dark else MUTED
        line_col = NAVY_L if dark else LINE
        rect(s, MARGIN, Inches(6.98), CONTENT_W, Pt(0.75), fill=line_col)
        text(s, MARGIN, Inches(7.02), Inches(8), Inches(0.35),
             f"{self.program_label}  ·  {self.agent_label}", size=10, color=col)
        text(s, W - MARGIN - Inches(1.5), Inches(7.02), Inches(1.5), Inches(0.35),
             str(self.n), size=10, color=col, align=PP_ALIGN.RIGHT)

    def title(self, s, title_txt, kicker=None, dark=False):
        col = WHITE if dark else INK
        kcol = TEAL if dark else TEAL_D
        y = Inches(0.42)
        if kicker:
            text(s, MARGIN, y, CONTENT_W, Inches(0.3), kicker.upper(), size=11,
                 color=kcol, bold=True, font=FONT, spacing=True)
            y = Inches(0.68)
        text(s, MARGIN, y, CONTENT_W, Inches(0.7), title_txt, size=28,
             color=col, bold=True, font=FONT_H)
        rect(s, MARGIN, Inches(1.32), Inches(1.1), Pt(3), fill=TEAL)

    def notes(self, s, txt):
        s.notes_slide.notes_text_frame.text = txt

    # ------------------------------------------------------------ slide types
    def cover(self, title_txt, subtitle, meta_lines, notes=None):
        s = self.slide(dark=True)
        rect(s, 0, 0, Inches(0.35), H, fill=TEAL)
        text(s, Inches(1.0), Inches(0.9), Inches(10), Inches(0.4),
             self.program_label.upper(), size=12, color=TEAL, bold=True, spacing=True)
        text(s, Inches(1.0), Inches(2.0), Inches(11.3), Inches(1.9), title_txt, size=44,
             color=WHITE, bold=True, font=FONT_H)
        text(s, Inches(1.0), Inches(4.0), Inches(11), Inches(1.0), subtitle, size=20,
             color=PALE, font=FONT_H)
        y = Inches(5.5)
        for ln in meta_lines:
            text(s, Inches(1.0), y, Inches(11), Inches(0.35), ln, size=13, color=PALE)
            y += Inches(0.36)
        if notes:
            self.notes(s, notes)
        return s

    def statement(self, kicker, big, small=None, notes=None):
        """One big idea per slide."""
        s = self.slide()
        self.title(s, kicker)
        text(s, MARGIN, Inches(1.7), CONTENT_W, Inches(2.8), big, size=26, color=NAVY,
             bold=True, font=FONT_H, anchor=MSO_ANCHOR.MIDDLE)
        if small:
            text(s, MARGIN, Inches(4.75), CONTENT_W, Inches(1.9), small, size=15, color=BODY)
        self.footer(s)
        if notes:
            self.notes(s, notes)
        return s

    def bullets_slide(self, title_txt, items, kicker=None, notes=None, takeaway=None,
                      size=17):
        s = self.slide()
        self.title(s, title_txt, kicker)
        bottom = BOTTOM - (Inches(0.85) if takeaway else 0)
        bullets(s, MARGIN, TOP, CONTENT_W, bottom - TOP, items, size=size)
        if takeaway:
            takeaway_bar(s, takeaway)
        self.footer(s)
        if notes:
            self.notes(s, notes)
        return s

    def two_col(self, title_txt, left_head, left_items, right_head, right_items,
                kicker=None, notes=None, takeaway=None, left_tone="grey",
                right_tone="teal", size=15):
        s = self.slide()
        self.title(s, title_txt, kicker)
        bottom = BOTTOM - (Inches(0.85) if takeaway else 0)
        gap = Inches(0.35)
        cw = (CONTENT_W - gap) / 2
        for i, (head, items, tone) in enumerate(
                [(left_head, left_items, left_tone), (right_head, right_items, right_tone)]):
            x = MARGIN + i * (cw + gap)
            panel(s, x, TOP, cw, bottom - TOP, head, items, tone=tone, size=size)
        if takeaway:
            takeaway_bar(s, takeaway)
        self.footer(s)
        if notes:
            self.notes(s, notes)
        return s

    def cards(self, title_txt, items, kicker=None, cols=3, notes=None, takeaway=None,
              body_size=13):
        """items: list of (heading, body[, tone])."""
        s = self.slide()
        self.title(s, title_txt, kicker)
        bottom = BOTTOM - (Inches(0.85) if takeaway else 0)
        rows = (len(items) + cols - 1) // cols
        gap = Inches(0.28)
        cw = (CONTENT_W - gap * (cols - 1)) / cols
        ch = (bottom - TOP - gap * (rows - 1)) / rows
        for i, it in enumerate(items):
            head, body = it[0], it[1]
            tone = it[2] if len(it) > 2 else "grey"
            r, c = divmod(i, cols)
            x = MARGIN + c * (cw + gap)
            y = TOP + r * (ch + gap)
            card(s, x, y, cw, ch, head, body, tone=tone, body_size=body_size)
        if takeaway:
            takeaway_bar(s, takeaway)
        self.footer(s)
        if notes:
            self.notes(s, notes)
        return s

    def kpis(self, title_txt, tiles, kicker=None, notes=None, source=None, below=None,
             cols=None, takeaway=None):
        """tiles: list of (big, label, sub). Optional `below` bullets under tiles."""
        s = self.slide()
        self.title(s, title_txt, kicker)
        cols = cols or min(len(tiles), 4)
        rows = (len(tiles) + cols - 1) // cols
        gap = Inches(0.25)
        cw = (CONTENT_W - gap * (cols - 1)) / cols
        th = Inches(1.9)
        bottom = BOTTOM - (Inches(0.85) if takeaway else 0)
        for i, (big, label, sub) in enumerate(tiles):
            r, c = divmod(i, cols)
            x = MARGIN + c * (cw + gap)
            y = TOP + r * (th + gap)
            kpi_tile(s, x, y, cw, th, big, label, sub)
        y_after = TOP + rows * (th + gap)
        if below:
            bullets(s, MARGIN, y_after + Inches(0.1), CONTENT_W,
                    bottom - y_after - Inches(0.5), below, size=15)
        if source:
            text(s, MARGIN, bottom - Inches(0.4), CONTENT_W, Inches(0.4),
                 "Source: " + source, size=10, color=MUTED, italic=True)
        if takeaway:
            takeaway_bar(s, takeaway)
        self.footer(s)
        if notes:
            self.notes(s, notes)
        return s

    def flow_slide(self, title_txt, steps, kicker=None, notes=None, below=None,
                   takeaway=None, legend=True, y=None, step_h=None):
        """steps: list of (label, sub, kind). Rendered as a left-to-right flow."""
        s = self.slide()
        self.title(s, title_txt, kicker)
        y = y or (TOP + Inches(0.25))
        step_h = step_h or (Inches(1.8) if len(steps) <= 5 else Inches(2.2))
        flow(s, MARGIN, y, CONTENT_W, step_h, steps)
        yy = y + step_h + Inches(0.15)
        if legend:
            legend_row(s, MARGIN, yy, [k for k in ("llm", "code", "human", "data")
                                       if any(st[2] == k for st in steps)])
            yy += Inches(0.45)
        bottom = BOTTOM - (Inches(0.85) if takeaway else 0)
        if below:
            bullets(s, MARGIN, yy + Inches(0.1), CONTENT_W, bottom - yy - Inches(0.1),
                    below, size=14)
        if takeaway:
            takeaway_bar(s, takeaway)
        self.footer(s)
        if notes:
            self.notes(s, notes)
        return s

    def table_slide(self, title_txt, headers, rows, kicker=None, notes=None,
                    col_widths=None, takeaway=None, source=None, size=12,
                    highlight_last=False):
        s = self.slide()
        self.title(s, title_txt, kicker)
        bottom = BOTTOM - (Inches(0.85) if takeaway else 0)
        if source:
            bottom -= Inches(0.6)
        table(s, MARGIN, TOP, CONTENT_W, bottom - TOP, headers, rows,
              col_widths=col_widths, size=size, highlight_last=highlight_last)
        if source:
            text(s, MARGIN, bottom + Inches(0.05), CONTENT_W, Inches(0.55),
                 "Source / assumptions: " + source, size=9.5, color=MUTED, italic=True)
        if takeaway:
            takeaway_bar(s, takeaway)
        self.footer(s)
        if notes:
            self.notes(s, notes)
        return s

    def program_map(self, title_txt, highlight, kicker=None, notes=None, below=None,
                    takeaway=None):
        """The five-agent program map with one agent highlighted."""
        s = self.slide()
        self.title(s, title_txt, kicker)
        agents = [
            ("BRD → FRD", "Business requirements →\nfunctional requirements", "brd"),
            ("FRD → STTM", "Functional requirements →\nsource-to-target mapping", "sttm"),
            ("CodeGen", "Mapping contract →\npipeline code + tests", "codegen"),
            ("Code Review", "PR / diff →\nreview findings + gate", "review"),
        ]
        y = TOP + Inches(0.3)
        bh = Inches(1.5)
        gap = Inches(0.55)
        bw = (CONTENT_W - gap * 3) / 4
        for i, (name, sub, key) in enumerate(agents):
            x = MARGIN + i * (bw + gap)
            hi = key == highlight
            box = rect(s, x, y, bw, bh, fill=(TEAL if hi else WHITE),
                       line=(TEAL if hi else PALE), lw=1.5,
                       shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.12)
            text(s, x, y + Inches(0.15), bw, Inches(0.45), name, size=16, bold=True,
                 color=(WHITE if hi else NAVY), align=PP_ALIGN.CENTER)
            text(s, x + Inches(0.1), y + Inches(0.6), bw - Inches(0.2), Inches(0.85), sub,
                 size=11, color=(WHITE if hi else BODY), align=PP_ALIGN.CENTER)
            if i < 3:
                arrow(s, x + bw + Inches(0.08), y + bh / 2, gap - Inches(0.16))
        # sidecar
        sy = y + bh + Inches(0.4)
        hi = highlight == "sql"
        sw = Inches(5.2)
        sx = MARGIN + (CONTENT_W - sw) / 2
        rect(s, sx, sy, sw, Inches(0.95), fill=(TEAL if hi else WHITE),
             line=(TEAL if hi else PALE), lw=1.5, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
             adj=0.2, dash=(not hi))
        text(s, sx, sy + Inches(0.08), sw, Inches(0.4), "SQL Optimization (standalone)",
             size=15, bold=True, color=(WHITE if hi else NAVY), align=PP_ALIGN.CENTER)
        text(s, sx, sy + Inches(0.48), sw, Inches(0.4),
             "Reads warehouse query telemetry, returns a ranked optimization backlog",
             size=11, color=(WHITE if hi else BODY), align=PP_ALIGN.CENTER)
        text(s, MARGIN, sy + Inches(1.05), CONTENT_W, Inches(0.35),
             "Human approval gate between every stage · shared hub + append-only ledger "
             "(agent-pipeline-orchestrator)", size=11, color=MUTED, align=PP_ALIGN.CENTER,
             italic=True)
        yy = sy + Inches(1.5)
        bottom = BOTTOM - (Inches(0.85) if takeaway else 0)
        if below:
            bullets(s, MARGIN, yy, CONTENT_W, bottom - yy, below, size=14)
        if takeaway:
            takeaway_bar(s, takeaway)
        self.footer(s)
        if notes:
            self.notes(s, notes)
        return s

    def closing(self, title_txt, lines, notes=None):
        s = self.slide(dark=True)
        rect(s, 0, 0, Inches(0.35), H, fill=TEAL)
        text(s, Inches(1.0), Inches(1.2), Inches(11), Inches(1.0), title_txt, size=36,
             color=WHITE, bold=True, font=FONT_H)
        bullets(s, Inches(1.0), Inches(2.5), Inches(11.3), Inches(4.0), lines, size=18,
                color=PALE, bullet_color=TEAL)
        self.footer(s, dark=True)
        if notes:
            self.notes(s, notes)
        return s


# ------------------------------------------------------------------ helpers
def rect(s, x, y, w, h, fill=None, line=None, lw=1.0, shape=MSO_SHAPE.RECTANGLE,
         adj=None, dash=False):
    sh = s.shapes.add_shape(shape, int(x), int(y), int(w), int(h))
    if adj is not None and sh.adjustments:
        sh.adjustments[0] = adj
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
        if dash:
            from pptx.enum.dml import MSO_LINE
            sh.line.dash_style = MSO_LINE.DASH
    sh.shadow.inherit = False
    return sh


def text(s, x, y, w, h, txt, size=14, color=INK, bold=False, italic=False, font=FONT,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=False, wrap=True):
    tb = s.shapes.add_textbox(int(x), int(y), int(w), int(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    tf.vertical_anchor = anchor
    lines = txt.split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = ln
        f = r.font
        f.size = Pt(size)
        f.bold = bold
        f.italic = italic
        f.name = font
        f.color.rgb = color
        if spacing:
            rPr = r._r.get_or_add_rPr()
            rPr.set("spc", "150")
    return tb


def bullets(s, x, y, w, h, items, size=15, color=INK, bullet_color=TEAL, sub_color=BODY):
    """items: str | (str, level) | (str, level, {'bold':..,'color':..}). level 0/1."""
    tb = s.shapes.add_textbox(int(x), int(y), int(w), int(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    first = True
    for it in items:
        if isinstance(it, str):
            it = (it, 0)
        txt, level = it[0], it[1]
        opts = it[2] if len(it) > 2 else {}
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(8 if level == 0 else 4)
        p.level = 0
        indent = "        " if level else ""
        r0 = p.add_run()
        r0.text = indent + ("▪  " if level == 0 else "–  ")
        r0.font.size = Pt(size if level == 0 else size - 2)
        r0.font.color.rgb = bullet_color if level == 0 else MUTED
        r0.font.name = FONT
        r0.font.bold = True
        r = p.add_run()
        r.text = txt
        r.font.size = Pt(opts.get("size", size if level == 0 else size - 2))
        r.font.bold = opts.get("bold", False)
        r.font.italic = opts.get("italic", False)
        r.font.name = FONT
        r.font.color.rgb = opts.get("color", color if level == 0 else sub_color)
    return tb


TONES = {
    "grey": (GREY_BG, MUTED, NAVY),
    "teal": (TEAL_BG, TEAL, TEAL_D),
    "blue": (BLUE_BG, BLUE, BLUE),
    "amber": (AMBER_BG, AMBER, AMBER),
    "navy": (NAVY, TEAL, WHITE),
    "white": (WHITE, LINE, NAVY),
}


def card(s, x, y, w, h, head, body, tone="grey", body_size=13):
    bg, accent, headcol = TONES[tone]
    rect(s, x, y, w, h, fill=bg, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.06)
    rect(s, x, y, Pt(4), h, fill=accent)
    pad = Inches(0.22)
    text(s, x + pad, y + Inches(0.15), w - 2 * pad, Inches(0.55), head, size=15,
         bold=True, color=headcol)
    body_col = PALE if tone == "navy" else BODY
    if isinstance(body, (list, tuple)):
        bullets(s, x + pad, y + Inches(0.7), w - 2 * pad, h - Inches(0.8), body,
                size=body_size, color=body_col, sub_color=body_col,
                bullet_color=accent)
    else:
        text(s, x + pad, y + Inches(0.7), w - 2 * pad, h - Inches(0.8), body,
             size=body_size, color=body_col)


def panel(s, x, y, w, h, head, items, tone="grey", size=15):
    bg, accent, headcol = TONES[tone]
    rect(s, x, y, w, h, fill=bg, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.04)
    rect(s, x, y, w, Inches(0.6), fill=accent, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
         adj=0.25)
    rect(s, x, y + Inches(0.3), w, Inches(0.3), fill=accent)  # square off bottom
    text(s, x + Inches(0.25), y + Inches(0.12), w - Inches(0.5), Inches(0.4), head,
         size=15, bold=True, color=WHITE)
    bullets(s, x + Inches(0.25), y + Inches(0.8), w - Inches(0.5), h - Inches(0.95),
            items, size=size, bullet_color=accent)


def kpi_tile(s, x, y, w, h, big, label, sub):
    rect(s, x, y, w, h, fill=WHITE, line=LINE, lw=1, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
         adj=0.08)
    rect(s, x, y + h - Pt(4), w, Pt(4), fill=TEAL)
    big_size = 32 if len(big) <= 8 else (26 if len(big) <= 11 else 21)
    text(s, x + Inches(0.2), y + Inches(0.15), w - Inches(0.4), Inches(0.75), big, size=big_size,
         bold=True, color=NAVY, font=FONT_H, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + Inches(0.2), y + Inches(0.9), w - Inches(0.4), Inches(0.35), label, size=13,
         bold=True, color=TEAL_D)
    text(s, x + Inches(0.2), y + Inches(1.22), w - Inches(0.4), Inches(0.5), sub, size=10.5,
         color=MUTED)


def arrow(s, x, y_center, length):
    a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, int(x), int(y_center - Inches(0.11)),
                           int(length), int(Inches(0.22)))
    a.fill.solid()
    a.fill.fore_color.rgb = PALE
    a.line.fill.background()
    a.shadow.inherit = False
    return a


def flow(s, x, y, w, h, steps):
    n = len(steps)
    gap = Inches(0.32) if n <= 5 else Inches(0.26)
    bw = (w - gap * (n - 1)) / n
    sub_size = 10.5 if n <= 5 else 10
    for i, (label, sub, kind) in enumerate(steps):
        accent, bg, _ = KIND_COLOR[kind]
        bx = x + i * (bw + gap)
        rect(s, bx, y, bw, h, fill=bg, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.1)
        rect(s, bx, y, bw, Pt(5), fill=accent)
        text(s, bx, y + Inches(0.15), bw, Inches(0.3), f"{i + 1}", size=11, bold=True,
             color=accent, align=PP_ALIGN.CENTER)
        text(s, bx + Inches(0.08), y + Inches(0.42), bw - Inches(0.16), Inches(0.5), label,
             size=13, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        text(s, bx + Inches(0.1), y + Inches(0.9), bw - Inches(0.2), h - Inches(0.95), sub,
             size=sub_size, color=BODY, align=PP_ALIGN.CENTER)
        if i < n - 1:
            arrow(s, bx + bw + Inches(0.04), y + h / 2, gap - Inches(0.08))


def legend_row(s, x, y, kinds):
    cx = x
    for k in kinds:
        accent, bg, label = KIND_COLOR[k]
        rect(s, cx, y + Inches(0.08), Inches(0.22), Inches(0.22), fill=accent,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.3)
        text(s, cx + Inches(0.3), y + Inches(0.03), Inches(2.4), Inches(0.35), label,
             size=11, color=BODY)
        cx += Inches(2.6)


def takeaway_bar(s, txt):
    y = BOTTOM - Inches(0.7)
    rect(s, MARGIN, y, CONTENT_W, Inches(0.7), fill=NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
         adj=0.2)
    rect(s, MARGIN + Inches(0.25), y + Inches(0.2), Pt(4), Inches(0.3), fill=TEAL)
    text(s, MARGIN + Inches(0.45), y, CONTENT_W - Inches(0.7), Inches(0.7), txt, size=14,
         color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)


def table(s, x, y, w, h, headers, rows, col_widths=None, size=12, highlight_last=False):
    nrows, ncols = len(rows) + 1, len(headers)
    shp = s.shapes.add_table(nrows, ncols, int(x), int(y), int(w),
                             int(min(h, Inches(0.62) * nrows)))
    tbl = shp.table
    if col_widths:
        tot = sum(col_widths)
        for i, cw in enumerate(col_widths):
            tbl.columns[i].width = int(w * cw / tot)
    row_h = int(min(h / nrows, Inches(0.62)))
    for r in range(nrows):
        tbl.rows[r].height = row_h

    def _cell(c, txt, bold=False, color=INK, fill=None, sz=size, align=PP_ALIGN.LEFT):
        c.text = ""
        tf = c.text_frame
        tf.word_wrap = True
        c.margin_left = c.margin_right = Inches(0.08)
        c.margin_top = c.margin_bottom = Inches(0.04)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        lines = str(txt).split("\n")
        for i, ln in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            rr = p.add_run()
            rr.text = ln
            rr.font.size = Pt(sz)
            rr.font.bold = bold
            rr.font.name = FONT
            rr.font.color.rgb = color
        if fill is not None:
            c.fill.solid()
            c.fill.fore_color.rgb = fill

    for ci, hd in enumerate(headers):
        _cell(tbl.cell(0, ci), hd, bold=True, color=WHITE, fill=NAVY, sz=size)
    for ri, row in enumerate(rows, start=1):
        last = highlight_last and ri == nrows - 1
        for ci, val in enumerate(row):
            fill = TEAL_BG if last else (WHITE if ri % 2 else PAPER)
            _cell(tbl.cell(ri, ci), val, bold=(ci == 0 or last),
                  color=(TEAL_D if last else INK), fill=fill)
    # kill the default table style banding look
    tblPr = shp._element.graphic.graphicData.tbl.tblPr
    tblPr.set("bandRow", "0")
    tblPr.set("firstRow", "0")
    return shp


# ---------------------------------------------------------------------------
# Content: FRD → STTM Agent — ACFC overview deck
# Facts sourced from README.md, CLAUDE.md, docs/LIVE_E2E_2026-08-07.md, docs/DEMO_RUNBOOK.md,
# docs/NATIVE_REBUILD_SPEC.md, reports/demo_frd.phase5.md (repo state 2026-08-10) and the
# sibling model-comparison-eval repo (SLIDES_INPUT.md, README.md, results/) for model economics.
# ---------------------------------------------------------------------------
import sys
from pathlib import Path

OUT = (Path(sys.argv[1]) if len(sys.argv) > 1
       else Path(__file__).resolve().parent.parent / "docs" / "frd-to-sttm-agent_overview_deck.pptx")

d = Deck("FRD → STTM Agent")

# 1 ── cover
d.cover(
    "FRD → STTM Agent",
    "From an approved Functional Requirements Document to the Source-to-Target Mapping workbook — "
    "in about thirty seconds, with every value checked back against the document.",
    ["Prepared for AmeriHealth Caritas Family of Companies (ACFC)",
     "Hexaware · AI-in-Engineering program · Agent 2 of 5 · August 2026",
     "Status: live end-to-end runs on an anonymised FRD (94.1% cell match vs. golden workbook); Databricks bundle built, Apps deploy pending"],
    notes="Opening. What the FRD→STTM agent does, why ACFC would want it, measured results from the "
          "7 Aug 2026 live runs, expected benefits, and an explicitly labelled illustrative savings model. "
          "All measurements are on a scrubbed practice document — a real run, not real client data.")

# 2 ── one sentence
d.statement(
    "What it does — in one sentence",
    "It reads an approved FRD, extracts every feed, file pattern, target table and rule with one Claude "
    "call, verifies every extracted value against the document text, asks a person about anything "
    "ambiguous, and renders the STTM workbook plus a machine-readable feed contract for the CodeGen agent.",
    "Input: FRD .docx (from the BRD→FRD agent or hand-authored).\n"
    "Output: STTM workbook .xlsx (FILE_DETAILS, VERSION_HISTORY, one MAPPING sheet per target table) + "
    "<doc_id>.contract.json + run report + cell-level score against a reference workbook.\n"
    "Exactly one LLM call in the middle; everything else is deterministic code.",
    notes="'The reading and the typing — not the thinking.' The agent copies out what the document says, "
          "checks every value back against the text, fills the workbook, and stops to ask when the FRD is unclear.")

# 3 ── problem
d.two_col(
    "Why ACFC needs it",
    "STTM authoring today",
    ["A mapping analyst reads the whole FRD, cross-checks the data dictionary and types every row — days per document",
     "Two analysts produce two different spreadsheets; names and rules drift between feeds and teams",
     "No record of which sentence a cell came from — hard to audit later",
     "It holds everything up: pipeline code, tests and reviews all wait on the workbook",
     "Hundreds of columns per feed (the demo FRD maps 411) — transcription errors are inevitable"],
    "With the FRD → STTM agent",
    ["Workbook drafted in ~33 s for a 3-feed FRD; the analyst reviews a short list of flagged items",
     "Same derivation rules every run — consistency is code, not habit",
     "Every extracted value is grounded: strict fields must appear verbatim, advisory fields ≥ 75% overlap",
     "Feed contract JSON is emitted alongside — CodeGen can start the moment the workbook is approved",
     "0 of 50 sampled identifiers invented; ≈ 96% of extracted field values kept as-is by a proxy SME"],
    kicker="The business problem",
    takeaway="Days of transcription become a 30-second run plus a short human review — and every cell can be traced back to the sentence it came from.",
    notes="ACFC's actual analyst effort per STTM is not in the repo; 'it takes days' is the deck's framing. "
          "Ask the client for their number to calibrate slide 9.")

# 4 ── how it works
d.flow_slide(
    "How it works — four notebooks, one LLM call",
    [("01 INGEST", ".docx → normalised markdown, requirement IDs preserved → frd_documents Delta table", "code"),
     ("02 EXTRACT", "One Claude call: schema-in-prompt, streamed JSON, validated client-side — truncation or refusal is a failure, never a guess", "llm"),
     ("03 CONTRACT BUILD", "Pydantic validation, regex enrichment, grounding audit (strict verbatim + advisory overlap), ambiguity gating → contract.json", "code"),
     ("REVIEW", "Analyst resolves gated ambiguities in the review app: pick a candidate, none-of-these, or free text", "human"),
     ("04 RENDER + EVAL", "Apply resolutions → dictionary cross-check → derive stage/standard mappings → .xlsx + cell-level score", "code")],
    kicker="LLM proposes, code verifies, human resolves",
    below=["Gate vocabulary shared across the program: FAIL (any strict grounding miss — no contract emitted) / PASS WITH FLAGS (advisory flags or ambiguities gated for review) / PASS",
           "Three ambiguity kinds only — attribution, disagreement, advisory grounding — each with a stable ID that survives re-runs so human decisions are never lost",
           "Runs as a Databricks Asset Bundle job (serverless) or the same four scripts on a laptop; artifacts in Unity Catalog volumes; key in a secret scope"],
    takeaway="The AI only reads. Plain, predictable code does everything else — and never writes a cell it cannot back up with the document.",
    notes="Model today: claude-opus-4-8 via the Anthropic SDK; the eval repo recommends serving on "
          "claude-sonnet-4-6 (same gate profile, half the cost). Server-side structured output was "
          "rejected by the API for this schema size, so the schema is rendered into the prompt and "
          "validated client-side with extra='forbid'.")

# 5 ── trust
d.cards(
    "Built for a regulated environment",
    [("Nothing invented",
      ["Strict fields (IDs, file patterns, tables, LOBs) must appear verbatim in the FRD", "45/45 strict + 30/30 advisory checks passed on the live run", "'Never relax the audit to make a run pass'"], "teal"),
     ("When unsure, it asks",
      ["Same rule on two feeds → attribution ambiguity with candidates", "Model vs regex disagreement → gated, not guessed", "Dictionary confirms what the data proves; the rest goes to a person"], "blue"),
     ("A person has the last word",
      ["Review app: candidate pick / none-of-these / free text", "Resolutions applied first and authoritatively at render", "Every application audited: applied / not applied / stale"], "amber"),
     ("Data stays in your workspace",
      ["UC volumes frd_raw / sttm_reference / sttm_out; Delta tables for documents, contracts, runs", "API key in a Databricks secret scope", "No real client documents ever uploaded during build"], "grey"),
     ("Fail loud, replay free",
      ["No silent retries; max_tokens or refusal raises", "Every live run saved as a replay set — demo without a key", "77 offline tests, no LLM / network / Spark"], "grey"),
     ("Scored, not asserted",
      ["Cell-level eval vs a golden workbook on every run", "Hallucination sweep: 0/50 identifiers ungrounded", "Gate shape may vary run to run; mapping quality did not"], "grey")],
    kicker="Governance by construction",
    cols=3, body_size=12,
    notes="The reference workbook plays two separated roles: its source-layout sections stand in for the "
          "vendor source dictionary (a production input); its stage/standard sections are ground truth for "
          "eval only and never an input to generation.")

# 6 ── where it fits
d.program_map(
    "Where it fits in the AI-in-Engineering pipeline", "sttm",
    kicker="Second of five agents",
    below=["Upstream is solved: the shared, versioned label contract with the BRD→FRD agent is byte-identical in both repos and tested for drift",
           "Downstream: <doc_id>.contract.json is consumed by CodeGen as its FRD feed contract today — the workbook-derived STTM mapping contract still needs a committed extractor (the program's biggest open gap)",
           "Run rows in frd_contracts / frd_sttm_runs give the orchestrator its status; humans approve before CodeGen is armed"],
    notes="Be candid about the CodeGen mapping-contract gap: half the hand-off works today, the other "
          "half was produced out-of-repo for the fixtures.")

# 7 ── measured
d.kpis(
    "What we have measured so far",
    [("94.1%", "cell-level match vs. golden", "3,094 / 3,288 target cells; all 194 misses are planted datatype divergences"),
     ("0 / 50", "identifiers invented", "Hallucination sweep; 45/45 strict + 30/30 advisory grounding"),
     ("~33 s · ~$0.15", "per FRD, one billed call", "14.6k in / 3.3k out tokens on claude-opus-4-8"),
     ("≈ 96%", "first-pass keep rate", "118 field values, ~4 proxy-SME edits — vs ≥ 80% acceptance bar")],
    kicker="Live end-to-end runs, 2026-08-07 (anonymised 3-feed FRD, 411 mapped columns)",
    below=["Three live runs produced three different gate shapes (1 blocking → 1 auto-confirmed → 0) at an identical 94.1% and 45/45 strict — non-determinism shows up in what is asked, not in what is mapped",
           "Output-token headroom: ~5% of the 64k budget used → ~75–90 feeds per document before chunking is needed",
           "Serving-model study (model-comparison-eval): cell accuracy is model-invariant; Sonnet 4.6 matches Opus 4.8's gate profile at ≈ $0.08/call vs $0.16 — Haiku is 6× cheaper but fabricated a catalog name and is not production-achievable"],
    source="docs/LIVE_E2E_2026-08-07.md, reports/demo_frd.phase5.md, docs/DEMO_RUNBOOK.md; model economics from model-comparison-eval/README.md and results/cost_projection.csv (n = 2 document families).",
    takeaway="A real run on a real-shaped document: nothing invented, ~94% of the workbook right first time, for 15 cents and half a minute.",
    notes="Never present a bare '100%' — cell accuracy on the demo pair is bounded by the derivation ceiling. "
          "The 194 misses are the fixture's deliberate string-vs-decimal divergences. Eval-repo caveats: "
          "n=2 document pairs; shipped harness was iterated on both.")

# 8 ── benefits
d.cards(
    "Expected benefits for ACFC",
    [("Days → minutes", "The workbook is drafted in ~30 s; the analyst's day goes to the handful of flagged items and QA.", "teal"),
     ("Easy to audit", "Every value points back to the sentence it came from; when an auditor asks 'why is this mapped that way?', the answer is already written down.", "teal"),
     ("Same result every time", "Derivation rules are code — no drift between analysts, feeds or teams.", "blue"),
     ("Fewer downstream defects", "A grounded, contract-conformant mapping means CodeGen builds from facts, not transcription errors.", "blue"),
     ("Your data stays with you", "Runs inside your Databricks workspace: UC volumes, secret scope, serverless job — nothing leaves.", "amber"),
     ("Predictable, tiny run cost", "~$0.15 per FRD today; ≈ $0.08 on the recommended serving model. Replays cost nothing.", "amber")],
    kicker="Why it is worth doing",
    cols=3, body_size=13,
    notes="Frame the review load honestly: on the demo document, at most a handful of gated items per run, often zero.")

# 9 ── savings model
d.table_slide(
    "Time & cost savings — illustrative model",
    ["Per FRD (≈3 feeds, ≈400 columns)", "Manual today (assumed)", "With agent (assumed)", "Saving"],
    [["Read FRD, identify feeds, files, tables, rules", "4 h", "0 h (extracted)", "4 h"],
     ["Type mapping rows: source / stage / standard for every column", "12 h", "0 h (rendered)", "12 h"],
     ["Cross-check against the source data dictionary", "3 h", "0 h (automatic)", "3 h"],
     ["Resolve gated ambiguities", "—", "0.5 h (0–1 items on the demo doc)", "—"],
     ["Analyst QA of the workbook", "3 h", "2 h", "1 h"],
     ["LLM cost", "—", "≈ $0.15 (measured; ≈ $0.08 on Sonnet 4.6)", "—"],
     ["Total per FRD", "22 h", "2.5 h + $0.15", "≈ 19.5 h (~89%)"]],
    kicker="Assumptions are placeholders — replace with ACFC's analyst effort and document volume",
    col_widths=[3.2, 1.6, 2.4, 1.4], size=12, highlight_last=True,
    source="Effort figures are illustrative assumptions, not measurements ('it takes days' is the only baseline stated). Only run time, LLM cost and review-item counts are measured (docs/LIVE_E2E_2026-08-07.md). "
           "Example scale-up: 3 FRDs/month × 19.5 h ≈ 700 h/yr; at an illustrative $100/h ≈ $70k/yr vs ≈ $5/yr of LLM spend.",
    takeaway="The LLM bill is a rounding error; the case rests on analyst hours and on downstream defects avoided. One real FRD in a pilot would replace every assumption here.",
    notes="Invite the client to replace each figure. QA time is deliberately kept: the agent removes typing, not accountability.")

# 10 ── deployment
d.two_col(
    "Deployment & integration",
    "What exists today",
    ["Databricks Asset Bundle: job frd_sttm_pipeline (ingest → extract → contract_build → render), serverless, manual trigger",
     "Unity Catalog layout: volumes frd_raw / sttm_reference / sttm_out; Delta tables frd_documents, frd_contracts, frd_sttm_runs; secret scope for the API key",
     "Review app (FastAPI + React, Databricks AppKit UI): existing-documents review, new upload, client demo with live runs and zero-call replay",
     "Local mode with the same four scripts and a mock provider — 77 offline tests",
     "Native rebuild spec written: Lakeflow Declarative Pipeline + Genie Code + Model Serving / AI Gateway"],
    "Pending ACFC decisions / work",
    ["Databricks Apps deployment of the review app (workspace access was blocked at build time)",
     "Wire the app's live run to the bundle job; API key via Apps secret; SSE through the Apps proxy",
     "Model Serving / AI Gateway endpoint (Claude via Foundation Model API or external model) instead of the SDK client",
     "Committed workbook → STTM-mapping-contract extractor for CodeGen",
     "A production source-dictionary artifact (the reference workbook stands in today)"],
    left_tone="teal", right_tone="grey", size=14,
    notes="'Built for the workspace, shown on a laptop.' The three plumbing items are small; the "
          "extractor for CodeGen is the one real engineering gap.")

# 11 ── limitations
d.two_col(
    "Known gaps & roadmap",
    "Honest limitations today",
    ["Tested at scale on one 3-feed anonymised FRD (plus a second document family in the eval repo)",
     "Datatypes default to String at both layers — an analyst decision the FRD does not state (the 194 planted misses)",
     "Gate shape varies run to run; demos must not depend on a specific count (replay sets are the safe path)",
     "Anthropic SDK today; Model Serving port is a spec, not code",
     "Test count and max_tokens in some docs/YAML are stale (77 tests; 64k default)"],
    "Proposed next steps",
    ["Pilot on one real (non-sensitive) FRD at full size — measure keep rate and analyst hours",
     "Databricks wiring: job trigger, key via secret store, live progress through Apps",
     "Build the workbook → mapping-contract extractor and close the CodeGen hand-off",
     "Move serving to Sonnet 4.6 via Model Serving / AI Gateway per the eval findings",
     "Chunked per-feed extraction for very large FRDs"],
    left_tone="amber", right_tone="teal", size=14,
    notes="Say these out loud; they are the asks. Eval-repo caveats apply to any model claim: n=2 pairs, "
          "shipped harness not held-out.")

# 12 ── closing
d.closing(
    "What we need from ACFC",
    ["One real (non-sensitive) FRD to test at full size, and its hand-built STTM for scoring",
     "A workspace to install into: catalog/schema, volumes, secret scope, Apps approval",
     "The production source data dictionary format",
     "Baseline numbers: STTMs per month, analyst hours per STTM — to replace the illustrative savings model",
     "A named mapping analyst to own ambiguity resolution in the pilot"],
    notes="Close on the pilot ask. Cost of the pilot in LLM spend is under a dollar.")

d.save(OUT)
print(f"wrote {OUT} ({d.n} slides)")
