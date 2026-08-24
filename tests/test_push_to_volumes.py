"""tools/push_local_source_to_volumes.py — local folder → Unity Catalog volumes.

Offline: the WorkspaceClient is a stub, so no workspace, no credential, no
upload. What matters and is tested here is that the tool reuses the pipeline's
own staging semantics (prefix split, pairing, corpus index) rather than
reimplementing them, and that it never creates volumes for an empty push.

All documents are SYNTHETIC, per the no-client-documents rule.
"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest
from openpyxl import Workbook

TOOLS = Path(__file__).resolve().parent.parent / "tools"
if str(TOOLS) not in sys.path:
    sys.path.insert(0, str(TOOLS))

import push_local_source_to_volumes as push  # noqa: E402

from frdsttm.corpus import CORPUS_INDEX_NAME  # noqa: E402


def _docx(path, paragraphs):
    from docx import Document

    d = Document()
    for p in paragraphs:
        d.add_paragraph(p)
    d.save(str(path))


def _xlsx(path, rows):
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    wb.save(str(path))


@pytest.fixture
def source(tmp_path):
    src = tmp_path / "frd-to-sttm-agent-documents"
    src.mkdir()
    _docx(src / "FRD_ClaimsFeed.docx",
          ["Claims feed", "Source table claims_raw", "Target CLAIM_FACT",
           "Claims arrive daily from the vendor extract."])
    _xlsx(src / "STTM_ClaimsFeed.xlsx",
          [["Source Table", "Source Column", "Target Table", "Target Column"],
           ["claims_raw", "clm_id", "CLAIM_FACT", "CLAIM_ID"]])
    return src


class FakeVolumes:
    def __init__(self, existing=()):
        self.existing = list(existing)
        self.created = []

    def list(self, catalog_name, schema_name):
        return [type("V", (), {"name": n})() for n in self.existing]

    def create(self, catalog_name, schema_name, name, volume_type, comment=None):
        self.created.append(name)
        self.existing.append(name)


class FakeFiles:
    def __init__(self):
        self.uploaded = []

    def upload(self, file_path, contents, overwrite=None):
        self.uploaded.append((file_path, len(contents.read())))


class FakeWorkspace:
    def __init__(self, existing=()):
        self.volumes = FakeVolumes(existing)
        self.files = FakeFiles()


# --------------------------------------------------------------------------- #
# staging — the part that must match the app exactly
# --------------------------------------------------------------------------- #
def test_staging_splits_pairs_and_builds_the_index(source, tmp_path):
    staging = tmp_path / "staging"
    result = push.stage(source, staging)

    assert result["frd_downloaded"] == 1
    assert result["reference_downloaded"] == 1
    assert list(result["index"]["pairs"]) == ["FRD_ClaimsFeed"]
    assert result["index"]["pairs"]["FRD_ClaimsFeed"]["matched_by"] == "name"
    # the index is what lets a deployed App show the corpus with no job running
    assert (staging / "sttm_reference" / CORPUS_INDEX_NAME).is_file()


def test_files_outside_the_convention_are_counted_not_uploaded(source, tmp_path):
    (source / "notes.docx").write_bytes((source / "FRD_ClaimsFeed.docx").read_bytes())
    result = push.stage(source, tmp_path / "staging")
    assert result["frd_ignored"] == 1
    assert result["frd_downloaded"] == 1


# --------------------------------------------------------------------------- #
# volumes
# --------------------------------------------------------------------------- #
def test_missing_volumes_are_created_existing_ones_left_alone(capsys):
    w = FakeWorkspace(existing=["frd_raw"])
    created = push.ensure_volumes(w, "cat", "sch", ["frd_raw", "sttm_reference"],
                                  dry_run=False)
    assert created == ["sttm_reference"]
    assert w.volumes.created == ["sttm_reference"]


def test_dry_run_creates_nothing():
    w = FakeWorkspace(existing=[])
    created = push.ensure_volumes(w, "cat", "sch", ["frd_raw"], dry_run=True)
    assert created == ["frd_raw"]      # reported...
    assert w.volumes.created == []     # ...but not actually created


def test_dry_run_uploads_nothing(source, tmp_path):
    staging = tmp_path / "staging"
    push.stage(source, staging)
    w = FakeWorkspace()
    n = push.upload_dir(w, staging / "frd_raw", "/Volumes/c/s/frd_raw", dry_run=True)
    assert n == 1 and w.files.uploaded == []


def test_upload_sends_every_staged_file_flat(source, tmp_path):
    staging = tmp_path / "staging"
    push.stage(source, staging)
    w = FakeWorkspace()
    push.upload_dir(w, staging / "sttm_reference", "/Volumes/c/s/sttm_reference",
                    dry_run=False)
    names = [p.rsplit("/", 1)[1] for p, _ in w.files.uploaded]
    assert "STTM_ClaimsFeed.xlsx" in names
    assert CORPUS_INDEX_NAME in names          # the index travels with the workbooks
    assert all(size > 0 for _, size in w.files.uploaded)


# --------------------------------------------------------------------------- #
# refusals
# --------------------------------------------------------------------------- #
def test_empty_folder_refuses_rather_than_creating_volumes(tmp_path, monkeypatch, capsys):
    """A folder whose files don't match the convention must not silently
    provision empty volumes — the operator has a naming problem to fix."""
    empty = tmp_path / "docs"
    empty.mkdir()
    (empty / "readme.txt").write_text("not an FRD")
    monkeypatch.setenv("STTM_LOCAL_SOURCE_DIR", str(empty))

    called = {}
    monkeypatch.setattr(push, "ensure_volumes",
                        lambda *a, **k: called.setdefault("ran", True))

    rc = push.main([str(empty)])

    assert rc == 1
    assert "ran" not in called


def test_a_missing_source_folder_exits_without_touching_the_workspace(tmp_path, monkeypatch):
    monkeypatch.delenv("STTM_LOCAL_SOURCE_DIR", raising=False)
    called = {}
    monkeypatch.setattr(push, "ensure_volumes",
                        lambda *a, **k: called.setdefault("ran", True))
    rc = push.main([str(tmp_path / "nope")])
    assert rc == 2
    assert "ran" not in called
