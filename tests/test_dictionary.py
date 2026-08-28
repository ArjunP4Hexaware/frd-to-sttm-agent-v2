import pytest
from openpyxl import Workbook

from conftest import make_vdd
from frdsttm.dictionary import DictionaryError, parse_dictionary_dir, parse_dictionary_workbook, source_layout


def test_parse_vdd(tmp_path):
    d = parse_dictionary_workbook(make_vdd(tmp_path / "VDD_x.xlsx"))
    assert d["n_files"] == 1 and d["n_fields"] == 3
    layout = source_layout(d)
    cols = layout["claims_YYYYMMDD.csv"]
    assert [c["name"] for c in cols] == ["CLAIM_ID", "MEMBER_ID", "AMOUNT"]
    assert cols[0]["required"] is True and cols[2]["required"] is False
    assert cols[1]["phi"] is True and cols[0]["phi"] is False
    assert cols[0]["position"] == 1


def test_template_example_rows_are_dropped_and_reported(tmp_path):
    d = parse_dictionary_workbook(make_vdd(tmp_path / "VDD_x.xlsx", marker_rows=True))
    assert d["n_files"] == 1
    assert any(p["kind"] == "template_example_rows" for p in d["problems"])


def test_missing_field_sheet_gates_not_raises(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "FILES"
    ws.append(["File Name Pattern", "File Title", "Format", "Delimiter", "Header Row", "Encoding",
               "Delivery Cadence", "Content Description", "Field Sheet"])
    ws.append(["a.csv", "A", "csv", ",", "Y", "UTF-8", "Daily", "d", "nope"])
    p = tmp_path / "VDD_a.xlsx"
    wb.save(str(p))
    d = parse_dictionary_workbook(p)
    assert d["n_files"] == 1 and d["n_fields"] == 0
    assert {q["kind"] for q in d["problems"]} == {"field_sheet_missing"}


def test_no_files_sheet_raises(tmp_path):
    wb = Workbook()
    wb.active.title = "Sheet1"
    p = tmp_path / "VDD_bad.xlsx"
    wb.save(str(p))
    with pytest.raises(DictionaryError):
        parse_dictionary_workbook(p)


def test_blank_flags_are_unknown_not_false(tmp_path):
    p = make_vdd(tmp_path / "VDD_x.xlsx", {"f.csv": [("A", "int", "", "", "d", "", "", "")]})
    d = parse_dictionary_workbook(p)
    f = d["fields"]["file1"][0]
    assert f["required"] is None and f["phi"] is None
    assert {q["kind"] for q in d["problems"]} >= {"missing_required_flags", "missing_phi_flags"}


def test_parse_dir_isolates_errors(tmp_path):
    make_vdd(tmp_path / "VDD_ok.xlsx")
    wb = Workbook()
    wb.active.title = "nothing"
    wb.save(str(tmp_path / "VDD_bad.xlsx"))
    out = parse_dictionary_dir(tmp_path)
    assert set(out["dictionaries"]) == {"VDD_ok.xlsx"} and set(out["errors"]) == {"VDD_bad.xlsx"}
