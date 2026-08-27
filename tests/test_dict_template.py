"""The ISSUED vendor data dictionary template must parse as nothing.

`tests/test_dictionary.py` pins the returned-template trap on synthetic
workbooks. This file pins it on the real thing: it runs
`scripts/build_dict_template.py` and parses what the script wrote.

Why it exists: the template's example rows are written by position and the
parser recognises an example only by the marker in the LAST column. Adding a
column without its value shifted the marker one cell left, un-marked the
row, and the blank template parsed as a real file (v1.3). The script now
keys rows by header and self-checks after saving; this test is the guard
that the guard is still there.

Run from the repo root — the script resolves `src/` relative to itself.
"""

from __future__ import annotations

import subprocess
import sys
from pathlib import Path

import pytest

from frdsttm.dictionary import TEMPLATE_EXAMPLE_MARKER, parse_dictionary_workbook

REPO = Path(__file__).resolve().parent.parent
SCRIPT = REPO / "scripts" / "build_dict_template.py"


@pytest.fixture(scope="module")
def built_template(tmp_path_factory) -> Path:
    out = tmp_path_factory.mktemp("template") / "DICT_TEMPLATE.xlsx"
    proc = subprocess.run([sys.executable, str(SCRIPT), str(out)],
                          cwd=REPO, capture_output=True, text=True)
    assert proc.returncode == 0, proc.stderr
    assert out.exists()
    return out


def test_blank_template_parses_as_zero_files(built_template):
    """No FILES example row may survive as a real file."""
    d = parse_dictionary_workbook(built_template)
    assert d["n_files"] == 0, [f["file_name_pattern"] for f in d["files"]]
    assert d["n_fields"] == 0


def test_every_files_example_row_is_dropped_and_counted(built_template):
    """The parser must report the examples it removed, and the count must
    equal the number of example rows the FILES sheet actually carries."""
    from openpyxl import load_workbook

    ws = load_workbook(built_template, read_only=True)["FILES"]
    rows = list(ws.iter_rows(values_only=True))
    notes_col = rows[0].index("Notes")
    marked = sum(1 for r in rows[1:]
                 if r[0] and TEMPLATE_EXAMPLE_MARKER in str(r[notes_col] or "").lower())
    unmarked = [r[0] for r in rows[1:]
                if r[0] and TEMPLATE_EXAMPLE_MARKER not in str(r[notes_col] or "").lower()]
    assert not unmarked, f"FILES example rows without the marker: {unmarked}"
    assert marked >= 2

    d = parse_dictionary_workbook(built_template)
    dropped = [p for p in d["problems"]
               if p["kind"] == "template_example_rows" and p["file"] is None]
    assert len(dropped) == 1
    assert dropped[0]["detail"].startswith(f"{marked} FILES row")


def test_marker_is_in_the_notes_column_not_elsewhere(built_template):
    """The trap this file guards against: the marker one column to the left.

    If the marker ever appears in a column other than Notes, a header was
    added without its value and every later cell shifted."""
    from openpyxl import load_workbook

    wb = load_workbook(built_template, read_only=True)
    for name in wb.sheetnames:
        rows = list(wb[name].iter_rows(values_only=True))
        if not rows or "Notes" not in rows[0]:
            continue
        notes_col = rows[0].index("Notes")
        for r in rows[1:]:
            for i, v in enumerate(r):
                if i != notes_col and TEMPLATE_EXAMPLE_MARKER in str(v or "").lower():
                    pytest.fail(f"{name}: marker found in column {rows[0][i]!r}, "
                                f"not in Notes — a column shifted")
