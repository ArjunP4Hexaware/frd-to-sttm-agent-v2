from openpyxl import load_workbook

from conftest import make_reference_sheet_per_table, make_reference_single_sheet, make_vdd, spec_for
from frdsttm import completeness as c
from frdsttm import render
from frdsttm.dictionary import parse_dictionary_workbook
from frdsttm.reference_layout import layout_of


def _ready(tmp_path, spec=None, vdd_files=None):
    spec = spec or spec_for()
    vdd = parse_dictionary_workbook(make_vdd(tmp_path / "v.xlsx", vdd_files))
    sources, _ = c.derive_targets(spec)
    pairing, _ = c.pair_files(spec, vdd)
    for s in sources:
        f = pairing.get(s["feed_index"])
        s["file"], s["field_sheet"] = (f["file_name_pattern"], f["field_sheet"]) if f else (None, None)
    return spec, vdd, sources


def test_build_rows_as_is_names_promoted_types_and_audit_rows(tmp_path):
    spec, vdd, sources = _ready(tmp_path)
    rows, notes = render.build_rows(sources[0], spec["feeds"][0], vdd["fields"]["file1"])
    data = [r for r in rows if not r["audit"]]
    assert [r["stage"]["column"] for r in data] == ["CLAIM_ID", "MEMBER_ID", "AMOUNT"]
    assert [r["standard"]["datatype"] for r in data] == ["Int", "String", "Decimal(10,2)"]
    assert all(r["stage"]["datatype"] == "String" for r in data)
    assert data[0]["stage"]["table"] == "clm_claims_stg" and data[0]["standard"]["table"] == "clm_claims"
    assert data[0]["standard"]["schema"] == "CLM" and data[0]["stage"]["catalog"] == "PR_DLK"
    audit = [r for r in rows if r["audit"]]
    assert [r["stage"]["column"] for r in audit] == ["SRC_FILE_NAME", "REC_CREATION_TIME", "REC_UPDATED_TIME", "LOB"]
    assert notes["n_audit_rows"] == 4


def test_segment_routing_and_no_lob_on_control_records(tmp_path):
    spec = spec_for(stage_table="X_HDR")
    spec["feeds"][0]["stage_target"]["tables"] = ["X_HDR", "X_DTL", "X_TRL"]
    spec["feeds"][0]["standard_target"]["tables"] = ["X_HDR", "X_DTL", "X_TRL"]
    files = {"claims_YYYYMMDD.csv": [("H1", "varchar", "Y", "N", "d", "", "", "Header"),
                                     ("D1", "int", "Y", "N", "d", "", "", "Detail"),
                                     ("T1", "int", "Y", "N", "d", "", "", "Trailer")]}
    spec, vdd, sources = _ready(tmp_path, spec, files)
    rows, notes = render.build_rows(sources[0], spec["feeds"][0], vdd["fields"]["file1"])
    by = {(r["source_column"], r["segment"]): r for r in rows if not r["audit"]}
    assert by[("H1", "Header")]["stage"]["table"] == "X_HDR" and by[("T1", "Trailer")]["stage"]["table"] == "X_TRL"
    lob_tables = {r["stage"]["table"] for r in rows if r["audit"] and r["stage"]["column"] == "LOB"}
    assert lob_tables == {"X_DTL"}


def test_rule_placement_names_the_column():
    feed = spec_for()["feeds"][0]
    feed["validation_rules"].append("Something with no column name in it.")
    feed["recycle_rule"] = "Recycle when MEMBER_ID is unknown for 7 days"
    rows = [{"source_column": n, "audit": False} for n in ("CLAIM_ID", "MEMBER_ID")] + \
           [{"source_column": "NA", "audit": True}]
    p = render.place_rules(feed, rows)
    assert list(p["by_column"]) == ["CLAIM_ID"]
    assert p["unattributed"] == ["Something with no column name in it."]
    assert p["recycle"]["column"] == "MEMBER_ID"


def test_render_builtin_single_source(tmp_path):
    spec, vdd, sources = _ready(tmp_path)
    out = tmp_path / "out.xlsx"
    info = render.render_workbook(spec, sources, vdd, out)
    assert info["dialect"] == "single_sheet" and info["layout_from"] is None and info["n_rows"] == 7
    wb = load_workbook(out)
    ws = wb.active
    header_row = next(r for r in range(1, 20) if ws.cell(r, 2).value == "Field Name")
    assert ws.cell(header_row + 1, 2).value == "CLAIM_ID"
    assert ws.cell(header_row + 1, 15).value == "CLAIM_ID"            # stage ColumnName as-is
    assert ws.cell(header_row + 1, 11).value.startswith("If the CLAIM_ID column")  # rule on its row


def test_render_into_sheet_per_table_layout(tmp_path):
    spec, vdd, sources = _ready(tmp_path)
    layout = layout_of(str(make_reference_sheet_per_table(tmp_path / "STTM_other.xlsx")))
    out = tmp_path / "out.xlsx"
    info = render.render_workbook(spec, sources, vdd, out, layout=layout)
    assert info["layout_from"] == "STTM_other.xlsx"
    wb = load_workbook(out)
    assert "MAPPING-CLM_CLAIMS_STG" in wb.sheetnames and "MAPPING-OTHER_TABLE" not in wb.sheetnames
    ws = wb["MAPPING-CLM_CLAIMS_STG"]
    assert ws.cell(2, 1).value == "Database column Name"        # template headers kept
    assert ws.cell(3, 1).value == "CLAIM_ID" and ws.cell(3, 11).value == "CLAIM_ID"
    assert ws.cell(3, 18).value is None                        # unmapped template column stays blank
    assert info["unfilled_columns"] == {"MAPPING-CLM_CLAIMS_STG": ["Owner"]}
    assert ws.cell(6, 1).value == "NA" and ws.cell(6, 11).value == "SRC_FILE_NAME"
    assert wb["FILE_DETAILS"].cell(2, 2).value == "claims_YYYYMMDD.csv"


def test_render_into_single_sheet_layout(tmp_path):
    spec, vdd, sources = _ready(tmp_path)
    layout = layout_of(str(make_reference_single_sheet(tmp_path / "STTM_other.xlsx")))
    out = tmp_path / "out.xlsx"
    info = render.render_workbook(spec, sources, vdd, out, layout=layout)
    assert info["layout_from"] == "STTM_other.xlsx"
    ws = load_workbook(out)["mapping"]
    assert ws.cell(1, 2).value == "claims_YYYYMMDD.csv"          # metadata block refilled
    assert ws.cell(11, 2).value == "CLAIM_ID" and ws.cell(11, 15).value == "CLAIM_ID"


def test_render_refuses_with_no_rows(tmp_path):
    spec, vdd, sources = _ready(tmp_path)
    sources[0]["file"] = None
    try:
        render.render_workbook(spec, sources, vdd, tmp_path / "o.xlsx")
    except ValueError as exc:
        assert "no rows" in str(exc)
    else:
        raise AssertionError("expected a refusal")


def test_standard_type_from_example_value_when_vendor_says_string(tmp_path):
    from frdsttm import standards as std
    assert std.type_from_example("0.56") == "Decimal(10,2)"
    assert std.type_from_example("1,234.50") == "Decimal(10,2)"
    assert std.type_from_example("46508") is None
    assert std.type_from_example("2024.01.05") is None and std.type_from_example("") is None
    files = {"claims_YYYYMMDD.csv": [
        ("PCT_A", "String", "Y", "N", "d", "", "0.56", ""),      # String declared, decimal example -> Decimal
        ("POP_B", "String", "Y", "N", "d", "", "46508", ""),     # String declared, integer example -> String
        ("CNT_C", "int", "Y", "N", "d", "", "12.5", ""),         # vendor type is specific -> vendor wins
        ("AMT_D", "", "Y", "N", "d", "", "3.14", ""),            # no type at all, decimal example -> Decimal
    ]}
    spec, vdd, sources = _ready(tmp_path, spec_for(), files)
    rows, notes = render.build_rows(sources[0], spec["feeds"][0], vdd["fields"]["file1"])
    by = {r["source_column"]: r for r in rows if not r["audit"]}
    assert by["PCT_A"]["standard"]["datatype"] == "Decimal(10,2)" and by["PCT_A"]["type_origin"] == "example value"
    assert by["POP_B"]["standard"]["datatype"] == "String" and by["POP_B"]["type_origin"] == "vendor type"
    assert by["CNT_C"]["standard"]["datatype"] == "Int" and by["CNT_C"]["type_origin"] == "vendor type"
    assert by["AMT_D"]["standard"]["datatype"] == "Decimal(10,2)"
    assert all(r["stage"]["datatype"] == "String" for r in rows if not r["audit"])   # stage is never inferred
    assert notes["inferred_from_example"] == 2


# --------------------------------------------------------------------------- #
# a type the vendor never gave must not read like one they did
# --------------------------------------------------------------------------- #
_NO_TYPE = {"claims_YYYYMMDD.csv": [
    ("CLAIM_ID", "int", "Y", "N", "Claim identifier", "", "1001", ""),
    ("MEMBER_ID", "", "Y", "Y", "Member identifier", "", "", ""),   # vendor gave nothing
]}


def _amber(cell):
    return (cell.fill is not None and cell.fill.fgColor.rgb == render._INFERRED_FILL.fgColor.rgb
            and bool(cell.font.italic))


def _two_sources(tmp_path):
    """Two feeds — what makes render_workbook choose the sheet_per_table dialect."""
    import json
    spec = spec_for()
    second = json.loads(json.dumps(spec["feeds"][0]))
    second.update(feed_name="Roster", file_name_patterns=["roster_YYYYMMDD.csv"],
                  validation_rules=[],
                  stage_target={**second["stage_target"], "tables": ["rst_roster_stg"]},
                  standard_target={**second["standard_target"], "tables": ["rst_roster"]})
    spec["feeds"].append(second)
    vdd = parse_dictionary_workbook(make_vdd(tmp_path / "v2.xlsx", {
        "claims_YYYYMMDD.csv": _NO_TYPE["claims_YYYYMMDD.csv"],
        "roster_YYYYMMDD.csv": [("ROSTER_ID", "int", "Y", "N", "Roster id", "", "1", "")]}))
    sources, _ = c.derive_targets(spec)
    pairing, _ = c.pair_files(spec, vdd)
    for src in sources:
        f = pairing.get(src["feed_index"])
        src["file"], src["field_sheet"] = (f["file_name_pattern"], f["field_sheet"]) if f else (None, None)
    return spec, vdd, sources


def test_single_sheet_blanks_the_source_type_and_marks_the_standard(tmp_path):
    spec, vdd, sources = _ready(tmp_path, spec_for(), _NO_TYPE)
    out = tmp_path / "o.xlsx"
    info = render.render_workbook(spec, sources, vdd, out)
    assert info["dialect"] == "single_sheet" and info["marked_types"] == 1
    ws = load_workbook(out).active
    hdr_r = next(r for r in range(1, 30) if ws.cell(r, 2).value == "Field Name")
    hdr = [ws.cell(hdr_r, col).value for col in range(1, ws.max_column + 1)]
    src_dt = hdr.index("Data Type") + 1
    std_dt = len(hdr) - hdr[::-1].index("DataType")        # last DataType = Standard Layer
    rows = {ws.cell(r, 2).value: r for r in range(hdr_r + 1, ws.max_row + 1)}

    claim, member = rows["CLAIM_ID"], rows["MEMBER_ID"]
    # the vendor's own word is transcribed and left alone
    assert ws.cell(claim, src_dt).value == "int" and not _amber(ws.cell(claim, src_dt))
    assert not _amber(ws.cell(claim, std_dt))
    # what the vendor never said is blank on the source side, marked on the target
    assert ws.cell(member, src_dt).value in (None, "")
    assert _amber(ws.cell(member, src_dt)) and _amber(ws.cell(member, std_dt))
    assert ws.cell(member, std_dt).value                   # still carries the ACFC default
    # and the legend sits below the data
    assert any(ws.cell(r, 1).value == render.TYPE_LEGEND for r in range(1, ws.max_row + 1))


def test_sheet_per_table_legend_lands_on_file_details(tmp_path):
    spec, vdd, sources = _two_sources(tmp_path)
    out = tmp_path / "two.xlsx"
    info = render.render_workbook(spec, sources, vdd, out)
    assert info["dialect"] == "sheet_per_table" and info["marked_types"] == 1
    fd = load_workbook(out)["FILE_DETAILS"]
    assert any(fd.cell(r, 1).value == render.TYPE_LEGEND for r in range(1, fd.max_row + 1))


def test_no_legend_when_every_type_came_from_the_vendor(tmp_path):
    spec, vdd, sources = _ready(tmp_path)                  # the default fixture states every type
    clean = tmp_path / "clean.xlsx"
    info = render.render_workbook(spec, sources, vdd, clean)
    ws = load_workbook(clean).active
    assert info["marked_types"] == 0
    assert not any(ws.cell(r, 1).value == render.TYPE_LEGEND for r in range(1, ws.max_row + 1))


def test_audit_rows_are_never_marked(tmp_path):
    """Audit columns are the standards' own — they were never the vendor's to state."""
    spec, vdd, sources = _ready(tmp_path, spec_for(), _NO_TYPE)
    rows, _ = render.build_rows(sources[0], spec["feeds"][0], vdd["fields"]["file1"])
    assert [render._type_marks(r) for r in rows if r["audit"]] == \
        [(False, False)] * len([r for r in rows if r["audit"]])


def test_marking_survives_a_borrowed_layouts_own_styles(tmp_path):
    """_fill_sheet restores the approved workbook's styles per row; the mark has
    to be applied after that or it is silently overwritten."""
    spec, vdd, sources = _ready(tmp_path, spec_for(), _NO_TYPE)
    ref = make_reference_single_sheet(tmp_path / "STTM_Other.xlsx")
    out = tmp_path / "borrowed.xlsx"
    info = render.render_workbook(spec, sources, vdd, out, layout=layout_of(str(ref)))
    assert info["layout_from"] == "STTM_Other.xlsx" and info["marked_types"] == 1
    ws = load_workbook(out)[load_workbook(out).sheetnames[0]]
    assert any(_amber(ws.cell(r, col))
               for r in range(1, ws.max_row + 1) for col in range(1, ws.max_column + 1))
