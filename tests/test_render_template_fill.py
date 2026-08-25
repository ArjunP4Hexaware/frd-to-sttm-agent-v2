"""Template fill (2026-08-22): 04 renders INTO the chosen reference workbook's
own layout — whatever dialect it is — instead of one of two hard-coded
header lists. These tests build templates in dialects the code has never
seen (different header names, an extra column the contract knows nothing
about, a different FILE_DETAILS column order) and check that the rendered
workbook keeps the template's headers, puts each value under the right
one, leaves unknown columns blank AND reports them, renames/creates/removes
sheets per feed, and still round-trips through our own parser.

Functions are lifted from 04_sttm_render.py by name (AST), as in
test_render_rules.py.
"""

import ast
import re
from copy import copy
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from frdsttm.reference_workbooks import _n, _nl, layout_of, parse_reference_workbook, match_feeds

RENDER_SRC = Path(__file__).resolve().parent.parent / "notebooks" / "04_sttm_render.py"
_NEEDED = {
    "GENERATOR", "_HDR_FILL", "_SEC_FILL", "_BOLD", "_style_row", "_quote_rule",
    "_column_mentions", "place_feed_rules", "_feed_level_rule_text",
    "_PER_TABLE_SRC_HEADERS", "_PER_TABLE_TGT_HEADERS", "_RECYCLE_HEADER",
    "render_sheet_per_table", "render_single_sheet", "render_contract",
    "_SEGMENT_SUFFIX", "_table_for_segment", "derive_field_mappings",
    "_copy_row_style", "_field_value", "_fill_sheet", "render_into_template",
    "render_into_single_sheet_template",
}


def _load():
    tree = ast.parse(RENDER_SRC.read_text(encoding="utf-8"))
    picked, found = [], set()
    for node in tree.body:
        if isinstance(node, ast.FunctionDef) and node.name in _NEEDED:
            picked.append(node); found.add(node.name)
        elif isinstance(node, ast.Assign):
            names = {t.id for t in node.targets if isinstance(t, ast.Name)} & _NEEDED
            if names:
                picked.append(node); found |= names
    assert found == _NEEDED, f"loader drifted: missing {_NEEDED - found}"
    ns = {"re": re, "copy": copy, "datetime": datetime, "timezone": timezone, "Path": Path,
          "Workbook": Workbook, "PatternFill": PatternFill, "Font": Font, "Alignment": Alignment,
          "get_column_letter": get_column_letter, "_n": _n, "_nl": _nl}
    exec(compile(ast.fix_missing_locations(ast.Module(body=picked, type_ignores=[])),
                 str(RENDER_SRC), "exec"), ns)
    return ns


R = _load()

RULE = "Reject the record when MEMBER_ID is NULL."
FEED_RULE = "Files arriving after 6 PM are processed next day."
RECYCLE = "Recycle Flag enabled for 7 days: check SUBS_ID against STG.SUBS; else recycle."


def _unusual_template(path, *, table="T_OLD", extra_sheet=None):
    """A dialect the code never hard-coded: different header spellings, a
    column ('Owner') the contract knows nothing about, FILE_DETAILS in a
    different column order, a recycle column beyond the bands."""
    wb = Workbook()
    fd = wb.active
    fd.title = "FILE_DETAILS"
    fd.append(["Frequency", "Vendor", "File Name", "Location", "File Description"])
    fd.append(["daily", "Old Vendor", "OLD_FILE.txt", "/old", "old rows"])
    vh = wb.create_sheet("VERSION_HISTORY")
    vh.append(["Version", "Date", "Author", "Change Description"])
    vh.append(["1.0", "2026-01-01", "someone", "old"])
    for name in [f"MAPPING-{table}"] + ([extra_sheet] if extra_sheet else []):
        ws = wb.create_sheet(name)
        ws.append(["Source File Layout", None, None, None, None, "Stage Layer", None, None, None,
                   "Standard Layer", None, None, None, None])
        ws.append(["Client Data Table Column Name", "Data Type", "Null Check", "Owner", "Comments",
                   "Schema", "Table Name", "Column Name", "Data Type",
                   "Schema", "Table Name", "Column Name", "Data Type", "Recycle Flag (Enabled for 7 Days)"])
        ws["A2"].font = Font(bold=True, color="FF0000")
        for col, dt in (("OLD_A", "String"), ("OLD_B", "Int")):
            ws.append([col, dt, "Not NULL", "team-x", "", "stg", table, col, dt, "std", table, col, dt, ""])
        ws.append(["NA", "timestamp", "", "", "", "stg", table, "LOAD_TS", "timestamp", "std", table, "LOAD_TS", "timestamp", ""])
        ws.column_dimensions["A"].width = 33
    wb.save(path)


def _contract(columns=("MEMBER_ID", "ZIP_CODE", "SUBS_ID"), table="T_MEMBER", rules=(RULE, FEED_RULE), recycle=RECYCLE):
    return {"status": "PASS", "generated_from_frd": "demo.docx", "_provenance": {},
            "feeds": [{"feed_name": "member feed", "source_system": "New Vendor",
                       "file_name_patterns": ["MEMBER_YYYYMMDD.txt"], "landing_location": "/landing/member",
                       "frequency": "weekly",
                       "stage_target": {"catalog": "c", "schema": "stg", "tables": [table]},
                       "standard_target": {"catalog": "c", "schema": "std", "tables": [table]},
                       "validation_rules": list(rules), "recycle_rule": recycle}]}


def _dictionary_from_template_with_our_columns(path, columns):
    """The template's dictionary but with OUR columns as the source fields
    (what a single-mode render on a same-table template does)."""
    d = parse_reference_workbook(str(path))
    (key, feed), = d["feeds"].items()
    real = [f for f in feed["fields"] if not f.get("audit")]
    audit = [f for f in feed["fields"] if f.get("audit")]
    proto = real[0]
    feed["fields"] = [{**proto, "source_column": c, "comment": ""} for c in columns] + audit
    feed["ref_targets"] = [{"stage": {"catalog": "", "schema": "stg", "table": "T_OLD", "column": c, "datatype": "String"},
                            "standard": {"catalog": "", "schema": "std", "table": "T_OLD", "column": c, "datatype": "String"}}
                           for c in columns] + feed["ref_targets"][len(real):]
    return d, key


def _rows(ws):
    return [list(r) for r in ws.iter_rows(values_only=True)]


def test_layout_of_recovers_roles_and_unmapped_columns(tmp_path):
    p = tmp_path / "tpl.xlsx"; _unusual_template(p)
    lay = layout_of(str(p))
    assert lay["dialect"] == "sheet_per_table"
    sheet = lay["sheets"]["MAPPING-T_OLD"]
    roles = {c["header"]: c["role"] for c in sheet["columns"]}
    assert roles["Client Data Table Column Name"] == ("source", "source_column")
    assert roles["Comments"] == ("source", "comment")
    assert roles["Owner"] is None and sheet["unmapped"] == ["Owner"]
    assert roles["Recycle Flag (Enabled for 7 Days)"] == ("recycle", None)
    assert sheet["first_data_row"] == 3
    assert lay["file_details"]["columns"] == {"frequency": 0, "vendor": 1, "file_name": 2, "location": 3, "description": 4}
    assert lay["version_history"]["columns"]["change"] == 3


def test_render_into_template_keeps_layout_fills_under_its_headers(tmp_path):
    p = tmp_path / "tpl.xlsx"; _unusual_template(p)
    d, key = _dictionary_from_template_with_our_columns(p, ["MEMBER_ID", "ZIP_CODE", "SUBS_ID"])
    contract = _contract()
    fm = {0: key}
    R["derive_field_mappings"](contract, d, fm)
    out = tmp_path / "out.xlsx"
    R["render_contract"](contract, d, str(out), layout=layout_of(str(p)), feed_match=fm)
    wb = load_workbook(out)
    # sheet renamed for OUR table; the template's header row kept verbatim (and its style)
    assert "MAPPING-T_MEMBER" in wb.sheetnames and "MAPPING-T_OLD" not in wb.sheetnames
    ws = wb["MAPPING-T_MEMBER"]
    rows = _rows(ws)
    assert rows[1][:5] == ["Client Data Table Column Name", "Data Type", "Null Check", "Owner", "Comments"]
    assert ws["A2"].font.bold and ws.column_dimensions["A"].width == 33
    data = rows[2:]
    assert [r[0] for r in data] == ["MEMBER_ID", "ZIP_CODE", "SUBS_ID", "NA"]
    # values under the template's own headers
    hdr = rows[1]
    i_owner, i_comm, i_rec = hdr.index("Owner"), hdr.index("Comments"), hdr.index("Recycle Flag (Enabled for 7 Days)")
    i_stage_tbl, i_std_col = 6, 11
    assert data[0][i_comm] == RULE and data[1][i_comm] is None
    assert all(r[i_owner] is None for r in data)                          # unknown column: blank
    assert data[0][i_stage_tbl] == "T_MEMBER" and data[0][i_std_col] == "MEMBER_ID"
    assert data[3][7] == "LOAD_TS" and data[3][0] == "NA"                 # audit row from the template
    assert data[2][i_rec].startswith("Y (") and RECYCLE in data[2][i_rec]  # recycle on SUBS_ID only
    assert all(r[i_rec] in (None, "") for r in data if r[0] != "SUBS_ID")
    # FILE_DETAILS under ITS column order; old row gone
    fd = _rows(wb["FILE_DETAILS"])
    assert fd[0] == ["Frequency", "Vendor", "File Name", "Location", "File Description"]
    assert fd[1] == ["weekly", "New Vendor", "MEMBER_YYYYMMDD.txt", "/landing/member", FEED_RULE]
    assert len(fd) == 2
    vh = _rows(wb["VERSION_HISTORY"])
    assert vh[1][0] == "0.1" and "tpl.xlsx" in vh[1][3] and len(vh) == 2
    # provenance says what the template had that we could not fill
    tf = contract["_provenance"]["template_fill"]
    assert tf["template"] == "tpl.xlsx"
    assert tf["unfilled_columns"] == {"MAPPING-T_MEMBER": ["Owner"]}
    assert tf["sheets"]["MAPPING-T_MEMBER"]["rows"] == 4


def test_own_template_render_keeps_its_datatypes_and_verbatim_flag_text(tmp_path):
    """The pinned own-template case (2026-08-25): the row set is the
    template's, so its per-row stage/standard datatypes and the analyst's
    exact cell spelling ("NOT NULL", "NA") must survive into the draft --
    not a "String" default and a re-spelled "Not NULL"."""
    p = tmp_path / "tpl.xlsx"; _unusual_template(p, table="T_MEMBER")
    wb = load_workbook(p); ws = wb["MAPPING-T_MEMBER"]
    ws["C3"] = "NOT NULL"; ws["C4"] = "NA"          # verbatim flag text, two spellings
    ws["I4"] = "Decimal(10,2)"; ws["M4"] = "Int"    # stage / standard datatypes differ
    wb.save(p)
    d = parse_reference_workbook(str(p))
    (key, feed), = d["feeds"].items()
    assert feed["fields"][0]["nullable_raw"] == "NOT NULL" and feed["fields"][1]["nullable_raw"] == "NA"
    contract = _contract(table="T_MEMBER")
    R["derive_field_mappings"](contract, d, {0: key})
    out = tmp_path / "out.xlsx"
    R["render_contract"](contract, d, str(out), layout=layout_of(str(p)), feed_match={0: key})
    data = _rows(load_workbook(out)["MAPPING-T_MEMBER"])[2:]
    assert [r[0] for r in data] == ["OLD_A", "OLD_B", "NA"]
    assert [r[2] for r in data[:2]] == ["NOT NULL", "NA"]                 # verbatim, not re-spelled
    assert (data[0][8], data[0][12]) == ("String", "String")
    assert (data[1][8], data[1][12]) == ("Decimal(10,2)", "Int")          # template datatypes, per layer
    assert data[2][12] == "timestamp"                                      # audit row unchanged


def test_rendered_template_fill_round_trips_through_our_parser(tmp_path):
    p = tmp_path / "tpl.xlsx"; _unusual_template(p)
    d, key = _dictionary_from_template_with_our_columns(p, ["MEMBER_ID", "ZIP_CODE"])
    contract = _contract(columns=("MEMBER_ID", "ZIP_CODE"))
    R["derive_field_mappings"](contract, d, {0: key})
    out = tmp_path / "out.xlsx"
    R["render_contract"](contract, d, str(out), layout=layout_of(str(p)), feed_match={0: key})
    back = parse_reference_workbook(str(out))
    (k, feed), = back["feeds"].items()
    assert k == "t_member"
    cols = [f["source_column"] for f in feed["fields"]]
    assert cols == ["MEMBER_ID", "ZIP_CODE", "NA"]
    assert feed["fields"][0]["comment"] == RULE and feed["recycle_note"]
    assert feed["fields"][2]["audit"] is True


def _two_feed_contract(d, key):
    contract = _contract(columns=("A1",), table="T_ONE", rules=(), recycle=None)
    contract["feeds"].append({**contract["feeds"][0], "feed_name": "second feed",
                              "stage_target": {"catalog": "c", "schema": "stg", "tables": ["T_TWO"]},
                              "standard_target": {"catalog": "c", "schema": "std", "tables": ["T_TWO"]}})
    d["feeds"]["t_two"] = {**d["feeds"][key], "sheet": "MAPPING-NOWHERE"}
    return contract, {0: key, 1: "t_two"}


def test_a_spare_template_sheet_is_reused_for_an_unmatched_feed(tmp_path):
    """Template has two mapping sheets; we render two feeds, the second with
    no sheet of its own → the spare sheet is reused (renamed), nothing is
    copied, nothing removed."""
    p = tmp_path / "tpl.xlsx"; _unusual_template(p, extra_sheet="MAPPING-T_OTHER")
    d, key = _dictionary_from_template_with_our_columns(p, ["A1"])
    contract, fm = _two_feed_contract(d, key)
    R["derive_field_mappings"](contract, d, fm)
    out = tmp_path / "out.xlsx"
    R["render_contract"](contract, d, str(out), layout=layout_of(str(p)), feed_match=fm)
    wb = load_workbook(out)
    assert {n for n in wb.sheetnames if n.startswith("MAPPING-")} == {"MAPPING-T_ONE", "MAPPING-T_TWO"}
    tf = contract["_provenance"]["template_fill"]
    assert tf["created_sheets"] == [] and tf["removed_sheets"] == []
    assert {info["from"] for info in tf["sheets"].values()} == {"MAPPING-T_OLD", "MAPPING-T_OTHER"}
    assert _rows(wb["MAPPING-T_TWO"])[1][0] == "Client Data Table Column Name"


def test_extra_feed_gets_a_copy_of_the_lead_sheet_and_spare_sheets_are_removed(tmp_path):
    # one-sheet template, two feeds → the second is a COPY of the lead sheet
    p = tmp_path / "tpl.xlsx"; _unusual_template(p)
    d, key = _dictionary_from_template_with_our_columns(p, ["A1"])
    contract, fm = _two_feed_contract(d, key)
    R["derive_field_mappings"](contract, d, fm)
    out = tmp_path / "out.xlsx"
    R["render_contract"](contract, d, str(out), layout=layout_of(str(p)), feed_match=fm)
    wb = load_workbook(out)
    assert {n for n in wb.sheetnames if n.startswith("MAPPING-")} == {"MAPPING-T_ONE", "MAPPING-T_TWO"}
    tf = contract["_provenance"]["template_fill"]
    assert any("MAPPING-T_TWO" in c and "copy of" in c for c in tf["created_sheets"])
    assert _rows(wb["MAPPING-T_TWO"])[1][0] == "Client Data Table Column Name"
    # two-sheet template, ONE feed → the spare sheet (another FRD's rows) is removed
    p2 = tmp_path / "tpl2.xlsx"; _unusual_template(p2, extra_sheet="MAPPING-T_OTHER")
    d2, key2 = _dictionary_from_template_with_our_columns(p2, ["A1"])
    c2 = _contract(columns=("A1",), table="T_ONE", rules=(), recycle=None)
    R["derive_field_mappings"](c2, d2, {0: key2})
    out2 = tmp_path / "out2.xlsx"
    R["render_contract"](c2, d2, str(out2), layout=layout_of(str(p2)), feed_match={0: key2})
    removed = c2["_provenance"]["template_fill"]["removed_sheets"]
    assert len(removed) == 1 and removed[0] in ("MAPPING-T_OLD", "MAPPING-T_OTHER")
    assert [n for n in load_workbook(out2).sheetnames if n.startswith("MAPPING-")] == ["MAPPING-T_ONE"]


def test_freeform_without_a_layout_still_uses_the_built_in_fallback(tmp_path):
    contract = _contract(rules=(RULE,), recycle=None)
    contract["feeds"][0]["fields"] = [{"source_column": "MEMBER_ID", "description": "", "sample": "",
                                       "datatype": "String", "nullable": False, "phi": False, "mandatory": False,
                                       "comment": "", "segment": "", "business_rule": "",
                                       "stage": {"catalog": "c", "schema": "stg", "table": "T_MEMBER", "column": "MEMBER_ID", "datatype": "String"},
                                       "standard": {"catalog": "c", "schema": "std", "table": "T_MEMBER", "column": "MEMBER_ID", "datatype": "String"}}]
    out = tmp_path / "out.xlsx"
    R["render_contract"](contract, {"dialect": "sheet_per_table", "feeds": {}, "meta": {}}, str(out))
    assert "template_fill" not in contract["_provenance"]
    assert _rows(load_workbook(out)["MAPPING-T_MEMBER"])[1][0] == "Database column Name"


def _single_sheet_template(path):
    wb = Workbook(); ws = wb.active; ws.title = "mapping"
    for k, v in (("File(s)", "OLD.txt"), ("File Generator", "Old"), ("File Location", "/old"),
                 ("LOB", "X"), ("File frequency", "daily"), ("Domain", "OLD"), ("Owner team", "team-x")):
        ws.append([k, v])
    ws.append(["Source Layout", None, None, None, "Stage Layer", None, None, None, "Standard Layer", None, None, None])
    ws.append(["#", "Field Name", "Data Type", "Business Rule", "Catalog", "Schema", "TableName", "ColumnName",
               "Catalog", "Schema", "TableName", "ColumnName"])
    ws.append([1, "OLD_A", "String", "", "c", "stg", "T_OLD", "OLD_A", "c", "std", "T_OLD", "OLD_A"])
    wb.save(path)


def test_single_sheet_template_fill_writes_metadata_and_business_rules(tmp_path):
    p = tmp_path / "caqh.xlsx"; _single_sheet_template(p)
    d = parse_reference_workbook(str(p))
    (key, feed), = d["feeds"].items()
    proto = feed["fields"][0]
    feed["fields"] = [{**proto, "source_column": c} for c in ("MEMBER_ID", "ZIP_CODE")]
    feed["ref_targets"] = feed["ref_targets"] * 2
    contract = _contract(columns=("MEMBER_ID", "ZIP_CODE"), rules=(RULE, FEED_RULE), recycle=None)
    contract["feeds"][0]["lobs"] = ["REG#1 Medicaid"]
    R["derive_field_mappings"](contract, d, {0: key})
    out = tmp_path / "out.xlsx"
    R["render_contract"](contract, d, str(out), layout=layout_of(str(p)), feed_match={0: key})
    rows = _rows(load_workbook(out)["mapping"])
    meta = {r[0]: r[1] for r in rows[:7]}
    assert meta["File(s)"] == "MEMBER_YYYYMMDD.txt" and meta["File Generator"] == "New Vendor"
    assert meta["LOB"] == "Medicaid" and meta["Owner team"] is None       # unknown key: cleared, reported
    hdr = rows[8]
    assert hdr[:4] == ["#", "Field Name", "Data Type", "Business Rule"]
    data = rows[9:]
    assert [r[1] for r in data] == ["MEMBER_ID", "ZIP_CODE"] and data[0][0] == 1
    assert data[0][3] == RULE and data[1][3] is None                     # rules → Business Rule (no Comment col)
    tf = contract["_provenance"]["template_fill"]
    assert tf["unfilled_meta"] == ["Owner team"] and FEED_RULE in tf["feed_level_rules_unplaced"]
