"""Synthetic fixtures built in-process — no client document is ever needed."""

from __future__ import annotations

import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook

FRD_TEXT = (
    "This Functional Requirements Document describes the onboarding of the vendor "
    "claims intake feed into the enterprise data lake, from the landing zone through "
    "the stage layer into the standard layer, with data quality rules applied at each "
    "layer as described in the sections below. The feed is delivered daily by the vendor."
)


def make_frd(path: Path, *, project_id="1234567", file_pattern="claims_YYYYMMDD.csv",
             stage_table="clm_claims_stg", standard_table="clm_claims", schema="stg_clm",
             extra_paragraphs=()) -> Path:
    import docx

    d = docx.Document()
    d.add_heading("Claims Intake FRD", 0)
    d.add_paragraph(f"Project ID: {project_id} Claims Intake Redesign")
    d.add_heading("In Scope", 1)
    d.add_paragraph(FRD_TEXT)
    d.add_paragraph("Ingest the vendor claims file into the data lake.", style="List Bullet")
    d.add_heading("Data Ingestion Requirements", 1)
    d.add_paragraph("SRQ226433 The process shall load the claims file daily.")
    t = d.add_table(rows=6, cols=2)
    rows = [("Object/data Format", "csv"), ("Target Schema", schema),
            ("Target Table Name", stage_table), ("Domain and Subdomain", "CLAIMS / Intake"),
            ("File Name", file_pattern), ("Standard Table", standard_table)]
    for r, (k, v) in zip(t.rows, rows):
        r.cells[0].text, r.cells[1].text = k, v
    d.add_heading("Data Quality", 2)
    d.add_paragraph("If the CLAIM_ID column is NULL, reject the record to the reject table.")
    d.add_paragraph("Process shall fail when file layout is not as per source dictionary.")
    for p in extra_paragraphs:
        d.add_paragraph(p)
    d.save(str(path))
    return path


def make_vdd(path: Path, files: dict | None = None, *, marker_rows=False) -> Path:
    """files: {pattern: [(name, type, required, phi, description, allowed, example, segment), ...]}"""
    files = files or {"claims_YYYYMMDD.csv": [
        ("CLAIM_ID", "int", "Y", "N", "Claim identifier", "", "1001", ""),
        ("MEMBER_ID", "varchar", "Y", "Y", "Member identifier", "", "M123", ""),
        ("AMOUNT", "decimal", "N", "N", "Billed amount", ">= 0", "12.50", ""),
    ]}
    wb = Workbook()
    ws = wb.active
    ws.title = "FILES"
    ws.append(["File Name Pattern", "File Title", "Format", "Delimiter", "Header Row", "Encoding",
               "Delivery Cadence", "Content Description", "Field Sheet", "Notes"])
    sheets = []
    for i, (pattern, cols) in enumerate(files.items(), 1):
        sheet = f"file{i}"
        sheets.append((sheet, cols))
        ws.append([pattern, f"File {i}", "csv", ",", "Y", "UTF-8", "Daily", "desc", sheet, ""])
    if marker_rows:
        ws.append(["example_*.csv", "Example", "csv", ",", "Y", "UTF-8", "Daily", "x", "fields_template",
                   "Example row — delete once replaced"])
    for sheet, cols in sheets:
        fs = wb.create_sheet(sheet)
        fs.append(["Position", "Field Name", "Data Type", "Length", "Required (Y/N)", "Description",
                   "Allowed Values / Range", "Example Value", "PHI/PII (Y/N)", "Segment", "Notes"])
        for pos, (name, typ, req, phi, desc, allowed, ex, seg) in enumerate(cols, 1):
            fs.append([pos, name, typ, "", req, desc, allowed, ex, phi, seg, ""])
    wb.save(str(path))
    return path


def make_reference_sheet_per_table(path: Path, table="other_table") -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "FILE_DETAILS"
    ws.append(["Vendor", "FileName", "File Description", "Location", "Frequency"])
    ws.append(["Other", "other_*.csv", "", "", "Monthly"])
    vh = wb.create_sheet("VERSION_HISTORY")
    vh.append(["Version", "Date", "Author", "Change Description"])
    vh.append(["1.0", "2026-01-01", "someone", "approved"])
    m = wb.create_sheet(f"MAPPING-{table.upper()}")
    m.append(["Source File Layout", "", "", "", "", "", "", "", "Stage Layer", "", "", "",
              "", "Standard Layer", "", "", "", "Owner"])
    m.append(["Database column Name", "NULL CHECK", "Description", "Sample Value", "DataType", "PHI Field",
              "Mandatory Field", "Comment", "Schema", "TableName", "ColumnName", "DataType", "",
              "Schema", "TableName", "ColumnName", "DataType", "Owner"])
    m.append(["other_col", "NULL", "d", "s", "String", "No", "No", "", "stg_x", table, "other_col",
              "String", "", "x", table, "other_col", "String", "bob"])
    m.append(["NA", "", "", "", "", "", "", "", "stg_x", table, "SRC_FILE_NAME", "String", "",
              "x", table, "SRC_FILE_NAME", "String", ""])
    wb.save(str(path))
    return path


def make_reference_single_sheet(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "mapping"
    for k, v in [("File(s)", "other.txt"), ("File Generator", "Other"), ("File Location", "x"),
                 ("LOB", "0100"), ("File frequency", "Weekly"), ("Domain", "MEMBER"),
                 ("Sub-Domain", "TPL"), ("File type", "txt")]:
        ws.append([k, v])
    ws.append(["Source Layout"] + [""] * 10 + ["Stage Layer"] + [""] * 7 + ["Standard Layer"])
    src = ["#", "Field Name", "Data Type", "Length", "Field Length (fixed width)",
           "Start position (fixed width)", "End Position (fixed width)", "Segment", "PII", "Comments",
           "Business Rule"]
    tgt = ["Catalog", "Schema", "TableName", "ColumnName", "DataType", "Mandatory Column", "Primary Key",
           "Field Description"]
    ws.append(src + tgt + tgt)
    ws.append([1, "Other Field", "varchar", "10", "", "", "", "Detail", "", "c", "r",
               "PR_DLK", "STG_MBR", "T", "OTHER_FIELD", "String", "", "", "d",
               "PR_STD", "MBR", "T", "OTHER_FIELD", "String", "", "", "d"])
    wb.save(str(path))
    return path


def spec_for(file_pattern="claims_YYYYMMDD.csv", stage_table="clm_claims_stg", standard_table="clm_claims",
             schema="stg_clm", project_id="1234567", **feed_overrides) -> dict:
    feed = {
        "feed_name": "Claims Intake", "source_system": "vendor", "file_name_patterns": [file_pattern],
        "file_format": "csv", "delimiter": ",", "record_segments": [], "frequency": "daily",
        "load_windows_sla": [], "lobs": [], "domain": "CLAIMS", "sub_domain": "Intake",
        "landing_location": None,
        "stage_target": {"catalog": None, "schema": schema, "tables": [stage_table], "load_strategy": None},
        "standard_target": {"catalog": None, "schema": None, "tables": [standard_table], "load_strategy": None},
        "validation_rules": ["If the CLAIM_ID column is NULL, reject the record to the reject table."],
        "recycle_rule": None, "history_backfill": None, "archive_retention": None, "phi_pii_notes": None,
        "sttm_reference": None, "requirement_ids": ["SRQ226433"],
    }
    feed.update(feed_overrides)
    return {"project": {"project_id": project_id, "project_name": None, "business_context_summary": None},
            "in_scope": [], "out_of_scope": [], "assumptions_constraints_dependencies": [],
            "feeds": [feed], "system_interfaces": [], "open_items": []}


class FakeClient:
    """Stands in for anthropic.Anthropic: returns a canned spec."""

    def __init__(self, spec: dict, stop_reason="end_turn"):
        self.spec, self.stop_reason, self.calls = spec, stop_reason, 0
        self.messages = self

    def stream(self, **kwargs):
        self.calls += 1
        self.last_kwargs = kwargs
        client = self

        class _Ctx:
            def __enter__(self_inner):
                return self_inner

            def __exit__(self_inner, *a):
                return False

            def get_final_message(self_inner):
                return SimpleNamespace(
                    stop_reason=client.stop_reason,
                    content=[SimpleNamespace(type="text", text=json.dumps(client.spec))],
                    usage=SimpleNamespace(input_tokens=10, output_tokens=5))
        return _Ctx()


@pytest.fixture
def data_root(tmp_path: Path) -> Path:
    """The four folders with one synthetic pair (FRD + VDD) and one other feed's approved STTM."""
    for d in ("frds", "vdds", "reference_sttms", "output_sttms"):
        (tmp_path / d).mkdir()
    make_frd(tmp_path / "frds" / "FRD_Claims_Intake.docx")
    make_vdd(tmp_path / "vdds" / "VDD_Claims_Intake.xlsx")
    make_reference_sheet_per_table(tmp_path / "reference_sttms" / "STTM_Other_Feed.xlsx")
    return tmp_path
