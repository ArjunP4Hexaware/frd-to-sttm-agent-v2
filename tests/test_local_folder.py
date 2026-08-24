"""frdsttm.local_folder — a directory standing in for the SharePoint library.

The point of this module is that it plugs into the EXISTING sync engine, so
these tests check the client contract (listing, change detection, path
safety) and then drive a real ``sync_from_sharepoint`` over a temp folder to
prove the downstream half — split by prefix, volumes, manifest, corpus index,
name pairing — behaves the same as it does over Graph.

All documents are SYNTHETIC, per the no-client-documents rule.
"""

from __future__ import annotations

import pytest
from openpyxl import Workbook

from frdsttm.corpus import CORPUS_INDEX_NAME, load_corpus_index
from frdsttm.local_folder import (
    SOURCE_DIR_VAR,
    LocalFolderClient,
    LocalFolderConfigError,
    load_local_folder_config,
)
from frdsttm.similarity import thresholds_from
from frdsttm.sync import MANIFEST_NAME, load_manifest, sync_from_sharepoint

NOW = "2026-08-24T12:00:00+00:00"
_thresholds = lambda: thresholds_from(lambda name, default: default)  # noqa: E731


def _docx(path, paragraphs):
    """A real .docx, so normalize_to_markdown parses it the way 01 does."""
    from docx import Document

    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    doc.save(str(path))


def _xlsx(path, rows):
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    wb.save(str(path))


@pytest.fixture
def source(tmp_path):
    """A flat folder in the demo's shape: two FRD/STTM pairs, matching names."""
    src = tmp_path / "frd-to-sttm-demo-document"
    src.mkdir()
    _docx(src / "FRD_ClaimsFeed.docx",
          ["Claims feed", "Source table claims_raw", "Target CLAIM_FACT",
           "Members are loaded daily from the claims extract."])
    _xlsx(src / "STTM_ClaimsFeed.xlsx",
          [["Source Table", "Source Column", "Target Table", "Target Column"],
           ["claims_raw", "clm_id", "CLAIM_FACT", "CLAIM_ID"]])
    _docx(src / "FRD_MemberFeed.docx",
          ["Member feed", "Source table member_raw", "Target MEMBER_DIM",
           "Enrollment records arrive weekly from the eligibility file."])
    _xlsx(src / "STTM_MemberFeed.xlsx",
          [["Source Table", "Source Column", "Target Table", "Target Column"],
           ["member_raw", "mbr_id", "MEMBER_DIM", "MEMBER_ID"]])
    return src


# --------------------------------------------------------------------------- #
# config
# --------------------------------------------------------------------------- #
def test_unset_variable_is_an_ordinary_state_not_a_crash():
    """Unset must raise the typed error the callers catch to fall back to
    SharePoint — never something that escapes as a 500."""
    with pytest.raises(LocalFolderConfigError) as exc:
        load_local_folder_config(lambda name, default: "")
    assert SOURCE_DIR_VAR in str(exc.value)


def test_missing_folder_names_the_path_it_looked_at(tmp_path):
    missing = tmp_path / "not-there"
    with pytest.raises(LocalFolderConfigError, match="does not exist"):
        load_local_folder_config(lambda name, default: str(missing))


def test_a_file_is_rejected_as_a_source_folder(tmp_path):
    f = tmp_path / "a.txt"
    f.write_text("x")
    with pytest.raises(LocalFolderConfigError, match="not a directory"):
        load_local_folder_config(lambda name, default: str(f))


def test_tilde_is_expanded(monkeypatch, tmp_path):
    monkeypatch.setenv("HOME", str(tmp_path))
    (tmp_path / "docs").mkdir()
    cfg = load_local_folder_config(lambda name, default: "~/docs")
    assert cfg.root == tmp_path / "docs"


def test_env_is_read_when_no_accessor_is_passed(monkeypatch, source):
    monkeypatch.setenv(SOURCE_DIR_VAR, str(source))
    assert load_local_folder_config().root == source


# --------------------------------------------------------------------------- #
# client contract
# --------------------------------------------------------------------------- #
def test_listing_filters_by_suffix_and_sorts_by_name(source):
    client = LocalFolderClient(load_local_folder_config(lambda n, d: str(source)))
    names = [i.name for i in client.list_documents(suffixes={".docx"})]
    assert names == ["FRD_ClaimsFeed.docx", "FRD_MemberFeed.docx"]


def test_listing_skips_directories_and_dotfiles(source):
    (source / "archive").mkdir()
    (source / ".DS_Store").write_bytes(b"junk")
    client = LocalFolderClient(load_local_folder_config(lambda n, d: str(source)))
    names = [i.name for i in client.list_documents()]
    assert "archive" not in names and ".DS_Store" not in names


def test_etag_changes_when_content_changes(source):
    """The incremental check leans on etag; an edited document must look
    different or a re-sync would skip the reviewer's change."""
    client = LocalFolderClient(load_local_folder_config(lambda n, d: str(source)))
    before = {i.name: i.etag for i in client.list_documents(suffixes={".xlsx"})}
    _xlsx(source / "STTM_ClaimsFeed.xlsx",
          [["Source Table", "Source Column", "Target Table", "Target Column"],
           ["claims_raw", "clm_id", "CLAIM_FACT", "CLAIM_ID"],
           ["claims_raw", "svc_dt", "CLAIM_FACT", "SERVICE_DATE"]])
    after = {i.name: i.etag for i in client.list_documents(suffixes={".xlsx"})}
    assert after["STTM_ClaimsFeed.xlsx"] != before["STTM_ClaimsFeed.xlsx"]
    assert after["STTM_MemberFeed.xlsx"] == before["STTM_MemberFeed.xlsx"]


def test_download_returns_the_bytes_on_disk(source):
    client = LocalFolderClient(load_local_folder_config(lambda n, d: str(source)))
    assert client.download_item("FRD_ClaimsFeed.docx") == \
        (source / "FRD_ClaimsFeed.docx").read_bytes()


def test_download_refuses_to_escape_the_source_folder(source, tmp_path):
    """item_id is treated as untrusted even though our own listing produced
    it — the same posture sync.safe_name takes on the write side."""
    secret = tmp_path / "outside.txt"
    secret.write_text("not yours")
    client = LocalFolderClient(load_local_folder_config(lambda n, d: str(source)))
    with pytest.raises((LocalFolderConfigError, FileNotFoundError)):
        client.download_item("../outside.txt")


def test_download_of_a_departed_file_raises_not_found(source):
    client = LocalFolderClient(load_local_folder_config(lambda n, d: str(source)))
    (source / "FRD_ClaimsFeed.docx").unlink()
    with pytest.raises(FileNotFoundError):
        client.download_item("FRD_ClaimsFeed.docx")


# --------------------------------------------------------------------------- #
# end to end through the real sync engine
# --------------------------------------------------------------------------- #
def _sync(source, tmp_path, now=NOW):
    client = LocalFolderClient(load_local_folder_config(lambda n, d: str(source)))
    return client, sync_from_sharepoint(
        client,
        frd_dir=tmp_path / "frd_raw",
        reference_dir=tmp_path / "sttm_reference",
        reference_folder=None,
        thresholds=_thresholds(),
        now_iso=now,
        frd_prefix="FRD_",
        reference_prefix="STTM_",
    )


def test_sync_splits_by_prefix_into_the_two_volumes(source, tmp_path):
    _, result = _sync(source, tmp_path)
    assert result["frd_downloaded"] == 2
    assert result["reference_downloaded"] == 2
    assert not result["skipped"]
    assert sorted(p.name for p in (tmp_path / "frd_raw").iterdir()) == [
        "FRD_ClaimsFeed.docx", "FRD_MemberFeed.docx"]
    refs = sorted(p.name for p in (tmp_path / "sttm_reference").iterdir())
    assert "STTM_ClaimsFeed.xlsx" in refs and "STTM_MemberFeed.xlsx" in refs


def test_both_pairs_match_by_name(source, tmp_path):
    """The demo's whole story: two FRDs, each already mapped to its STTM."""
    _, result = _sync(source, tmp_path)
    pairs = result["index"]["pairs"]
    assert set(pairs) == {"FRD_ClaimsFeed", "FRD_MemberFeed"}
    assert all(p["matched_by"] == "name" for p in pairs.values())
    assert result["index"]["unmapped"] == []


def test_files_outside_the_convention_are_counted_not_dropped(source, tmp_path):
    (source / "notes.docx").write_bytes((source / "FRD_ClaimsFeed.docx").read_bytes())
    _, result = _sync(source, tmp_path)
    assert result["frd_ignored"] == 1
    assert not (tmp_path / "frd_raw" / "notes.docx").exists()


def test_second_sync_downloads_nothing_when_nothing_changed(source, tmp_path):
    _sync(source, tmp_path)
    _, again = _sync(source, tmp_path, now="2026-08-24T13:00:00+00:00")
    assert again["frd_downloaded"] == 0 and again["reference_downloaded"] == 0
    assert again["frd_unchanged"] == 2 and again["reference_unchanged"] == 2


def test_an_edited_document_is_picked_up_on_the_next_sync(source, tmp_path):
    _sync(source, tmp_path)
    _docx(source / "FRD_ClaimsFeed.docx",
          ["Claims feed", "Source table claims_raw", "Target CLAIM_FACT",
           "Members are loaded daily.", "New: rejects route to CLAIM_REJECT."])
    _, again = _sync(source, tmp_path, now="2026-08-24T13:00:00+00:00")
    assert again["frd_downloaded"] == 1
    # .docx is a zip — assert on the parsed text, which is what 01 consumes.
    from frdsttm.frd_parsing import normalize_to_markdown

    assert "CLAIM_REJECT" in normalize_to_markdown(
        str(tmp_path / "frd_raw" / "FRD_ClaimsFeed.docx"))


def test_a_removed_document_leaves_the_volume(source, tmp_path):
    _sync(source, tmp_path)
    (source / "FRD_MemberFeed.docx").unlink()
    _, again = _sync(source, tmp_path, now="2026-08-24T13:00:00+00:00")
    assert again["frd_removed"] == 1
    assert not (tmp_path / "frd_raw" / "FRD_MemberFeed.docx").exists()
    assert again["index"]["unpaired_references"]  # its STTM is now on its own


def test_sync_writes_the_index_and_manifest_where_the_notebooks_read_them(source, tmp_path):
    _sync(source, tmp_path)
    ref = tmp_path / "sttm_reference"
    assert (ref / CORPUS_INDEX_NAME).is_file()
    assert (ref / MANIFEST_NAME).is_file()
    assert load_corpus_index(ref)["pairs"]
    assert load_manifest(ref)["synced_at"] == NOW


def test_nothing_under_the_source_folder_is_written(source, tmp_path):
    """Read-only, mirroring the 'never writes to SharePoint' rule it stands
    in for: no .part files, no renames, no new entries."""
    before = {p.name: p.read_bytes() for p in source.iterdir()}
    _sync(source, tmp_path)
    after = {p.name: p.read_bytes() for p in source.iterdir()}
    assert after == before
