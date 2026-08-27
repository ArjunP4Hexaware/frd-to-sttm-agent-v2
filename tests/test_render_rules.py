"""Rule placement in the rendered STTM (2026-08-22).

Until this change neither renderer wrote a feed's `validation_rules` /
`recycle_rule` anywhere in the workbook — the rules survived only in the
contract JSON. These tests pin that every FRD-stated rule lands somewhere a
human sees, that attributed rules land on the row of the column they name in
the exact headers CodeGen's `extract-sttm` reads (`Comment` → value_spec,
`Recycle Flag` → `Y ( verbatim )`), and that an unattributable rule is
placed in a source-level cell and recorded in provenance — never pinned to an
arbitrary row, never dropped.

04_sttm_render.py runs driver code at import, so (like
test_render_attribution.py) the functions under test are lifted from the
committed source by name via the AST and exec'd.
"""

import ast
import re
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path

from frdsttm import standards as _std
from frdsttm import term_catalog

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from frdsttm.reference_workbooks import _n, _nl, parse_reference_workbook

RENDER_SRC = Path(__file__).resolve().parent.parent / "notebooks" / "04_sttm_render.py"
_NEEDED = {
    "GENERATOR", "_HDR_FILL", "_SEC_FILL", "_BOLD", "_style_row", "_quote_rule",
    "_column_mentions", "place_feed_rules", "_feed_level_rule_text",
    "_PER_TABLE_SRC_HEADERS", "_PER_TABLE_TGT_HEADERS", "_RECYCLE_HEADER",
    "render_sheet_per_table", "render_single_sheet", "render_contract",
    "_SEGMENT_SUFFIX", "_table_for_segment", "derive_field_mappings",
    # 2026-08-27 evening: the eval is functional; the two old names are
    # aliases of evaluate_functional, which needs its two helpers.
    "evaluate_cross_reference", "evaluate_functional", "_type_family", "_norm_ident",
    # 2026-08-27: target names come from the harvested term catalog, so
    # derive_field_mappings calls target_column, which calls _ambiguity_id.
    "target_column", "_ambiguity_id",
}


def _load():
    tree = ast.parse(RENDER_SRC.read_text(encoding="utf-8"))
    picked, found = [], set()
    for node in tree.body:
        if isinstance(node, ast.FunctionDef) and node.name in _NEEDED:
            picked.append(node)
            found.add(node.name)
        elif isinstance(node, ast.Assign):
            names = {t.id for t in node.targets if isinstance(t, ast.Name)} & _NEEDED
            if names:
                picked.append(node)
                found |= names
    assert found == _NEEDED, f"loader drifted from source: missing {_NEEDED - found}"
    ns = {"re": re, "datetime": datetime, "timezone": timezone, "Workbook": Workbook,
          "PatternFill": PatternFill, "Font": Font, "Alignment": Alignment,
          "get_column_letter": get_column_letter, "_n": _n, "_nl": _nl,
          "_tc": term_catalog, "hashlib": hashlib, "json": json, "_std": _std}
    exec(compile(ast.fix_missing_locations(ast.Module(body=picked, type_ignores=[])),
                 str(RENDER_SRC), "exec"), ns)
    return ns


R = _load()

# CodeGen's `extract-sttm` raises on any source-band header outside its
# configured synonyms (code-gen-agent config/config.yaml › extractor.header_synonyms,
# normalized: lowercased, whitespace-collapsed). Pinned here because that repo
# is not importable from this one; update BOTH when the dialect changes.
CODEGEN_SOURCE_HEADER_SYNONYMS = {
    "database column name", "database name", "client data table column name",
    "description", "sample value", "example values", "datatype", "data type",
    "null check", "phi field", "phi/pii field", "mandatory field", "mandatory",
    "comment",
}
CODEGEN_RECYCLE_HEADER_PREFIX = "recycle flag"
CODEGEN_RECYCLE_CELL_RE = re.compile(r"\bY\s*\(")

MEMBER_RULE = ("If the MEMBER_ID column is NULL, then we are rejecting the record "
               "and moving it to the reject table.")
ZIP_RULE = "If the ZIP_CODE column is NULL, then we are rejecting the record."
FEED_RULE = "Files arriving after 6 PM ET are processed on the next business day."
RECYCLE_RULE = ("Recycle Flag enabled for 7 days: check SUBS_ID against "
                "STG.SUBSCRIBER_MASTER; if available process it, else load to the "
                "Recycle table with Recycle Flag enabled for 7 days.")


def _field(col, comment=""):
    return {"source_column": col, "description": f"{col} desc", "sample": "", "datatype": "String",
            "nullable": False, "phi": False, "mandatory": True, "comment": comment,
            "segment": "", "business_rule": "",
            "stage": {"catalog": "c", "schema": "stg", "table": "T_MEMBER", "column": col,
                      "datatype": "String"},
            "standard": {"catalog": "c", "schema": "std", "table": "T_MEMBER", "column": col,
                         "datatype": "String"}}


def _feed(rules, recycle, columns=("MEMBER_ID", "ZIP_CODE", "SUBS_ID", "MEMBER_ID_2")):
    return {"feed_name": "member feed", "source_system": "Vendor",
            "file_name_patterns": ["MEMBER_YYYYMMDD.txt"], "landing_location": "/landing",
            "frequency": "daily",
            "stage_target": {"catalog": "c", "schema": "stg", "tables": ["T_MEMBER"]},
            "standard_target": {"catalog": "c", "schema": "std", "tables": ["T_MEMBER"]},
            "validation_rules": list(rules), "recycle_rule": recycle,
            "fields": [_field(c) for c in columns]}


def _contract(feed):
    return {"status": "PASS", "generated_from_frd": "demo.docx", "feeds": [feed],
            "_provenance": {}}


def _rows(ws):
    return [list(r) for r in ws.iter_rows(values_only=True)]


# --------------------------------------------------------------------------- #
# place_feed_rules — the pure decision
# --------------------------------------------------------------------------- #

def test_rules_attributed_to_every_row_whose_column_they_name():
    p = R["place_feed_rules"](_feed([MEMBER_RULE, ZIP_RULE, FEED_RULE], RECYCLE_RULE))
    assert p["by_column"] == {"MEMBER_ID": [MEMBER_RULE], "ZIP_CODE": [ZIP_RULE]}
    assert p["unattributed"] == [FEED_RULE]
    assert p["recycle"] == {"text": RECYCLE_RULE, "column": "SUBS_ID"}


def test_whole_word_match_does_not_hit_longer_column_names():
    # MEMBER_ID must not attribute to MEMBER_ID_2, and a rule naming only
    # MEMBER_ID_2 must not attribute to MEMBER_ID.
    p = R["place_feed_rules"](_feed(["Reject when MEMBER_ID_2 is blank."], None))
    assert p["by_column"] == {"MEMBER_ID_2": ["Reject when MEMBER_ID_2 is blank."]}
    assert R["_column_mentions"]("member_id is null", ["MEMBER_ID", "MEMBER_ID_2"]) == ["MEMBER_ID"]


def test_recycle_rule_naming_no_column_is_not_pinned_to_a_row():
    p = R["place_feed_rules"](_feed([], "Recycle unmatched records for 7 days."))
    assert p["recycle"] == {"text": "Recycle unmatched records for 7 days.", "column": None}
    assert "Recycle rule:" in R["_feed_level_rule_text"](p)


def test_no_rules_no_placement():
    p = R["place_feed_rules"](_feed([], None))
    assert p == {"by_column": {}, "recycle": None, "unattributed": []}
    assert R["_feed_level_rule_text"](p) == ""


# --------------------------------------------------------------------------- #
# sheet-per-table dialect — the one CodeGen's extractor reads
# --------------------------------------------------------------------------- #

def test_per_table_source_headers_are_all_known_to_codegen():
    assert {_nl(h) for h in R["_PER_TABLE_SRC_HEADERS"]} <= CODEGEN_SOURCE_HEADER_SYNONYMS
    assert "comment" in {_nl(h) for h in R["_PER_TABLE_SRC_HEADERS"]}  # → value_spec
    assert _nl(R["_RECYCLE_HEADER"]).startswith(CODEGEN_RECYCLE_HEADER_PREFIX)


def test_per_table_places_rules_on_the_named_rows(tmp_path):
    out = tmp_path / "out.xlsx"
    contract = _contract(_feed([MEMBER_RULE, ZIP_RULE, FEED_RULE], RECYCLE_RULE))
    R["render_sheet_per_table"](contract, str(out))
    wb = load_workbook(out)
    ws = wb["MAPPING-T_MEMBER"]
    rows = _rows(ws)
    bands, headers, data = rows[0], rows[1], rows[2:]

    # bands where CodeGen expects them: row 1, source | stage | standard in order
    n_src = len(R["_PER_TABLE_SRC_HEADERS"])
    n_tgt = len(R["_PER_TABLE_TGT_HEADERS"])
    assert bands[0] == "Source File Layout"
    assert bands[n_src] == "Stage Layer"
    assert bands[n_src + n_tgt + 1] == "Standard Layer"
    # every header inside the source band is one CodeGen knows
    assert {_nl(h) for h in headers[:n_src]} <= CODEGEN_SOURCE_HEADER_SYNONYMS

    comment_i = headers.index("Comment")
    recycle_i = headers.index(R["_RECYCLE_HEADER"])
    assert recycle_i > n_src + 2 * n_tgt  # beyond the standard band, outside every band
    by_col = {r[0]: r for r in data}
    assert by_col["MEMBER_ID"][comment_i] == MEMBER_RULE
    assert by_col["ZIP_CODE"][comment_i] == ZIP_RULE
    assert not by_col["SUBS_ID"][comment_i]
    assert not by_col["MEMBER_ID_2"][comment_i]

    # exactly one flagged recycle row, in the Y ( verbatim ) shape, on SUBS_ID
    flagged = [r for r in data if r[recycle_i]]
    assert [r[0] for r in flagged] == ["SUBS_ID"]
    cell = flagged[0][recycle_i]
    assert CODEGEN_RECYCLE_CELL_RE.search(cell)
    assert RECYCLE_RULE in cell

    # the rule naming no column is in the feed-level cell, not lost
    fd = _rows(wb["FILE_DETAILS"])
    desc_i = fd[0].index("File Description")
    assert fd[1][desc_i] == FEED_RULE


def test_per_table_no_recycle_rule_means_no_recycle_column(tmp_path):
    out = tmp_path / "out.xlsx"
    R["render_sheet_per_table"](_contract(_feed([MEMBER_RULE], None)), str(out))
    headers = _rows(load_workbook(out)["MAPPING-T_MEMBER"])[1]
    assert R["_RECYCLE_HEADER"] not in headers
    assert "Comment" in headers


def test_per_table_unattributed_recycle_goes_to_feed_level_cell_only(tmp_path):
    out = tmp_path / "out.xlsx"
    rec = "Recycle unmatched records for 7 days."
    R["render_sheet_per_table"](_contract(_feed([], rec)), str(out))
    wb = load_workbook(out)
    rows = _rows(wb["MAPPING-T_MEMBER"])
    recycle_i = rows[1].index(R["_RECYCLE_HEADER"])
    assert all(not r[recycle_i] for r in rows[2:])  # no row is flagged — no guess
    fd = _rows(wb["FILE_DETAILS"])
    assert rec in fd[1][fd[0].index("File Description")]


def test_per_table_keeps_template_comment_behind_the_rule(tmp_path):
    out = tmp_path / "out.xlsx"
    feed = _feed([MEMBER_RULE], None)
    feed["fields"][0]["comment"] = "template note"
    R["render_sheet_per_table"](_contract(feed), str(out))
    rows = _rows(load_workbook(out)["MAPPING-T_MEMBER"])
    cell = rows[2][rows[1].index("Comment")]
    assert cell == f"{MEMBER_RULE}\ntemplate note"


def test_rendered_per_table_workbook_round_trips_through_our_own_parser(tmp_path):
    """A rendered workbook uploaded back as an approved STTM must parse as a
    dictionary whose fields carry the rule (comment) and whose recycle marker
    is detected — the same signals resolve_attribution() reads from golden
    workbooks."""
    out = tmp_path / "out.xlsx"
    R["render_sheet_per_table"](_contract(_feed([MEMBER_RULE], RECYCLE_RULE)), str(out))
    d = parse_reference_workbook(str(out))
    assert d["dialect"] == "sheet_per_table"
    (key, feed), = d["feeds"].items()
    assert feed["recycle_note"]
    by_col = {f["source_column"]: f for f in feed["fields"]}
    assert by_col["MEMBER_ID"]["comment"] == MEMBER_RULE
    assert by_col["SUBS_ID"]["comment"] == ""


# --------------------------------------------------------------------------- #
# render_contract — provenance
# --------------------------------------------------------------------------- #

def test_render_contract_records_placement_in_provenance(tmp_path):
    out = tmp_path / "out.xlsx"
    contract = _contract(_feed([MEMBER_RULE, FEED_RULE], RECYCLE_RULE))
    dictionary = {"dialect": "sheet_per_table", "feeds": {}, "meta": {}, "feed_sources": {}}
    R["render_contract"](contract, dictionary, str(out))
    (p,) = contract["_provenance"]["rule_placement"]
    assert p["feed"] == "member feed" and p["dialect"] == "sheet_per_table"
    assert p["by_column"] == {"MEMBER_ID": [MEMBER_RULE]}
    assert p["unattributed"] == [FEED_RULE]
    assert p["recycle"]["column"] == "SUBS_ID"
    assert out.exists()


# --------------------------------------------------------------------------- #
# single-sheet (CAQH) dialect
# --------------------------------------------------------------------------- #

def test_single_sheet_places_rules_in_business_rule_and_metadata(tmp_path):
    out = tmp_path / "out.xlsx"
    contract = _contract(_feed([MEMBER_RULE, FEED_RULE], RECYCLE_RULE))
    R["render_single_sheet"](contract, str(out), sheet_name="mapping")
    rows = _rows(load_workbook(out)["mapping"])
    meta = {r[0]: r[1] for r in rows if r and r[0] and r[1] is not None}
    assert meta["Recycle rule"] == RECYCLE_RULE
    assert meta["Validation rules (source-level)"] == FEED_RULE
    header_r = next(i for i, r in enumerate(rows) if r and r[0] == "#")
    headers = rows[header_r]
    br_i, name_i = headers.index("Business Rule"), headers.index("Field Name")
    by_col = {r[name_i]: r for r in rows[header_r + 1:]}
    assert by_col["MEMBER_ID"][br_i] == MEMBER_RULE
    assert not by_col["ZIP_CODE"][br_i]


def test_single_sheet_without_rules_adds_no_metadata_rows(tmp_path):
    out = tmp_path / "out.xlsx"
    R["render_single_sheet"](_contract(_feed([], None)), str(out), sheet_name="mapping")
    keys = {r[0] for r in _rows(load_workbook(out)["mapping"]) if r}
    assert "Recycle rule" not in keys and "Validation rules (source-level)" not in keys


# --------------------------------------------------------------------------- #
# audit rows (source "NA") — template convention, carried through end to end
# --------------------------------------------------------------------------- #

def _template_workbook(path, audit=(("SRC_FILE", "string"), ("LOAD_TS", "timestamp"))):
    """A sheet-per-table reference workbook with two real columns and
    trailing NA audit rows, the shape of every client workbook in this dialect."""
    from openpyxl import Workbook as _WB
    wb = _WB()
    ws = wb.active
    ws.title = "MAPPING-T_MEMBER"
    src = ["Database Column Name", "Description", "Datatype", "Null Check", "Comment"]
    tgt = ["Catalog", "Schema", "TableName", "ColumnName", "Datatype"]
    ws.append(["Source File Layout"] + [""] * 4 + ["Stage Layer"] + [""] * 4 + ["Standard Layer"] + [""] * 4)
    ws.append(src + tgt + tgt)
    for col in ("MEMBER_ID", "ZIP_CODE"):
        ws.append([col, f"{col} desc", "String", "Not Null", ""]
                  + ["c", "stg", "T_MEMBER", col, "String"] + ["c", "std", "T_MEMBER", col, "String"])
    for col, dt in audit:
        ws.append(["NA", "audit", dt, "", ""]
                  + ["c", "stg", "T_MEMBER", col, dt] + ["c", "std", "T_MEMBER", col, dt])
    wb.save(path)


def test_parser_flags_na_rows_as_audit_and_keeps_them(tmp_path):
    p = tmp_path / "ref.xlsx"
    _template_workbook(p)
    (feed,) = parse_reference_workbook(str(p))["feeds"].values()
    assert [f["source_column"] for f in feed["fields"]] == ["MEMBER_ID", "ZIP_CODE", "NA", "NA"]
    assert [f["audit"] for f in feed["fields"]] == [False, False, True, True]
    assert feed["ref_targets"][2]["stage"]["column"] == "SRC_FILE"


def test_derive_takes_audit_identity_from_template_not_1_to_1(tmp_path):
    p = tmp_path / "ref.xlsx"
    _template_workbook(p)
    dictionary = parse_reference_workbook(str(p))
    (key,) = dictionary["feeds"]
    contract = _contract({**_feed([], None), "fields": []})
    R["derive_field_mappings"](contract, dictionary, {0: key})
    fields = contract["feeds"][0]["fields"]
    real = [f for f in fields if not f.get("audit")]
    audit = [f for f in fields if f.get("audit")]
    assert [f["stage"]["column"] for f in real] == ["MEMBER_ID", "ZIP_CODE"]   # 1:1 rule, unchanged
    assert [(f["stage"]["column"], f["stage"]["datatype"]) for f in audit] == \
        [("SRC_FILE", "string"), ("LOAD_TS", "timestamp")]
    assert [f["standard"]["column"] for f in audit] == ["SRC_FILE", "LOAD_TS"]
    # tables still come from the contract's own targets, not the template
    assert {f["stage"]["table"] for f in fields} == {"T_MEMBER"}
    assert {f["stage"]["schema"] for f in fields} == {"stg"}


def test_audit_rows_render_with_target_identity_and_take_no_rules(tmp_path):
    p = tmp_path / "ref.xlsx"
    _template_workbook(p)
    dictionary = parse_reference_workbook(str(p))
    (key,) = dictionary["feeds"]
    contract = _contract({**_feed(["NA values are rejected; MEMBER_ID must be present."], None),
                          "fields": []})
    R["derive_field_mappings"](contract, dictionary, {0: key})
    out = tmp_path / "out.xlsx"
    R["render_contract"](contract, dictionary, str(out))
    rows = _rows(load_workbook(out)["MAPPING-T_MEMBER"])
    headers = rows[1]
    comment_i = headers.index("Comment")
    stage_col_i = len(R["_PER_TABLE_SRC_HEADERS"]) + 2  # Schema, TableName, ColumnName
    tail = rows[-2:]
    assert [r[0] for r in tail] == ["NA", "NA"]
    assert [r[stage_col_i] for r in tail] == ["SRC_FILE", "LOAD_TS"]
    assert all(not r[comment_i] for r in tail)          # never a rule target
    (p,) = contract["_provenance"]["rule_placement"]
    assert list(p["by_column"]) == ["MEMBER_ID"]


def test_cross_eval_aligns_audit_rows_by_target_column(tmp_path):
    p = tmp_path / "ref.xlsx"
    _template_workbook(p)
    dictionary = parse_reference_workbook(str(p))
    (key,) = dictionary["feeds"]
    contract = _contract({**_feed([], None), "fields": []})
    R["derive_field_mappings"](contract, dictionary, {0: key})
    ev = R["evaluate_cross_reference"](contract, dictionary, {0: key})
    # every reference ROW (2 real + 2 audit) is structurally correct — the eval
    # is functional since 2026-08-27 evening: rows, not cells, and audit rows
    # still align by their target column rather than colliding on source "NA"
    assert ev["totals"]["cells"] == 4
    assert ev["totals"]["pct"] == 100.0, ev["feeds"][0]["sample_diffs"]
    assert ev["totals"]["structural"] == 0
